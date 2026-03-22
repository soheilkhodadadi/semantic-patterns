"""Build standalone journal-style DOCX tables for preliminary delivery."""

from __future__ import annotations

import argparse
import csv
import math
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from semantic_ai_washing.analysis.delivery_table_payloads import (
    sig_stars,
    summarize_table_1,
    summarize_table_2_timing_focus,
    summarize_table_2_timing_focus_counts,
    summarize_table_3_timing_composition,
    summarize_table_3_timing_composition_counts,
    summarize_table_4_actionable_patent_timing,
    summarize_table_4b_speculative_patent_timing,
    summarize_table_5_credibility_metrics_tplus1,
    summarize_table_5b_credibility_metrics_tplus2,
    summarize_table_6_as_patent_mismatch_tplus1,
    summarize_table_6b_as_patent_mismatch_tplus2,
)

REPO_ROOT = Path(__file__).resolve().parents[3]

DEFAULT_REG_READY_PANEL = "data/processed/panel/panel_reg_ready_ever_speaker_2016_2024_v1.csv"
DEFAULT_PORTFOLIO_COEFFS = (
    "results/01_baseline/tables_2016_2024_applied_v2_legalnorm_unique/portfolio_coefficients.csv"
)
DEFAULT_OUTPUT_DIR = "output/doc/delivery_tables_v1"


def _split_csv_arg(raw_value: str) -> set[str]:
    return {item.strip().lower() for item in raw_value.split(",") if item.strip()}


def _load_coeff_lookup(
    coeff_path: str | Path,
) -> dict[tuple[str, str], tuple[float, float, float | None, int | None]]:
    lookup: dict[tuple[str, str], tuple[float, float, float | None, int | None]] = {}
    with Path(coeff_path).open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            model = row.get("model", "")
            term = row.get("term", "")
            try:
                coef = float(row["coef"])
                se = float(row["se"])
            except (TypeError, ValueError):
                continue
            try:
                pvalue = float(row["p"]) if row.get("p", "") not in {"", None} else None
            except ValueError:
                pvalue = None
            try:
                nobs = int(float(row["N"])) if row.get("N", "") not in {"", None} else None
            except ValueError:
                nobs = None
            lookup[(model, term)] = (coef, se, pvalue, nobs)
    return lookup


def _summarize_conditional_appendix_table(coeff_path: str | Path) -> dict[str, object]:
    lookup = _load_coeff_lookup(coeff_path)
    models = [
        {"number": "(1)", "label": "Actionable-only"},
        {"number": "(2)", "label": "Speculative-only"},
        {"number": "(3)", "label": "Share model"},
    ]
    row_specs = [
        ("Actionable disclosure", ["has_actionable", None, None]),
        ("Speculative-only disclosure", [None, "has_spec_only", None]),
        ("Actionable share", [None, None, "ActShare"]),
        ("Speculative share", [None, None, "SpecShare"]),
    ]
    body_rows: list[dict[str, object]] = []
    footer_n: list[str] = []
    model_ids = [
        "portfolio_lpm_anypat_k1_actionable_only_fe",
        "portfolio_lpm_anypat_k1_speculative_only_fe",
        "portfolio_lpm_anypat_k1_shares_fe",
    ]
    for label, terms in row_specs:
        coef_cells: list[str] = []
        se_cells: list[str] = []
        for model_id, term in zip(model_ids, terms, strict=True):
            if term is None:
                coef_cells.append("")
                se_cells.append("")
                continue
            coef, se, pvalue, nobs = lookup.get(
                (model_id, term), (float("nan"), float("nan"), None, None)
            )
            coef_cells.append(f"{coef:.3f}{sig_stars(pvalue)}" if not math.isnan(coef) else "")
            se_cells.append(f"({se:.3f})" if not math.isnan(se) else "")
            footer_n.append(f"{nobs:,}" if nobs is not None else "")
        body_rows.append({"label": label, "cells": coef_cells, "kind": "coef"})
        body_rows.append({"label": "", "cells": se_cells, "kind": "se"})

    note = (
        "This appendix table presents the earlier conditional validation result estimated on the narrower AI-speaking panel. "
        "The dependent variable is an indicator for whether the firm records any AI patent in t+1. Control variables include size, leverage, cash/assets, R&D/assets, CAPX/assets, ROA, sales growth, and employees. "
        "Firm and year fixed effects are included in all columns. Standard errors clustered at the firm level are shown in parentheses. Constants are omitted. (* p<0.1, ** p<0.05, *** p<0.01)."
    )
    return {
        "title": "Appendix Table. Conditional Disclosure Composition and Future AI Patenting",
        "note": note,
        "dependent_label": "Any AI patent in t+1",
        "models": models,
        "body_rows": body_rows,
        "footer_rows": [
            {"label": "Controls", "cells": ["Y", "Y", "Y"]},
            {"label": "Firm FE", "cells": ["Y", "Y", "Y"]},
            {"label": "Year FE", "cells": ["Y", "Y", "Y"]},
            {"label": "Observations", "cells": footer_n[:3] if footer_n else ["", "", ""]},
        ],
    }


def _set_run_font(run, *, size: float = 11, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def _clear_cell(cell) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)


def _write_cell(
    cell,
    text: str,
    *,
    align: WD_ALIGN_PARAGRAPH,
    size: float = 10.5,
    bold: bool = False,
    italic: bool = False,
) -> None:
    _clear_cell(cell)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run(text)
    _set_run_font(run, size=size, bold=bold, italic=italic)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _set_cell_border(cell, **kwargs: dict[str, str]) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("left", "top", "right", "bottom"):
        edge_data = kwargs.get(edge)
        if edge_data is None:
            continue
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key, value in edge_data.items():
            element.set(qn(f"w:{key}"), str(value))


def _set_table_no_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_borders = tbl_pr.first_child_found_in("w:tblBorders")
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = tbl_borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tbl_borders.append(element)
        element.set(qn("w:val"), "nil")


def _apply_row_rule(row, *, top: bool = False, bottom: bool = False, size: str = "8") -> None:
    border = {"val": "single", "sz": size, "space": "0", "color": "000000"}
    for cell in row.cells:
        kwargs: dict[str, dict[str, str]] = {}
        if top:
            kwargs["top"] = border
        if bottom:
            kwargs["bottom"] = border
        _set_cell_border(cell, **kwargs)


def _set_document_defaults(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)


def _add_title(document: Document, title: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(title)
    _set_run_font(run, size=16, bold=True)


def _add_note(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run(text)
    _set_run_font(run, size=11)


def _build_table_1_doc(panel_path: str | Path, output_path: str | Path) -> None:
    payload = summarize_table_1(panel_path)
    document = Document()
    _set_document_defaults(document)
    _add_title(document, str(payload["title"]))
    _add_note(document, str(payload["note"]))

    headers: list[str] = payload["headers"]  # type: ignore[assignment]
    rows: list[dict[str, str]] = payload["rows"]  # type: ignore[assignment]
    table = document.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_no_borders(table)

    widths = [2.6, 0.58, 0.72, 0.52, 0.52, 0.52, 0.52, 0.52, 0.58]
    for index, width in enumerate(widths):
        for row in table.rows:
            row.cells[index].width = Inches(width)

    for idx, header in enumerate(headers):
        _write_cell(
            table.rows[0].cells[idx],
            header,
            align=WD_ALIGN_PARAGRAPH.CENTER if idx > 0 else WD_ALIGN_PARAGRAPH.LEFT,
            size=10.5,
            bold=True,
        )
    _apply_row_rule(table.rows[0], top=True, bottom=True)

    for row_idx, row_payload in enumerate(rows, start=1):
        for col_idx, header in enumerate(headers):
            align = WD_ALIGN_PARAGRAPH.LEFT if col_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            _write_cell(
                table.rows[row_idx].cells[col_idx], row_payload[header], align=align, size=10.5
            )
    _apply_row_rule(table.rows[-1], bottom=True)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def _build_panel_timing_doc(payload: dict[str, object], output_path: str | Path) -> None:
    document = Document()
    _set_document_defaults(document)
    _add_title(document, str(payload["title"]))
    _add_note(document, str(payload["note"]))

    models: list[dict[str, str]] = payload["models"]  # type: ignore[assignment]
    panels: list[dict[str, object]] = payload["panels"]  # type: ignore[assignment]
    total_rows = 3
    for panel in panels:
        total_rows += (1 if panel.get("heading") else 0) + 2 + len(panel["footer_rows"])

    cols = len(models) + 1
    table = document.add_table(rows=total_rows, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_no_borders(table)

    widths = [2.55] + [0.85] * len(models)
    for index, width in enumerate(widths):
        for row in table.rows:
            row.cells[index].width = Inches(width)

    _write_cell(
        table.rows[0].cells[0], "Variable", align=WD_ALIGN_PARAGRAPH.LEFT, size=10.5, bold=True
    )
    merged = table.rows[0].cells[1].merge(table.rows[0].cells[-1])
    _write_cell(
        merged,
        str(payload["dependent_label"]),
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=10.5,
        bold=True,
    )
    _apply_row_rule(table.rows[0], top=True)

    _write_cell(table.rows[1].cells[0], "", align=WD_ALIGN_PARAGRAPH.LEFT, size=10.5)
    for idx, model in enumerate(models, start=1):
        _write_cell(
            table.rows[1].cells[idx], model["label"], align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5
        )

    _write_cell(table.rows[2].cells[0], "", align=WD_ALIGN_PARAGRAPH.LEFT, size=10.5)
    for idx, model in enumerate(models, start=1):
        _write_cell(
            table.rows[2].cells[idx], model["number"], align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5
        )
    _apply_row_rule(table.rows[2], bottom=True)

    cursor = 3
    for panel in panels:
        heading = panel.get("heading")
        if heading:
            merged = table.rows[cursor].cells[0].merge(table.rows[cursor].cells[-1])
            _write_cell(merged, str(heading), align=WD_ALIGN_PARAGRAPH.LEFT, size=10.5, bold=True)
            _apply_row_rule(table.rows[cursor], top=True)
            cursor += 1

        _write_cell(
            table.rows[cursor].cells[0],
            str(panel["label"]),
            align=WD_ALIGN_PARAGRAPH.LEFT,
            size=10.5,
        )
        for idx, value in enumerate(panel["coef_cells"], start=1):
            _write_cell(
                table.rows[cursor].cells[idx],
                str(value),
                align=WD_ALIGN_PARAGRAPH.CENTER,
                size=10.5,
            )
        cursor += 1

        _write_cell(
            table.rows[cursor].cells[0], "", align=WD_ALIGN_PARAGRAPH.LEFT, size=10.5, italic=True
        )
        for idx, value in enumerate(panel["se_cells"], start=1):
            _write_cell(
                table.rows[cursor].cells[idx],
                str(value),
                align=WD_ALIGN_PARAGRAPH.CENTER,
                size=10.5,
                italic=True,
            )
        cursor += 1

        for footer in panel["footer_rows"]:  # type: ignore[index]
            _write_cell(
                table.rows[cursor].cells[0],
                str(footer["label"]),
                align=WD_ALIGN_PARAGRAPH.LEFT,
                size=10.5,
            )
            for idx, value in enumerate(footer["cells"], start=1):
                _write_cell(
                    table.rows[cursor].cells[idx],
                    str(value),
                    align=WD_ALIGN_PARAGRAPH.CENTER,
                    size=10.5,
                )
            cursor += 1

        _apply_row_rule(table.rows[cursor - 1], bottom=True)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def _build_conditional_appendix_doc(coeff_path: str | Path, output_path: str | Path) -> None:
    payload = _summarize_conditional_appendix_table(coeff_path)
    document = Document()
    _set_document_defaults(document)
    _add_title(document, str(payload["title"]))
    _add_note(document, str(payload["note"]))

    body_rows: list[dict[str, object]] = payload["body_rows"]  # type: ignore[assignment]
    footer_rows: list[dict[str, object]] = payload["footer_rows"]  # type: ignore[assignment]
    models: list[dict[str, str]] = payload["models"]  # type: ignore[assignment]

    total_rows = 2 + len(body_rows) + len(footer_rows)
    table = document.add_table(rows=total_rows, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_no_borders(table)

    widths = [2.85, 1.05, 1.05, 1.05]
    for index, width in enumerate(widths):
        for row in table.rows:
            row.cells[index].width = Inches(width)

    _write_cell(
        table.rows[0].cells[0], "Variable", align=WD_ALIGN_PARAGRAPH.LEFT, size=10.5, bold=True
    )
    merged = table.rows[0].cells[1].merge(table.rows[0].cells[3])
    _write_cell(
        merged,
        str(payload["dependent_label"]),
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=10.5,
        bold=True,
    )
    _apply_row_rule(table.rows[0], top=True)

    _write_cell(table.rows[1].cells[0], "", align=WD_ALIGN_PARAGRAPH.LEFT, size=10.5)
    for idx, model in enumerate(models, start=1):
        _write_cell(
            table.rows[1].cells[idx], model["number"], align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5
        )
    _apply_row_rule(table.rows[1], bottom=True)

    cursor = 2
    for row_payload in body_rows:
        is_se = row_payload.get("kind") == "se"
        _write_cell(
            table.rows[cursor].cells[0],
            str(row_payload["label"]),
            align=WD_ALIGN_PARAGRAPH.LEFT,
            size=10.5,
            italic=is_se,
        )
        for idx, value in enumerate(row_payload["cells"], start=1):
            _write_cell(
                table.rows[cursor].cells[idx],
                str(value),
                align=WD_ALIGN_PARAGRAPH.CENTER,
                size=10.5,
                italic=is_se,
            )
        cursor += 1
    _apply_row_rule(table.rows[cursor - 1], bottom=True)

    for footer in footer_rows:
        _write_cell(
            table.rows[cursor].cells[0],
            str(footer["label"]),
            align=WD_ALIGN_PARAGRAPH.LEFT,
            size=10.5,
        )
        for idx, value in enumerate(footer["cells"], start=1):
            _write_cell(
                table.rows[cursor].cells[idx],
                str(value),
                align=WD_ALIGN_PARAGRAPH.CENTER,
                size=10.5,
            )
        cursor += 1
    _apply_row_rule(table.rows[cursor - 1], bottom=True)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def _build_row_matrix_doc(payload: dict[str, object], output_path: str | Path) -> None:
    document = Document()
    _set_document_defaults(document)
    _add_title(document, str(payload["title"]))
    _add_note(document, str(payload["note"]))

    body_rows: list[dict[str, object]] = payload["body_rows"]  # type: ignore[assignment]
    footer_rows: list[dict[str, object]] = payload["footer_rows"]  # type: ignore[assignment]
    models: list[dict[str, str]] = payload["models"]  # type: ignore[assignment]

    total_rows = 3 + len(body_rows) + len(footer_rows)
    table = document.add_table(rows=total_rows, cols=len(models) + 1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_no_borders(table)

    widths = [2.8] + [0.95] * len(models)
    for index, width in enumerate(widths):
        for row in table.rows:
            row.cells[index].width = Inches(width)

    _write_cell(
        table.rows[0].cells[0], "Variable", align=WD_ALIGN_PARAGRAPH.LEFT, size=10.5, bold=True
    )
    merged = table.rows[0].cells[1].merge(table.rows[0].cells[-1])
    _write_cell(
        merged,
        str(payload["dependent_label"]),
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=10.5,
        bold=True,
    )
    _apply_row_rule(table.rows[0], top=True)

    _write_cell(table.rows[1].cells[0], "", align=WD_ALIGN_PARAGRAPH.LEFT, size=10.5)
    for idx, model in enumerate(models, start=1):
        _write_cell(
            table.rows[1].cells[idx], model["label"], align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5
        )

    _write_cell(table.rows[2].cells[0], "", align=WD_ALIGN_PARAGRAPH.LEFT, size=10.5)
    for idx, model in enumerate(models, start=1):
        _write_cell(
            table.rows[2].cells[idx], model["number"], align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5
        )
    _apply_row_rule(table.rows[2], bottom=True)

    cursor = 3
    for row_payload in body_rows:
        is_se = row_payload.get("kind") == "se"
        _write_cell(
            table.rows[cursor].cells[0],
            str(row_payload["label"]),
            align=WD_ALIGN_PARAGRAPH.LEFT,
            size=10.5,
            italic=is_se,
        )
        for idx, value in enumerate(row_payload["cells"], start=1):
            _write_cell(
                table.rows[cursor].cells[idx],
                str(value),
                align=WD_ALIGN_PARAGRAPH.CENTER,
                size=10.5,
                italic=is_se,
            )
        cursor += 1
    _apply_row_rule(table.rows[cursor - 1], bottom=True)

    for footer in footer_rows:
        _write_cell(
            table.rows[cursor].cells[0],
            str(footer["label"]),
            align=WD_ALIGN_PARAGRAPH.LEFT,
            size=10.5,
        )
        for idx, value in enumerate(footer["cells"], start=1):
            _write_cell(
                table.rows[cursor].cells[idx],
                str(value),
                align=WD_ALIGN_PARAGRAPH.CENTER,
                size=10.5,
            )
        cursor += 1
    _apply_row_rule(table.rows[cursor - 1], bottom=True)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--reg-ready-panel", default=DEFAULT_REG_READY_PANEL)
    parser.add_argument("--portfolio-coeffs", default=DEFAULT_PORTFOLIO_COEFFS)
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR)
    parser.add_argument(
        "--tables",
        default="table1,table2_timing_focus,table3_timing_composition",
        help=(
            "Comma-separated list of standalone tables to build: table1, table2_timing_focus, "
            "table2_timing_focus_count, table3_timing_composition, "
            "table3_timing_composition_count, table4_actionable_patent_timing, "
            "table4b_speculative_patent_timing, table5_credibility_metrics_tplus1, "
            "table5b_credibility_metrics_tplus2, table6_as_patent_mismatch_tplus1, "
            "table6b_as_patent_mismatch_tplus2, table2_conditional_appendix"
        ),
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    selected = _split_csv_arg(args.tables)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    if "table1" in selected:
        _build_table_1_doc(
            args.reg_ready_panel, output_dir / "table_1_summary_statistics_prelim_v1.docx"
        )

    if "table2_timing_focus" in selected:
        _build_panel_timing_doc(
            summarize_table_2_timing_focus(args.reg_ready_panel),
            output_dir / "table_2_ai_focus_timing_prelim_v1.docx",
        )

    if "table2_timing_focus_count" in selected:
        _build_panel_timing_doc(
            summarize_table_2_timing_focus_counts(args.reg_ready_panel),
            output_dir / "table_2b_ai_focus_timing_counts_prelim_v1.docx",
        )

    if "table3_timing_composition" in selected:
        _build_panel_timing_doc(
            summarize_table_3_timing_composition(args.reg_ready_panel),
            output_dir / "table_3_disclosure_composition_timing_prelim_v1.docx",
        )

    if "table3_timing_composition_count" in selected:
        _build_panel_timing_doc(
            summarize_table_3_timing_composition_counts(args.reg_ready_panel),
            output_dir / "table_3b_disclosure_composition_timing_counts_prelim_v1.docx",
        )

    if "table4_actionable_patent_timing" in selected:
        _build_row_matrix_doc(
            summarize_table_4_actionable_patent_timing(args.reg_ready_panel),
            output_dir / "table_4_actionable_patent_timing_prelim_v1.docx",
        )

    if "table4b_speculative_patent_timing" in selected:
        _build_row_matrix_doc(
            summarize_table_4b_speculative_patent_timing(args.reg_ready_panel),
            output_dir / "table_4b_speculative_patent_timing_prelim_v1.docx",
        )

    if "table5_credibility_metrics_tplus1" in selected:
        _build_row_matrix_doc(
            summarize_table_5_credibility_metrics_tplus1(args.reg_ready_panel),
            output_dir / "table_5_credibility_metrics_tplus1_prelim_v1.docx",
        )

    if "table5b_credibility_metrics_tplus2" in selected:
        _build_row_matrix_doc(
            summarize_table_5b_credibility_metrics_tplus2(args.reg_ready_panel),
            output_dir / "table_5b_credibility_metrics_tplus2_prelim_v1.docx",
        )

    if "table6_as_patent_mismatch_tplus1" in selected:
        _build_row_matrix_doc(
            summarize_table_6_as_patent_mismatch_tplus1(args.reg_ready_panel),
            output_dir / "table_6_as_patent_mismatch_tplus1_prelim_v1.docx",
        )

    if "table6b_as_patent_mismatch_tplus2" in selected:
        _build_row_matrix_doc(
            summarize_table_6b_as_patent_mismatch_tplus2(args.reg_ready_panel),
            output_dir / "table_6b_as_patent_mismatch_tplus2_prelim_v1.docx",
        )

    if "table2_conditional_appendix" in selected:
        _build_conditional_appendix_doc(
            args.portfolio_coeffs,
            output_dir / "table_2_core_patent_validation_prelim_v1.docx",
        )

    print(f"[delivery-table-docs] wrote selected DOCX tables under {output_dir}")


if __name__ == "__main__":
    main()
