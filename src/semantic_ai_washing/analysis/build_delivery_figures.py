"""Build standalone delivery-phase figures and review DOCX wrappers."""

from __future__ import annotations

import argparse
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

REPO_ROOT = Path(__file__).resolve().parents[3]
DEFAULT_PANEL = "data/processed/panel/panel_ai_patents_controls_ever_speaker_2016_2024_v1.csv"
DEFAULT_FIG_DIR = "output/figures/delivery_figures_v1"
DEFAULT_DOC_DIR = "output/doc/delivery_figures_v1"


def _set_doc_defaults(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(11)


def _add_title(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    r.font.size = Pt(16)
    r.bold = True


def _add_note(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    r.font.size = Pt(11)


def _save_figure_docx(title: str, note: str, image_path: Path, output_path: Path) -> None:
    doc = Document()
    _set_doc_defaults(doc)
    _add_title(doc, title)
    _add_note(doc, note)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image_path), width=Inches(6.6))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))


def _base_style() -> None:
    plt.style.use("seaborn-v0_8-whitegrid")
    plt.rcParams.update(
        {
            "font.family": "serif",
            "font.serif": ["Times New Roman", "Times", "DejaVu Serif"],
            "axes.titlesize": 14,
            "axes.labelsize": 11,
            "xtick.labelsize": 10,
            "ytick.labelsize": 10,
            "legend.fontsize": 10,
        }
    )


def _style_axes(ax) -> None:
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.grid(True, axis="y", color="#d9d9d9", linewidth=0.7)
    ax.grid(False, axis="x")


def _load_panel(panel_path: str | Path) -> pd.DataFrame:
    return pd.read_csv(panel_path, low_memory=False)


def build_figure_1(panel_path: str | Path, output_dir: str | Path, doc_dir: str | Path) -> tuple[Path, Path]:
    df = _load_panel(panel_path)
    yearly = (
        df.groupby("year", as_index=False)
        .agg(
            ai_total_mean=("ai_total", "mean"),
            n_A_mean=("n_A", "mean"),
            n_S_mean=("n_S", "mean"),
            any_ai_share=("any_ai_talk", "mean"),
        )
        .sort_values("year")
    )
    talker = df.loc[df["any_ai_talk"] > 0].copy()
    comp = (
        talker.groupby("year", as_index=False)
        .agg(act_share_mean=("ActShare", "mean"), spec_share_mean=("SpecShare", "mean"))
        .sort_values("year")
    )

    _base_style()
    fig, axes = plt.subplots(2, 1, figsize=(8.2, 8.8), constrained_layout=True)

    ax = axes[0]
    ax.plot(yearly["year"], yearly["ai_total_mean"], color="#22333b", marker="o", linewidth=2.3, label="AI sentences")
    ax.plot(yearly["year"], yearly["n_A_mean"], color="#2a9d8f", marker="o", linewidth=2.0, label="Actionable")
    ax.plot(yearly["year"], yearly["n_S_mean"], color="#c46b48", marker="o", linewidth=2.0, label="Speculative")
    ax.set_title("Panel A. Mean AI Disclosure Counts per Firm-Year")
    ax.set_ylabel("Mean count")
    ax.legend(loc="upper left", frameon=False)
    _style_axes(ax)

    ax2 = axes[1]
    ax2.plot(comp["year"], comp["act_share_mean"], color="#2a9d8f", marker="o", linewidth=2.2, label="Actionable share")
    ax2.plot(comp["year"], comp["spec_share_mean"], color="#c46b48", marker="o", linewidth=2.2, label="Speculative share")
    ax2.plot(yearly["year"], yearly["any_ai_share"], color="#6c757d", marker="o", linewidth=1.8, linestyle="--", label="Any AI talk share")
    ax2.set_title("Panel B. Disclosure Composition Over Time")
    ax2.set_xlabel("Year")
    ax2.set_ylabel("Share")
    ax2.set_ylim(bottom=0)
    ax2.legend(loc="upper left", frameon=False)
    _style_axes(ax2)

    fig_dir = Path(output_dir)
    fig_dir.mkdir(parents=True, exist_ok=True)
    image_path = fig_dir / "figure_1_disclosure_volume_composition_prelim_v1.png"
    fig.savefig(image_path, dpi=220, bbox_inches="tight")
    plt.close(fig)

    note = (
        "This figure uses the merged ever-speaker annual panel. Panel A plots mean AI sentence counts per firm-year, separating total AI sentences from the actionable and speculative subsets. "
        "Panel B plots mean actionable and speculative shares among AI-talking firm-years, together with the share of ever-speaker firm-years that contain any AI disclosure."
    )
    doc_path = Path(doc_dir) / "figure_1_disclosure_volume_composition_prelim_v1.docx"
    _save_figure_docx("Figure 1. AI Disclosure Volume and Composition Over Time", note, image_path, doc_path)
    return image_path, doc_path


def build_figure_2(panel_path: str | Path, output_dir: str | Path, doc_dir: str | Path) -> tuple[Path, Path]:
    df = _load_panel(panel_path)
    yearly = (
        df.groupby("year", as_index=False)
        .agg(
            ai_patent_share=("patents_ai", lambda s: (pd.to_numeric(s, errors="coerce").fillna(0) > 0).mean()),
            ai_patent_mean=("patents_ai", "mean"),
            log_ai_patent_mean=("patents_ai", lambda s: np.log1p(pd.to_numeric(s, errors="coerce").fillna(0)).mean()),
        )
        .sort_values("year")
    )
    yearly["log_ai_patent_mean"] = yearly["log_ai_patent_mean"].astype(float)

    _base_style()
    fig, axes = plt.subplots(2, 1, figsize=(8.2, 8.4), constrained_layout=True)

    ax = axes[0]
    ax.plot(yearly["year"], yearly["ai_patent_share"], color="#1d3557", marker="o", linewidth=2.3)
    ax.set_title("Panel A. Share of Firm-Years with Any AI Patent")
    ax.set_ylabel("Share")
    ax.set_ylim(bottom=0)
    _style_axes(ax)

    ax2 = axes[1]
    ax2.plot(yearly["year"], yearly["ai_patent_mean"], color="#457b9d", marker="o", linewidth=2.1, label="Mean AI patents")
    ax2.plot(yearly["year"], yearly["log_ai_patent_mean"], color="#a8dadc", marker="o", linewidth=2.1, label="Mean log(1 + AI patents)")
    ax2.set_title("Panel B. Mean AI Patent Intensity")
    ax2.set_xlabel("Year")
    ax2.set_ylabel("Mean level")
    ax2.legend(loc="upper left", frameon=False)
    _style_axes(ax2)

    fig_dir = Path(output_dir)
    fig_dir.mkdir(parents=True, exist_ok=True)
    image_path = fig_dir / "figure_2_ai_patent_coverage_prelim_v1.png"
    fig.savefig(image_path, dpi=220, bbox_inches="tight")
    plt.close(fig)

    note = (
        "This figure uses the merged ever-speaker annual panel. Panel A plots the share of firm-years with any AI patent. "
        "Panel B plots the mean AI patent count and the mean log-transformed AI patent intensity, showing both sparsity and the growth of patenting in the upper tail."
    )
    doc_path = Path(doc_dir) / "figure_2_ai_patent_coverage_prelim_v1.docx"
    _save_figure_docx("Figure 2. AI Patent Coverage Over Time", note, image_path, doc_path)
    return image_path, doc_path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--panel", default=DEFAULT_PANEL)
    parser.add_argument("--figure-dir", default=DEFAULT_FIG_DIR)
    parser.add_argument("--doc-dir", default=DEFAULT_DOC_DIR)
    parser.add_argument("--figures", default="figure1,figure2")
    return parser.parse_args()


def _selected(raw: str) -> set[str]:
    return {part.strip().lower() for part in raw.split(",") if part.strip()}


def main() -> None:
    args = parse_args()
    selected = _selected(args.figures)
    if "figure1" in selected:
        build_figure_1(args.panel, args.figure_dir, args.doc_dir)
    if "figure2" in selected:
        build_figure_2(args.panel, args.figure_dir, args.doc_dir)
    print(f"[delivery-figures] wrote selected figures under {args.figure_dir} and {args.doc_dir}")


if __name__ == "__main__":
    main()
