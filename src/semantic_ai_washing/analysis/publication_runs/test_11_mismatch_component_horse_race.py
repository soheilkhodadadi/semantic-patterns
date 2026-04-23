"""Publication run driver for Test 11: horse-race decomposition of PatentMismatch."""

from __future__ import annotations

import argparse
import json
import math
import shutil
from datetime import UTC, date, datetime
from pathlib import Path

from docx import Document
import pandas as pd
import statsmodels.formula.api as smf

from semantic_ai_washing.analysis.delivery_table_payloads import _add_patent_mismatch
from semantic_ai_washing.analysis.publication_runs.test_03_post_filing_drift import (
    _add_note,
    _add_title,
    _build_panel_table,
    _set_document_defaults,
    _set_landscape,
)

REPO_ROOT = Path(__file__).resolve().parents[4]
DEFAULT_EVENT_PANEL = (
    REPO_ROOT
    / "data/processed/panel/filing_event_estimation_sample_hybrid_api_a_conf49_v1.parquet"
)
DEFAULT_ANNUAL_PANEL = (
    REPO_ROOT
    / "data/processed/panel/canonical/ever_speaker_panel_2016_2025_hybrid_api_a_conf49_v1.parquet"
)
DEFAULT_TEST_ROOT = Path(
    "/Users/soheilkhodadadi/DataWork/semantic-patterns/ai_washing/derived/v3_1/test_runs/test_11_mismatch_component_horse_race"
)
DEFAULT_PAPER_ROOT = REPO_ROOT / "paper/generated/v3_1"
DEFAULT_RUN_ID = f"{date.today():%Y%m%d}_aiw_v3_1_test_11_mismatch_component_horse_race_main_v1"
TEST_ID = "test_11_mismatch_component_horse_race"
MODULE_PATH = "semantic_ai_washing.analysis.publication_runs.test_11_mismatch_component_horse_race"
CONTROL_TERMS = ["ln_assets", "leverage", "cash", "roa"]
OUTCOME_SPECS = [
    ("car_m1_p1", "CAR[-1,+1]", "car_m1_p1_complete"),
    ("bhar_3m", "BHAR[+2,+63]", "bhar_3m_complete"),
]
MODEL_SPECS = [
    ("mismatch_only", "Mismatch-only", "PatentMismatch + AI_Focus + {controls} + C(sic2) + C(filing_year)"),
    ("components", "Components", "LowCredibility + WeakPatentRelative + AI_Focus + {controls} + C(sic2) + C(filing_year)"),
    ("interaction", "Components + interaction", "LowCredibility + WeakPatentRelative + LowCredibility:WeakPatentRelative + AI_Focus + {controls} + C(sic2) + C(filing_year)"),
]
DISPLAY_TERMS = [
    ("PatentMismatch", "PatentMismatch"),
    ("LowCredibility", "LowCredibility"),
    ("WeakPatentRelative", "WeakPatentRelative"),
    ("LowCredibility:WeakPatentRelative", "LowCred × WeakPatent"),
    ("AI_Focus", "AI_Focus"),
    ("Outcome mean", "Outcome mean"),
    ("N", "N"),
    ("Adj. R-squared", "Adj. R-squared"),
]


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--event-panel", type=Path, default=DEFAULT_EVENT_PANEL)
    parser.add_argument("--annual-panel", type=Path, default=DEFAULT_ANNUAL_PANEL)
    parser.add_argument("--test-root", type=Path, default=DEFAULT_TEST_ROOT)
    parser.add_argument("--paper-root", type=Path, default=DEFAULT_PAPER_ROOT)
    parser.add_argument("--run-id", default=DEFAULT_RUN_ID)
    return parser.parse_args()


def _normalize_id(series: pd.Series) -> pd.Series:
    return (
        series.astype(str)
        .str.replace(r"\.0$", "", regex=True)
        .str.strip()
        .replace({"nan": "", "None": ""})
    )


def _ensure_sic2(panel: pd.DataFrame) -> pd.DataFrame:
    out = panel.copy()
    if "sic2" not in out.columns or out["sic2"].isna().all():
        if "sic" in out.columns:
            sic_raw = pd.to_numeric(out["sic"], errors="coerce")
            out["sic2"] = (sic_raw // 100).astype("Int64")
        else:
            out["sic2"] = pd.Series(pd.NA, index=out.index, dtype="Int64")
    return out


def _load_annual_component_panel(path: Path) -> pd.DataFrame:
    panel = pd.read_parquet(path).copy()
    panel = _ensure_sic2(panel)
    panel = _add_patent_mismatch(panel)
    keep = [
        "cik",
        "year",
        "sic2",
        "PatentMismatch",
        "LowCredibility",
        "WeakPatentRelative",
        "AI_Focus",
        "A_S",
        "SpecShare",
        "any_ai_talk",
    ]
    available = [column for column in keep if column in panel.columns]
    out = panel[available].copy()
    out["cik"] = _normalize_id(out["cik"])
    out["year"] = pd.to_numeric(out["year"], errors="coerce").astype("Int64")
    out["sic2"] = pd.to_numeric(out["sic2"], errors="coerce").astype("Int64")
    return out


def _build_analysis_sample(event_panel: Path, annual_panel: Path) -> tuple[pd.DataFrame, dict[str, object]]:
    event = pd.read_parquet(event_panel).copy()
    event["cik"] = _normalize_id(event["cik"])
    event["gvkey"] = _normalize_id(event["gvkey"])
    event["filing_year"] = pd.to_numeric(event["filing_year"], errors="coerce").astype("Int64")
    event["n_ai_total"] = pd.to_numeric(event["n_ai_total"], errors="coerce")

    annual = _load_annual_component_panel(annual_panel)
    sample = event.merge(
        annual,
        left_on=["cik", "filing_year"],
        right_on=["cik", "year"],
        how="left",
        validate="many_to_one",
        suffixes=("", "_annual"),
    )
    sample["is_ai_filing"] = sample["n_ai_total"].fillna(0).gt(0)
    sample = sample.loc[sample["is_ai_filing"]].copy()

    numeric_cols = [
        "PatentMismatch",
        "LowCredibility",
        "WeakPatentRelative",
        "AI_Focus",
        *CONTROL_TERMS,
        *[spec[0] for spec in OUTCOME_SPECS],
    ]
    for column in numeric_cols:
        if column in sample.columns:
            sample[column] = pd.to_numeric(sample[column], errors="coerce")
    for complete_flag in [spec[2] for spec in OUTCOME_SPECS]:
        sample[complete_flag] = pd.to_numeric(sample[complete_flag], errors="coerce").fillna(0).astype(int)

    summary = {
        "ai_filing_rows": int(len(sample)),
        "unique_gvkey": int(sample["gvkey"].replace("", pd.NA).dropna().nunique()),
        "filing_year_min": int(sample["filing_year"].min()),
        "filing_year_max": int(sample["filing_year"].max()),
        "mismatch_share": float(sample["PatentMismatch"].fillna(0).mean()),
        "low_cred_share": float(sample["LowCredibility"].fillna(0).mean()),
        "weak_patent_share": float(sample["WeakPatentRelative"].fillna(0).mean()),
    }
    return sample, summary


def _fit_model(sample: pd.DataFrame, outcome: str, outcome_label: str, complete_flag: str, model_id: str, model_label: str, rhs_formula: str) -> dict[str, object]:
    needed = [outcome, complete_flag, "gvkey", "sic2", "filing_year", "AI_Focus", *CONTROL_TERMS]
    components = ["PatentMismatch", "LowCredibility", "WeakPatentRelative"]
    use = sample.dropna(subset=needed + components).copy()
    use = use.loc[use[complete_flag].eq(1)].copy()
    use = use.loc[use["gvkey"].str.len().gt(0)].copy()
    use["sic2"] = pd.to_numeric(use["sic2"], errors="coerce")
    use = use.loc[use["sic2"].notna()].copy()
    use["sic2"] = use["sic2"].astype(int)
    use["filing_year"] = pd.to_numeric(use["filing_year"], errors="coerce").astype(int)
    for col in [outcome, "AI_Focus", *CONTROL_TERMS, *components]:
        use[col] = pd.to_numeric(use[col], errors="coerce")
    formula = f"{outcome} ~ " + rhs_formula.format(controls=" + ".join(CONTROL_TERMS))
    model = smf.ols(formula=formula, data=use)
    result = model.fit(cov_type="cluster", cov_kwds={"groups": use["gvkey"].astype(str)})
    return {
        "outcome": outcome,
        "outcome_label": outcome_label,
        "model_id": model_id,
        "model_label": model_label,
        "formula": formula,
        "nobs": int(result.nobs),
        "adj_r_squared": float(result.rsquared_adj),
        "outcome_mean": float(use[outcome].mean()),
        "params": result.params.to_dict(),
        "bse": result.bse.to_dict(),
        "pvalues": result.pvalues.to_dict(),
    }


def _stars(p_value: float | None) -> str:
    if p_value is None or not math.isfinite(p_value):
        return ""
    if p_value < 0.01:
        return "***"
    if p_value < 0.05:
        return "**"
    if p_value < 0.10:
        return "*"
    return ""


def _coef_cell(result: dict[str, object], term: str) -> str:
    if term not in result.get("params", {}):
        return ""
    coef = result["params"].get(term)
    se = result["bse"].get(term)
    p_value = result["pvalues"].get(term)
    if coef is None or se is None or not math.isfinite(coef) or not math.isfinite(se):
        return ""
    return f"{coef:.4f}{_stars(p_value)} ({se:.4f})"


def _build_results_table(model_rows: list[dict[str, object]]) -> tuple[pd.DataFrame, dict[tuple[str, str], dict[str, object]]]:
    keyed = {(row["outcome"], row["model_id"]): row for row in model_rows}
    rows: list[dict[str, object]] = []
    for outcome, outcome_label, _ in OUTCOME_SPECS:
        for term, row_label in DISPLAY_TERMS:
            row = {"panel": outcome_label, "row_label": row_label}
            for model_id, model_label, _ in MODEL_SPECS:
                result = keyed[(outcome, model_id)]
                if term in {"Outcome mean", "N", "Adj. R-squared"}:
                    if term == "Outcome mean":
                        row[model_label] = f"{result['outcome_mean']:.4f}"
                    elif term == "N":
                        row[model_label] = str(result["nobs"])
                    else:
                        row[model_label] = f"{result['adj_r_squared']:.3f}"
                else:
                    row[model_label] = _coef_cell(result, term)
            rows.append(row)
    return pd.DataFrame(rows), keyed


def _markdown_table(headers: list[str], rows: list[list[object]]) -> str:
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join(["---"] * len(headers)) + " |",
    ]
    for row in rows:
        lines.append("| " + " | ".join(str(item) for item in row) + " |")
    return "\n".join(lines)


def _render_table_outputs(table_df: pd.DataFrame) -> tuple[str, str]:
    headers = ["Row", *[label for _, label, _ in MODEL_SPECS]]
    md_lines = ["# Table Main", ""]
    latex_lines = [
        "\\begin{table}[!htbp]",
        "\\centering",
        "\\caption{Horse-race decomposition of PatentMismatch}",
        "\\begin{tabular}{l" + "c" * len(MODEL_SPECS) + "}",
        "\\hline",
    ]
    for _, panel_label, _ in OUTCOME_SPECS:
        panel_rows = table_df.loc[table_df["panel"].eq(panel_label)].copy()
        rendered = panel_rows[["row_label", *[label for _, label, _ in MODEL_SPECS]]].values.tolist()
        md_lines.extend([f"## Panel {panel_label}", _markdown_table(headers, rendered), ""])
        latex_lines.append(f"\\multicolumn{{{len(headers)}}}{{l}}{{\\textit{{Panel {panel_label}}}}} \\")
        latex_lines.append(" & ".join(headers) + " \\")
        for row in rendered:
            latex_lines.append(" & ".join(str(item) for item in row) + " \\")
    latex_lines.extend(["\\hline", "\\end{tabular}", "\\end{table}"])
    return "\n".join(md_lines), "\n".join(latex_lines) + "\n"


def _docx_panel_rows(table_df: pd.DataFrame, panel_name: str) -> list[list[tuple[str, bool]]]:
    rows = []
    subset = table_df.loc[table_df["panel"].eq(panel_name)].copy()
    for row in subset[["row_label", *[label for _, label, _ in MODEL_SPECS]]].itertuples(index=False):
        rows.append([(str(row[0]), True), *[(str(value), False) for value in row[1:]]])
    return rows


def _build_table_docx(table_df: pd.DataFrame, output_path: Path) -> None:
    document = Document()
    _set_document_defaults(document)
    _set_landscape(document)
    _add_title(document, "Table 11. Horse-Race Decomposition of PatentMismatch")
    note = (
        "This table decomposes PatentMismatch into its disclosure-side and patent-side ingredients. Panel A uses filing-date CAR[-1,+1] as the dependent variable and Panel B uses post-filing BHAR[+2,+63]. Column (1) uses PatentMismatch alone, Column (2) replaces it with LowCredibility and WeakPatentRelative, and Column (3) adds their interaction. All specifications include AI Focus, lagged firm controls, SIC2 fixed effects, filing-year fixed effects, and standard errors clustered by gvkey."
    )
    _add_note(document, note)
    headers = ["", *[label for _, label, _ in MODEL_SPECS]]
    for _, panel_label, _ in OUTCOME_SPECS:
        p = document.add_paragraph()
        p.add_run(f"Panel {panel_label}").bold = True
        _build_panel_table(document, headers, _docx_panel_rows(table_df, panel_label))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def _result_notes(model_keyed: dict[tuple[str, str], dict[str, object]], sample_summary: dict[str, object]) -> str:
    car_components = model_keyed[("car_m1_p1", "components")]
    car_interaction = model_keyed[("car_m1_p1", "interaction")]
    bhar_components = model_keyed[("bhar_3m", "components")]
    bhar_interaction = model_keyed[("bhar_3m", "interaction")]
    return "\n".join(
        [
            f"# Result Notes: {TEST_ID}",
            "",
            f"- AI-filing sample: `{sample_summary['ai_filing_rows']}` rows across `{sample_summary['unique_gvkey']}` firms.",
            f"- Filing-date CAR components model: `LowCredibility = {car_components['params'].get('LowCredibility', float('nan')):.4f}` (p=`{car_components['pvalues'].get('LowCredibility', float('nan')):.3f}`), `WeakPatentRelative = {car_components['params'].get('WeakPatentRelative', float('nan')):.4f}` (p=`{car_components['pvalues'].get('WeakPatentRelative', float('nan')):.3f}`).",
            f"- Filing-date CAR interaction model: `LowCred × WeakPatent = {car_interaction['params'].get('LowCredibility:WeakPatentRelative', float('nan')):.4f}` (p=`{car_interaction['pvalues'].get('LowCredibility:WeakPatentRelative', float('nan')):.3f}`).",
            f"- Post-filing BHAR components model: `LowCredibility = {bhar_components['params'].get('LowCredibility', float('nan')):.4f}` (p=`{bhar_components['pvalues'].get('LowCredibility', float('nan')):.3f}`), `WeakPatentRelative = {bhar_components['params'].get('WeakPatentRelative', float('nan')):.4f}` (p=`{bhar_components['pvalues'].get('WeakPatentRelative', float('nan')):.3f}`).",
            f"- Post-filing BHAR interaction model: `LowCred × WeakPatent = {bhar_interaction['params'].get('LowCredibility:WeakPatentRelative', float('nan')):.4f}` (p=`{bhar_interaction['pvalues'].get('LowCredibility:WeakPatentRelative', float('nan')):.3f}`).",
            "- Interpretation discipline: the horse race is informative if the interaction survives more strongly than either ingredient alone; if the patent-side component dominates, the market story has to be narrowed toward weaker capability rather than low-credibility disclosure per se.",
            "",
        ]
    )


def _writer_packet(args: argparse.Namespace, sample_summary: dict[str, object], model_keyed: dict[tuple[str, str], dict[str, object]]) -> str:
    bhar_interaction = model_keyed[("bhar_3m", "interaction")]
    return "\n".join(
        [
            f"# Writer Packet: {TEST_ID}",
            "",
            "## Purpose",
            "- This run tests whether the market signal is driven by the combined mismatch state or by one of its two ingredients separately.",
            "- It is designed to answer the most important identification question left after factor-adjusted alpha survives in the equal-weight 3-month portfolio.",
            "",
            "## Sample Definition",
            f"- Event panel: `{args.event_panel}`",
            f"- Annual panel: `{args.annual_panel}`",
            "- Unit of observation: `AI-talking annual filing event`",
            "- Outcomes: `CAR[-1,+1]` and `BHAR[+2,+63]`",
            "- Treatment variables: `PatentMismatch`, `LowCredibility`, `WeakPatentRelative`, and `LowCredibility × WeakPatentRelative`",
            "- Controls: `ln_assets`, `leverage`, `cash`, `roa`, plus `AI_Focus`",
            "- Fixed effects: `SIC2` and `filing-year`",
            "- Inference: `clustered by gvkey`",
            "",
            "## Sample Counts",
            f"- AI-filing rows: `{sample_summary['ai_filing_rows']}`",
            f"- Unique firms: `{sample_summary['unique_gvkey']}`",
            f"- Mismatch share: `{sample_summary['mismatch_share']:.3f}`",
            f"- LowCredibility share: `{sample_summary['low_cred_share']:.3f}`",
            f"- WeakPatentRelative share: `{sample_summary['weak_patent_share']:.3f}`",
            "",
            "## Most Important Coefficient",
            f"- BHAR interaction coefficient: `{bhar_interaction['params'].get('LowCredibility:WeakPatentRelative', float('nan')):.4f}` with p=`{bhar_interaction['pvalues'].get('LowCredibility:WeakPatentRelative', float('nan')):.3f}`",
            "",
            "## Caption Draft",
            "This table decomposes PatentMismatch into its disclosure-side and patent-side ingredients. Panel A uses filing-date CAR[-1,+1] as the dependent variable and Panel B uses post-filing BHAR[+2,+63]. Column (1) uses PatentMismatch alone, Column (2) replaces it with LowCredibility and WeakPatentRelative, and Column (3) adds their interaction. All specifications include AI Focus, lagged firm controls, SIC2 fixed effects, filing-year fixed effects, and standard errors clustered by gvkey.",
            "",
        ]
    )


def _dataset_summary(sample_summary: dict[str, object], model_rows: list[dict[str, object]], *, args: argparse.Namespace) -> dict[str, object]:
    return {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "inputs": {
            "event_panel": str(args.event_panel),
            "annual_panel": str(args.annual_panel),
        },
        "sample_summary": sample_summary,
        "model_rows": model_rows,
    }


def _copy_exports(run_dir: Path, paper_root: Path, run_id: str) -> dict[str, str]:
    exports = {
        "table_csv": paper_root / "tables" / f"{TEST_ID}_{run_id}.csv",
        "table_md": paper_root / "tables" / f"{TEST_ID}_{run_id}.md",
        "table_tex": paper_root / "latex" / f"{TEST_ID}_{run_id}.tex",
        "table_docx": paper_root / "docx" / f"{TEST_ID}_{run_id}.docx",
        "writer_packet": paper_root / "writer_packets" / f"{TEST_ID}_{run_id}.md",
        "result_notes": paper_root / "snippets" / f"{TEST_ID}_{run_id}_result_notes.md",
    }
    for path in exports.values():
        path.parent.mkdir(parents=True, exist_ok=True)
    mapping = {
        run_dir / "table_main.csv": exports["table_csv"],
        run_dir / "table_main.md": exports["table_md"],
        run_dir / "table_main.tex": exports["table_tex"],
        run_dir / "table_main.docx": exports["table_docx"],
        run_dir / "writer_packet.md": exports["writer_packet"],
        run_dir / "result_notes.md": exports["result_notes"],
    }
    for src, dst in mapping.items():
        shutil.copy2(src, dst)
    return {key: str(path) for key, path in exports.items()}


def main() -> None:
    args = _parse_args()
    run_dir = args.test_root / args.run_id
    run_dir.mkdir(parents=True, exist_ok=True)

    sample, sample_summary = _build_analysis_sample(args.event_panel, args.annual_panel)
    model_rows: list[dict[str, object]] = []
    for outcome, outcome_label, complete_flag in OUTCOME_SPECS:
        for model_id, model_label, rhs_formula in MODEL_SPECS:
            model_rows.append(
                _fit_model(sample, outcome, outcome_label, complete_flag, model_id, model_label, rhs_formula)
            )

    table_df, model_keyed = _build_results_table(model_rows)
    table_df.to_csv(run_dir / "table_main.csv", index=False)
    table_md, table_tex = _render_table_outputs(table_df)
    (run_dir / "table_main.md").write_text(table_md, encoding="utf-8")
    (run_dir / "table_main.tex").write_text(table_tex, encoding="utf-8")
    _build_table_docx(table_df, run_dir / "table_main.docx")
    (run_dir / "result_notes.md").write_text(_result_notes(model_keyed, sample_summary), encoding="utf-8")
    (run_dir / "writer_packet.md").write_text(_writer_packet(args, sample_summary, model_keyed), encoding="utf-8")
    dataset_summary = _dataset_summary(sample_summary, model_rows, args=args)
    (run_dir / "dataset_summary.json").write_text(json.dumps(dataset_summary, indent=2), encoding="utf-8")

    paper_exports = _copy_exports(run_dir, args.paper_root, args.run_id)
    manifest = {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "module_path": MODULE_PATH,
        "run_dir": str(run_dir),
        "inputs": {
            "event_panel": str(args.event_panel),
            "annual_panel": str(args.annual_panel),
        },
        "outputs": {
            "dataset_summary": str(run_dir / "dataset_summary.json"),
            "table_csv": str(run_dir / "table_main.csv"),
            "table_md": str(run_dir / "table_main.md"),
            "table_tex": str(run_dir / "table_main.tex"),
            "table_docx": str(run_dir / "table_main.docx"),
            "writer_packet": str(run_dir / "writer_packet.md"),
            "result_notes": str(run_dir / "result_notes.md"),
        },
        "paper_exports": paper_exports,
    }
    (run_dir / "run_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    print(f"[{TEST_ID}] wrote run bundle to {run_dir}")
    print(f"[{TEST_ID}] paper table: {paper_exports['table_docx']}")


if __name__ == "__main__":
    main()
