"""Publication run driver for Test 12: predictive return regressions with richer observables."""

from __future__ import annotations

import argparse
import json
import math
import shutil
from datetime import UTC, date, datetime
from pathlib import Path

from docx import Document
import numpy as np
import pandas as pd
import statsmodels.formula.api as smf

from semantic_ai_washing.analysis.delivery_table_payloads import _add_patent_mismatch
from semantic_ai_washing.analysis.publication_runs.test_03_post_filing_drift import (
    _add_note,
    _add_title,
    _build_panel_table,
    _set_document_defaults,
    _set_landscape,
)

REPO_ROOT = Path(__file__).resolve().parents[4]
DEFAULT_EVENT_PANEL = (
    REPO_ROOT
    / "data/processed/panel/filing_event_estimation_sample_hybrid_api_a_conf49_v1.parquet"
)
DEFAULT_ANNUAL_PANEL = (
    REPO_ROOT
    / "data/processed/panel/canonical/ever_speaker_panel_2016_2025_hybrid_api_a_conf49_v1.parquet"
)
DEFAULT_TEST_ROOT = Path(
    "/Users/soheilkhodadadi/DataWork/semantic-patterns/ai_washing/derived/v3_1/test_runs/test_12_predictive_return_controls"
)
DEFAULT_PAPER_ROOT = REPO_ROOT / "paper/generated/v3_1"
DEFAULT_RUN_ID = f"{date.today():%Y%m%d}_aiw_v3_1_test_12_predictive_return_controls_main_v1"
TEST_ID = "test_12_predictive_return_controls"
MODULE_PATH = "semantic_ai_washing.analysis.publication_runs.test_12_predictive_return_controls"
FOCAL_TERMS = ["PatentMismatch", "A_S", "AI_Focus"]
CORE_CONTROLS = ["ln_assets", "leverage", "cash", "roa"]
OPERATING_CONTROLS = ["rd_intensity", "capx_at", "sales_growth"]
MARKET_CONTROLS = ["ln_mktcap_year_end", "annual_ret", "annual_bhar_vw", "ln_vol", "firm_age_market"]
PATENT_CONTROLS = ["log_patents_ai_lag1"]
OUTCOME_SPECS = [
    ("bhar_3m", "BHAR[+2,+63]", "bhar_3m_complete"),
    ("bhar_12m", "BHAR[+2,+252]", "bhar_12m_complete"),
]
MODEL_SPECS = [
    ("core", "Core controls", CORE_CONTROLS),
    ("operating", "+ operating controls", [*CORE_CONTROLS, *OPERATING_CONTROLS]),
    ("market", "+ market characteristics", [*CORE_CONTROLS, *OPERATING_CONTROLS, *MARKET_CONTROLS]),
    ("market_patent", "+ market + patent history", [*CORE_CONTROLS, *OPERATING_CONTROLS, *MARKET_CONTROLS, *PATENT_CONTROLS]),
]
DISPLAY_TERMS = [
    ("PatentMismatch", "PatentMismatch"),
    ("A_S", "A_S"),
    ("AI_Focus", "AI_Focus"),
    ("Outcome mean", "Outcome mean"),
    ("N", "N"),
    ("Adj. R-squared", "Adj. R-squared"),
]


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--event-panel", type=Path, default=DEFAULT_EVENT_PANEL)
    parser.add_argument("--annual-panel", type=Path, default=DEFAULT_ANNUAL_PANEL)
    parser.add_argument("--test-root", type=Path, default=DEFAULT_TEST_ROOT)
    parser.add_argument("--paper-root", type=Path, default=DEFAULT_PAPER_ROOT)
    parser.add_argument("--run-id", default=DEFAULT_RUN_ID)
    return parser.parse_args()


def _normalize_id(series: pd.Series) -> pd.Series:
    return (
        series.astype(str)
        .str.replace(r"\.0$", "", regex=True)
        .str.strip()
        .replace({"nan": "", "None": ""})
    )


def _ensure_sic2(panel: pd.DataFrame) -> pd.DataFrame:
    out = panel.copy()
    if "sic2" not in out.columns or out["sic2"].isna().all():
        if "sic" in out.columns:
            sic_raw = pd.to_numeric(out["sic"], errors="coerce")
            out["sic2"] = (sic_raw // 100).astype("Int64")
        else:
            out["sic2"] = pd.Series(pd.NA, index=out.index, dtype="Int64")
    return out


def _winsorize(series: pd.Series, lower_q: float = 0.01, upper_q: float = 0.99) -> pd.Series:
    valid = pd.to_numeric(series, errors="coerce").replace([np.inf, -np.inf], np.nan)
    if valid.notna().sum() == 0:
        return valid
    lower = valid.quantile(lower_q)
    upper = valid.quantile(upper_q)
    return valid.clip(lower=lower, upper=upper)


def _load_annual_controls(path: Path) -> pd.DataFrame:
    panel = pd.read_parquet(path).copy()
    panel = _ensure_sic2(panel)
    panel = _add_patent_mismatch(panel)
    keep = [
        "cik",
        "year",
        "sic2",
        "PatentMismatch",
        "A_S",
        "AI_Focus",
        "rd_intensity",
        "capx_at",
        "sales_growth",
        "annual_ret",
        "annual_bhar_vw",
        "market_cap_year_end",
        "firm_age_market",
        "vol",
        "log_patents_ai_lag1",
    ]
    available = [column for column in keep if column in panel.columns]
    out = panel[available].copy()
    out["cik"] = _normalize_id(out["cik"])
    out["year"] = pd.to_numeric(out["year"], errors="coerce").astype("Int64")
    out["sic2"] = pd.to_numeric(out["sic2"], errors="coerce").astype("Int64")
    return out


def _build_analysis_sample(event_panel: Path, annual_panel: Path) -> tuple[pd.DataFrame, dict[str, object]]:
    event = pd.read_parquet(event_panel).copy()
    event["cik"] = _normalize_id(event["cik"])
    event["gvkey"] = _normalize_id(event["gvkey"])
    event["filing_year"] = pd.to_numeric(event["filing_year"], errors="coerce").astype("Int64")
    event["n_ai_total"] = pd.to_numeric(event["n_ai_total"], errors="coerce")

    annual = _load_annual_controls(annual_panel)
    sample = event.merge(
        annual,
        left_on=["cik", "filing_year"],
        right_on=["cik", "year"],
        how="left",
        validate="many_to_one",
        suffixes=("", "_annual"),
    )
    sample["is_ai_filing"] = sample["n_ai_total"].fillna(0).gt(0)
    sample = sample.loc[sample["is_ai_filing"]].copy()

    for outcome, _, complete_flag in OUTCOME_SPECS:
        sample[outcome] = _winsorize(sample[outcome])
        sample[complete_flag] = pd.to_numeric(sample[complete_flag], errors="coerce").fillna(0).astype(int)

    sample["ln_mktcap_year_end"] = np.log1p(pd.to_numeric(sample.get("market_cap_year_end"), errors="coerce"))
    sample["ln_vol"] = np.log1p(pd.to_numeric(sample.get("vol"), errors="coerce"))
    for column in [
        *FOCAL_TERMS,
        *CORE_CONTROLS,
        *OPERATING_CONTROLS,
        *MARKET_CONTROLS,
        *PATENT_CONTROLS,
    ]:
        if column in sample.columns:
            sample[column] = pd.to_numeric(sample[column], errors="coerce").replace([np.inf, -np.inf], np.nan)

    summary = {
        "ai_filing_rows": int(len(sample)),
        "unique_gvkey": int(sample["gvkey"].replace("", pd.NA).dropna().nunique()),
        "filing_year_min": int(sample["filing_year"].min()),
        "filing_year_max": int(sample["filing_year"].max()),
        "winsorization": "1st/99th percentile on BHAR outcomes and heavy-tailed market controls transformed with log1p where applicable",
    }
    return sample, summary


def _fit_model(sample: pd.DataFrame, outcome: str, outcome_label: str, complete_flag: str, model_id: str, model_label: str, controls: list[str]) -> dict[str, object]:
    needed = [outcome, complete_flag, "gvkey", "sic2", "filing_year", *FOCAL_TERMS, *controls]
    use = sample.dropna(subset=needed).copy()
    use = use.loc[use[complete_flag].eq(1)].copy()
    use = use.loc[use["gvkey"].str.len().gt(0)].copy()
    use["sic2"] = pd.to_numeric(use["sic2"], errors="coerce")
    use = use.loc[use["sic2"].notna()].copy()
    use["sic2"] = use["sic2"].astype(int)
    use["filing_year"] = pd.to_numeric(use["filing_year"], errors="coerce").astype(int)

    formula = (
        f"{outcome} ~ "
        + " + ".join([*FOCAL_TERMS, *controls])
        + " + C(sic2) + C(filing_year)"
    )
    model = smf.ols(formula=formula, data=use)
    result = model.fit(cov_type="cluster", cov_kwds={"groups": use["gvkey"].astype(str)})
    return {
        "outcome": outcome,
        "outcome_label": outcome_label,
        "model_id": model_id,
        "model_label": model_label,
        "controls": controls,
        "formula": formula,
        "nobs": int(result.nobs),
        "adj_r_squared": float(result.rsquared_adj),
        "outcome_mean": float(use[outcome].mean()),
        "params": result.params.to_dict(),
        "bse": result.bse.to_dict(),
        "pvalues": result.pvalues.to_dict(),
    }


def _stars(p_value: float | None) -> str:
    if p_value is None or not math.isfinite(p_value):
        return ""
    if p_value < 0.01:
        return "***"
    if p_value < 0.05:
        return "**"
    if p_value < 0.10:
        return "*"
    return ""


def _coef_cell(result: dict[str, object], term: str) -> str:
    coef = result["params"].get(term)
    se = result["bse"].get(term)
    p_value = result["pvalues"].get(term)
    if coef is None or se is None or not math.isfinite(coef) or not math.isfinite(se):
        return ""
    return f"{coef:.4f}{_stars(p_value)} ({se:.4f})"


def _build_results_table(model_rows: list[dict[str, object]]) -> tuple[pd.DataFrame, dict[tuple[str, str], dict[str, object]]]:
    keyed = {(row["outcome"], row["model_id"]): row for row in model_rows}
    rows: list[dict[str, object]] = []
    for outcome, outcome_label, _ in OUTCOME_SPECS:
        for term, row_label in DISPLAY_TERMS:
            row = {"panel": outcome_label, "row_label": row_label}
            for model_id, model_label, _ in MODEL_SPECS:
                result = keyed[(outcome, model_id)]
                if term in {"Outcome mean", "N", "Adj. R-squared"}:
                    if term == "Outcome mean":
                        row[model_label] = f"{result['outcome_mean']:.4f}"
                    elif term == "N":
                        row[model_label] = str(result["nobs"])
                    else:
                        row[model_label] = f"{result['adj_r_squared']:.3f}"
                else:
                    row[model_label] = _coef_cell(result, term)
            rows.append(row)
    return pd.DataFrame(rows), keyed


def _markdown_table(headers: list[str], rows: list[list[object]]) -> str:
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join(["---"] * len(headers)) + " |",
    ]
    for row in rows:
        lines.append("| " + " | ".join(str(item) for item in row) + " |")
    return "\n".join(lines)


def _render_table_outputs(table_df: pd.DataFrame) -> tuple[str, str]:
    headers = ["Row", *[label for _, label, _ in MODEL_SPECS]]
    md_lines = ["# Table Main", ""]
    latex_lines = [
        "\\begin{table}[!htbp]",
        "\\centering",
        "\\caption{Predictive post-filing return regressions with richer observables}",
        "\\begin{tabular}{l" + "c" * len(MODEL_SPECS) + "}",
        "\\hline",
    ]
    for _, panel_label, _ in OUTCOME_SPECS:
        panel_rows = table_df.loc[table_df["panel"].eq(panel_label)].copy()
        rendered = panel_rows[["row_label", *[label for _, label, _ in MODEL_SPECS]]].values.tolist()
        md_lines.extend([f"## Panel {panel_label}", _markdown_table(headers, rendered), ""])
        latex_lines.append(f"\\multicolumn{{{len(headers)}}}{{l}}{{\\textit{{Panel {panel_label}}}}} \\")
        latex_lines.append(" & ".join(headers) + " \\")
        for row in rendered:
            latex_lines.append(" & ".join(str(item) for item in row) + " \\")
    latex_lines.extend(["\\hline", "\\end{tabular}", "\\end{table}"])
    return "\n".join(md_lines), "\n".join(latex_lines) + "\n"


def _docx_panel_rows(table_df: pd.DataFrame, panel_name: str) -> list[list[tuple[str, bool]]]:
    rows = []
    subset = table_df.loc[table_df["panel"].eq(panel_name)].copy()
    for row in subset[["row_label", *[label for _, label, _ in MODEL_SPECS]]].itertuples(index=False):
        rows.append([(str(row[0]), True), *[(str(value), False) for value in row[1:]]])
    return rows


def _build_table_docx(table_df: pd.DataFrame, output_path: Path) -> None:
    document = Document()
    _set_document_defaults(document)
    _set_landscape(document)
    _add_title(document, "Table 12. Predictive Return Regressions with Richer Observables")
    note = (
        "This table tests whether the post-filing return pattern is subsumed by observable firm characteristics. The dependent variables are BHAR[+2,+63] in Panel A and BHAR[+2,+252] in Panel B. Each column adds progressively richer controls: core balance-sheet controls, operating controls, market characteristics, and prior AI patent history. All specifications include PatentMismatch, the A/S ratio, AI Focus, SIC2 fixed effects, filing-year fixed effects, and standard errors clustered by gvkey."
    )
    _add_note(document, note)
    headers = ["", *[label for _, label, _ in MODEL_SPECS]]
    for _, panel_label, _ in OUTCOME_SPECS:
        p = document.add_paragraph()
        p.add_run(f"Panel {panel_label}").bold = True
        _build_panel_table(document, headers, _docx_panel_rows(table_df, panel_label))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def _result_notes(model_keyed: dict[tuple[str, str], dict[str, object]], sample_summary: dict[str, object]) -> str:
    base_3m = model_keyed[("bhar_3m", "core")]
    full_3m = model_keyed[("bhar_3m", "market_patent")]
    full_12m = model_keyed[("bhar_12m", "market_patent")]
    return "\n".join(
        [
            f"# Result Notes: {TEST_ID}",
            "",
            f"- AI-filing sample: `{sample_summary['ai_filing_rows']}` rows across `{sample_summary['unique_gvkey']}` firms.",
            f"- BHAR[+2,+63] core-controls PatentMismatch coefficient: `{base_3m['params'].get('PatentMismatch', float('nan')):.4f}` (p=`{base_3m['pvalues'].get('PatentMismatch', float('nan')):.3f}`).",
            f"- BHAR[+2,+63] full-controls PatentMismatch coefficient: `{full_3m['params'].get('PatentMismatch', float('nan')):.4f}` (p=`{full_3m['pvalues'].get('PatentMismatch', float('nan')):.3f}`).",
            f"- BHAR[+2,+252] full-controls PatentMismatch coefficient: `{full_12m['params'].get('PatentMismatch', float('nan')):.4f}` (p=`{full_12m['pvalues'].get('PatentMismatch', float('nan')):.3f}`).",
            "- Interpretation discipline: the key question is whether PatentMismatch remains economically and statistically material once balance-sheet controls, operating variables, market characteristics, and prior AI patent history are all absorbed.",
            "",
        ]
    )


def _writer_packet(args: argparse.Namespace, sample_summary: dict[str, object], model_keyed: dict[tuple[str, str], dict[str, object]]) -> str:
    full_3m = model_keyed[("bhar_3m", "market_patent")]
    return "\n".join(
        [
            f"# Writer Packet: {TEST_ID}",
            "",
            "## Purpose",
            "- This run asks whether the post-filing return pattern survives richer observable controls instead of being subsumed by simple firm characteristics.",
            "- It is the direct follow-up to the factor-adjusted alpha and horse-race results.",
            "",
            "## Sample Definition",
            f"- Event panel: `{args.event_panel}`",
            f"- Annual panel: `{args.annual_panel}`",
            "- Unit of observation: `AI-talking annual filing event`",
            "- Outcomes: `BHAR[+2,+63]` and `BHAR[+2,+252]`",
            "- Fixed effects: `SIC2` and `filing-year`",
            "- Inference: `clustered by gvkey`",
            "",
            "## Control Ladder",
            "- Core: `ln_assets`, `leverage`, `cash`, `roa`",
            "- Operating: `rd_intensity`, `capx_at`, `sales_growth`",
            "- Market characteristics: `ln_mktcap_year_end`, `annual_ret`, `annual_bhar_vw`, `ln_vol`, `firm_age_market`",
            "- Patent history: `log_patents_ai_lag1`",
            "",
            "## Sample Counts",
            f"- AI-filing rows: `{sample_summary['ai_filing_rows']}`",
            f"- Unique firms: `{sample_summary['unique_gvkey']}`",
            "",
            "## Most Important Coefficient",
            f"- Full-controls BHAR[+2,+63] PatentMismatch coefficient: `{full_3m['params'].get('PatentMismatch', float('nan')):.4f}` with p=`{full_3m['pvalues'].get('PatentMismatch', float('nan')):.3f}`",
            "",
            "## Caption Draft",
            "This table tests whether the post-filing return pattern is subsumed by observable firm characteristics. The dependent variables are BHAR[+2,+63] in Panel A and BHAR[+2,+252] in Panel B. Each column adds progressively richer controls: core balance-sheet controls, operating controls, market characteristics, and prior AI patent history. All specifications include PatentMismatch, the A/S ratio, AI Focus, SIC2 fixed effects, filing-year fixed effects, and standard errors clustered by gvkey.",
            "",
        ]
    )


def _dataset_summary(sample_summary: dict[str, object], model_rows: list[dict[str, object]], *, args: argparse.Namespace) -> dict[str, object]:
    return {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "inputs": {
            "event_panel": str(args.event_panel),
            "annual_panel": str(args.annual_panel),
        },
        "sample_summary": sample_summary,
        "model_rows": model_rows,
    }


def _copy_exports(run_dir: Path, paper_root: Path, run_id: str) -> dict[str, str]:
    exports = {
        "table_csv": paper_root / "tables" / f"{TEST_ID}_{run_id}.csv",
        "table_md": paper_root / "tables" / f"{TEST_ID}_{run_id}.md",
        "table_tex": paper_root / "latex" / f"{TEST_ID}_{run_id}.tex",
        "table_docx": paper_root / "docx" / f"{TEST_ID}_{run_id}.docx",
        "writer_packet": paper_root / "writer_packets" / f"{TEST_ID}_{run_id}.md",
        "result_notes": paper_root / "snippets" / f"{TEST_ID}_{run_id}_result_notes.md",
    }
    for path in exports.values():
        path.parent.mkdir(parents=True, exist_ok=True)
    mapping = {
        run_dir / "table_main.csv": exports["table_csv"],
        run_dir / "table_main.md": exports["table_md"],
        run_dir / "table_main.tex": exports["table_tex"],
        run_dir / "table_main.docx": exports["table_docx"],
        run_dir / "writer_packet.md": exports["writer_packet"],
        run_dir / "result_notes.md": exports["result_notes"],
    }
    for src, dst in mapping.items():
        shutil.copy2(src, dst)
    return {key: str(path) for key, path in exports.items()}


def main() -> None:
    args = _parse_args()
    run_dir = args.test_root / args.run_id
    run_dir.mkdir(parents=True, exist_ok=True)

    sample, sample_summary = _build_analysis_sample(args.event_panel, args.annual_panel)
    model_rows: list[dict[str, object]] = []
    for outcome, outcome_label, complete_flag in OUTCOME_SPECS:
        for model_id, model_label, controls in MODEL_SPECS:
            model_rows.append(
                _fit_model(sample, outcome, outcome_label, complete_flag, model_id, model_label, controls)
            )

    table_df, model_keyed = _build_results_table(model_rows)
    table_df.to_csv(run_dir / "table_main.csv", index=False)
    table_md, table_tex = _render_table_outputs(table_df)
    (run_dir / "table_main.md").write_text(table_md, encoding="utf-8")
    (run_dir / "table_main.tex").write_text(table_tex, encoding="utf-8")
    _build_table_docx(table_df, run_dir / "table_main.docx")
    (run_dir / "result_notes.md").write_text(_result_notes(model_keyed, sample_summary), encoding="utf-8")
    (run_dir / "writer_packet.md").write_text(_writer_packet(args, sample_summary, model_keyed), encoding="utf-8")
    dataset_summary = _dataset_summary(sample_summary, model_rows, args=args)
    (run_dir / "dataset_summary.json").write_text(json.dumps(dataset_summary, indent=2), encoding="utf-8")

    paper_exports = _copy_exports(run_dir, args.paper_root, args.run_id)
    manifest = {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "module_path": MODULE_PATH,
        "run_dir": str(run_dir),
        "inputs": {
            "event_panel": str(args.event_panel),
            "annual_panel": str(args.annual_panel),
        },
        "outputs": {
            "dataset_summary": str(run_dir / "dataset_summary.json"),
            "table_csv": str(run_dir / "table_main.csv"),
            "table_md": str(run_dir / "table_main.md"),
            "table_tex": str(run_dir / "table_main.tex"),
            "table_docx": str(run_dir / "table_main.docx"),
            "writer_packet": str(run_dir / "writer_packet.md"),
            "result_notes": str(run_dir / "result_notes.md"),
        },
        "paper_exports": paper_exports,
    }
    (run_dir / "run_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    print(f"[{TEST_ID}] wrote run bundle to {run_dir}")
    print(f"[{TEST_ID}] paper table: {paper_exports['table_docx']}")


if __name__ == "__main__":
    main()
