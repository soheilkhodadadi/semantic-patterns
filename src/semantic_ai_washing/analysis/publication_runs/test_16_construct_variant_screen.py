"""Publication run driver for Test 16: construct-variant validation screen."""

from __future__ import annotations

import argparse
import json
import math
import shutil
from datetime import UTC, date, datetime
from pathlib import Path

from docx import Document
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from semantic_ai_washing.analysis.delivery_table_payloads import _fit_absorbed_ols
from semantic_ai_washing.analysis.publication_runs.test_03_post_filing_drift import (
    _add_note,
    _add_title,
    _build_panel_table,
    _set_document_defaults,
    _set_landscape,
)

REPO_ROOT = Path(__file__).resolve().parents[4]
DEFAULT_ANNUAL_PANEL = (
    REPO_ROOT
    / "data/processed/panel/canonical/ever_speaker_panel_2016_2025_hybrid_api_a_conf49_v1.parquet"
)
DEFAULT_TEST_ROOT = Path(
    "/Users/soheilkhodadadi/DataWork/semantic-patterns/ai_washing/derived/v3_1/test_runs/test_16_construct_variant_screen"
)
DEFAULT_PAPER_ROOT = REPO_ROOT / "paper/generated/v3_1"
DEFAULT_RUN_ID = f"{date.today():%Y%m%d}_aiw_v3_1_test_16_construct_variant_screen_main_v1"
TEST_ID = "test_16_construct_variant_screen"
MODULE_PATH = "semantic_ai_washing.analysis.publication_runs.test_16_construct_variant_screen"
CONTROL_TERMS = ["AI_Focus", "ln_assets", "leverage", "cash", "roa"]
VARIANT_SPECS: list[tuple[str, str]] = [
    ("PatentMismatch", "Canonical grant mismatch"),
    ("StrictPatentMismatch", "Strict grant mismatch"),
    ("ApplicationMismatch", "Application mismatch"),
    ("LowCredibility", "Low credibility only"),
    ("WeakPatentRelative", "Weak patent only"),
]
OUTCOME_SPECS: list[tuple[str, str, str]] = [
    ("log_patents_ai_lead1", "Log(1 + AI grants) t+1", "Panel A. Future AI grant outcomes"),
    ("log_patents_ai_lead2", "Log(1 + AI grants) t+2", "Panel A. Future AI grant outcomes"),
    (
        "log_applications_ai_lead1",
        "Log(1 + AI applications) t+1",
        "Panel B. Future AI application outcomes",
    ),
    (
        "log_applications_ai_lead2",
        "Log(1 + AI applications) t+2",
        "Panel B. Future AI application outcomes",
    ),
]


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--annual-panel", type=Path, default=DEFAULT_ANNUAL_PANEL)
    parser.add_argument("--test-root", type=Path, default=DEFAULT_TEST_ROOT)
    parser.add_argument("--paper-root", type=Path, default=DEFAULT_PAPER_ROOT)
    parser.add_argument("--run-id", default=DEFAULT_RUN_ID)
    return parser.parse_args()


def _base_style() -> None:
    plt.style.use("seaborn-v0_8-whitegrid")
    plt.rcParams.update(
        {
            "font.family": "serif",
            "font.serif": ["Times New Roman", "Times", "DejaVu Serif"],
            "axes.titlesize": 12,
            "axes.labelsize": 10,
            "xtick.labelsize": 9,
            "ytick.labelsize": 9,
        }
    )


def _normalize_id(series: pd.Series) -> pd.Series:
    return (
        series.astype(str)
        .str.replace(r"\.0$", "", regex=True)
        .str.strip()
        .replace({"nan": "", "None": ""})
    )


def _ensure_sic2(panel: pd.DataFrame) -> pd.DataFrame:
    out = panel.copy()
    if "sic2" not in out.columns or out["sic2"].isna().all():
        if "sic" in out.columns:
            sic_raw = pd.to_numeric(out["sic"], errors="coerce")
            out["sic2"] = (sic_raw // 100).astype("Int64")
        else:
            out["sic2"] = pd.Series(pd.NA, index=out.index, dtype="Int64")
    return out


def _add_construct_variants(panel: pd.DataFrame) -> pd.DataFrame:
    out = _ensure_sic2(panel)
    if "SpecShare" not in out.columns and "share_S" in out.columns:
        out["SpecShare"] = out["share_S"]
    for column in [
        "A_S",
        "SpecShare",
        "AI_Focus",
        "ln_assets",
        "leverage",
        "cash",
        "roa",
        "log_patents_ai_lead0",
        "log_patents_ai_lead1",
        "log_patents_ai_lead2",
        "log_applications_ai_lead0",
        "log_applications_ai_lead1",
        "log_applications_ai_lead2",
    ]:
        if column in out.columns:
            out[column] = pd.to_numeric(out[column], errors="coerce").replace(
                [np.inf, -np.inf], np.nan
            )

    out["cik"] = _normalize_id(out["cik"])
    out["year"] = pd.to_numeric(out["year"], errors="coerce").astype("Int64")

    valid_industry = out["sic2"].notna() & out["year"].notna()
    out["industry_year"] = pd.Series(pd.NA, index=out.index, dtype="object")
    out.loc[valid_industry, "industry_year"] = (
        out.loc[valid_industry, "sic2"]
        .astype(int)
        .astype(str)
        .str.cat(out.loc[valid_industry, "year"].astype(int).astype(str), sep="_")
    )

    talk_mask = out["any_ai_talk"].fillna(0).astype(int).eq(1)
    low_cred_mask = talk_mask & out["A_S"].notna() & out["SpecShare"].notna() & out["year"].notna()
    out["LowCredibility"] = 0
    out["StrictLowCred"] = 0
    if low_cred_mask.any():
        talk = out.loc[low_cred_mask, ["year", "A_S", "SpecShare"]].copy()
        as_cut = talk.groupby("year")["A_S"].transform(lambda s: s.quantile(0.25))
        spec_cut = talk.groupby("year")["SpecShare"].transform(lambda s: s.quantile(0.75))
        low_cred = (talk["A_S"] <= as_cut) | (talk["SpecShare"] >= spec_cut)
        strict_low_cred = (talk["A_S"] <= as_cut) & (talk["SpecShare"] >= spec_cut)
        out.loc[low_cred_mask, "LowCredibility"] = low_cred.astype(int).to_numpy()
        out.loc[low_cred_mask, "StrictLowCred"] = strict_low_cred.astype(int).to_numpy()

    grant_industry_mean = out.groupby("industry_year")["log_patents_ai_lead0"].transform("mean")
    out["WeakPatentRelative"] = (
        (out["log_patents_ai_lead0"] < grant_industry_mean).fillna(False).astype(int)
    )
    app_industry_mean = out.groupby("industry_year")["log_applications_ai_lead0"].transform("mean")
    out["WeakAppRelative"] = (
        (out["log_applications_ai_lead0"] < app_industry_mean).fillna(False).astype(int)
    )

    out["PatentMismatch"] = (
        talk_mask & out["LowCredibility"].astype(bool) & out["WeakPatentRelative"].astype(bool)
    ).astype(int)
    out["StrictPatentMismatch"] = (
        talk_mask & out["StrictLowCred"].astype(bool) & out["WeakPatentRelative"].astype(bool)
    ).astype(int)
    out["ApplicationMismatch"] = (
        talk_mask & out["LowCredibility"].astype(bool) & out["WeakAppRelative"].astype(bool)
    ).astype(int)
    return out


def _prepare_sample(panel: pd.DataFrame) -> tuple[pd.DataFrame, dict[str, object]]:
    use = _add_construct_variants(panel)
    use = use.loc[use["any_ai_talk"].fillna(0).astype(int).eq(1)].copy()
    needed = [label for label, _ in VARIANT_SPECS] + CONTROL_TERMS
    for column in needed:
        use[column] = pd.to_numeric(use[column], errors="coerce").replace(
            [np.inf, -np.inf], np.nan
        )
    use = use.loc[use["cik"].str.len().gt(0) & use["year"].notna()].copy()
    summary = {
        "ai_talking_rows": int(len(use)),
        "unique_firms": int(use["cik"].nunique()),
        "year_min": int(use["year"].min()),
        "year_max": int(use["year"].max()),
        "variant_shares": {
            column: float(pd.to_numeric(use[column], errors="coerce").fillna(0).mean())
            for column, _label in VARIANT_SPECS
        },
    }
    return use, summary


def _fit_variant_model(sample: pd.DataFrame, variant: str, dependent: str) -> dict[str, object]:
    result, use, adj_r2 = _fit_absorbed_ols(
        sample,
        dependent=dependent,
        rhs_terms=[variant],
        absorb_col="cik",
        include_year=True,
        controls=CONTROL_TERMS,
    )
    return {
        "variant": variant,
        "dependent": dependent,
        "nobs": int(result.nobs),
        "adj_r2": float(adj_r2) if adj_r2 is not None else math.nan,
        "params": result.params.to_dict(),
        "bse": result.bse.to_dict(),
        "pvalues": result.pvalues.to_dict(),
        "outcome_mean": float(use[dependent].mean()),
    }


def _stars(p_value: float | None) -> str:
    if p_value is None or not math.isfinite(p_value):
        return ""
    if p_value < 0.01:
        return "***"
    if p_value < 0.05:
        return "**"
    if p_value < 0.10:
        return "*"
    return ""


def _coef_cell(coef: float | None, se: float | None, p_value: float | None) -> str:
    if coef is None or se is None or not math.isfinite(coef) or not math.isfinite(se):
        return ""
    return f"{coef:.4f}{_stars(p_value)} ({se:.4f})"


def _build_results_table(
    model_rows: list[dict[str, object]],
) -> tuple[pd.DataFrame, dict[tuple[str, str], dict[str, object]]]:
    keyed = {(row["dependent"], row["variant"]): row for row in model_rows}
    rows: list[dict[str, object]] = []
    for panel_label in [
        "Panel A. Future AI grant outcomes",
        "Panel B. Future AI application outcomes",
    ]:
        panel_outcomes = [spec for spec in OUTCOME_SPECS if spec[2] == panel_label]
        for variant, variant_label in VARIANT_SPECS:
            row = {"panel": panel_label, "row_label": variant_label}
            for dependent, dependent_label, _panel in panel_outcomes:
                result = keyed[(dependent, variant)]
                row[dependent_label] = _coef_cell(
                    result["params"].get(variant),
                    result["bse"].get(variant),
                    result["pvalues"].get(variant),
                )
            rows.append(row)
        n_row = {"panel": panel_label, "row_label": "Observations"}
        mean_row = {"panel": panel_label, "row_label": "Outcome mean"}
        for dependent, dependent_label, _panel in panel_outcomes:
            result = keyed[(dependent, VARIANT_SPECS[0][0])]
            n_row[dependent_label] = str(result["nobs"])
            mean_row[dependent_label] = f"{result['outcome_mean']:.3f}"
        rows.extend([mean_row, n_row])
    return pd.DataFrame(rows), keyed


def _render_table_outputs(table_df: pd.DataFrame) -> tuple[str, str]:
    panel_headers = {
        "Panel A. Future AI grant outcomes": [
            spec[1] for spec in OUTCOME_SPECS if spec[2] == "Panel A. Future AI grant outcomes"
        ],
        "Panel B. Future AI application outcomes": [
            spec[1]
            for spec in OUTCOME_SPECS
            if spec[2] == "Panel B. Future AI application outcomes"
        ],
    }

    def md_table(panel: str) -> str:
        headers = ["Variant"] + panel_headers[panel]
        rows = table_df.loc[table_df["panel"].eq(panel)]
        lines = [
            "| " + " | ".join(headers) + " |",
            "| " + " | ".join(["---"] * len(headers)) + " |",
        ]
        for _, row in rows.iterrows():
            cells = [str(row["row_label"])] + [str(row[h]) for h in panel_headers[panel]]
            lines.append("| " + " | ".join(cells) + " |")
        return "\n".join(lines)

    md = "\n".join(
        [
            "# Table Main",
            "",
            "## Panel A. Future AI grant outcomes",
            md_table("Panel A. Future AI grant outcomes"),
            "",
            "## Panel B. Future AI application outcomes",
            md_table("Panel B. Future AI application outcomes"),
            "",
        ]
    )

    latex_lines = [
        "\\begin{table}[!htbp]",
        "\\centering",
        "\\caption{Construct-variant screen for future AI realization}",
        "\\begin{tabular}{l" + "c" * 2 + "}",
        "\\hline",
        "Variant & Log(1 + AI grants) t+1 & Log(1 + AI grants) t+2 \\\\",
        "\\hline",
    ]
    for _, row in table_df.loc[
        table_df["panel"].eq("Panel A. Future AI grant outcomes")
    ].iterrows():
        latex_lines.append(
            " & ".join(
                [str(row["row_label"])]
                + [str(row[h]) for h in panel_headers["Panel A. Future AI grant outcomes"]]
            )
            + " \\\\"
        )
    latex_lines.extend(
        [
            "\\hline",
            "\\multicolumn{3}{l}{\\textit{Panel B. Future AI application outcomes}} \\\\",
            "Variant & Log(1 + AI applications) t+1 & Log(1 + AI applications) t+2 \\\\",
            "\\hline",
        ]
    )
    for _, row in table_df.loc[
        table_df["panel"].eq("Panel B. Future AI application outcomes")
    ].iterrows():
        latex_lines.append(
            " & ".join(
                [str(row["row_label"])]
                + [str(row[h]) for h in panel_headers["Panel B. Future AI application outcomes"]]
            )
            + " \\\\"
        )
    latex_lines.extend(["\\hline", "\\end{tabular}", "\\end{table}"])
    return md, "\n".join(latex_lines) + "\n"


def _build_table_docx(table_df: pd.DataFrame, output_path: Path) -> None:
    document = Document()
    _set_document_defaults(document)
    _set_landscape(document)
    _add_title(document, "Table 16. Construct-Variant Screen for Future AI Realization")
    _add_note(
        document,
        "This table screens nearby AI-washing constructs inside the AI-talking annual sample. Panel A asks which variants predict future AI grant realization at horizons t+1 and t+2. Panel B asks the same for future AI application realization. All specifications absorb firm and year fixed effects, control for AI Focus, log assets, leverage, cash, and ROA, and use firm-clustered standard errors. The goal is not to replace the canonical construct mechanically, but to learn whether nearby variants tell a cleaner capability-realization story.",
    )
    for panel in ["Panel A. Future AI grant outcomes", "Panel B. Future AI application outcomes"]:
        p = document.add_paragraph()
        p.add_run(panel).bold = True
        headers = [""] + [spec[1] for spec in OUTCOME_SPECS if spec[2] == panel]
        panel_rows: list[list[tuple[str, bool]]] = []
        for _, row in table_df.loc[table_df["panel"].eq(panel)].iterrows():
            panel_rows.append(
                [(str(row["row_label"]), True)] + [(str(row[h]), False) for h in headers[1:]]
            )
        _build_panel_table(document, headers, panel_rows)
    document.save(output_path)


def _plot_heatmap(
    model_rows: list[dict[str, object]], output_stem: Path
) -> tuple[Path, Path, pd.DataFrame]:
    _base_style()
    records = []
    for dependent, dependent_label, panel in OUTCOME_SPECS:
        for variant, variant_label in VARIANT_SPECS:
            row = next(
                r for r in model_rows if r["dependent"] == dependent and r["variant"] == variant
            )
            records.append(
                {
                    "panel": panel,
                    "variant": variant_label,
                    "outcome": dependent_label,
                    "coef": float(row["params"].get(variant, math.nan)),
                    "p_value": float(row["pvalues"].get(variant, math.nan)),
                }
            )
    figure_df = pd.DataFrame(records)
    outcome_order = [label for _dependent, label, _panel in OUTCOME_SPECS]
    variant_order = [label for _variant, label in VARIANT_SPECS]
    pivot = figure_df.pivot(index="variant", columns="outcome", values="coef").reindex(
        index=variant_order, columns=outcome_order
    )
    ann = (
        figure_df.assign(
            label=lambda x: x.apply(lambda r: f"{r['coef']:.3f}{_stars(r['p_value'])}", axis=1)
        )
        .pivot(index="variant", columns="outcome", values="label")
        .reindex(index=variant_order, columns=outcome_order)
    )
    vmax = float(np.nanmax(np.abs(pivot.to_numpy(dtype=float)))) if not pivot.empty else 0.1
    vmax = max(vmax, 0.05)
    fig, ax = plt.subplots(figsize=(8.6, 3.8))
    im = ax.imshow(
        pivot.to_numpy(dtype=float), cmap="RdBu_r", vmin=-vmax, vmax=vmax, aspect="auto"
    )
    ax.set_xticks(range(len(pivot.columns)), list(pivot.columns), rotation=18, ha="right")
    ax.set_yticks(range(len(pivot.index)), list(pivot.index))
    ax.set_title("Construct-Variant Screen: Future AI Realization Coefficients")
    for i, row_label in enumerate(pivot.index):
        for j, col_label in enumerate(pivot.columns):
            ax.text(j, i, ann.loc[row_label, col_label], ha="center", va="center", fontsize=8)
    cbar = fig.colorbar(im, ax=ax, shrink=0.9)
    cbar.set_label("Coefficient")
    fig.tight_layout()
    png_path = output_stem.with_suffix(".png")
    pdf_path = output_stem.with_suffix(".pdf")
    fig.savefig(png_path, dpi=240)
    fig.savefig(pdf_path)
    plt.close(fig)
    return png_path, pdf_path, figure_df


def _result_notes(model_rows: list[dict[str, object]], sample_summary: dict[str, object]) -> str:
    keyed = {(row["dependent"], row["variant"]): row for row in model_rows}
    patent_main = keyed[("log_patents_ai_lead1", "PatentMismatch")]
    app_main = keyed[("log_applications_ai_lead2", "PatentMismatch")]
    weak_pat = keyed[("log_patents_ai_lead1", "WeakPatentRelative")]
    return "\n".join(
        [
            f"# Result Notes: {TEST_ID}",
            "",
            f"- Sample: `{sample_summary['ai_talking_rows']}` AI-talking firm-years across `{sample_summary['unique_firms']}` firms.",
            f"- Canonical PatentMismatch on future grants t+1: `{patent_main['params'].get('PatentMismatch', math.nan):.4f}` (p=`{patent_main['pvalues'].get('PatentMismatch', math.nan):.3f}`).",
            f"- Canonical PatentMismatch on future applications t+2: `{app_main['params'].get('PatentMismatch', math.nan):.4f}` (p=`{app_main['pvalues'].get('PatentMismatch', math.nan):.3f}`).",
            f"- WeakPatentRelative on future grants t+1: `{weak_pat['params'].get('WeakPatentRelative', math.nan):.4f}` (p=`{weak_pat['pvalues'].get('WeakPatentRelative', math.nan):.3f}`).",
            "- Interpretation discipline: this is a construct-screening table. The purpose is to see which nearby variants best line up with later AI realization, not to retroactively replace the audited canonical construct with whichever column is numerically strongest.",
            "",
        ]
    )


def _writer_packet(
    args: argparse.Namespace,
    model_rows: list[dict[str, object]],
    sample_summary: dict[str, object],
) -> str:
    keyed = {(row["dependent"], row["variant"]): row for row in model_rows}
    return "\n".join(
        [
            f"# Writer Packet: {TEST_ID}",
            "",
            "## Role In The Paper",
            "- This run opens the post-market non-risky lane: a structured screen of nearby AI-washing constructs against later AI realization outcomes.",
            "- It helps decide whether the canonical PatentMismatch construct remains the right main text measure and which nearby variants deserve appendix or robustness status.",
            "",
            "## Sample Definition",
            f"- Annual panel: `{args.annual_panel}`",
            "- Working sample: `AI-talking firm-years only`",
            f"- Rows: `{sample_summary['ai_talking_rows']}`",
            f"- Firms: `{sample_summary['unique_firms']}`",
            "- Fixed effects: `firm + year`",
            "- Inference: `firm-clustered standard errors`",
            "",
            "## Variant Definitions",
            "- Canonical grant mismatch: low-credibility disclosure plus weak contemporaneous grant-side capability relative to the industry-year mean.",
            "- Strict grant mismatch: same as above, but the low-credibility leg requires both low A/S and high speculative share instead of either condition alone.",
            "- Application mismatch: low-credibility disclosure plus weak contemporaneous application-side capability relative to the industry-year mean.",
            "- Low credibility only: disclosure-side construct without the capability filter.",
            "- Weak patent only: capability-side construct without the disclosure filter.",
            "",
            "## Headline Read",
            f"- Canonical PatentMismatch on future grants t+1: `{keyed[('log_patents_ai_lead1', 'PatentMismatch')]['params'].get('PatentMismatch', math.nan):.4f}` (p=`{keyed[('log_patents_ai_lead1', 'PatentMismatch')]['pvalues'].get('PatentMismatch', math.nan):.3f}`)",
            f"- Canonical PatentMismatch on future applications t+2: `{keyed[('log_applications_ai_lead2', 'PatentMismatch')]['params'].get('PatentMismatch', math.nan):.4f}` (p=`{keyed[('log_applications_ai_lead2', 'PatentMismatch')]['pvalues'].get('PatentMismatch', math.nan):.3f}`)",
            f"- WeakPatentRelative on future grants t+1: `{keyed[('log_patents_ai_lead1', 'WeakPatentRelative')]['params'].get('WeakPatentRelative', math.nan):.4f}` (p=`{keyed[('log_patents_ai_lead1', 'WeakPatentRelative')]['pvalues'].get('WeakPatentRelative', math.nan):.3f}`)",
            "",
            "## Caption Draft",
            "This table screens nearby AI-washing constructs inside the AI-talking annual sample. Panel A relates each construct variant to future AI grant realization, and Panel B relates the same variants to future AI application realization. All specifications absorb firm and year fixed effects, include AI Focus and core firm controls, and cluster standard errors by firm. The table is meant to rank interpretability and predictive content across nearby constructs rather than to mechanically replace the audited canonical measure.",
            "",
        ]
    )


def _dataset_summary(
    args: argparse.Namespace,
    sample: pd.DataFrame,
    sample_summary: dict[str, object],
    model_rows: list[dict[str, object]],
    figure_df: pd.DataFrame,
) -> dict[str, object]:
    figure_payload = figure_df.copy()
    return {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "inputs": {"annual_panel": str(args.annual_panel)},
        "sample_summary": sample_summary,
        "variant_counts": {
            column: int(pd.to_numeric(sample[column], errors="coerce").fillna(0).sum())
            for column, _label in VARIANT_SPECS
        },
        "model_rows": model_rows,
        "figure_series": figure_payload.to_dict(orient="records"),
    }


def _copy_exports(run_dir: Path, paper_root: Path, run_id: str) -> dict[str, str]:
    exports = {
        "table_csv": paper_root / "tables" / f"{TEST_ID}_{run_id}.csv",
        "table_md": paper_root / "tables" / f"{TEST_ID}_{run_id}.md",
        "table_tex": paper_root / "latex" / f"{TEST_ID}_{run_id}.tex",
        "table_docx": paper_root / "docx" / f"{TEST_ID}_{run_id}.docx",
        "figure_png": paper_root / "figures" / f"{TEST_ID}_{run_id}.png",
        "figure_pdf": paper_root / "figures" / f"{TEST_ID}_{run_id}.pdf",
        "figure_series": paper_root / "tables" / f"{TEST_ID}_{run_id}_figure_series.csv",
        "writer_packet": paper_root / "writer_packets" / f"{TEST_ID}_{run_id}.md",
        "result_notes": paper_root / "snippets" / f"{TEST_ID}_{run_id}_result_notes.md",
    }
    for path in exports.values():
        path.parent.mkdir(parents=True, exist_ok=True)
    mapping = {
        run_dir / "table_main.csv": exports["table_csv"],
        run_dir / "table_main.md": exports["table_md"],
        run_dir / "table_main.tex": exports["table_tex"],
        run_dir / "table_main.docx": exports["table_docx"],
        run_dir / "figure_main.png": exports["figure_png"],
        run_dir / "figure_main.pdf": exports["figure_pdf"],
        run_dir / "figure_series.csv": exports["figure_series"],
        run_dir / "writer_packet.md": exports["writer_packet"],
        run_dir / "result_notes.md": exports["result_notes"],
    }
    for src, dst in mapping.items():
        shutil.copy2(src, dst)
    return {key: str(path) for key, path in exports.items()}


def main() -> None:
    args = _parse_args()
    run_dir = args.test_root / args.run_id
    run_dir.mkdir(parents=True, exist_ok=True)

    panel = pd.read_parquet(args.annual_panel)
    sample, sample_summary = _prepare_sample(panel)

    model_rows: list[dict[str, object]] = []
    for dependent, _label, _panel in OUTCOME_SPECS:
        for variant, _variant_label in VARIANT_SPECS:
            model_rows.append(_fit_variant_model(sample, variant, dependent))

    table_df, keyed = _build_results_table(model_rows)
    table_df.to_csv(run_dir / "table_main.csv", index=False)
    table_md, table_tex = _render_table_outputs(table_df)
    (run_dir / "table_main.md").write_text(table_md, encoding="utf-8")
    (run_dir / "table_main.tex").write_text(table_tex, encoding="utf-8")
    _build_table_docx(table_df, run_dir / "table_main.docx")

    png_path, pdf_path, figure_df = _plot_heatmap(model_rows, run_dir / "figure_main")
    figure_df.to_csv(run_dir / "figure_series.csv", index=False)

    (run_dir / "result_notes.md").write_text(
        _result_notes(model_rows, sample_summary), encoding="utf-8"
    )
    (run_dir / "writer_packet.md").write_text(
        _writer_packet(args, model_rows, sample_summary), encoding="utf-8"
    )
    dataset_summary = _dataset_summary(args, sample, sample_summary, model_rows, figure_df)
    (run_dir / "dataset_summary.json").write_text(
        json.dumps(dataset_summary, indent=2), encoding="utf-8"
    )

    paper_exports = _copy_exports(run_dir, args.paper_root, args.run_id)
    manifest = {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "module_path": MODULE_PATH,
        "run_dir": str(run_dir),
        "inputs": {"annual_panel": str(args.annual_panel)},
        "outputs": {
            "dataset_summary": str(run_dir / "dataset_summary.json"),
            "table_csv": str(run_dir / "table_main.csv"),
            "table_md": str(run_dir / "table_main.md"),
            "table_tex": str(run_dir / "table_main.tex"),
            "table_docx": str(run_dir / "table_main.docx"),
            "figure_png": str(png_path),
            "figure_pdf": str(pdf_path),
            "figure_series": str(run_dir / "figure_series.csv"),
            "writer_packet": str(run_dir / "writer_packet.md"),
            "result_notes": str(run_dir / "result_notes.md"),
        },
        "paper_exports": paper_exports,
        "sample_summary": sample_summary,
    }
    (run_dir / "run_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    print(f"[{TEST_ID}] wrote run bundle to {run_dir}")
    print(f"[{TEST_ID}] paper table: {paper_exports['table_docx']}")


if __name__ == "__main__":
    main()
