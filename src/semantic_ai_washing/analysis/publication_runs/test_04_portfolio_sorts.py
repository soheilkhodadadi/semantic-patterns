"""Publication run driver for Test 04: portfolio sorts and size-split alpha."""

from __future__ import annotations

import argparse
import json
import math
import shutil
from datetime import UTC, date, datetime
from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
import statsmodels.api as sm
from docx import Document

from semantic_ai_washing.analysis.publication_runs.test_03_post_filing_drift import (
    _add_note,
    _add_title,
    _build_analysis_sample,
    _build_panel_table,
    _set_document_defaults,
    _set_landscape,
)
from semantic_ai_washing.analysis.publication_runs.test_03_post_filing_portfolio_alpha import (
    _base_style,
    _load_market_index,
    _load_monthly_returns,
    _weighted_average,
)

REPO_ROOT = Path(__file__).resolve().parents[4]
DEFAULT_EVENT_PANEL = (
    REPO_ROOT
    / "data/processed/panel/filing_event_estimation_sample_hybrid_api_a_conf49_v1.parquet"
)
DEFAULT_ANNUAL_PANEL = (
    REPO_ROOT
    / "data/processed/panel/canonical/ever_speaker_panel_2016_2025_hybrid_api_a_conf49_v1.parquet"
)
DEFAULT_MONTHLY_RETURNS = REPO_ROOT / "data/interim/market/wrds_crsp_msf_full_sample_v1.parquet"
DEFAULT_MARKET_INDEX = REPO_ROOT / "data/interim/market/wrds_crsp_msi_full_sample_v1.parquet"
DEFAULT_TEST_ROOT = Path(
    "/Users/soheilkhodadadi/DataWork/semantic-patterns/ai_washing/derived/test_runs/test_04_portfolio_sorts"
)
DEFAULT_PAPER_ROOT = REPO_ROOT / "paper/generated"
DEFAULT_RUN_ID = f"{date.today():%Y%m%d}_hybrid_api_a_conf49_main_v1"
TEST_ID = "test_04_portfolio_sorts"
MODULE_PATH = "semantic_ai_washing.analysis.publication_runs.test_04_portfolio_sorts"
HORIZONS = {1: "1m", 3: "3m"}
AS_PORTFOLIOS = {1: "Q1 (low A/S)", 2: "Q2", 3: "Q3 (high A/S)"}
MISMATCH_ROWS = [
    ("overall", "ew", "All firms, equal-weight"),
    ("overall", "vw", "All firms, value-weight"),
    ("small", "ew", "Small firms, equal-weight"),
    ("big", "ew", "Big firms, equal-weight"),
    ("small", "vw", "Small firms, value-weight"),
    ("big", "vw", "Big firms, value-weight"),
]


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--event-panel", type=Path, default=DEFAULT_EVENT_PANEL)
    parser.add_argument("--annual-panel", type=Path, default=DEFAULT_ANNUAL_PANEL)
    parser.add_argument("--monthly-returns", type=Path, default=DEFAULT_MONTHLY_RETURNS)
    parser.add_argument("--market-index", type=Path, default=DEFAULT_MARKET_INDEX)
    parser.add_argument("--test-root", type=Path, default=DEFAULT_TEST_ROOT)
    parser.add_argument("--paper-root", type=Path, default=DEFAULT_PAPER_ROOT)
    parser.add_argument("--run-id", default=DEFAULT_RUN_ID)
    return parser.parse_args()


def _prepare_signal_sample(event_panel: Path, annual_panel: Path) -> pd.DataFrame:
    sample = _build_analysis_sample(event_panel, annual_panel)
    sample = sample.loc[sample["is_ai_filing"]].copy()
    sample["permno"] = sample["permno"].astype(str)
    sample["filing_date"] = pd.to_datetime(sample["filing_date"])
    sample["filing_month"] = sample["filing_date"].dt.to_period("M")
    for column in ["PatentMismatch", "A_S", "AI_Focus"]:
        sample[column] = pd.to_numeric(sample[column], errors="coerce")
    sample = sample.loc[sample["PatentMismatch"].notna() & sample["A_S"].notna()].copy()
    sample["PatentMismatch"] = sample["PatentMismatch"].fillna(0).astype(int)
    return sample


def _build_active_positions(sample: pd.DataFrame, horizon_months: int) -> pd.DataFrame:
    positions: list[tuple[object, ...]] = []
    for row in sample[
        [
            "permno",
            "filing_month",
            "filing_date",
            "A_S",
            "PatentMismatch",
            "AI_Focus",
            "post_chatgpt",
        ]
    ].itertuples(index=False):
        for k in range(1, horizon_months + 1):
            positions.append(
                (
                    row.permno,
                    row.filing_month + k,
                    row.filing_date,
                    row.A_S,
                    row.PatentMismatch,
                    row.AI_Focus,
                    row.post_chatgpt,
                )
            )
    out = pd.DataFrame(
        positions,
        columns=[
            "permno",
            "month",
            "filing_date",
            "A_S",
            "PatentMismatch",
            "AI_Focus",
            "post_chatgpt",
        ],
    )
    out = out.sort_values(["permno", "month", "filing_date"]).drop_duplicates(
        ["permno", "month"], keep="last"
    )
    return out.reset_index(drop=True)


def _month_terciles(values: pd.Series) -> pd.Series:
    out = pd.Series(pd.NA, index=values.index, dtype="Int64")
    valid = values.dropna()
    if len(valid) < 3:
        return out
    ranked = valid.rank(method="first")
    try:
        buckets = pd.qcut(ranked, 3, labels=[1, 2, 3])
    except ValueError:
        return out
    out.loc[valid.index] = buckets.astype(int).to_numpy()
    return out


def _merge_returns(active_positions: pd.DataFrame, monthly_returns: pd.DataFrame) -> pd.DataFrame:
    merged = active_positions.merge(monthly_returns, on=["permno", "month"], how="left")
    merged = merged.dropna(subset=["ret"]).copy()
    merged["A_S_bucket"] = (
        merged.groupby("month", group_keys=False)["A_S"].apply(_month_terciles).astype("Int64")
    )
    month_median_size = merged.groupby("month")["lag_mcap"].transform("median")
    merged["size_bucket"] = pd.Series(pd.NA, index=merged.index, dtype="object")
    has_size = merged["lag_mcap"].notna() & month_median_size.notna()
    merged.loc[has_size & merged["lag_mcap"].le(month_median_size), "size_bucket"] = "small"
    merged.loc[has_size & merged["lag_mcap"].gt(month_median_size), "size_bucket"] = "big"
    return merged


def _portfolio_series(
    merged: pd.DataFrame,
    group_cols: list[str],
    *,
    return_col: str = "ret",
    weight_col: str | None = None,
) -> pd.DataFrame:
    use = merged.dropna(subset=[return_col]).copy()
    if weight_col is None:
        grouped = (
            use.groupby(["month", *group_cols], as_index=False)[return_col]
            .mean()
            .rename(columns={return_col: "portfolio_ret"})
        )
    else:
        grouped = (
            use.groupby(["month", *group_cols])
            .apply(lambda g: _weighted_average(g[return_col], g[weight_col]))
            .reset_index(name="portfolio_ret")
        )
    return grouped


def _alpha_summary(series_df: pd.DataFrame) -> dict[str, float | int]:
    use = series_df.dropna(subset=["portfolio_ret", "vwretd"]).copy()
    if use.empty:
        return {
            "mean_monthly_return": math.nan,
            "annualized_simple_return": math.nan,
            "market_alpha": math.nan,
            "market_alpha_p": math.nan,
            "beta_to_market": math.nan,
            "n_months": 0,
        }
    monthly_mean = float(use["portfolio_ret"].mean())
    X = sm.add_constant(use["vwretd"])
    result = sm.OLS(use["portfolio_ret"], X).fit()
    return {
        "mean_monthly_return": monthly_mean,
        "annualized_simple_return": float(monthly_mean * 12.0),
        "market_alpha": float(result.params.get("const", math.nan)),
        "market_alpha_p": float(result.pvalues.get("const", math.nan)),
        "beta_to_market": float(result.params.get("vwretd", math.nan)),
        "n_months": int(len(use)),
    }


def _as_sort_summaries(
    merged: pd.DataFrame, market_index: pd.DataFrame, *, horizon_label: str
) -> tuple[list[dict[str, object]], pd.DataFrame | None]:
    rows: list[dict[str, object]] = []
    figure_series = None
    for weighting, weight_col in [("ew", None), ("vw", "lag_mcap")]:
        series = _portfolio_series(
            merged.dropna(subset=["A_S_bucket"]), ["A_S_bucket"], weight_col=weight_col
        )
        pivot = series.pivot(index="month", columns="A_S_bucket", values="portfolio_ret")
        pivot = pivot.rename(columns={1: "Q1 (low A/S)", 2: "Q2", 3: "Q3 (high A/S)"})
        portfolio = pd.DataFrame(index=pivot.index)
        for label in AS_PORTFOLIOS.values():
            portfolio[label] = pivot.get(label)
        portfolio["Q3 - Q1"] = portfolio["Q3 (high A/S)"] - portfolio["Q1 (low A/S)"]
        portfolio = portfolio.reset_index().merge(market_index, on="month", how="left")
        for portfolio_label in ["Q1 (low A/S)", "Q2", "Q3 (high A/S)", "Q3 - Q1"]:
            subset = portfolio[["month", portfolio_label, "vwretd"]].rename(
                columns={portfolio_label: "portfolio_ret"}
            )
            stats = _alpha_summary(subset)
            rows.append(
                {
                    "panel": "a_s_sorts",
                    "weighting": weighting,
                    "portfolio": portfolio_label,
                    "horizon": horizon_label,
                    **stats,
                }
            )
        if horizon_label == "3m" and weighting == "ew":
            figure_series = portfolio[["month", "Q3 - Q1"]].rename(
                columns={"Q3 - Q1": "as_q3_q1_ew"}
            )
    return rows, figure_series


def _mismatch_spread_series(
    merged: pd.DataFrame, market_index: pd.DataFrame, *, weight_col: str | None = None
) -> pd.DataFrame:
    series = _portfolio_series(merged, ["PatentMismatch"], weight_col=weight_col)
    pivot = series.pivot(index="month", columns="PatentMismatch", values="portfolio_ret").rename(
        columns={0: "credible", 1: "mismatch"}
    )
    out = pd.DataFrame(index=pivot.index)
    out.index.name = "month"
    out["portfolio_ret"] = pivot["credible"] - pivot["mismatch"]
    return out.reset_index().merge(market_index, on="month", how="left")


def _size_split_summaries(
    merged: pd.DataFrame, market_index: pd.DataFrame, *, horizon_label: str
) -> tuple[list[dict[str, object]], pd.DataFrame | None]:
    rows: list[dict[str, object]] = []
    figure_parts: dict[str, pd.DataFrame] = {}

    for sample_key, weight_col in [("ew", None), ("vw", "lag_mcap")]:
        overall_series = _mismatch_spread_series(merged, market_index, weight_col=weight_col)
        overall_stats = _alpha_summary(overall_series)
        rows.append(
            {
                "panel": "mismatch_size",
                "sample": "overall",
                "weighting": sample_key,
                "horizon": horizon_label,
                **overall_stats,
            }
        )

        for size_bucket in ["small", "big"]:
            sub = merged.loc[merged["size_bucket"].eq(size_bucket)].copy()
            series = _mismatch_spread_series(sub, market_index, weight_col=weight_col)
            stats = _alpha_summary(series)
            rows.append(
                {
                    "panel": "mismatch_size",
                    "sample": size_bucket,
                    "weighting": sample_key,
                    "horizon": horizon_label,
                    **stats,
                }
            )
            if horizon_label == "3m" and sample_key == "ew":
                figure_parts[size_bucket] = series.rename(
                    columns={"portfolio_ret": f"{size_bucket}_mismatch_ew"}
                )[["month", f"{size_bucket}_mismatch_ew"]]

    figure_series = None
    if horizon_label == "3m" and figure_parts:
        figure_series = figure_parts["small"].merge(figure_parts["big"], on="month", how="outer")
    return rows, figure_series


def _plot_cumulative_series(figure_df: pd.DataFrame, output_base: Path) -> tuple[Path, Path]:
    _base_style()
    plot_df = figure_df.sort_values("month").copy()
    plot_df["month_end"] = plot_df["month"].dt.to_timestamp("M")
    for column in ["as_q3_q1_ew", "small_mismatch_ew", "big_mismatch_ew"]:
        plot_df[f"cum_{column}"] = (1.0 + plot_df[column].fillna(0)).cumprod() - 1.0

    fig, ax = plt.subplots(figsize=(8.6, 4.9), constrained_layout=True)
    ax.plot(
        plot_df["month_end"],
        100 * plot_df["cum_as_q3_q1_ew"],
        color="#1d3557",
        linewidth=2.1,
        label="A/S Q3 - Q1 (EW, 3m)",
    )
    ax.plot(
        plot_df["month_end"],
        100 * plot_df["cum_small_mismatch_ew"],
        color="#c46b48",
        linewidth=2.1,
        label="Small-firm credible - mismatch (EW, 3m)",
    )
    ax.plot(
        plot_df["month_end"],
        100 * plot_df["cum_big_mismatch_ew"],
        color="#6d597a",
        linewidth=2.1,
        label="Big-firm credible - mismatch (EW, 3m)",
    )
    ax.axhline(0, color="#6c757d", linewidth=0.8)
    ax.set_ylabel("Cumulative portfolio return (pct)")
    ax.set_xlabel("Calendar month")
    ax.set_title("Portfolio-Sort Return Paths After AI Filing Signals")
    ax.legend(frameon=False, loc="best")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.grid(True, axis="y", color="#d9d9d9", linewidth=0.7)
    ax.grid(False, axis="x")

    png_path = output_base.with_suffix(".png")
    pdf_path = output_base.with_suffix(".pdf")
    fig.savefig(png_path, dpi=220, bbox_inches="tight")
    fig.savefig(pdf_path, bbox_inches="tight")
    plt.close(fig)
    return png_path, pdf_path


def _format_pct(value: float) -> str:
    return "" if pd.isna(value) else f"{100 * value:.3f}"


def _format_float(value: float) -> str:
    return "" if pd.isna(value) else f"{value:.3f}"


def _panel_rows_as(
    as_rows: list[dict[str, object]], weighting: str
) -> list[list[tuple[str, bool]]]:
    keyed = {
        (row["portfolio"], row["horizon"]): row for row in as_rows if row["weighting"] == weighting
    }
    rendered: list[list[tuple[str, bool]]] = []
    for portfolio in ["Q1 (low A/S)", "Q2", "Q3 (high A/S)", "Q3 - Q1"]:
        row = [(portfolio, True)]
        for horizon in HORIZONS.values():
            stats = keyed[(portfolio, horizon)]
            row.extend(
                [
                    (_format_pct(stats["mean_monthly_return"]), False),
                    (_format_pct(stats["market_alpha"]), False),
                    (_format_float(stats["market_alpha_p"]), False),
                ]
            )
        rendered.append(row)
    return rendered


def _panel_rows_size(size_rows: list[dict[str, object]]) -> list[list[tuple[str, bool]]]:
    keyed = {(row["sample"], row["weighting"], row["horizon"]): row for row in size_rows}
    rendered: list[list[tuple[str, bool]]] = []
    for sample_key, weighting, label in MISMATCH_ROWS:
        row = [(label, True)]
        for horizon in HORIZONS.values():
            stats = keyed[(sample_key, weighting, horizon)]
            row.extend(
                [
                    (_format_pct(stats["mean_monthly_return"]), False),
                    (_format_pct(stats["market_alpha"]), False),
                    (_format_float(stats["market_alpha_p"]), False),
                ]
            )
        rendered.append(row)
    return rendered


def _markdown_table(headers: list[str], rows: list[list[str]]) -> str:
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join(["---"] * len(headers)) + " |",
    ]
    for row in rows:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def _render_table_outputs(
    as_rows: list[dict[str, object]], size_rows: list[dict[str, object]]
) -> tuple[str, str]:
    headers = [
        "Row",
        "1m mean",
        "1m alpha",
        "1m p",
        "3m mean",
        "3m alpha",
        "3m p",
    ]
    panel_a = [[cell for cell, _ in row] for row in _panel_rows_as(as_rows, "ew")]
    panel_b = [[cell for cell, _ in row] for row in _panel_rows_as(as_rows, "vw")]
    panel_c = [[cell for cell, _ in row] for row in _panel_rows_size(size_rows)]

    md = "\n".join(
        [
            "# Table Main",
            "",
            "## Panel A. A/S tercile sorts, equal-weight",
            _markdown_table(headers, panel_a),
            "",
            "## Panel B. A/S tercile sorts, value-weight",
            _markdown_table(headers, panel_b),
            "",
            "## Panel C. Credible-minus-mismatch spread, overall and by size",
            _markdown_table(headers, panel_c),
            "",
        ]
    )

    latex_lines = [
        "\\begin{table}[!htbp]",
        "\\centering",
        "\\caption{Portfolio sorts on AI disclosure credibility after annual filings}",
        "\\begin{tabular}{l" + "c" * (len(headers) - 1) + "}",
        "\\hline",
        " & " + " & ".join(headers[1:]) + " \\\\",
        "\\hline",
        "\\multicolumn{7}{l}{\\textit{Panel A. A/S tercile sorts, equal-weight}} \\\\",
    ]
    for row in panel_a:
        latex_lines.append(row[0] + " & " + " & ".join(row[1:]) + " \\\\")
    latex_lines.append(
        "\\multicolumn{7}{l}{\\textit{Panel B. A/S tercile sorts, value-weight}} \\\\"
    )
    for row in panel_b:
        latex_lines.append(row[0] + " & " + " & ".join(row[1:]) + " \\\\")
    latex_lines.append(
        "\\multicolumn{7}{l}{\\textit{Panel C. Credible-minus-mismatch spread, overall and by size}} \\\\"
    )
    for row in panel_c:
        latex_lines.append(row[0] + " & " + " & ".join(row[1:]) + " \\\\")
    latex_lines.extend(["\\hline", "\\end{tabular}", "\\end{table}"])
    return md, "\n".join(latex_lines) + "\n"


def _build_table_docx(
    as_rows: list[dict[str, object]], size_rows: list[dict[str, object]], output_path: Path
) -> None:
    document = Document()
    _set_document_defaults(document)
    _set_landscape(document)
    _add_title(document, "Table 4. Portfolio Sorts on AI Disclosure Credibility")
    note = (
        "This table reports monthly portfolio returns sorted on filing-based AI disclosure credibility signals. Panel A and Panel B sort active stock-month observations into terciles on the A/S ratio using the most recent annual filing signal carried forward for one or three months after the filing month. "
        "Panel A reports equal-weight portfolios and Panel B reports value-weight portfolios using lagged market capitalization. Panel C reports the credible-minus-mismatch spread overall and separately for small and big firms, where size uses the active-sample monthly median lagged market cap because NYSE breakpoints are not available in the local CRSP monthly extract. Market alpha is the intercept from a regression of the monthly portfolio return on the CRSP value-weighted market return."
    )
    _add_note(document, note)

    headers = ["", "1m mean", "1m alpha", "1m p", "3m mean", "3m alpha", "3m p"]

    p = document.add_paragraph()
    p.add_run("Panel A. A/S tercile sorts, equal-weight").bold = True
    _build_panel_table(document, headers, _panel_rows_as(as_rows, "ew"))

    p = document.add_paragraph()
    p.add_run("Panel B. A/S tercile sorts, value-weight").bold = True
    _build_panel_table(document, headers, _panel_rows_as(as_rows, "vw"))

    p = document.add_paragraph()
    p.add_run("Panel C. Credible-minus-mismatch spread, overall and by size").bold = True
    _build_panel_table(document, headers, _panel_rows_size(size_rows))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def _result_notes(as_rows: list[dict[str, object]], size_rows: list[dict[str, object]]) -> str:
    as_key = {(row["weighting"], row["portfolio"], row["horizon"]): row for row in as_rows}
    size_key = {(row["sample"], row["weighting"], row["horizon"]): row for row in size_rows}
    return "\n".join(
        [
            "# Result Notes",
            "",
            f"- Equal-weight A/S high-minus-low spread at 3m: mean monthly return `{100 * as_key[('ew', 'Q3 - Q1', '3m')]['mean_monthly_return']:.3f}` percentage points; market alpha `{100 * as_key[('ew', 'Q3 - Q1', '3m')]['market_alpha']:.3f}` pp with p=`{as_key[('ew', 'Q3 - Q1', '3m')]['market_alpha_p']:.3f}`.",
            f"- Value-weight A/S high-minus-low spread at 3m: market alpha `{100 * as_key[('vw', 'Q3 - Q1', '3m')]['market_alpha']:.3f}` pp with p=`{as_key[('vw', 'Q3 - Q1', '3m')]['market_alpha_p']:.3f}`.",
            f"- Small-firm credible-minus-mismatch spread at 3m (EW): market alpha `{100 * size_key[('small', 'ew', '3m')]['market_alpha']:.3f}` pp with p=`{size_key[('small', 'ew', '3m')]['market_alpha_p']:.3f}`.",
            f"- Big-firm credible-minus-mismatch spread at 3m (EW): market alpha `{100 * size_key[('big', 'ew', '3m')]['market_alpha']:.3f}` pp with p=`{size_key[('big', 'ew', '3m')]['market_alpha_p']:.3f}`.",
            "",
        ]
    )


def _writer_packet(
    args: argparse.Namespace,
    signal_sample: pd.DataFrame,
    as_rows: list[dict[str, object]],
    size_rows: list[dict[str, object]],
) -> str:
    as_key = {(row["weighting"], row["portfolio"], row["horizon"]): row for row in as_rows}
    size_key = {(row["sample"], row["weighting"], row["horizon"]): row for row in size_rows}
    return "\n".join(
        [
            "# Writer Packet",
            "",
            "## Metadata",
            f"- Test id: `{TEST_ID}`",
            f"- Run id: `{args.run_id}`",
            f"- Date run: `{date.today().isoformat()}`",
            f"- Script/module path: `{MODULE_PATH}`",
            f"- Input files: `{args.event_panel}`, `{args.annual_panel}`, `{args.monthly_returns}`, `{args.market_index}`",
            "- Unit of observation: `calendar month portfolio return`",
            "- Sample filters: `AI-talking annual filers with merged grant-based PatentMismatch and valid CRSP monthly return`",
            "",
            "## Portfolio Construction",
            "- Sort variable 1: `A_S` terciles assigned each calendar month among active stock-month signals",
            "- Sort variable 2: `PatentMismatch` binary split, summarized as credible minus mismatch",
            "- Rebalance timing: `monthly, using the latest filing signal mapped to each active stock-month`",
            "- Holding periods in main table: `1 month and 3 months after filing month, starting in month +1`",
            "- Equal-weight version: `simple average across active stock-month constituents`",
            "- Value-weight version: `lagged market-cap weights from CRSP MSF`",
            "- Size split: `monthly active-sample median lagged market cap (NYSE breakpoints unavailable locally)`",
            "",
            "## Results",
            f"- Strongest A/S sort result: `EW 3m Q3-Q1 alpha = {100 * as_key[('ew', 'Q3 - Q1', '3m')]['market_alpha']:.3f} pp/month`, p = `{as_key[('ew', 'Q3 - Q1', '3m')]['market_alpha_p']:.3f}`",
            f"- Small-firm mismatch spread: `EW 3m alpha = {100 * size_key[('small', 'ew', '3m')]['market_alpha']:.3f} pp/month`, p = `{size_key[('small', 'ew', '3m')]['market_alpha_p']:.3f}`",
            f"- Big-firm mismatch spread: `EW 3m alpha = {100 * size_key[('big', 'ew', '3m')]['market_alpha']:.3f} pp/month`, p = `{size_key[('big', 'ew', '3m')]['market_alpha_p']:.3f}`",
            f"- Distinct permnos represented in the filing signal sample: `{signal_sample['permno'].nunique()}`",
            "",
            "## Caption Draft",
            "This table reports monthly portfolio returns sorted on filing-based AI disclosure credibility signals. Panels A and B sort active stock-month observations into terciles on the A/S ratio using the most recent annual filing signal carried forward for one or three months after the filing month. Panel A reports equal-weight portfolios and Panel B reports value-weight portfolios using lagged market capitalization. Panel C reports the credible-minus-mismatch spread overall and separately for small and big firms, where size uses the monthly active-sample median lagged market cap because NYSE breakpoints are not available in the local CRSP monthly extract. Market alpha is the intercept from a regression of monthly portfolio returns on the CRSP value-weighted market return.",
            "",
        ]
    )


def _dataset_summary(
    args: argparse.Namespace,
    signal_sample: pd.DataFrame,
    as_rows: list[dict[str, object]],
    size_rows: list[dict[str, object]],
    figure_df: pd.DataFrame,
) -> dict[str, object]:
    figure_payload = figure_df.copy()
    figure_payload["month"] = figure_payload["month"].astype(str)
    return {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "event_panel": str(args.event_panel),
        "annual_panel": str(args.annual_panel),
        "monthly_returns": str(args.monthly_returns),
        "market_index": str(args.market_index),
        "signal_filing_count": int(len(signal_sample)),
        "unique_permno": int(signal_sample["permno"].nunique()),
        "filing_year_min": int(signal_sample["filing_year"].min()),
        "filing_year_max": int(signal_sample["filing_year"].max()),
        "as_sort_summaries": as_rows,
        "mismatch_size_summaries": size_rows,
        "figure_series": figure_payload.to_dict(orient="records"),
    }


def _copy_exports(run_dir: Path, paper_root: Path, run_id: str) -> dict[str, str]:
    exports = {
        "figure_png": paper_root / "figures" / f"{TEST_ID}_{run_id}.png",
        "figure_pdf": paper_root / "figures" / f"{TEST_ID}_{run_id}.pdf",
        "table_csv": paper_root / "tables" / f"{TEST_ID}_{run_id}.csv",
        "table_md": paper_root / "tables" / f"{TEST_ID}_{run_id}.md",
        "table_tex": paper_root / "latex" / f"{TEST_ID}_{run_id}.tex",
        "table_docx": paper_root / "docx" / f"{TEST_ID}_{run_id}.docx",
        "figure_series": paper_root / "tables" / f"{TEST_ID}_{run_id}_figure_series.csv",
        "writer_packet": paper_root / "writer_packets" / f"{TEST_ID}_{run_id}.md",
        "result_notes": paper_root / "snippets" / f"{TEST_ID}_{run_id}_result_notes.md",
    }
    for path in exports.values():
        path.parent.mkdir(parents=True, exist_ok=True)
    mapping = {
        run_dir / "figure_main.png": exports["figure_png"],
        run_dir / "figure_main.pdf": exports["figure_pdf"],
        run_dir / "table_main.csv": exports["table_csv"],
        run_dir / "table_main.md": exports["table_md"],
        run_dir / "table_main.tex": exports["table_tex"],
        run_dir / "table_main.docx": exports["table_docx"],
        run_dir / "figure_series.csv": exports["figure_series"],
        run_dir / "writer_packet.md": exports["writer_packet"],
        run_dir / "result_notes.md": exports["result_notes"],
    }
    for src, dst in mapping.items():
        shutil.copy2(src, dst)
    return {key: str(path) for key, path in exports.items()}


def main() -> None:
    args = _parse_args()
    run_dir = args.test_root / args.run_id
    run_dir.mkdir(parents=True, exist_ok=True)

    signal_sample = _prepare_signal_sample(args.event_panel, args.annual_panel)
    monthly_returns = _load_monthly_returns(args.monthly_returns)
    market_index = _load_market_index(args.market_index)

    as_rows: list[dict[str, object]] = []
    size_rows: list[dict[str, object]] = []
    figure_df = None
    for horizon_months, horizon_label in HORIZONS.items():
        active_positions = _build_active_positions(signal_sample, horizon_months)
        merged = _merge_returns(active_positions, monthly_returns)

        as_part, as_figure = _as_sort_summaries(merged, market_index, horizon_label=horizon_label)
        size_part, size_figure = _size_split_summaries(
            merged, market_index, horizon_label=horizon_label
        )
        as_rows.extend(as_part)
        size_rows.extend(size_part)
        if horizon_label == "3m":
            figure_df = as_figure.merge(size_figure, on="month", how="outer")

    table_df = pd.concat([pd.DataFrame(as_rows), pd.DataFrame(size_rows)], ignore_index=True)
    table_df.to_csv(run_dir / "table_main.csv", index=False)
    figure_df.to_csv(run_dir / "figure_series.csv", index=False)

    table_md, table_tex = _render_table_outputs(as_rows, size_rows)
    (run_dir / "table_main.md").write_text(table_md, encoding="utf-8")
    (run_dir / "table_main.tex").write_text(table_tex, encoding="utf-8")
    _build_table_docx(as_rows, size_rows, run_dir / "table_main.docx")
    png_path, pdf_path = _plot_cumulative_series(figure_df, run_dir / "figure_main")

    (run_dir / "result_notes.md").write_text(_result_notes(as_rows, size_rows), encoding="utf-8")
    (run_dir / "writer_packet.md").write_text(
        _writer_packet(args, signal_sample, as_rows, size_rows), encoding="utf-8"
    )
    dataset_summary = _dataset_summary(args, signal_sample, as_rows, size_rows, figure_df)
    (run_dir / "dataset_summary.json").write_text(
        json.dumps(dataset_summary, indent=2), encoding="utf-8"
    )
    paper_exports = _copy_exports(run_dir, args.paper_root, args.run_id)
    manifest = {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "module_path": MODULE_PATH,
        "inputs": {
            "event_panel": str(args.event_panel),
            "annual_panel": str(args.annual_panel),
            "monthly_returns": str(args.monthly_returns),
            "market_index": str(args.market_index),
        },
        "run_dir": str(run_dir),
        "outputs": {
            "dataset_summary": str(run_dir / "dataset_summary.json"),
            "table_csv": str(run_dir / "table_main.csv"),
            "table_md": str(run_dir / "table_main.md"),
            "table_tex": str(run_dir / "table_main.tex"),
            "table_docx": str(run_dir / "table_main.docx"),
            "figure_series": str(run_dir / "figure_series.csv"),
            "figure_png": str(png_path),
            "figure_pdf": str(pdf_path),
            "writer_packet": str(run_dir / "writer_packet.md"),
            "result_notes": str(run_dir / "result_notes.md"),
        },
        "paper_exports": paper_exports,
    }
    (run_dir / "run_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    print(f"[{TEST_ID}] wrote run bundle to {run_dir}")
    print(f"[{TEST_ID}] paper table: {paper_exports['table_docx']}")


if __name__ == "__main__":
    main()
