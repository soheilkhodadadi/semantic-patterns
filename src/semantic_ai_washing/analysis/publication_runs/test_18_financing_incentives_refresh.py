"""Publication run driver for Test 18: financing and valuation refresh in the AI-talking sample."""

from __future__ import annotations

import argparse
import json
import math
import shutil
from datetime import UTC, date, datetime
from pathlib import Path

from docx import Document
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from semantic_ai_washing.analysis.delivery_table_payloads import _fit_absorbed_ols
from semantic_ai_washing.analysis.publication_runs.test_03_post_filing_drift import (
    _add_note,
    _add_title,
    _build_panel_table,
    _set_document_defaults,
    _set_landscape,
)
from semantic_ai_washing.analysis.publication_runs.test_16_construct_variant_screen import (
    _add_construct_variants,
)

REPO_ROOT = Path(__file__).resolve().parents[4]
DEFAULT_ANNUAL_PANEL = (
    REPO_ROOT
    / "data/processed/panel/canonical/ever_speaker_panel_2016_2025_hybrid_api_a_conf49_v1.parquet"
)
DEFAULT_TEST_ROOT = Path(
    "/Users/soheilkhodadadi/DataWork/semantic-patterns/ai_washing/derived/v3_1/test_runs/test_18_financing_incentives_refresh"
)
DEFAULT_PAPER_ROOT = REPO_ROOT / "paper/generated/v3_1"
DEFAULT_RUN_ID = f"{date.today():%Y%m%d}_aiw_v3_1_test_18_financing_incentives_refresh_main_v1"
TEST_ID = "test_18_financing_incentives_refresh"
MODULE_PATH = "semantic_ai_washing.analysis.publication_runs.test_18_financing_incentives_refresh"
VARIANT_SPECS: list[tuple[str, str]] = [
    ("PatentMismatch", "Canonical grant mismatch"),
    ("ApplicationMismatch", "Application mismatch"),
    ("LowCredibility", "Low credibility only"),
]
OUTCOME_SPECS: list[tuple[str, str, list[str], str]] = [
    (
        "log_mktcap_assets",
        "Log MktCap/assets",
        ["AI_Focus", "ln_assets", "cash", "roa"],
        "valuation",
    ),
    ("log_q_proxy", "Log Q proxy", ["AI_Focus", "ln_assets", "cash", "roa"], "valuation"),
    (
        "delta_log_q_proxy_lead1",
        "Delta log Q t+1",
        ["AI_Focus", "ln_assets", "cash", "roa"],
        "valuation",
    ),
    (
        "share_growth_lead1",
        "Delta Shares t+1",
        ["AI_Focus", "ln_assets", "leverage", "cash", "roa"],
        "financing",
    ),
    (
        "equity_issue_lead1",
        "Issue>5% t+1",
        ["AI_Focus", "ln_assets", "leverage", "cash", "roa"],
        "financing",
    ),
]
SUBSET_SPECS = [
    ("all", "Panel A. Full AI-talking sample"),
    ("nonbig", "Panel B. Non-big AI-talking sample"),
]


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--annual-panel", type=Path, default=DEFAULT_ANNUAL_PANEL)
    parser.add_argument("--test-root", type=Path, default=DEFAULT_TEST_ROOT)
    parser.add_argument("--paper-root", type=Path, default=DEFAULT_PAPER_ROOT)
    parser.add_argument("--run-id", default=DEFAULT_RUN_ID)
    return parser.parse_args()


def _base_style() -> None:
    plt.style.use("seaborn-v0_8-whitegrid")
    plt.rcParams.update(
        {
            "font.family": "serif",
            "font.serif": ["Times New Roman", "Times", "DejaVu Serif"],
            "axes.titlesize": 12,
            "axes.labelsize": 10,
            "xtick.labelsize": 9,
            "ytick.labelsize": 9,
        }
    )


def _winsorize(series: pd.Series, lower_q: float = 0.01, upper_q: float = 0.99) -> pd.Series:
    valid = pd.to_numeric(series, errors="coerce").replace([np.inf, -np.inf], np.nan)
    if valid.notna().sum() == 0:
        return valid
    lower = valid.quantile(lower_q)
    upper = valid.quantile(upper_q)
    return valid.clip(lower=lower, upper=upper)


def _nonbig_mask(series: pd.Series) -> pd.Series:
    return pd.Series(series, copy=False).eq(True).fillna(False)


def _prepare_sample(panel: pd.DataFrame) -> tuple[pd.DataFrame, dict[str, object]]:
    use = _add_construct_variants(panel).copy()
    use = use.loc[use["any_ai_talk"].fillna(0).astype(int).eq(1)].copy()
    for column in [
        "AI_Focus",
        "ln_assets",
        "leverage",
        "cash",
        "roa",
        "market_cap_year_end",
        "shrout",
        *[variant for variant, _label in VARIANT_SPECS],
    ]:
        if column in use.columns:
            use[column] = pd.to_numeric(use[column], errors="coerce").replace(
                [np.inf, -np.inf], np.nan
            )

    use["assets_proxy"] = np.exp(use["ln_assets"])
    use.loc[~np.isfinite(use["assets_proxy"]) | use["assets_proxy"].le(0), "assets_proxy"] = np.nan
    use["mktcap_assets"] = use["market_cap_year_end"] / use["assets_proxy"]
    use["q_proxy"] = use["mktcap_assets"] + use["leverage"]
    use["log_mktcap_assets"] = np.log1p(use["mktcap_assets"].clip(lower=0))
    use["log_q_proxy"] = np.log1p(use["q_proxy"].clip(lower=0))

    use = use.sort_values(["permno", "year", "cik"]).reset_index(drop=True)
    shrout_lead1 = use.groupby("permno", sort=False)["shrout"].shift(-1)
    log_q_lead1 = use.groupby("cik", sort=False)["log_q_proxy"].shift(-1)
    use["share_growth_lead1"] = shrout_lead1 / use["shrout"] - 1.0
    use["equity_issue_lead1"] = np.where(
        use["share_growth_lead1"].notna(), (use["share_growth_lead1"] > 0.05).astype(float), np.nan
    )
    use["delta_log_q_proxy_lead1"] = log_q_lead1 - use["log_q_proxy"]

    for column in [
        "log_mktcap_assets",
        "log_q_proxy",
        "delta_log_q_proxy_lead1",
        "share_growth_lead1",
    ]:
        use[column] = _winsorize(use[column])

    yearly_median_market_cap = use.groupby("year")["market_cap_year_end"].transform("median")
    use["nonbig"] = np.where(
        use["market_cap_year_end"].notna() & yearly_median_market_cap.notna(),
        use["market_cap_year_end"].le(yearly_median_market_cap),
        pd.NA,
    )
    use = use.loc[use["cik"].astype(str).str.len().gt(0) & use["year"].notna()].copy()
    summary = {
        "ai_talking_rows": int(len(use)),
        "unique_firms": int(use["cik"].nunique()),
        "year_min": int(pd.to_numeric(use["year"], errors="coerce").min()),
        "year_max": int(pd.to_numeric(use["year"], errors="coerce").max()),
        "nonbig_rows": int(_nonbig_mask(use["nonbig"]).sum()),
        "nonbig_firms": int(use.loc[_nonbig_mask(use["nonbig"]), "cik"].nunique()),
        "outcome_nonmissing": {
            outcome: int(pd.to_numeric(use[outcome], errors="coerce").notna().sum())
            for outcome, _label, _controls, _family in OUTCOME_SPECS
        },
    }
    return use, summary


def _subset_frame(sample: pd.DataFrame, subset_key: str) -> pd.DataFrame:
    if subset_key == "all":
        return sample.copy()
    if subset_key == "nonbig":
        return sample.loc[_nonbig_mask(sample["nonbig"])].copy()
    raise ValueError(f"Unknown subset: {subset_key}")


def _fit_model(
    sample: pd.DataFrame, subset_key: str, variant: str, outcome: str, controls: list[str]
) -> dict[str, object]:
    subset = _subset_frame(sample, subset_key)
    result, use, adj_r2 = _fit_absorbed_ols(
        subset,
        dependent=outcome,
        rhs_terms=[variant],
        absorb_col="cik",
        include_year=True,
        controls=controls,
    )
    return {
        "subset": subset_key,
        "variant": variant,
        "outcome": outcome,
        "nobs": int(result.nobs),
        "adj_r2": float(adj_r2) if adj_r2 is not None else math.nan,
        "params": result.params.to_dict(),
        "bse": result.bse.to_dict(),
        "pvalues": result.pvalues.to_dict(),
        "outcome_mean": float(use[outcome].mean()),
    }


def _stars(p_value: float | None) -> str:
    if p_value is None or not math.isfinite(p_value):
        return ""
    if p_value < 0.01:
        return "***"
    if p_value < 0.05:
        return "**"
    if p_value < 0.10:
        return "*"
    return ""


def _coef_cell(coef: float | None, se: float | None, p_value: float | None) -> str:
    if coef is None or se is None or not math.isfinite(coef) or not math.isfinite(se):
        return ""
    return f"{coef:.4f}{_stars(p_value)} ({se:.4f})"


def _build_results_table(
    model_rows: list[dict[str, object]],
) -> tuple[pd.DataFrame, dict[tuple[str, str, str], dict[str, object]]]:
    keyed = {(row["subset"], row["outcome"], row["variant"]): row for row in model_rows}
    rows: list[dict[str, object]] = []
    for subset_key, panel_label in SUBSET_SPECS:
        for family in ["valuation", "financing"]:
            for outcome, outcome_label, _controls, outcome_family in OUTCOME_SPECS:
                if outcome_family != family:
                    continue
                row = {"panel": panel_label, "family": family, "row_label": outcome_label}
                for variant, variant_label in VARIANT_SPECS:
                    result = keyed[(subset_key, outcome, variant)]
                    row[variant_label] = _coef_cell(
                        result["params"].get(variant),
                        result["bse"].get(variant),
                        result["pvalues"].get(variant),
                    )
                base = keyed[(subset_key, outcome, VARIANT_SPECS[0][0])]
                row["Outcome mean"] = f"{base['outcome_mean']:.3f}"
                row["Observations"] = str(base["nobs"])
                rows.append(row)
    return pd.DataFrame(rows), keyed


def _render_table_outputs(table_df: pd.DataFrame) -> tuple[str, str]:
    headers = (
        ["Outcome"]
        + [label for _variant, label in VARIANT_SPECS]
        + ["Outcome mean", "Observations"]
    )

    def md_table(panel: str) -> str:
        lines = [
            "| " + " | ".join(headers) + " |",
            "| " + " | ".join(["---"] * len(headers)) + " |",
        ]
        panel_df = table_df.loc[table_df["panel"].eq(panel)]
        for family in ["valuation", "financing"]:
            fam = panel_df.loc[panel_df["family"].eq(family)]
            if fam.empty:
                continue
            lines.append(f"| **{family.title()} outcomes** |  |  |  |  |  |")
            for _, row in fam.iterrows():
                cells = [str(row["row_label"])] + [str(row[h]) for h in headers[1:]]
                lines.append("| " + " | ".join(cells) + " |")
        return "\n".join(lines)

    md = "\n".join(
        [
            "# Table Main",
            "",
            "## Panel A. Full AI-talking sample",
            md_table("Panel A. Full AI-talking sample"),
            "",
            "## Panel B. Non-big AI-talking sample",
            md_table("Panel B. Non-big AI-talking sample"),
            "",
        ]
    )

    latex_lines = [
        "\\begin{table}[!htbp]",
        "\\centering",
        "\\caption{Financing and valuation refresh in the AI-talking sample}",
        "\\begin{tabular}{l" + "c" * (len(headers) - 1) + "}",
        "\\hline",
        "Outcome & Canonical grant mismatch & Application mismatch & Low credibility only & Outcome mean & Observations \\\\",
        "\\hline",
    ]
    for panel in ["Panel A. Full AI-talking sample", "Panel B. Non-big AI-talking sample"]:
        latex_lines.append(f"\\multicolumn{{6}}{{l}}{{\\textit{{{panel}}}}} \\\\")
        panel_df = table_df.loc[table_df["panel"].eq(panel)]
        for family in ["valuation", "financing"]:
            fam = panel_df.loc[panel_df["family"].eq(family)]
            if fam.empty:
                continue
            latex_lines.append(f"\\multicolumn{{6}}{{l}}{{{family.title()} outcomes}} \\\\")
            for _, row in fam.iterrows():
                latex_lines.append(
                    " & ".join([str(row["row_label"])] + [str(row[h]) for h in headers[1:]])
                    + " \\\\"
                )
        latex_lines.append("\\hline")
    latex_lines.extend(["\\end{tabular}", "\\end{table}"])
    return md, "\n".join(latex_lines) + "\n"


def _build_table_docx(table_df: pd.DataFrame, output_path: Path) -> None:
    document = Document()
    _set_document_defaults(document)
    _set_landscape(document)
    _add_title(document, "Table 18. Financing and Valuation Refresh in the AI-Talking Sample")
    _add_note(
        document,
        "This table refreshes the financing and valuation lane inside the AI-talking sample. Panel A uses the full AI-talking sample; Panel B restricts to non-big firms using the yearly median market-cap split available in the local panel. Valuation outcomes are log market-capitalization-to-assets, a log Q-style proxy, and next-year change in that Q proxy. Financing outcomes are next-year share growth and an indicator for share growth above 5 percent. All specifications absorb firm and year fixed effects, use firm-clustered standard errors, and compare the canonical grant-based mismatch construct, an application-side mismatch variant, and a disclosure-only low-credibility variant.",
    )
    headers = (
        [""] + [label for _variant, label in VARIANT_SPECS] + ["Outcome mean", "Observations"]
    )
    for panel in ["Panel A. Full AI-talking sample", "Panel B. Non-big AI-talking sample"]:
        p = document.add_paragraph()
        p.add_run(panel).bold = True
        panel_df = table_df.loc[table_df["panel"].eq(panel)]
        rows: list[list[tuple[str, bool]]] = []
        for family in ["valuation", "financing"]:
            fam = panel_df.loc[panel_df["family"].eq(family)]
            if fam.empty:
                continue
            rows.append(
                [(f"{family.title()} outcomes", True)] + [("", False)] * (len(headers) - 1)
            )
            for _, row in fam.iterrows():
                rows.append(
                    [(str(row["row_label"]), True)] + [(str(row[h]), False) for h in headers[1:]]
                )
        _build_panel_table(document, headers, rows)
    document.save(output_path)


def _plot_heatmap(
    model_rows: list[dict[str, object]], output_stem: Path
) -> tuple[Path, Path, pd.DataFrame]:
    _base_style()
    records = []
    row_order: list[str] = []
    for subset_key, panel_label in SUBSET_SPECS:
        for family in ["valuation", "financing"]:
            for outcome, outcome_label, _controls, outcome_family in OUTCOME_SPECS:
                if outcome_family != family:
                    continue
                row_name = f"{panel_label.replace('Panel A. ', '').replace('Panel B. ', '')}: {outcome_label}"
                row_order.append(row_name)
                for variant, variant_label in VARIANT_SPECS:
                    result = next(
                        r
                        for r in model_rows
                        if r["subset"] == subset_key
                        and r["outcome"] == outcome
                        and r["variant"] == variant
                    )
                    records.append(
                        {
                            "row_name": row_name,
                            "variant": variant_label,
                            "coef": float(result["params"].get(variant, math.nan)),
                            "p_value": float(result["pvalues"].get(variant, math.nan)),
                        }
                    )
    figure_df = pd.DataFrame(records)
    variant_order = [label for _variant, label in VARIANT_SPECS]
    pivot = figure_df.pivot(index="row_name", columns="variant", values="coef").reindex(
        index=row_order, columns=variant_order
    )
    ann = (
        figure_df.assign(
            label=lambda x: x.apply(lambda r: f"{r['coef']:.3f}{_stars(r['p_value'])}", axis=1)
        )
        .pivot(index="row_name", columns="variant", values="label")
        .reindex(index=row_order, columns=variant_order)
    )
    vmax = float(np.nanmax(np.abs(pivot.to_numpy(dtype=float)))) if not pivot.empty else 0.1
    vmax = max(vmax, 0.05)
    fig, ax = plt.subplots(figsize=(8.8, 5.6))
    im = ax.imshow(
        pivot.to_numpy(dtype=float), cmap="RdBu_r", vmin=-vmax, vmax=vmax, aspect="auto"
    )
    ax.set_xticks(range(len(pivot.columns)), list(pivot.columns), rotation=12, ha="right")
    ax.set_yticks(range(len(pivot.index)), list(pivot.index))
    ax.set_title("Financing and Valuation Refresh: AI-Talking Sample")
    for i, row_name in enumerate(pivot.index):
        for j, col_name in enumerate(pivot.columns):
            ax.text(j, i, ann.loc[row_name, col_name], ha="center", va="center", fontsize=8)
    cbar = fig.colorbar(im, ax=ax, shrink=0.9)
    cbar.set_label("Coefficient")
    fig.tight_layout()
    png_path = output_stem.with_suffix(".png")
    pdf_path = output_stem.with_suffix(".pdf")
    fig.savefig(png_path, dpi=240)
    fig.savefig(pdf_path)
    plt.close(fig)
    return png_path, pdf_path, figure_df


def _data_gap_note() -> str:
    return "\n".join(
        [
            f"# Data Gap Note: {TEST_ID}",
            "",
            "The current canonical panel is strong enough for a financing and valuation refresh, but it does not yet support a true governance extension.",
            "",
            "Missing from the current panel are the fields usually needed for a stricter incentives / governance section:",
            "- institutional ownership or 13F concentration;",
            "- analyst coverage or forecast dispersion;",
            "- board independence, CEO-chair duality, or governance-score fields;",
            "- richer executive incentive variables such as option intensity or performance pay.",
            "",
            "That means Packet C can credibly refresh valuation and issuance-style outcomes now, but a genuine governance packet should wait for a separate data pull rather than being improvised from the current panel.",
            "",
        ]
    )


def _result_notes(model_rows: list[dict[str, object]], sample_summary: dict[str, object]) -> str:
    keyed = {(row["subset"], row["outcome"], row["variant"]): row for row in model_rows}
    full_q = keyed[("all", "log_q_proxy", "PatentMismatch")]
    nonbig_q = keyed[("nonbig", "log_q_proxy", "PatentMismatch")]
    nonbig_val = keyed[("nonbig", "log_mktcap_assets", "PatentMismatch")]
    financing = keyed[("all", "equity_issue_lead1", "PatentMismatch")]
    return "\n".join(
        [
            f"# Result Notes: {TEST_ID}",
            "",
            f"- Sample: `{sample_summary['ai_talking_rows']}` AI-talking firm-years across `{sample_summary['unique_firms']}` firms; non-big subset = `{sample_summary['nonbig_rows']}` rows.",
            f"- Full-sample Log Q proxy with canonical PatentMismatch: `{full_q['params'].get('PatentMismatch', math.nan):.4f}` (p=`{full_q['pvalues'].get('PatentMismatch', math.nan):.3f}`).",
            f"- Non-big Log Q proxy with canonical PatentMismatch: `{nonbig_q['params'].get('PatentMismatch', math.nan):.4f}` (p=`{nonbig_q['pvalues'].get('PatentMismatch', math.nan):.3f}`).",
            f"- Non-big Log MktCap/assets with canonical PatentMismatch: `{nonbig_val['params'].get('PatentMismatch', math.nan):.4f}` (p=`{nonbig_val['pvalues'].get('PatentMismatch', math.nan):.3f}`).",
            f"- Full-sample Issue>5% t+1 with canonical PatentMismatch: `{financing['params'].get('PatentMismatch', math.nan):.4f}` (p=`{financing['pvalues'].get('PatentMismatch', math.nan):.3f}`).",
            "- Read this table as a refresh of the incentives lane. The valuation signal is concentrated in the non-big subset, while issuance-style outcomes remain mostly null in the current panel.",
            "",
        ]
    )


def _writer_packet(
    args: argparse.Namespace,
    model_rows: list[dict[str, object]],
    sample_summary: dict[str, object],
) -> str:
    keyed = {(row["subset"], row["outcome"], row["variant"]): row for row in model_rows}
    return "\n".join(
        [
            f"# Writer Packet: {TEST_ID}",
            "",
            "## Role In The Paper",
            "- This run refreshes the financing / valuation lane under the new discipline rather than assuming the older valuation table still carries the same interpretation.",
            "- The table is useful mainly because it shows where the signal does and does not survive: valuation is concentrated in non-big firms, while issuance outcomes remain weak.",
            "",
            "## Sample Definition",
            f"- Annual panel: `{args.annual_panel}`",
            "- Working sample: `AI-talking firm-years only`",
            f"- Rows: `{sample_summary['ai_talking_rows']}`",
            f"- Non-big rows: `{sample_summary['nonbig_rows']}`",
            "- Fixed effects: `firm + year`",
            "",
            "## Headline Read",
            f"- Full-sample Log Q proxy, canonical PatentMismatch: `{keyed[('all', 'log_q_proxy', 'PatentMismatch')]['params'].get('PatentMismatch', math.nan):.4f}` (p=`{keyed[('all', 'log_q_proxy', 'PatentMismatch')]['pvalues'].get('PatentMismatch', math.nan):.3f}`)",
            f"- Non-big Log Q proxy, canonical PatentMismatch: `{keyed[('nonbig', 'log_q_proxy', 'PatentMismatch')]['params'].get('PatentMismatch', math.nan):.4f}` (p=`{keyed[('nonbig', 'log_q_proxy', 'PatentMismatch')]['pvalues'].get('PatentMismatch', math.nan):.3f}`)",
            f"- Non-big Log MktCap/assets, canonical PatentMismatch: `{keyed[('nonbig', 'log_mktcap_assets', 'PatentMismatch')]['params'].get('PatentMismatch', math.nan):.4f}` (p=`{keyed[('nonbig', 'log_mktcap_assets', 'PatentMismatch')]['pvalues'].get('PatentMismatch', math.nan):.3f}`)",
            f"- Full-sample Issue>5% t+1, canonical PatentMismatch: `{keyed[('all', 'equity_issue_lead1', 'PatentMismatch')]['params'].get('PatentMismatch', math.nan):.4f}` (p=`{keyed[('all', 'equity_issue_lead1', 'PatentMismatch')]['pvalues'].get('PatentMismatch', math.nan):.3f}`)",
            "",
            "## Caption Draft",
            "This table refreshes the financing and valuation lane inside the AI-talking sample. Panel A uses the full AI-talking sample, while Panel B restricts to non-big firms using the yearly median market-cap split available in the local panel. Valuation outcomes are log market-capitalization-to-assets, a log Q-style proxy, and the next-year change in that Q proxy. Financing outcomes are next-year share growth and an indicator for substantial share issuance. All specifications absorb firm and year fixed effects, cluster standard errors by firm, and compare the canonical grant-based mismatch construct, an application-side mismatch variant, and a disclosure-only low-credibility variant.",
            "",
        ]
    )


def _dataset_summary(
    args: argparse.Namespace,
    sample_summary: dict[str, object],
    model_rows: list[dict[str, object]],
    figure_df: pd.DataFrame,
) -> dict[str, object]:
    return {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "inputs": {"annual_panel": str(args.annual_panel)},
        "sample_summary": sample_summary,
        "model_rows": model_rows,
        "figure_series": figure_df.to_dict(orient="records"),
    }


def _copy_exports(run_dir: Path, paper_root: Path, run_id: str) -> dict[str, str]:
    exports = {
        "table_csv": paper_root / "tables" / f"{TEST_ID}_{run_id}.csv",
        "table_md": paper_root / "tables" / f"{TEST_ID}_{run_id}.md",
        "table_tex": paper_root / "latex" / f"{TEST_ID}_{run_id}.tex",
        "table_docx": paper_root / "docx" / f"{TEST_ID}_{run_id}.docx",
        "figure_png": paper_root / "figures" / f"{TEST_ID}_{run_id}.png",
        "figure_pdf": paper_root / "figures" / f"{TEST_ID}_{run_id}.pdf",
        "figure_series": paper_root / "tables" / f"{TEST_ID}_{run_id}_figure_series.csv",
        "writer_packet": paper_root / "writer_packets" / f"{TEST_ID}_{run_id}.md",
        "result_notes": paper_root / "snippets" / f"{TEST_ID}_{run_id}_result_notes.md",
        "data_gap_note": paper_root / "snippets" / f"{TEST_ID}_{run_id}_data_gap_note.md",
    }
    for path in exports.values():
        path.parent.mkdir(parents=True, exist_ok=True)
    mapping = {
        run_dir / "table_main.csv": exports["table_csv"],
        run_dir / "table_main.md": exports["table_md"],
        run_dir / "table_main.tex": exports["table_tex"],
        run_dir / "table_main.docx": exports["table_docx"],
        run_dir / "figure_main.png": exports["figure_png"],
        run_dir / "figure_main.pdf": exports["figure_pdf"],
        run_dir / "figure_series.csv": exports["figure_series"],
        run_dir / "writer_packet.md": exports["writer_packet"],
        run_dir / "result_notes.md": exports["result_notes"],
        run_dir / "data_gap_note.md": exports["data_gap_note"],
    }
    for src, dst in mapping.items():
        shutil.copy2(src, dst)
    return {key: str(path) for key, path in exports.items()}


def main() -> None:
    args = _parse_args()
    run_dir = args.test_root / args.run_id
    run_dir.mkdir(parents=True, exist_ok=True)

    panel = pd.read_parquet(args.annual_panel)
    sample, sample_summary = _prepare_sample(panel)

    model_rows: list[dict[str, object]] = []
    for subset_key, _panel_label in SUBSET_SPECS:
        for outcome, _label, controls, _family in OUTCOME_SPECS:
            for variant, _variant_label in VARIANT_SPECS:
                model_rows.append(_fit_model(sample, subset_key, variant, outcome, controls))

    table_df, _keyed = _build_results_table(model_rows)
    table_df.to_csv(run_dir / "table_main.csv", index=False)
    table_md, table_tex = _render_table_outputs(table_df)
    (run_dir / "table_main.md").write_text(table_md, encoding="utf-8")
    (run_dir / "table_main.tex").write_text(table_tex, encoding="utf-8")
    _build_table_docx(table_df, run_dir / "table_main.docx")

    png_path, pdf_path, figure_df = _plot_heatmap(model_rows, run_dir / "figure_main")
    figure_df.to_csv(run_dir / "figure_series.csv", index=False)

    (run_dir / "data_gap_note.md").write_text(_data_gap_note(), encoding="utf-8")
    (run_dir / "result_notes.md").write_text(
        _result_notes(model_rows, sample_summary), encoding="utf-8"
    )
    (run_dir / "writer_packet.md").write_text(
        _writer_packet(args, model_rows, sample_summary), encoding="utf-8"
    )
    dataset_summary = _dataset_summary(args, sample_summary, model_rows, figure_df)
    (run_dir / "dataset_summary.json").write_text(
        json.dumps(dataset_summary, indent=2), encoding="utf-8"
    )

    paper_exports = _copy_exports(run_dir, args.paper_root, args.run_id)
    manifest = {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "module_path": MODULE_PATH,
        "run_dir": str(run_dir),
        "inputs": {"annual_panel": str(args.annual_panel)},
        "outputs": {
            "dataset_summary": str(run_dir / "dataset_summary.json"),
            "table_csv": str(run_dir / "table_main.csv"),
            "table_md": str(run_dir / "table_main.md"),
            "table_tex": str(run_dir / "table_main.tex"),
            "table_docx": str(run_dir / "table_main.docx"),
            "figure_png": str(png_path),
            "figure_pdf": str(pdf_path),
            "figure_series": str(run_dir / "figure_series.csv"),
            "writer_packet": str(run_dir / "writer_packet.md"),
            "result_notes": str(run_dir / "result_notes.md"),
            "data_gap_note": str(run_dir / "data_gap_note.md"),
        },
        "paper_exports": paper_exports,
        "sample_summary": sample_summary,
    }
    (run_dir / "run_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    print(f"[{TEST_ID}] wrote run bundle to {run_dir}")
    print(f"[{TEST_ID}] paper table: {paper_exports['table_docx']}")


if __name__ == "__main__":
    main()
