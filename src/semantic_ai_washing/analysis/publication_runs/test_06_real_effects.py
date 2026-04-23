"""Publication run driver for Test 06: real effects beyond patents."""

from __future__ import annotations

import argparse
import json
import math
import shutil
from datetime import UTC, date, datetime
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
import statsmodels.api as sm

from semantic_ai_washing.analysis.delivery_table_payloads import _add_patent_mismatch
from semantic_ai_washing.analysis.publication_runs.test_03_post_filing_drift import (
    _add_note,
    _add_title,
    _build_panel_table,
    _set_document_defaults,
    _set_landscape,
)

REPO_ROOT = Path(__file__).resolve().parents[4]
DEFAULT_ANNUAL_PANEL = (
    REPO_ROOT
    / "data/processed/panel/canonical/ever_speaker_panel_2016_2025_hybrid_api_a_conf49_v1.parquet"
)
DEFAULT_TEST_ROOT = Path(
    "/Users/soheilkhodadadi/DataWork/semantic-patterns/ai_washing/derived/test_runs/test_06_real_effects"
)
DEFAULT_PAPER_ROOT = REPO_ROOT / "paper/generated"
DEFAULT_RUN_ID = f"{date.today():%Y%m%d}_hybrid_api_a_conf49_main_v1"
TEST_ID = "test_06_real_effects"
MODULE_PATH = "semantic_ai_washing.analysis.publication_runs.test_06_real_effects"
CONTROL_TERMS = ["ln_assets", "leverage", "cash", "roa"]
FOCAL_TERMS = ["PatentMismatch", "A_S", "AI_Focus"]
OUTCOME_SPECS = [
    ("roa", "ROA"),
    ("sales_growth", "Sales growth"),
    ("capx_at", "CAPX/assets"),
    ("rd_intensity", "R&D/assets"),
]
HORIZONS = [(1, "t+1"), (2, "t+2")]


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--annual-panel", type=Path, default=DEFAULT_ANNUAL_PANEL)
    parser.add_argument("--test-root", type=Path, default=DEFAULT_TEST_ROOT)
    parser.add_argument("--paper-root", type=Path, default=DEFAULT_PAPER_ROOT)
    parser.add_argument("--run-id", default=DEFAULT_RUN_ID)
    return parser.parse_args()


def _base_style() -> None:
    plt.style.use("seaborn-v0_8-whitegrid")
    plt.rcParams.update(
        {
            "font.family": "serif",
            "font.serif": ["Times New Roman", "Times", "DejaVu Serif"],
            "axes.titlesize": 13,
            "axes.labelsize": 11,
            "xtick.labelsize": 10,
            "ytick.labelsize": 10,
            "legend.fontsize": 9,
        }
    )


def _ensure_sic2(df: pd.DataFrame) -> pd.DataFrame:
    panel = df.copy()
    if "sic2" not in panel.columns or panel["sic2"].isna().all():
        if "sic" in panel.columns:
            sic_raw = pd.to_numeric(panel["sic"], errors="coerce")
            panel["sic2"] = (sic_raw // 100).astype("Int64")
        else:
            panel["sic2"] = pd.Series(pd.NA, index=panel.index, dtype="Int64")
    return panel


def _load_panel(path: Path) -> pd.DataFrame:
    panel = pd.read_parquet(path).copy()
    panel = _ensure_sic2(panel)
    panel = _add_patent_mismatch(panel)
    panel["cik"] = panel["cik"].astype(str)
    panel["year"] = pd.to_numeric(panel["year"], errors="coerce").astype("Int64")
    panel = panel.sort_values(["cik", "year"]).reset_index(drop=True)
    return panel


def _winsorize(series: pd.Series, lower_q: float = 0.01, upper_q: float = 0.99) -> pd.Series:
    valid = pd.to_numeric(series, errors="coerce").replace([np.inf, -np.inf], np.nan)
    if valid.notna().sum() == 0:
        return valid
    lower = valid.quantile(lower_q)
    upper = valid.quantile(upper_q)
    return valid.clip(lower=lower, upper=upper)


def _prepare_panel(panel: pd.DataFrame) -> tuple[pd.DataFrame, dict[str, object]]:
    use = panel.copy()
    use["PatentMismatch"] = (
        pd.to_numeric(use["PatentMismatch"], errors="coerce").fillna(0).astype(int)
    )
    use["A_S"] = pd.to_numeric(use["A_S"], errors="coerce").replace([np.inf, -np.inf], np.nan)
    use["AI_Focus"] = pd.to_numeric(use["AI_Focus"], errors="coerce").replace(
        [np.inf, -np.inf], np.nan
    )
    for control in CONTROL_TERMS:
        use[control] = pd.to_numeric(use[control], errors="coerce").replace(
            [np.inf, -np.inf], np.nan
        )

    for outcome, _ in OUTCOME_SPECS:
        use[outcome] = pd.to_numeric(use[outcome], errors="coerce").replace(
            [np.inf, -np.inf], np.nan
        )
        for horizon, _ in HORIZONS:
            lead = use.groupby("cik", sort=False)[outcome].shift(-horizon)
            use[f"{outcome}_lead{horizon}"] = _winsorize(lead)

    summary = {
        "firm_year_rows": int(len(use)),
        "unique_firms": int(use["cik"].nunique()),
        "year_min": int(use["year"].min()),
        "year_max": int(use["year"].max()),
        "winsorization": "1st/99th percentile on each future outcome distribution",
    }
    return use, summary


def _two_way_absorb(
    use: pd.DataFrame, columns: list[str], *, max_iter: int = 200, tol: float = 1e-10
) -> pd.DataFrame:
    entity = use["cik"]
    year = use["year"]
    absorbed = pd.DataFrame(index=use.index)
    for col in columns:
        z = pd.to_numeric(use[col], errors="coerce").astype(float).copy()
        for _ in range(max_iter):
            old = z.to_numpy(copy=True)
            z = z - z.groupby(entity).transform("mean")
            z = z - z.groupby(year).transform("mean")
            if len(z):
                delta_arr = np.abs(z.to_numpy() - old)
                delta = 0.0 if np.isnan(delta_arr).all() else float(np.nanmax(delta_arr))
            else:
                delta = 0.0
            if delta < tol:
                break
        absorbed[col] = z
    return absorbed


def _fit_absorbed_ols(
    use: pd.DataFrame, dependent: str, rhs: list[str]
) -> sm.regression.linear_model.RegressionResultsWrapper:
    absorbed = _two_way_absorb(use, [dependent, *rhs])
    y = absorbed[dependent]
    X = absorbed[rhs]
    return sm.OLS(y, X).fit(cov_type="cluster", cov_kwds={"groups": use["cik"].astype(str)})


def _fit_model(panel: pd.DataFrame, outcome: str, horizon: int, label: str) -> dict[str, object]:
    dependent = f"{outcome}_lead{horizon}"
    needed = [dependent, "cik", "year", *FOCAL_TERMS, *CONTROL_TERMS]
    use = panel.dropna(subset=needed).copy()
    finite_mask = np.isfinite(
        use[[dependent, *FOCAL_TERMS, *CONTROL_TERMS]].to_numpy(dtype=float)
    ).all(axis=1)
    use = use.loc[finite_mask].copy()
    use = use.loc[use["cik"].str.len().gt(0)].copy()
    use["year"] = pd.to_numeric(use["year"], errors="coerce").astype(int)
    rhs = [*FOCAL_TERMS, *CONTROL_TERMS]
    result = _fit_absorbed_ols(use, dependent, rhs)
    outcome_mean = float(use[dependent].mean())
    return {
        "outcome": outcome,
        "outcome_label": label,
        "horizon": horizon,
        "horizon_label": f"t+{horizon}",
        "dependent": dependent,
        "nobs": int(result.nobs),
        "adj_r_squared": float(result.rsquared_adj),
        "outcome_mean": outcome_mean,
        "params": result.params.to_dict(),
        "bse": result.bse.to_dict(),
        "pvalues": result.pvalues.to_dict(),
    }


def _stars(p_value: float | None) -> str:
    if p_value is None or not math.isfinite(p_value):
        return ""
    if p_value < 0.01:
        return "***"
    if p_value < 0.05:
        return "**"
    if p_value < 0.10:
        return "*"
    return ""


def _coef_cell(coef: float | None, se: float | None, p_value: float | None) -> str:
    if coef is None or se is None or not math.isfinite(coef) or not math.isfinite(se):
        return ""
    return f"{coef:.4f}{_stars(p_value)} ({se:.4f})"


def _build_results_table(
    model_rows: list[dict[str, object]],
) -> tuple[pd.DataFrame, dict[tuple[int, str], dict[str, object]]]:
    keyed = {(row["horizon"], row["outcome"]): row for row in model_rows}
    rows: list[dict[str, object]] = []
    for horizon, horizon_label in HORIZONS:
        for term in ["PatentMismatch", "A_S", "AI_Focus", "Outcome mean", "N", "Adj. R-squared"]:
            row = {"panel": horizon_label, "row_label": term}
            for outcome, label in OUTCOME_SPECS:
                cell_key = label
                result = keyed[(horizon, outcome)]
                if term in FOCAL_TERMS:
                    row[cell_key] = _coef_cell(
                        result["params"].get(term),
                        result["bse"].get(term),
                        result["pvalues"].get(term),
                    )
                elif term == "Outcome mean":
                    row[cell_key] = f"{result['outcome_mean']:.4f}"
                elif term == "N":
                    row[cell_key] = str(result["nobs"])
                else:
                    row[cell_key] = f"{result['adj_r_squared']:.3f}"
            rows.append(row)
    return pd.DataFrame(rows), keyed


def _markdown_table(headers: list[str], rows: list[list[object]]) -> str:
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join(["---"] * len(headers)) + " |",
    ]
    for row in rows:
        lines.append("| " + " | ".join(str(item) for item in row) + " |")
    return "\n".join(lines)


def _render_table_outputs(table_df: pd.DataFrame) -> tuple[str, str]:
    headers = ["Row"] + [label for _, label in OUTCOME_SPECS]
    md_lines = ["# Table Main", ""]
    latex_lines = [
        "\\begin{table}[!htbp]",
        "\\centering",
        "\\caption{Future real outcomes and disclosure credibility}",
        "\\begin{tabular}{l" + "c" * len(OUTCOME_SPECS) + "}",
        "\\hline",
    ]
    for _, horizon_label in HORIZONS:
        panel_rows = table_df.loc[table_df["panel"].eq(horizon_label)].copy()
        rendered = panel_rows[
            ["row_label", *[label for _, label in OUTCOME_SPECS]]
        ].values.tolist()
        md_lines.extend([f"## Panel {horizon_label}", _markdown_table(headers, rendered), ""])
        latex_lines.append(
            f"\\multicolumn{{{len(headers)}}}{{l}}{{\\textit{{Panel {horizon_label}}}}} \\\\"
        )
        latex_lines.append(" & ".join(headers) + " \\\\")
        for row in rendered:
            latex_lines.append(" & ".join(str(item) for item in row) + " \\\\")
    latex_lines.extend(["\\hline", "\\end{tabular}", "\\end{table}"])
    return "\n".join(md_lines), "\n".join(latex_lines) + "\n"


def _docx_panel_rows(table_df: pd.DataFrame, panel_name: str) -> list[list[tuple[str, bool]]]:
    rows = []
    subset = table_df.loc[table_df["panel"].eq(panel_name)].copy()
    for row in subset[["row_label", *[label for _, label in OUTCOME_SPECS]]].itertuples(
        index=False
    ):
        rows.append([(str(row[0]), True), *[(str(value), False) for value in row[1:]]])
    return rows


def _build_table_docx(table_df: pd.DataFrame, output_path: Path) -> None:
    document = Document()
    _set_document_defaults(document)
    _set_landscape(document)
    _add_title(document, "Table 6. Future Real Outcomes and Disclosure Credibility")
    note = (
        "This table reports future real outcomes as a function of disclosure credibility. Dependent variables are future ROA, sales growth, CAPX/assets, and R&D/assets at horizons t+1 and t+2. The key regressors are PatentMismatch, the A/S ratio, and AI Focus. "
        "All specifications absorb firm and year fixed effects, use firm-clustered standard errors, and apply 1st/99th percentile winsorization to each future outcome distribution."
    )
    _add_note(document, note)
    headers = ["", *[label for _, label in OUTCOME_SPECS]]
    for _, horizon_label in HORIZONS:
        p = document.add_paragraph()
        p.add_run(f"Panel {horizon_label}").bold = True
        _build_panel_table(document, headers, _docx_panel_rows(table_df, horizon_label))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def _plot_patent_mismatch_coefficients(
    model_rows: list[dict[str, object]], output_base: Path
) -> tuple[Path, Path]:
    _base_style()
    fig, ax = plt.subplots(figsize=(8.2, 4.8), constrained_layout=True)
    color_map = {1: "#c46b48", 2: "#6d597a"}
    y_positions = {}
    current = len(OUTCOME_SPECS) - 1
    plotted = []
    for outcome, label in OUTCOME_SPECS:
        for horizon, horizon_label in HORIZONS:
            y_positions[(outcome, horizon)] = current + (0.15 if horizon == 1 else -0.15)
        current -= 1
    for row in model_rows:
        coef = row["params"].get("PatentMismatch", math.nan)
        se = row["bse"].get("PatentMismatch", math.nan)
        if not (math.isfinite(coef) and math.isfinite(se)):
            continue
        y = y_positions[(row["outcome"], row["horizon"])]
        ax.errorbar(
            coef,
            y,
            xerr=1.96 * se,
            fmt="o",
            color=color_map[row["horizon"]],
            capsize=4,
            linewidth=1.6,
            label=f"{row['horizon_label']}",
        )
        plotted.append((row["outcome_label"], y))
    ax.axvline(0, color="#6c757d", linewidth=0.9)
    ax.set_yticks(range(len(OUTCOME_SPECS)))
    ax.set_yticklabels([label for _, label in OUTCOME_SPECS][::-1])
    ax.set_xlabel("PatentMismatch coefficient")
    ax.set_title("PatentMismatch and Future Real Outcomes")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.grid(True, axis="x", color="#d9d9d9", linewidth=0.7)
    ax.grid(False, axis="y")
    handles, labels = ax.get_legend_handles_labels()
    dedup: dict[str, object] = {}
    for handle, label in zip(handles, labels, strict=False):
        dedup.setdefault(label, handle)
    ax.legend(dedup.values(), dedup.keys(), frameon=False, loc="best")

    png_path = output_base.with_suffix(".png")
    pdf_path = output_base.with_suffix(".pdf")
    fig.savefig(png_path, dpi=220, bbox_inches="tight")
    fig.savefig(pdf_path, bbox_inches="tight")
    plt.close(fig)
    return png_path, pdf_path


def _result_notes(
    model_keyed: dict[tuple[int, str], dict[str, object]], panel_summary: dict[str, object]
) -> str:
    t1_roa = model_keyed[(1, "roa")]
    t1_sales = model_keyed[(1, "sales_growth")]
    t2_roa = model_keyed[(2, "roa")]
    return "\n".join(
        [
            "# Result Notes",
            "",
            f"- Future ROA t+1: PatentMismatch `{t1_roa['params'].get('PatentMismatch', math.nan):.4f}` (p=`{t1_roa['pvalues'].get('PatentMismatch', math.nan):.3f}`), A/S `{t1_roa['params'].get('A_S', math.nan):.4f}` (p=`{t1_roa['pvalues'].get('A_S', math.nan):.3f}`).",
            f"- Future sales growth t+1: PatentMismatch `{t1_sales['params'].get('PatentMismatch', math.nan):.4f}` (p=`{t1_sales['pvalues'].get('PatentMismatch', math.nan):.3f}`), A/S `{t1_sales['params'].get('A_S', math.nan):.4f}` (p=`{t1_sales['pvalues'].get('A_S', math.nan):.3f}`).",
            f"- Future ROA t+2: PatentMismatch `{t2_roa['params'].get('PatentMismatch', math.nan):.4f}` (p=`{t2_roa['pvalues'].get('PatentMismatch', math.nan):.3f}`).",
            f"- Annual panel coverage: `{panel_summary['firm_year_rows']}` rows across `{panel_summary['unique_firms']}` firms.",
            "",
        ]
    )


def _writer_packet(
    args: argparse.Namespace, panel_summary: dict[str, object], model_rows: list[dict[str, object]]
) -> str:
    keyed = {(row["horizon"], row["outcome"]): row for row in model_rows}
    return "\n".join(
        [
            "# Writer Packet",
            "",
            "## Metadata",
            f"- Test id: `{TEST_ID}`",
            f"- Run id: `{args.run_id}`",
            f"- Date run: `{date.today().isoformat()}`",
            f"- Script/module path: `{MODULE_PATH}`",
            f"- Input file: `{args.annual_panel}`",
            "",
            "## Design",
            "- Core regressors: `PatentMismatch`, `A_S`, `AI_Focus`",
            "- Controls: `ln_assets`, `leverage`, `cash`, `roa`",
            "- Fixed effects: `firm and year`",
            "- Clustering: `firm (cik)`",
            f"- Winsorization: `{panel_summary['winsorization']}`",
            "",
            "## Main Results",
            f"- ROA t+1 PatentMismatch: `{keyed[(1, 'roa')]['params'].get('PatentMismatch', math.nan):.4f}` (p=`{keyed[(1, 'roa')]['pvalues'].get('PatentMismatch', math.nan):.3f}`)",
            f"- Sales growth t+1 PatentMismatch: `{keyed[(1, 'sales_growth')]['params'].get('PatentMismatch', math.nan):.4f}` (p=`{keyed[(1, 'sales_growth')]['pvalues'].get('PatentMismatch', math.nan):.3f}`)",
            f"- CAPX/assets t+1 PatentMismatch: `{keyed[(1, 'capx_at')]['params'].get('PatentMismatch', math.nan):.4f}` (p=`{keyed[(1, 'capx_at')]['pvalues'].get('PatentMismatch', math.nan):.3f}`)",
            f"- R&D/assets t+1 PatentMismatch: `{keyed[(1, 'rd_intensity')]['params'].get('PatentMismatch', math.nan):.4f}` (p=`{keyed[(1, 'rd_intensity')]['pvalues'].get('PatentMismatch', math.nan):.3f}`)",
            "",
            "## Caption Draft",
            "This table reports future real outcomes as a function of disclosure credibility. The dependent variables are future ROA, sales growth, CAPX/assets, and R&D/assets at horizons t+1 and t+2. The key regressors are PatentMismatch, the A/S ratio, and AI Focus. All specifications absorb firm and year fixed effects, use firm-clustered standard errors, and apply 1st/99th percentile winsorization to each future outcome distribution.",
            "",
        ]
    )


def _dataset_summary(
    args: argparse.Namespace, panel_summary: dict[str, object], model_rows: list[dict[str, object]]
) -> dict[str, object]:
    return {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "annual_panel": str(args.annual_panel),
        "panel_summary": panel_summary,
        "model_rows": model_rows,
    }


def _copy_exports(run_dir: Path, paper_root: Path, run_id: str) -> dict[str, str]:
    exports = {
        "figure_png": paper_root / "figures" / f"{TEST_ID}_{run_id}.png",
        "figure_pdf": paper_root / "figures" / f"{TEST_ID}_{run_id}.pdf",
        "table_csv": paper_root / "tables" / f"{TEST_ID}_{run_id}.csv",
        "table_md": paper_root / "tables" / f"{TEST_ID}_{run_id}.md",
        "table_tex": paper_root / "latex" / f"{TEST_ID}_{run_id}.tex",
        "table_docx": paper_root / "docx" / f"{TEST_ID}_{run_id}.docx",
        "figure_series": paper_root / "tables" / f"{TEST_ID}_{run_id}_figure_series.csv",
        "writer_packet": paper_root / "writer_packets" / f"{TEST_ID}_{run_id}.md",
        "result_notes": paper_root / "snippets" / f"{TEST_ID}_{run_id}_result_notes.md",
    }
    for path in exports.values():
        path.parent.mkdir(parents=True, exist_ok=True)
    mapping = {
        run_dir / "figure_main.png": exports["figure_png"],
        run_dir / "figure_main.pdf": exports["figure_pdf"],
        run_dir / "table_main.csv": exports["table_csv"],
        run_dir / "table_main.md": exports["table_md"],
        run_dir / "table_main.tex": exports["table_tex"],
        run_dir / "table_main.docx": exports["table_docx"],
        run_dir / "figure_series.csv": exports["figure_series"],
        run_dir / "writer_packet.md": exports["writer_packet"],
        run_dir / "result_notes.md": exports["result_notes"],
    }
    for src, dst in mapping.items():
        shutil.copy2(src, dst)
    return {key: str(path) for key, path in exports.items()}


def main() -> None:
    args = _parse_args()
    run_dir = args.test_root / args.run_id
    run_dir.mkdir(parents=True, exist_ok=True)

    panel = _load_panel(args.annual_panel)
    panel, panel_summary = _prepare_panel(panel)

    model_rows = []
    for horizon, _ in HORIZONS:
        for outcome, label in OUTCOME_SPECS:
            model_rows.append(_fit_model(panel, outcome, horizon, label))

    table_df, model_keyed = _build_results_table(model_rows)
    table_df.to_csv(run_dir / "table_main.csv", index=False)
    figure_df = pd.DataFrame(
        [
            {
                "horizon": row["horizon_label"],
                "outcome": row["outcome_label"],
                "coef": row["params"].get("PatentMismatch", math.nan),
                "se": row["bse"].get("PatentMismatch", math.nan),
                "p_value": row["pvalues"].get("PatentMismatch", math.nan),
            }
            for row in model_rows
        ]
    )
    figure_df.to_csv(run_dir / "figure_series.csv", index=False)

    table_md, table_tex = _render_table_outputs(table_df)
    (run_dir / "table_main.md").write_text(table_md, encoding="utf-8")
    (run_dir / "table_main.tex").write_text(table_tex, encoding="utf-8")
    _build_table_docx(table_df, run_dir / "table_main.docx")
    png_path, pdf_path = _plot_patent_mismatch_coefficients(model_rows, run_dir / "figure_main")

    (run_dir / "result_notes.md").write_text(
        _result_notes(model_keyed, panel_summary), encoding="utf-8"
    )
    (run_dir / "writer_packet.md").write_text(
        _writer_packet(args, panel_summary, model_rows), encoding="utf-8"
    )
    summary = _dataset_summary(args, panel_summary, model_rows)
    (run_dir / "dataset_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")

    paper_exports = _copy_exports(run_dir, args.paper_root, args.run_id)
    manifest = {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "module_path": MODULE_PATH,
        "inputs": {"annual_panel": str(args.annual_panel)},
        "run_dir": str(run_dir),
        "outputs": {
            "dataset_summary": str(run_dir / "dataset_summary.json"),
            "table_csv": str(run_dir / "table_main.csv"),
            "table_md": str(run_dir / "table_main.md"),
            "table_tex": str(run_dir / "table_main.tex"),
            "table_docx": str(run_dir / "table_main.docx"),
            "figure_series": str(run_dir / "figure_series.csv"),
            "figure_png": str(png_path),
            "figure_pdf": str(pdf_path),
            "writer_packet": str(run_dir / "writer_packet.md"),
            "result_notes": str(run_dir / "result_notes.md"),
        },
        "paper_exports": paper_exports,
    }
    (run_dir / "run_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    print(f"[{TEST_ID}] wrote run bundle to {run_dir}")
    print(f"[{TEST_ID}] paper table: {paper_exports['table_docx']}")


if __name__ == "__main__":
    main()
