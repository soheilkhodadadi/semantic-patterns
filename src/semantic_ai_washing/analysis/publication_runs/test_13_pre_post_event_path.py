"""Publication run driver for Test 13: pre/post monthly event path around the filing."""

from __future__ import annotations

import argparse
import json
import shutil
from dataclasses import dataclass
from datetime import UTC, date, datetime
from pathlib import Path

from docx import Document
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from scipy import stats

from semantic_ai_washing.analysis.delivery_table_payloads import _add_patent_mismatch
from semantic_ai_washing.analysis.publication_runs.test_03_post_filing_drift import (
    _add_note,
    _add_title,
    _build_panel_table,
    _set_document_defaults,
    _set_landscape,
)

REPO_ROOT = Path(__file__).resolve().parents[4]
DEFAULT_EVENT_PANEL = (
    REPO_ROOT
    / "data/processed/panel/filing_event_estimation_sample_hybrid_api_a_conf49_v1.parquet"
)
DEFAULT_MONTHLY_RETURNS = REPO_ROOT / "data/interim/market/wrds_crsp_msf_full_sample_v1.parquet"
DEFAULT_MONTHLY_MARKET = REPO_ROOT / "data/interim/market/wrds_crsp_msi_full_sample_v1.parquet"
DEFAULT_ANNUAL_PANEL = (
    REPO_ROOT
    / "data/processed/panel/canonical/ever_speaker_panel_2016_2025_hybrid_api_a_conf49_v1.parquet"
)
DEFAULT_TEST_ROOT = Path(
    "/Users/soheilkhodadadi/DataWork/semantic-patterns/ai_washing/derived/v3_1/test_runs/test_13_pre_post_event_path"
)
DEFAULT_PAPER_ROOT = REPO_ROOT / "paper/generated/v3_1"
DEFAULT_RUN_ID = f"{date.today():%Y%m%d}_aiw_v3_1_test_13_pre_post_event_path_main_v1"
TEST_ID = "test_13_pre_post_event_path"
MODULE_PATH = "semantic_ai_washing.analysis.publication_runs.test_13_pre_post_event_path"
FULL_WINDOW = tuple(range(-12, 13))


@dataclass(frozen=True)
class WindowSpec:
    panel: str
    start_month: int
    end_month: int
    label: str


WINDOW_SPECS = [
    WindowSpec("Pre-filing windows", -12, -2, "BHAR[-12,-2]"),
    WindowSpec("Pre-filing windows", -6, -2, "BHAR[-6,-2]"),
    WindowSpec("Pre-filing windows", -3, -1, "BHAR[-3,-1]"),
    WindowSpec("Filing and post-filing windows", 0, 0, "BHAR[0,0]"),
    WindowSpec("Filing and post-filing windows", 1, 3, "BHAR[+1,+3]"),
    WindowSpec("Filing and post-filing windows", 1, 6, "BHAR[+1,+6]"),
    WindowSpec("Filing and post-filing windows", 1, 12, "BHAR[+1,+12]"),
]


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--event-panel", type=Path, default=DEFAULT_EVENT_PANEL)
    parser.add_argument("--monthly-returns", type=Path, default=DEFAULT_MONTHLY_RETURNS)
    parser.add_argument("--monthly-market", type=Path, default=DEFAULT_MONTHLY_MARKET)
    parser.add_argument("--annual-panel", type=Path, default=DEFAULT_ANNUAL_PANEL)
    parser.add_argument("--test-root", type=Path, default=DEFAULT_TEST_ROOT)
    parser.add_argument("--paper-root", type=Path, default=DEFAULT_PAPER_ROOT)
    parser.add_argument("--run-id", default=DEFAULT_RUN_ID)
    return parser.parse_args()


def _normalize_id(series: pd.Series) -> pd.Series:
    return (
        series.astype(str)
        .str.replace(r"\.0$", "", regex=True)
        .str.strip()
        .replace({"nan": "", "None": ""})
    )


def _base_style() -> None:
    plt.style.use("seaborn-v0_8-whitegrid")
    plt.rcParams.update(
        {
            "font.family": "serif",
            "font.serif": ["Times New Roman", "Times", "DejaVu Serif"],
            "axes.titlesize": 13,
            "axes.labelsize": 11,
            "xtick.labelsize": 10,
            "ytick.labelsize": 10,
            "legend.fontsize": 9,
        }
    )


def _ensure_sic2(panel: pd.DataFrame) -> pd.DataFrame:
    out = panel.copy()
    if "sic2" not in out.columns or out["sic2"].isna().all():
        if "sic" in out.columns:
            sic_raw = pd.to_numeric(out["sic"], errors="coerce")
            out["sic2"] = (sic_raw // 100).astype("Int64")
        else:
            out["sic2"] = pd.Series(pd.NA, index=out.index, dtype="Int64")
    return out


def _load_annual_backbone(path: Path) -> pd.DataFrame:
    annual = pd.read_parquet(path).copy()
    annual = _ensure_sic2(annual)
    annual = _add_patent_mismatch(annual)
    keep = ["cik", "year", "PatentMismatch", "A_S", "AI_Focus"]
    annual = annual[keep].copy()
    annual["cik"] = _normalize_id(annual["cik"])
    annual["year"] = pd.to_numeric(annual["year"], errors="coerce").astype("Int64")
    return annual


def _load_ai_filing_backbone(event_panel: Path, annual_panel: Path) -> pd.DataFrame:
    event = pd.read_parquet(
        event_panel,
        columns=[
            "filing_id",
            "cik",
            "gvkey",
            "permno",
            "anchor_trading_date",
            "filing_year",
            "n_ai_total",
        ],
    ).copy()
    event["cik"] = _normalize_id(event["cik"])
    event["gvkey"] = _normalize_id(event["gvkey"])
    event["permno"] = _normalize_id(event["permno"])
    event["filing_year"] = pd.to_numeric(event["filing_year"], errors="coerce").astype("Int64")
    event["n_ai_total"] = pd.to_numeric(event["n_ai_total"], errors="coerce")
    event = event.loc[event["n_ai_total"].fillna(0).gt(0)].copy()
    event["anchor_trading_date"] = pd.to_datetime(event["anchor_trading_date"], errors="coerce")
    event["filing_month_end"] = event["anchor_trading_date"].dt.to_period("M").dt.to_timestamp("M")

    annual = _load_annual_backbone(annual_panel)
    merged = event.merge(
        annual,
        left_on=["cik", "filing_year"],
        right_on=["cik", "year"],
        how="left",
        validate="many_to_one",
    )
    merged["PatentMismatch"] = pd.to_numeric(merged["PatentMismatch"], errors="coerce")
    merged = merged.loc[
        merged["permno"].str.len().gt(0)
        & merged["filing_month_end"].notna()
        & merged["PatentMismatch"].notna()
    ].copy()
    merged["PatentMismatch"] = merged["PatentMismatch"].astype(int)
    return merged[
        [
            "filing_id",
            "cik",
            "gvkey",
            "permno",
            "filing_year",
            "filing_month_end",
            "PatentMismatch",
            "A_S",
            "AI_Focus",
        ]
    ].copy()


def _load_monthly_returns(monthly_returns: Path, monthly_market: Path) -> pd.DataFrame:
    firm = pd.read_parquet(
        monthly_returns, columns=["permno", "date", "ret", "retx", "prc", "shrout", "vol"]
    ).copy()
    market = pd.read_parquet(monthly_market, columns=["date", "vwretd", "ewretd", "sprtrn"]).copy()
    firm["permno"] = _normalize_id(firm["permno"])
    firm["date"] = pd.to_datetime(firm["date"], errors="coerce")
    market["date"] = pd.to_datetime(market["date"], errors="coerce")
    for column in ["ret", "retx", "prc", "shrout", "vol"]:
        firm[column] = pd.to_numeric(firm[column], errors="coerce")
    for column in ["vwretd", "ewretd", "sprtrn"]:
        market[column] = pd.to_numeric(market[column], errors="coerce")
    out = firm.merge(market, on="date", how="left", validate="many_to_one")
    out["month_end"] = out["date"].dt.to_period("M").dt.to_timestamp("M")
    return out


def _months_between(left: pd.Series, right: pd.Series) -> pd.Series:
    return (left.dt.year - right.dt.year) * 12 + (left.dt.month - right.dt.month)


def _build_monthly_event_panel(
    event_panel: Path, annual_panel: Path, monthly_returns: Path, monthly_market: Path
) -> tuple[pd.DataFrame, dict[str, object]]:
    filings = _load_ai_filing_backbone(event_panel, annual_panel)
    monthly = _load_monthly_returns(monthly_returns, monthly_market)
    merged = filings.merge(monthly, on="permno", how="left", validate="many_to_many")
    merged = merged.loc[merged["month_end"].notna()].copy()
    merged["rel_month"] = _months_between(merged["month_end"], merged["filing_month_end"])
    merged = merged.loc[merged["rel_month"].between(min(FULL_WINDOW), max(FULL_WINDOW))].copy()

    needed_months = set(FULL_WINDOW)
    complete = merged.groupby("filing_id").apply(
        lambda frame: needed_months.issubset(set(frame["rel_month"].tolist()))
        and frame["ret"].notna().all()
        and frame["vwretd"].notna().all(),
        include_groups=False,
    )
    keep_ids = complete.loc[complete].index
    balanced = merged.loc[merged["filing_id"].isin(keep_ids)].copy()
    balanced["group_label"] = balanced["PatentMismatch"].map({0: "No mismatch", 1: "Mismatch"})

    summary = {
        "ai_filing_rows": int(filings["filing_id"].nunique()),
        "balanced_filing_rows": int(balanced["filing_id"].nunique()),
        "balanced_unique_firms": int(balanced["gvkey"].replace("", pd.NA).dropna().nunique()),
        "window_min": int(min(FULL_WINDOW)),
        "window_max": int(max(FULL_WINDOW)),
        "balanced_mismatch_share": float(
            balanced[["filing_id", "PatentMismatch"]].drop_duplicates()["PatentMismatch"].mean()
        ),
    }
    return balanced, summary


def _build_cumulative_path_series(monthly_panel: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    ordered = monthly_panel.sort_values(["filing_id", "rel_month"]).copy()
    ordered["cum_firm"] = ordered.groupby("filing_id")["ret"].transform(lambda s: (1.0 + s).cumprod() - 1.0)
    ordered["cum_market"] = ordered.groupby("filing_id")["vwretd"].transform(lambda s: (1.0 + s).cumprod() - 1.0)
    ordered["cum_bhar_vw"] = ordered["cum_firm"] - ordered["cum_market"]

    grouped = (
        ordered.groupby(["PatentMismatch", "group_label", "rel_month"], as_index=False)
        .agg(
            mean_cum_bhar_vw=("cum_bhar_vw", "mean"),
            std_cum_bhar_vw=("cum_bhar_vw", "std"),
            n_filings=("filing_id", "nunique"),
        )
        .sort_values(["PatentMismatch", "rel_month"])
        .reset_index(drop=True)
    )
    grouped["se_cum_bhar_vw"] = grouped["std_cum_bhar_vw"] / np.sqrt(grouped["n_filings"])

    pivot_mean = grouped.pivot(index="rel_month", columns="PatentMismatch", values="mean_cum_bhar_vw")
    pivot_n = grouped.pivot(index="rel_month", columns="PatentMismatch", values="n_filings")
    pivot_var = grouped.pivot(index="rel_month", columns="PatentMismatch", values="std_cum_bhar_vw") ** 2
    diff = pd.DataFrame({"rel_month": pivot_mean.index.to_numpy()})
    diff["diff_mismatch_minus_non"] = (pivot_mean[1] - pivot_mean[0]).to_numpy()
    diff["se_diff"] = np.sqrt((pivot_var[1] / pivot_n[1]) + (pivot_var[0] / pivot_n[0])).to_numpy()
    diff["ci_low"] = diff["diff_mismatch_minus_non"] - 1.96 * diff["se_diff"]
    diff["ci_high"] = diff["diff_mismatch_minus_non"] + 1.96 * diff["se_diff"]

    figure_series = diff.merge(
        grouped.loc[grouped["PatentMismatch"].eq(0), ["rel_month", "mean_cum_bhar_vw", "n_filings"]]
        .rename(columns={"mean_cum_bhar_vw": "mean_non_mismatch", "n_filings": "n_non_mismatch"}),
        on="rel_month",
        how="left",
    ).merge(
        grouped.loc[grouped["PatentMismatch"].eq(1), ["rel_month", "mean_cum_bhar_vw", "n_filings"]]
        .rename(columns={"mean_cum_bhar_vw": "mean_mismatch", "n_filings": "n_mismatch"}),
        on="rel_month",
        how="left",
    )
    return grouped, figure_series


def _window_bhar(frame: pd.DataFrame) -> float:
    return float((1.0 + frame["ret"]).prod() - (1.0 + frame["vwretd"]).prod())


def _build_window_table(monthly_panel: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for spec in WINDOW_SPECS:
        sub = monthly_panel.loc[
            monthly_panel["rel_month"].between(spec.start_month, spec.end_month)
        ].copy()
        per_filing = (
            sub.groupby(["filing_id", "PatentMismatch"], as_index=False)
            .apply(_window_bhar, include_groups=False)
            .rename(columns={None: "window_bhar"})
        )
        mismatch = per_filing.loc[per_filing["PatentMismatch"].eq(1), "window_bhar"].astype(float)
        non = per_filing.loc[per_filing["PatentMismatch"].eq(0), "window_bhar"].astype(float)
        t_stat, p_value = stats.ttest_ind(mismatch, non, equal_var=False, nan_policy="omit")
        rows.append(
            {
                "panel": spec.panel,
                "window_label": spec.label,
                "mean_mismatch": float(mismatch.mean()),
                "mean_non_mismatch": float(non.mean()),
                "diff_mismatch_minus_non": float(mismatch.mean() - non.mean()),
                "p_value": float(p_value),
                "n_mismatch": int(len(mismatch)),
                "n_non_mismatch": int(len(non)),
            }
        )
    return pd.DataFrame(rows)


def _markdown_table(headers: list[str], rows: list[list[object]]) -> str:
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join(["---"] * len(headers)) + " |",
    ]
    for row in rows:
        lines.append("| " + " | ".join(str(item) for item in row) + " |")
    return "\n".join(lines)


def _render_table_outputs(table_df: pd.DataFrame) -> tuple[str, str]:
    headers = ["Window", "Mean mismatch", "Mean non-mismatch", "Diff", "p-value", "N mismatch", "N non-mismatch"]
    md_lines = ["# Table Main", ""]
    latex_lines = [
        "\\begin{table}[!htbp]",
        "\\centering",
        "\\caption{Monthly pre/post event path around the filing by PatentMismatch}",
        "\\begin{tabular}{l" + "c" * (len(headers) - 1) + "}",
        "\\hline",
    ]
    for panel in table_df["panel"].drop_duplicates():
        subset = table_df.loc[table_df["panel"].eq(panel)].copy()
        rendered = []
        for row in subset.itertuples(index=False):
            rendered.append(
                [
                    row.window_label,
                    f"{100 * row.mean_mismatch:.3f}",
                    f"{100 * row.mean_non_mismatch:.3f}",
                    f"{100 * row.diff_mismatch_minus_non:.3f}",
                    f"{row.p_value:.3f}",
                    row.n_mismatch,
                    row.n_non_mismatch,
                ]
            )
        md_lines.extend([f"## Panel {panel}", _markdown_table(headers, rendered), ""])
        latex_lines.append(f"\\multicolumn{{{len(headers)}}}{{l}}{{\\textit{{Panel {panel}}}}} \\\\")
        latex_lines.append(" & ".join(headers) + " \\\\")
        for row in rendered:
            latex_lines.append(" & ".join(str(item) for item in row) + " \\\\")
    latex_lines.extend(["\\hline", "\\end{tabular}", "\\end{table}"])
    return "\n".join(md_lines), "\n".join(latex_lines) + "\n"


def _docx_panel_rows(table_df: pd.DataFrame, panel_name: str) -> list[list[tuple[str, bool]]]:
    rows = []
    subset = table_df.loc[table_df["panel"].eq(panel_name)].copy()
    for row in subset.itertuples(index=False):
        rows.append(
            [
                (str(row.window_label), True),
                (f"{100 * row.mean_mismatch:.3f}", False),
                (f"{100 * row.mean_non_mismatch:.3f}", False),
                (f"{100 * row.diff_mismatch_minus_non:.3f}", False),
                (f"{row.p_value:.3f}", False),
                (str(row.n_mismatch), False),
                (str(row.n_non_mismatch), False),
            ]
        )
    return rows


def _build_table_docx(table_df: pd.DataFrame, output_path: Path) -> None:
    document = Document()
    _set_document_defaults(document)
    _set_landscape(document)
    _add_title(document, "Table 13. Monthly Pre/Post Event Path Around the Filing")
    note = (
        "This table reports monthly buy-and-hold abnormal return windows around the AI-related annual filing. The sample is restricted to the balanced set of filings with complete monthly return coverage from month -12 to month +12 around the filing month. Means are reported separately for mismatch and non-mismatch filings, and p-values are from unequal-variance t-tests."
    )
    _add_note(document, note)
    headers = ["", "Mismatch", "No mismatch", "Diff", "p-value", "N mismatch", "N non-mismatch"]
    for panel in table_df["panel"].drop_duplicates():
        p = document.add_paragraph()
        p.add_run(f"Panel {panel}").bold = True
        _build_panel_table(document, headers, _docx_panel_rows(table_df, panel))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def _plot_event_path(
    grouped: pd.DataFrame, figure_series: pd.DataFrame, output_base: Path
) -> tuple[Path, Path]:
    _base_style()
    fig, axes = plt.subplots(2, 1, figsize=(8.5, 7.2), sharex=True, constrained_layout=True)
    color_map = {0: "#1d3557", 1: "#c46b48"}

    ax = axes[0]
    for mismatch_value in [0, 1]:
        sub = grouped.loc[grouped["PatentMismatch"].eq(mismatch_value)].copy()
        ax.plot(
            sub["rel_month"],
            100 * sub["mean_cum_bhar_vw"],
            linewidth=2.1,
            color=color_map[mismatch_value],
            label=sub["group_label"].iloc[0],
        )
    ax.axvline(0, color="#6c757d", linewidth=0.9, linestyle="--")
    ax.axhline(0, color="#6c757d", linewidth=0.8)
    ax.set_ylabel("Mean cumulative BHAR (pct)")
    ax.set_title("Monthly Event-Time Path Around the AI-Related Filing")
    ax.legend(frameon=False, loc="best")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)

    ax = axes[1]
    ax.plot(
        figure_series["rel_month"],
        100 * figure_series["diff_mismatch_minus_non"],
        color="#3a3a3a",
        linewidth=2.0,
        label="Mismatch minus no mismatch",
    )
    ax.fill_between(
        figure_series["rel_month"],
        100 * figure_series["ci_low"],
        100 * figure_series["ci_high"],
        color="#b8c1cc",
        alpha=0.35,
        linewidth=0,
    )
    ax.axvline(0, color="#6c757d", linewidth=0.9, linestyle="--")
    ax.axhline(0, color="#6c757d", linewidth=0.8)
    ax.set_xlabel("Months relative to filing month")
    ax.set_ylabel("Diff in cumulative BHAR (pct)")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.legend(frameon=False, loc="best")

    png_path = output_base.with_suffix(".png")
    pdf_path = output_base.with_suffix(".pdf")
    fig.savefig(png_path, dpi=220, bbox_inches="tight")
    fig.savefig(pdf_path, bbox_inches="tight")
    plt.close(fig)
    return png_path, pdf_path


def _result_notes(
    figure_series: pd.DataFrame, window_table: pd.DataFrame, sample_summary: dict[str, object]
) -> str:
    diff_lookup = figure_series.set_index("rel_month")
    pre_window = window_table.loc[window_table["window_label"].eq("BHAR[-12,-2]")].iloc[0]
    early_post = window_table.loc[window_table["window_label"].eq("BHAR[+1,+3]")].iloc[0]
    post_window = window_table.loc[window_table["window_label"].eq("BHAR[+1,+12]")].iloc[0]
    return "\n".join(
        [
            f"# Result Notes: {TEST_ID}",
            "",
            f"- Balanced event-time sample: `{sample_summary['balanced_filing_rows']}` filings across `{sample_summary['balanced_unique_firms']}` firms.",
            f"- Diff in cumulative BHAR at month `-1`: `{100 * diff_lookup.loc[-1, 'diff_mismatch_minus_non']:.3f}` pct.",
            f"- Diff in cumulative BHAR at month `0`: `{100 * diff_lookup.loc[0, 'diff_mismatch_minus_non']:.3f}` pct.",
            f"- Diff in cumulative BHAR at month `+3`: `{100 * diff_lookup.loc[3, 'diff_mismatch_minus_non']:.3f}` pct.",
            f"- Diff in cumulative BHAR at month `+12`: `{100 * diff_lookup.loc[12, 'diff_mismatch_minus_non']:.3f}` pct.",
            f"- Pre-filing BHAR[-12,-2] diff: `{100 * pre_window['diff_mismatch_minus_non']:.3f}` pct (p=`{pre_window['p_value']:.3f}`).",
            f"- Early post-filing BHAR[+1,+3] diff: `{100 * early_post['diff_mismatch_minus_non']:.3f}` pct (p=`{early_post['p_value']:.3f}`).",
            f"- Post-filing BHAR[+1,+12] diff: `{100 * post_window['diff_mismatch_minus_non']:.3f}` pct (p=`{post_window['p_value']:.3f}`).",
            "- Interpretation discipline: if the pre-filing window is flat while the early post-filing window turns negative, that is more consistent with a short-horizon filing-related differentiation than with a deep pre-existing trend. If the gap is already open before month 0, the delayed-correction language should be dropped.",
            "",
        ]
    )


def _writer_packet(
    args: argparse.Namespace,
    figure_series: pd.DataFrame,
    window_table: pd.DataFrame,
    sample_summary: dict[str, object],
) -> str:
    diff_lookup = figure_series.set_index("rel_month")
    early_post = window_table.loc[window_table["window_label"].eq("BHAR[+1,+3]")].iloc[0]
    return "\n".join(
        [
            f"# Writer Packet: {TEST_ID}",
            "",
            "## Purpose",
            "- This run tests whether the mismatch return gap opens after the filing or is already visible in the year before the filing month.",
            "- Because the daily event-return file begins only at trading day -2, this diagnostic is built from monthly CRSP event time instead of the daily file.",
            "",
            "## Sample Definition",
            f"- Event panel: `{args.event_panel}`",
            f"- Monthly firm returns: `{args.monthly_returns}`",
            f"- Monthly market index: `{args.monthly_market}`",
            "- Unit of observation: `AI-talking annual filing event × event month`",
            "- Balanced window: `-12` to `+12` months around the filing month",
            "",
            "## Sample Counts",
            f"- AI-filing events before balancing: `{sample_summary['ai_filing_rows']}`",
            f"- Balanced filing events: `{sample_summary['balanced_filing_rows']}`",
            f"- Balanced unique firms: `{sample_summary['balanced_unique_firms']}`",
            "",
            "## Key Diagnostics",
            f"- Cumulative BHAR diff at month -1: `{100 * diff_lookup.loc[-1, 'diff_mismatch_minus_non']:.3f}` pct",
            f"- Cumulative BHAR diff at month +3: `{100 * diff_lookup.loc[3, 'diff_mismatch_minus_non']:.3f}` pct",
            f"- Cumulative BHAR diff at month +12: `{100 * diff_lookup.loc[12, 'diff_mismatch_minus_non']:.3f}` pct",
            f"- Pre window BHAR[-12,-2] diff p-value: `{window_table.loc[window_table['window_label'].eq('BHAR[-12,-2]'), 'p_value'].iloc[0]:.3f}`",
            f"- Early post window BHAR[+1,+3] diff p-value: `{early_post['p_value']:.3f}`",
            f"- Post window BHAR[+1,+12] diff p-value: `{window_table.loc[window_table['window_label'].eq('BHAR[+1,+12]'), 'p_value'].iloc[0]:.3f}`",
            "",
            "## Caption Draft",
            "This figure and table trace the event-time return path around the AI-related annual filing using monthly CRSP returns. The sample is restricted to filings with complete return coverage from month -12 to month +12 around the filing month. The key diagnostic is whether the mismatch-minus-non-mismatch gap is already open in the pre-filing window or instead appears mainly in the first few months after the filing month.",
            "",
        ]
    )


def _dataset_summary(
    sample_summary: dict[str, object],
    window_table: pd.DataFrame,
    figure_series: pd.DataFrame,
    *,
    args: argparse.Namespace,
) -> dict[str, object]:
    return {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "inputs": {
            "event_panel": str(args.event_panel),
            "monthly_returns": str(args.monthly_returns),
            "monthly_market": str(args.monthly_market),
            "annual_panel": str(args.annual_panel),
        },
        "sample_summary": sample_summary,
        "window_table": window_table.to_dict(orient="records"),
        "event_path_snapshot": figure_series.to_dict(orient="records"),
    }


def _copy_exports(run_dir: Path, paper_root: Path, run_id: str) -> dict[str, str]:
    exports = {
        "table_csv": paper_root / "tables" / f"{TEST_ID}_{run_id}.csv",
        "table_md": paper_root / "tables" / f"{TEST_ID}_{run_id}.md",
        "table_tex": paper_root / "latex" / f"{TEST_ID}_{run_id}.tex",
        "table_docx": paper_root / "docx" / f"{TEST_ID}_{run_id}.docx",
        "figure_png": paper_root / "figures" / f"{TEST_ID}_{run_id}.png",
        "figure_pdf": paper_root / "figures" / f"{TEST_ID}_{run_id}.pdf",
        "figure_series": paper_root / "tables" / f"{TEST_ID}_{run_id}_figure_series.csv",
        "writer_packet": paper_root / "writer_packets" / f"{TEST_ID}_{run_id}.md",
        "result_notes": paper_root / "snippets" / f"{TEST_ID}_{run_id}_result_notes.md",
    }
    for path in exports.values():
        path.parent.mkdir(parents=True, exist_ok=True)
    mapping = {
        run_dir / "table_main.csv": exports["table_csv"],
        run_dir / "table_main.md": exports["table_md"],
        run_dir / "table_main.tex": exports["table_tex"],
        run_dir / "table_main.docx": exports["table_docx"],
        run_dir / "main_figure.png": exports["figure_png"],
        run_dir / "main_figure.pdf": exports["figure_pdf"],
        run_dir / "figure_series.csv": exports["figure_series"],
        run_dir / "writer_packet.md": exports["writer_packet"],
        run_dir / "result_notes.md": exports["result_notes"],
    }
    for src, dst in mapping.items():
        shutil.copy2(src, dst)
    return {key: str(path) for key, path in exports.items()}


def main() -> None:
    args = _parse_args()
    run_dir = args.test_root / args.run_id
    run_dir.mkdir(parents=True, exist_ok=True)

    monthly_panel, sample_summary = _build_monthly_event_panel(
        args.event_panel,
        args.annual_panel,
        args.monthly_returns,
        args.monthly_market,
    )
    grouped, figure_series = _build_cumulative_path_series(monthly_panel)
    window_table = _build_window_table(monthly_panel)

    window_table.to_csv(run_dir / "table_main.csv", index=False)
    table_md, table_tex = _render_table_outputs(window_table)
    (run_dir / "table_main.md").write_text(table_md, encoding="utf-8")
    (run_dir / "table_main.tex").write_text(table_tex, encoding="utf-8")
    _build_table_docx(window_table, run_dir / "table_main.docx")
    figure_series.to_csv(run_dir / "figure_series.csv", index=False)
    fig_png, fig_pdf = _plot_event_path(grouped, figure_series, run_dir / "main_figure")
    (run_dir / "result_notes.md").write_text(
        _result_notes(figure_series, window_table, sample_summary), encoding="utf-8"
    )
    (run_dir / "writer_packet.md").write_text(
        _writer_packet(args, figure_series, window_table, sample_summary), encoding="utf-8"
    )
    dataset_summary = _dataset_summary(sample_summary, window_table, figure_series, args=args)
    (run_dir / "dataset_summary.json").write_text(json.dumps(dataset_summary, indent=2), encoding="utf-8")

    paper_exports = _copy_exports(run_dir, args.paper_root, args.run_id)
    manifest = {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "module_path": MODULE_PATH,
        "run_dir": str(run_dir),
        "inputs": {
            "event_panel": str(args.event_panel),
            "monthly_returns": str(args.monthly_returns),
            "monthly_market": str(args.monthly_market),
            "annual_panel": str(args.annual_panel),
        },
        "outputs": {
            "dataset_summary": str(run_dir / "dataset_summary.json"),
            "table_csv": str(run_dir / "table_main.csv"),
            "table_md": str(run_dir / "table_main.md"),
            "table_tex": str(run_dir / "table_main.tex"),
            "table_docx": str(run_dir / "table_main.docx"),
            "figure_png": str(fig_png),
            "figure_pdf": str(fig_pdf),
            "figure_series": str(run_dir / "figure_series.csv"),
            "writer_packet": str(run_dir / "writer_packet.md"),
            "result_notes": str(run_dir / "result_notes.md"),
        },
        "paper_exports": paper_exports,
    }
    (run_dir / "run_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    print(f"[{TEST_ID}] wrote run bundle to {run_dir}")
    print(f"[{TEST_ID}] paper table: {paper_exports['table_docx']}")
    print(f"[{TEST_ID}] paper figure: {paper_exports['figure_png']}")


if __name__ == "__main__":
    main()
