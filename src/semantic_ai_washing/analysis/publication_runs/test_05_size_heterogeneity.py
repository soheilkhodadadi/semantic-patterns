"""Publication run driver for Test 05: size heterogeneity."""

from __future__ import annotations

import argparse
import json
import math
import shutil
from datetime import UTC, date, datetime
from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from scipy.stats import norm
import statsmodels.formula.api as smf

from semantic_ai_washing.analysis.publication_runs.test_03_post_filing_drift import (
    _add_note,
    _add_title,
    _build_analysis_sample,
    _build_panel_table,
    _set_document_defaults,
    _set_landscape,
)
from semantic_ai_washing.analysis.publication_runs.test_03_post_filing_portfolio_alpha import (
    _base_style,
    _load_monthly_returns,
)

REPO_ROOT = Path(__file__).resolve().parents[4]
DEFAULT_EVENT_PANEL = (
    REPO_ROOT
    / "data/processed/panel/filing_event_estimation_sample_hybrid_api_a_conf49_v1.parquet"
)
DEFAULT_ANNUAL_PANEL = (
    REPO_ROOT
    / "data/processed/panel/canonical/ever_speaker_panel_2016_2025_hybrid_api_a_conf49_v1.parquet"
)
DEFAULT_MONTHLY_RETURNS = REPO_ROOT / "data/interim/market/wrds_crsp_msf_full_sample_v1.parquet"
DEFAULT_TEST_ROOT = Path(
    "/Users/soheilkhodadadi/DataWork/semantic-patterns/ai_washing/derived/test_runs/test_05_size_heterogeneity"
)
DEFAULT_PAPER_ROOT = REPO_ROOT / "paper/generated"
DEFAULT_RUN_ID = f"{date.today():%Y%m%d}_hybrid_api_a_conf49_main_v1"
TEST_ID = "test_05_size_heterogeneity"
MODULE_PATH = "semantic_ai_washing.analysis.publication_runs.test_05_size_heterogeneity"
CONTROL_TERMS = ["ln_assets", "leverage", "cash", "roa"]
OUTCOME_SPECS = [
    {
        "dependent": "car_m1_p1",
        "title": "CAR[-1,+1]",
        "short_label": "CAR[-1,+1]",
    },
    {
        "dependent": "bhar_3m",
        "title": "BHAR[+2,+63]",
        "short_label": "BHAR[+2,+63]",
    },
]


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--event-panel", type=Path, default=DEFAULT_EVENT_PANEL)
    parser.add_argument("--annual-panel", type=Path, default=DEFAULT_ANNUAL_PANEL)
    parser.add_argument("--monthly-returns", type=Path, default=DEFAULT_MONTHLY_RETURNS)
    parser.add_argument("--test-root", type=Path, default=DEFAULT_TEST_ROOT)
    parser.add_argument("--paper-root", type=Path, default=DEFAULT_PAPER_ROOT)
    parser.add_argument("--run-id", default=DEFAULT_RUN_ID)
    return parser.parse_args()


def _prepare_sample(
    event_panel: Path, annual_panel: Path, monthly_returns: Path
) -> tuple[pd.DataFrame, dict[str, object]]:
    sample = _build_analysis_sample(event_panel, annual_panel)
    sample = sample.loc[sample["is_ai_filing"]].copy()
    sample["permno"] = sample["permno"].astype(str)
    sample["filing_date"] = pd.to_datetime(sample["filing_date"])
    sample["filing_month"] = sample["filing_date"].dt.to_period("M")

    monthly = _load_monthly_returns(monthly_returns)[["permno", "month", "lag_mcap"]].copy()
    sample = sample.merge(
        monthly,
        left_on=["permno", "filing_month"],
        right_on=["permno", "month"],
        how="left",
    )
    sample = sample.drop(columns=["month"], errors="ignore")
    sample["lag_mcap"] = pd.to_numeric(sample["lag_mcap"], errors="coerce")
    sample["month_size_break"] = sample.groupby("filing_month")["lag_mcap"].transform("median")
    sample["Small"] = pd.NA
    has_size = sample["lag_mcap"].notna() & sample["month_size_break"].notna()
    sample.loc[has_size, "Small"] = (
        sample.loc[has_size, "lag_mcap"].le(sample.loc[has_size, "month_size_break"]).astype(int)
    )

    summary = {
        "filing_count_total": int(len(sample)),
        "filing_count_with_size": int(sample["Small"].notna().sum()),
        "small_count": int(sample["Small"].fillna(-1).eq(1).sum()),
        "big_count": int(sample["Small"].fillna(-1).eq(0).sum()),
        "permno_with_size": int(sample.loc[sample["Small"].notna(), "permno"].nunique()),
        "filing_year_min": int(sample["filing_year"].min()),
        "filing_year_max": int(sample["filing_year"].max()),
        "size_break_method": "monthly active-sample median lagged market cap from CRSP MSF",
    }
    return sample, summary


def _model_sample(
    sample: pd.DataFrame, dependent: str, *, require_small: bool = True
) -> pd.DataFrame:
    needed = [
        dependent,
        "PatentMismatch",
        "A_S",
        "AI_Focus",
        "sic2",
        "filing_year",
        "gvkey",
        *CONTROL_TERMS,
    ]
    if require_small:
        needed.append("Small")
    use = sample.copy()
    use = use.dropna(subset=needed)
    use = use.loc[use["gvkey"].astype(str).str.len().gt(0)].copy()
    for column in ["PatentMismatch", "A_S", "AI_Focus", "Small"]:
        if column in use.columns:
            use[column] = pd.to_numeric(use[column], errors="coerce").astype(float)
    use["sic2"] = pd.to_numeric(use["sic2"], errors="coerce").astype(int)
    use["filing_year"] = pd.to_numeric(use["filing_year"], errors="coerce").astype(int)
    for control in CONTROL_TERMS:
        use[control] = pd.to_numeric(use[control], errors="coerce").astype(float)
    return use


def _fit_subsample_regression(
    sample: pd.DataFrame, dependent: str, *, small_flag: int
) -> dict[str, object]:
    use = _model_sample(sample, dependent).loc[lambda df: df["Small"].eq(float(small_flag))].copy()
    formula = (
        f"{dependent} ~ PatentMismatch + A_S + AI_Focus + "
        + " + ".join(CONTROL_TERMS)
        + " + C(sic2) + C(filing_year)"
    )
    result = smf.ols(formula=formula, data=use).fit(
        cov_type="cluster", cov_kwds={"groups": use["gvkey"].astype(str)}
    )
    return {
        "formula": formula,
        "nobs": int(result.nobs),
        "adj_r_squared": float(result.rsquared_adj),
        "params": result.params.to_dict(),
        "bse": result.bse.to_dict(),
        "pvalues": result.pvalues.to_dict(),
    }


def _fit_interaction_regression(sample: pd.DataFrame, dependent: str) -> dict[str, object]:
    use = _model_sample(sample, dependent).copy()
    formula = (
        f"{dependent} ~ PatentMismatch + Small + PatentMismatch:Small + A_S + AI_Focus + "
        + " + ".join(CONTROL_TERMS)
        + " + C(sic2) + C(filing_year)"
    )
    result = smf.ols(formula=formula, data=use).fit(
        cov_type="cluster", cov_kwds={"groups": use["gvkey"].astype(str)}
    )

    big_term = {
        "coef": float(result.params.get("PatentMismatch", math.nan)),
        "se": float(result.bse.get("PatentMismatch", math.nan)),
        "p_value": float(result.pvalues.get("PatentMismatch", math.nan)),
    }
    small_test = result.t_test("PatentMismatch + PatentMismatch:Small = 0")
    small_coef = float(small_test.effect[0])
    small_se = float(
        small_test.sd[0][0] if hasattr(small_test.sd, "__getitem__") else small_test.sd
    )
    small_p = (
        float(small_test.pvalue)
        if hasattr(small_test, "pvalue")
        else (
            2 * (1 - norm.cdf(abs(small_coef / small_se)))
            if small_se and math.isfinite(small_se)
            else math.nan
        )
    )

    return {
        "formula": formula,
        "nobs": int(result.nobs),
        "adj_r_squared": float(result.rsquared_adj),
        "params": result.params.to_dict(),
        "bse": result.bse.to_dict(),
        "pvalues": result.pvalues.to_dict(),
        "big_mismatch_effect": big_term,
        "small_mismatch_effect": {
            "coef": small_coef,
            "se": small_se,
            "p_value": float(small_p),
        },
    }


def _stars(p_value: float | None) -> str:
    if p_value is None or not math.isfinite(p_value):
        return ""
    if p_value < 0.01:
        return "***"
    if p_value < 0.05:
        return "**"
    if p_value < 0.10:
        return "*"
    return ""


def _coef_cell(coef: float | None, se: float | None, p_value: float | None) -> str:
    if coef is None or se is None or not math.isfinite(coef) or not math.isfinite(se):
        return ""
    return f"{coef:.4f}{_stars(p_value)} ({se:.4f})"


def _table_bundle(
    outcome_results: list[dict[str, object]],
) -> tuple[list[dict[str, object]], pd.DataFrame]:
    rows: list[dict[str, object]] = []
    figure_rows: list[dict[str, object]] = []
    for outcome in outcome_results:
        small = outcome["small"]
        big = outcome["big"]
        pooled = outcome["pooled"]
        rows.extend(
            [
                {
                    "panel": outcome["title"],
                    "row_label": "PatentMismatch",
                    "small_subsample": _coef_cell(
                        small["params"].get("PatentMismatch"),
                        small["bse"].get("PatentMismatch"),
                        small["pvalues"].get("PatentMismatch"),
                    ),
                    "big_subsample": _coef_cell(
                        big["params"].get("PatentMismatch"),
                        big["bse"].get("PatentMismatch"),
                        big["pvalues"].get("PatentMismatch"),
                    ),
                    "pooled_interaction": _coef_cell(
                        pooled["params"].get("PatentMismatch"),
                        pooled["bse"].get("PatentMismatch"),
                        pooled["pvalues"].get("PatentMismatch"),
                    ),
                },
                {
                    "panel": outcome["title"],
                    "row_label": "PatentMismatch × Small",
                    "small_subsample": "",
                    "big_subsample": "",
                    "pooled_interaction": _coef_cell(
                        pooled["params"].get("PatentMismatch:Small"),
                        pooled["bse"].get("PatentMismatch:Small"),
                        pooled["pvalues"].get("PatentMismatch:Small"),
                    ),
                },
                {
                    "panel": outcome["title"],
                    "row_label": "Implied mismatch effect, small firms",
                    "small_subsample": "",
                    "big_subsample": "",
                    "pooled_interaction": _coef_cell(
                        pooled["small_mismatch_effect"]["coef"],
                        pooled["small_mismatch_effect"]["se"],
                        pooled["small_mismatch_effect"]["p_value"],
                    ),
                },
                {
                    "panel": outcome["title"],
                    "row_label": "Small",
                    "small_subsample": "",
                    "big_subsample": "",
                    "pooled_interaction": _coef_cell(
                        pooled["params"].get("Small"),
                        pooled["bse"].get("Small"),
                        pooled["pvalues"].get("Small"),
                    ),
                },
                {
                    "panel": outcome["title"],
                    "row_label": "A/S ratio",
                    "small_subsample": _coef_cell(
                        small["params"].get("A_S"),
                        small["bse"].get("A_S"),
                        small["pvalues"].get("A_S"),
                    ),
                    "big_subsample": _coef_cell(
                        big["params"].get("A_S"),
                        big["bse"].get("A_S"),
                        big["pvalues"].get("A_S"),
                    ),
                    "pooled_interaction": _coef_cell(
                        pooled["params"].get("A_S"),
                        pooled["bse"].get("A_S"),
                        pooled["pvalues"].get("A_S"),
                    ),
                },
                {
                    "panel": outcome["title"],
                    "row_label": "AI Focus",
                    "small_subsample": _coef_cell(
                        small["params"].get("AI_Focus"),
                        small["bse"].get("AI_Focus"),
                        small["pvalues"].get("AI_Focus"),
                    ),
                    "big_subsample": _coef_cell(
                        big["params"].get("AI_Focus"),
                        big["bse"].get("AI_Focus"),
                        big["pvalues"].get("AI_Focus"),
                    ),
                    "pooled_interaction": _coef_cell(
                        pooled["params"].get("AI_Focus"),
                        pooled["bse"].get("AI_Focus"),
                        pooled["pvalues"].get("AI_Focus"),
                    ),
                },
                {
                    "panel": outcome["title"],
                    "row_label": "Controls",
                    "small_subsample": "Y",
                    "big_subsample": "Y",
                    "pooled_interaction": "Y",
                },
                {
                    "panel": outcome["title"],
                    "row_label": "Industry FE",
                    "small_subsample": "Y",
                    "big_subsample": "Y",
                    "pooled_interaction": "Y",
                },
                {
                    "panel": outcome["title"],
                    "row_label": "Filing-year FE",
                    "small_subsample": "Y",
                    "big_subsample": "Y",
                    "pooled_interaction": "Y",
                },
                {
                    "panel": outcome["title"],
                    "row_label": "N",
                    "small_subsample": str(small["nobs"]),
                    "big_subsample": str(big["nobs"]),
                    "pooled_interaction": str(pooled["nobs"]),
                },
                {
                    "panel": outcome["title"],
                    "row_label": "Adj. R-squared",
                    "small_subsample": f"{small['adj_r_squared']:.3f}",
                    "big_subsample": f"{big['adj_r_squared']:.3f}",
                    "pooled_interaction": f"{pooled['adj_r_squared']:.3f}",
                },
            ]
        )
        for group_label, effect in [
            ("Big firms", pooled["big_mismatch_effect"]),
            ("Small firms", pooled["small_mismatch_effect"]),
        ]:
            figure_rows.append(
                {
                    "outcome": outcome["short_label"],
                    "group": group_label,
                    "coef": effect["coef"],
                    "se": effect["se"],
                    "ci_low": effect["coef"] - 1.96 * effect["se"],
                    "ci_high": effect["coef"] + 1.96 * effect["se"],
                    "p_value": effect["p_value"],
                }
            )
    return rows, pd.DataFrame(figure_rows)


def _markdown_table(headers: list[str], rows: list[list[object]]) -> str:
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join(["---"] * len(headers)) + " |",
    ]
    for row in rows:
        lines.append("| " + " | ".join(str(item) for item in row) + " |")
    return "\n".join(lines)


def _render_table_outputs(table_rows: list[dict[str, object]]) -> tuple[str, str]:
    headers = ["Row", "Small subsample", "Big subsample", "Pooled interaction"]
    panels: list[tuple[str, list[list[object]]]] = []
    for panel_name in [spec["title"] for spec in OUTCOME_SPECS]:
        panel_rows = []
        for row in table_rows:
            if row["panel"] != panel_name:
                continue
            panel_rows.append(
                [
                    row["row_label"],
                    row["small_subsample"],
                    row["big_subsample"],
                    row["pooled_interaction"],
                ]
            )
        panels.append((panel_name, panel_rows))

    md_lines = ["# Table Main", ""]
    latex_lines = [
        "\\begin{table}[!htbp]",
        "\\centering",
        "\\caption{Size heterogeneity in filing-date and post-filing mismatch effects}",
        "\\begin{tabular}{lccc}",
        "\\hline",
        " & Small subsample & Big subsample & Pooled interaction \\\\",
        "\\hline",
    ]
    for panel_name, panel_rows in panels:
        md_lines.extend([f"## {panel_name}", _markdown_table(headers, panel_rows), ""])
        latex_lines.append(f"\\multicolumn{{4}}{{l}}{{\\textit{{{panel_name}}}}} \\\\")
        for row in panel_rows:
            latex_lines.append(
                row[0] + " & " + " & ".join(str(item) for item in row[1:]) + " \\\\"
            )
    latex_lines.extend(["\\hline", "\\end{tabular}", "\\end{table}"])
    return "\n".join(md_lines), "\n".join(latex_lines) + "\n"


def _panel_rows(
    table_rows: list[dict[str, object]], panel_name: str
) -> list[list[tuple[str, bool]]]:
    rows = []
    for row in table_rows:
        if row["panel"] != panel_name:
            continue
        rows.append(
            [
                (row["row_label"], True),
                (row["small_subsample"], False),
                (row["big_subsample"], False),
                (row["pooled_interaction"], False),
            ]
        )
    return rows


def _build_table_docx(table_rows: list[dict[str, object]], output_path: Path) -> None:
    document = Document()
    _set_document_defaults(document)
    _set_landscape(document)
    _add_title(document, "Table 5. Small-Firm Heterogeneity in AI-Washing Pricing Effects")
    note = (
        "This table reports heterogeneity by firm size for the filing-date and post-filing mismatch effects. Small firms are defined using the monthly active-sample median lagged market capitalization from CRSP MSF because NYSE breakpoints are not available in the local extract. "
        "For each outcome, the first two columns re-estimate the baseline regression separately in the small-firm and big-firm subsamples. The third column estimates a pooled interaction regression with PatentMismatch, Small, and PatentMismatch × Small. Controls include lagged size, leverage, cash/assets, and ROA, with industry fixed effects, filing-year fixed effects, and firm-clustered standard errors."
    )
    _add_note(document, note)
    headers = ["", "Small subsample", "Big subsample", "Pooled interaction"]
    for panel_name in [spec["title"] for spec in OUTCOME_SPECS]:
        p = document.add_paragraph()
        p.add_run(panel_name).bold = True
        _build_panel_table(document, headers, _panel_rows(table_rows, panel_name))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def _plot_heterogeneity_figure(figure_df: pd.DataFrame, output_base: Path) -> tuple[Path, Path]:
    _base_style()
    fig, ax = plt.subplots(figsize=(8.0, 4.6), constrained_layout=True)
    color_map = {"Big firms": "#6d597a", "Small firms": "#c46b48"}
    marker_map = {"CAR[-1,+1]": "o", "BHAR[+2,+63]": "D"}
    y_positions = {
        ("CAR[-1,+1]", "Big firms"): 1.15,
        ("CAR[-1,+1]", "Small firms"): 0.85,
        ("BHAR[+2,+63]", "Big firms"): 0.15,
        ("BHAR[+2,+63]", "Small firms"): -0.15,
    }
    for row in figure_df.itertuples(index=False):
        y = y_positions[(row.outcome, row.group)]
        ax.errorbar(
            100 * row.coef,
            y,
            xerr=[[100 * (row.coef - row.ci_low)], [100 * (row.ci_high - row.coef)]],
            fmt=marker_map[row.outcome],
            color=color_map[row.group],
            markersize=7,
            capsize=4,
            linewidth=1.6,
            label=f"{row.outcome} / {row.group}",
        )
    ax.axvline(0, color="#6c757d", linewidth=0.9)
    ax.set_yticks([1.0, 0.0])
    ax.set_yticklabels(["CAR[-1,+1]", "BHAR[+2,+63]"])
    ax.set_xlabel("Implied PatentMismatch effect from pooled interaction (pct)")
    ax.set_title("Size Heterogeneity in Mismatch Pricing Effects")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.grid(True, axis="x", color="#d9d9d9", linewidth=0.7)
    ax.grid(False, axis="y")
    handles, labels = ax.get_legend_handles_labels()
    dedup: dict[str, object] = {}
    for handle, label in zip(handles, labels, strict=False):
        dedup.setdefault(label, handle)
    ax.legend(dedup.values(), dedup.keys(), frameon=False, loc="best", fontsize=9)

    png_path = output_base.with_suffix(".png")
    pdf_path = output_base.with_suffix(".pdf")
    fig.savefig(png_path, dpi=220, bbox_inches="tight")
    fig.savefig(pdf_path, bbox_inches="tight")
    plt.close(fig)
    return png_path, pdf_path


def _result_notes(
    outcome_results: list[dict[str, object]], size_summary: dict[str, object]
) -> str:
    by_outcome = {row["short_label"]: row for row in outcome_results}
    car = by_outcome["CAR[-1,+1]"]["pooled"]
    bhar = by_outcome["BHAR[+2,+63]"]["pooled"]
    return "\n".join(
        [
            "# Result Notes",
            "",
            f"- Filing-date CAR pooled interaction: big-firm mismatch effect `{car['big_mismatch_effect']['coef']:.4f}` (p=`{car['big_mismatch_effect']['p_value']:.3f}`); small-firm implied mismatch effect `{car['small_mismatch_effect']['coef']:.4f}` (p=`{car['small_mismatch_effect']['p_value']:.3f}`).",
            f"- Post-filing BHAR pooled interaction: big-firm mismatch effect `{bhar['big_mismatch_effect']['coef']:.4f}` (p=`{bhar['big_mismatch_effect']['p_value']:.3f}`); small-firm implied mismatch effect `{bhar['small_mismatch_effect']['coef']:.4f}` (p=`{bhar['small_mismatch_effect']['p_value']:.3f}`).",
            f"- Filing sample with usable size assignment: `{size_summary['filing_count_with_size']}` filings; small = `{size_summary['small_count']}`, big = `{size_summary['big_count']}`.",
            "",
        ]
    )


def _writer_packet(
    args: argparse.Namespace,
    size_summary: dict[str, object],
    outcome_results: list[dict[str, object]],
) -> str:
    by_outcome = {row["short_label"]: row for row in outcome_results}
    car = by_outcome["CAR[-1,+1]"]["pooled"]
    bhar = by_outcome["BHAR[+2,+63]"]["pooled"]
    return "\n".join(
        [
            "# Writer Packet",
            "",
            "## Metadata",
            f"- Test id: `{TEST_ID}`",
            f"- Run id: `{args.run_id}`",
            f"- Date run: `{date.today().isoformat()}`",
            f"- Script/module path: `{MODULE_PATH}`",
            f"- Input files: `{args.event_panel}`, `{args.annual_panel}`, `{args.monthly_returns}`",
            "- Unit of observation: `annual AI filing event`",
            "",
            "## Size Design",
            f"- Breakpoint method: `{size_summary['size_break_method']}`",
            f"- Filings with size assignment: `{size_summary['filing_count_with_size']}`",
            f"- Small-firm filings: `{size_summary['small_count']}`",
            f"- Big-firm filings: `{size_summary['big_count']}`",
            "",
            "## Outcomes",
            "- Short-run outcome: `CAR[-1,+1]`",
            "- Drift outcome: `BHAR[+2,+63]`",
            "- Main specification: `small subsample, big subsample, and pooled interaction with PatentMismatch × Small`",
            "",
            "## Results",
            f"- CAR pooled interaction differential: `PatentMismatch × Small = {car['params'].get('PatentMismatch:Small', math.nan):.4f}` with p = `{car['pvalues'].get('PatentMismatch:Small', math.nan):.3f}`",
            f"- BHAR pooled interaction differential: `PatentMismatch × Small = {bhar['params'].get('PatentMismatch:Small', math.nan):.4f}` with p = `{bhar['pvalues'].get('PatentMismatch:Small', math.nan):.3f}`",
            f"- Implied small-firm mismatch effect in BHAR: `{bhar['small_mismatch_effect']['coef']:.4f}` with p = `{bhar['small_mismatch_effect']['p_value']:.3f}`",
            "",
            "## Caption Draft",
            "This table reports heterogeneity by firm size for the filing-date and post-filing mismatch effects. Small firms are defined using the monthly active-sample median lagged market capitalization from CRSP MSF because NYSE breakpoints are not available in the local extract. For each outcome, the first two columns re-estimate the baseline regression separately in the small-firm and big-firm subsamples. The third column estimates a pooled interaction regression with PatentMismatch, Small, and PatentMismatch × Small, alongside the same disclosure controls, industry fixed effects, filing-year fixed effects, and firm-clustered standard errors.",
            "",
        ]
    )


def _dataset_summary(
    args: argparse.Namespace,
    size_summary: dict[str, object],
    outcome_results: list[dict[str, object]],
    figure_df: pd.DataFrame,
) -> dict[str, object]:
    payload = figure_df.copy()
    return {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "event_panel": str(args.event_panel),
        "annual_panel": str(args.annual_panel),
        "monthly_returns": str(args.monthly_returns),
        "size_summary": size_summary,
        "outcome_results": outcome_results,
        "figure_series": payload.to_dict(orient="records"),
    }


def _copy_exports(run_dir: Path, paper_root: Path, run_id: str) -> dict[str, str]:
    exports = {
        "figure_png": paper_root / "figures" / f"{TEST_ID}_{run_id}.png",
        "figure_pdf": paper_root / "figures" / f"{TEST_ID}_{run_id}.pdf",
        "table_csv": paper_root / "tables" / f"{TEST_ID}_{run_id}.csv",
        "table_md": paper_root / "tables" / f"{TEST_ID}_{run_id}.md",
        "table_tex": paper_root / "latex" / f"{TEST_ID}_{run_id}.tex",
        "table_docx": paper_root / "docx" / f"{TEST_ID}_{run_id}.docx",
        "figure_series": paper_root / "tables" / f"{TEST_ID}_{run_id}_figure_series.csv",
        "writer_packet": paper_root / "writer_packets" / f"{TEST_ID}_{run_id}.md",
        "result_notes": paper_root / "snippets" / f"{TEST_ID}_{run_id}_result_notes.md",
    }
    for path in exports.values():
        path.parent.mkdir(parents=True, exist_ok=True)
    mapping = {
        run_dir / "figure_main.png": exports["figure_png"],
        run_dir / "figure_main.pdf": exports["figure_pdf"],
        run_dir / "table_main.csv": exports["table_csv"],
        run_dir / "table_main.md": exports["table_md"],
        run_dir / "table_main.tex": exports["table_tex"],
        run_dir / "table_main.docx": exports["table_docx"],
        run_dir / "figure_series.csv": exports["figure_series"],
        run_dir / "writer_packet.md": exports["writer_packet"],
        run_dir / "result_notes.md": exports["result_notes"],
    }
    for src, dst in mapping.items():
        shutil.copy2(src, dst)
    return {key: str(path) for key, path in exports.items()}


def main() -> None:
    args = _parse_args()
    run_dir = args.test_root / args.run_id
    run_dir.mkdir(parents=True, exist_ok=True)

    sample, size_summary = _prepare_sample(
        args.event_panel, args.annual_panel, args.monthly_returns
    )

    outcome_results: list[dict[str, object]] = []
    for spec in OUTCOME_SPECS:
        outcome_results.append(
            {
                **spec,
                "small": _fit_subsample_regression(sample, spec["dependent"], small_flag=1),
                "big": _fit_subsample_regression(sample, spec["dependent"], small_flag=0),
                "pooled": _fit_interaction_regression(sample, spec["dependent"]),
            }
        )

    table_rows, figure_df = _table_bundle(outcome_results)
    pd.DataFrame(table_rows).to_csv(run_dir / "table_main.csv", index=False)
    figure_df.to_csv(run_dir / "figure_series.csv", index=False)

    table_md, table_tex = _render_table_outputs(table_rows)
    (run_dir / "table_main.md").write_text(table_md, encoding="utf-8")
    (run_dir / "table_main.tex").write_text(table_tex, encoding="utf-8")
    _build_table_docx(table_rows, run_dir / "table_main.docx")
    png_path, pdf_path = _plot_heterogeneity_figure(figure_df, run_dir / "figure_main")

    (run_dir / "result_notes.md").write_text(
        _result_notes(outcome_results, size_summary), encoding="utf-8"
    )
    (run_dir / "writer_packet.md").write_text(
        _writer_packet(args, size_summary, outcome_results), encoding="utf-8"
    )

    summary = _dataset_summary(args, size_summary, outcome_results, figure_df)
    (run_dir / "dataset_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")

    paper_exports = _copy_exports(run_dir, args.paper_root, args.run_id)
    manifest = {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "module_path": MODULE_PATH,
        "inputs": {
            "event_panel": str(args.event_panel),
            "annual_panel": str(args.annual_panel),
            "monthly_returns": str(args.monthly_returns),
        },
        "run_dir": str(run_dir),
        "outputs": {
            "dataset_summary": str(run_dir / "dataset_summary.json"),
            "table_csv": str(run_dir / "table_main.csv"),
            "table_md": str(run_dir / "table_main.md"),
            "table_tex": str(run_dir / "table_main.tex"),
            "table_docx": str(run_dir / "table_main.docx"),
            "figure_series": str(run_dir / "figure_series.csv"),
            "figure_png": str(png_path),
            "figure_pdf": str(pdf_path),
            "writer_packet": str(run_dir / "writer_packet.md"),
            "result_notes": str(run_dir / "result_notes.md"),
        },
        "paper_exports": paper_exports,
    }
    (run_dir / "run_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    print(f"[{TEST_ID}] wrote run bundle to {run_dir}")
    print(f"[{TEST_ID}] paper table: {paper_exports['table_docx']}")


if __name__ == "__main__":
    main()
