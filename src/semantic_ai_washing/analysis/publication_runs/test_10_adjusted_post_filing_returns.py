"""Publication run driver for Test 10: industry-adjusted and characteristic-adjusted post-filing returns."""

from __future__ import annotations

import argparse
import json
import shutil
from datetime import UTC, date, datetime
from pathlib import Path

from docx import Document
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from scipy import stats

from semantic_ai_washing.analysis.delivery_table_payloads import _add_patent_mismatch
from semantic_ai_washing.analysis.publication_runs.test_03_post_filing_drift import (
    _add_note,
    _add_title,
    _build_panel_table,
    _set_document_defaults,
    _set_landscape,
)

REPO_ROOT = Path(__file__).resolve().parents[4]
DEFAULT_EVENT_PANEL = (
    REPO_ROOT
    / "data/processed/panel/filing_event_estimation_sample_hybrid_api_a_conf49_v1.parquet"
)
DEFAULT_ANNUAL_PANEL = (
    REPO_ROOT
    / "data/processed/panel/canonical/ever_speaker_panel_2016_2025_hybrid_api_a_conf49_v1.parquet"
)
DEFAULT_TEST_ROOT = Path(
    "/Users/soheilkhodadadi/DataWork/semantic-patterns/ai_washing/derived/v3_1/test_runs/test_10_adjusted_post_filing_returns"
)
DEFAULT_PAPER_ROOT = REPO_ROOT / "paper/generated/v3_1"
DEFAULT_RUN_ID = f"{date.today():%Y%m%d}_aiw_v3_1_test_10_adjusted_post_filing_returns_main_v1"
TEST_ID = "test_10_adjusted_post_filing_returns"
MODULE_PATH = "semantic_ai_washing.analysis.publication_runs.test_10_adjusted_post_filing_returns"

OUTCOME_SPECS = [
    ("bhar_1m", "BHAR[+2,+21]", 1),
    ("bhar_3m", "BHAR[+2,+63]", 3),
    ("bhar_6m", "BHAR[+2,+126]", 6),
    ("bhar_12m", "BHAR[+2,+252]", 12),
]

BENCHMARK_SPECS = [
    ("raw", "Raw BHAR", None),
    ("industry", "Industry-adjusted", ["sic2"]),
    ("industry_year", "Industry-year-adjusted", ["filing_year", "sic2"]),
    ("characteristic", "Characteristic-adjusted", ["filing_year", "size_bucket", "mom_bucket"]),
]


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--event-panel", type=Path, default=DEFAULT_EVENT_PANEL)
    parser.add_argument("--annual-panel", type=Path, default=DEFAULT_ANNUAL_PANEL)
    parser.add_argument("--test-root", type=Path, default=DEFAULT_TEST_ROOT)
    parser.add_argument("--paper-root", type=Path, default=DEFAULT_PAPER_ROOT)
    parser.add_argument("--run-id", default=DEFAULT_RUN_ID)
    return parser.parse_args()


def _normalize_id(series: pd.Series) -> pd.Series:
    return (
        series.astype(str)
        .str.replace(r"\.0$", "", regex=True)
        .str.strip()
        .replace({"nan": "", "None": ""})
    )


def _ensure_sic2(panel: pd.DataFrame) -> pd.DataFrame:
    out = panel.copy()
    if "sic2" not in out.columns or out["sic2"].isna().all():
        if "sic" in out.columns:
            sic_raw = pd.to_numeric(out["sic"], errors="coerce")
            out["sic2"] = (sic_raw // 100).astype("Int64")
        else:
            out["sic2"] = pd.Series(pd.NA, index=out.index, dtype="Int64")
    return out


def _base_style() -> None:
    plt.style.use("seaborn-v0_8-whitegrid")
    plt.rcParams.update(
        {
            "font.family": "serif",
            "font.serif": ["Times New Roman", "Times", "DejaVu Serif"],
            "axes.titlesize": 13,
            "axes.labelsize": 11,
            "xtick.labelsize": 10,
            "ytick.labelsize": 10,
            "legend.fontsize": 9,
        }
    )


def _load_annual_backbone(path: Path) -> pd.DataFrame:
    annual = pd.read_parquet(path).copy()
    annual = _ensure_sic2(annual)
    annual = _add_patent_mismatch(annual)
    annual["cik"] = _normalize_id(annual["cik"])
    annual["year"] = pd.to_numeric(annual["year"], errors="coerce").astype("Int64")
    return annual


def _rank_bucket(series: pd.Series, labels: list[str]) -> pd.Series:
    valid = pd.to_numeric(series, errors="coerce")
    out = pd.Series(pd.NA, index=series.index, dtype="object")
    if valid.notna().sum() < len(labels):
        return out
    ranked = valid.rank(method="first")
    buckets = pd.qcut(ranked, len(labels), labels=labels, duplicates="drop")
    out.loc[buckets.index] = buckets.astype("object")
    return out


def _build_analysis_sample(event_panel: Path, annual_panel: Path) -> tuple[pd.DataFrame, dict[str, object]]:
    event = pd.read_parquet(event_panel).copy()
    event["cik"] = _normalize_id(event["cik"])
    event["filing_year"] = pd.to_numeric(event["filing_year"], errors="coerce").astype("Int64")
    event["n_ai_total"] = pd.to_numeric(event["n_ai_total"], errors="coerce")
    event = event.loc[event["n_ai_total"].fillna(0).gt(0)].copy()

    annual = _load_annual_backbone(annual_panel)
    current = annual[
        ["cik", "year", "sic2", "PatentMismatch", "A_S", "AI_Focus"]
    ].copy()
    lagged = annual[
        ["cik", "year", "market_cap_year_end", "annual_bhar_vw", "annual_ret"]
    ].copy()
    lagged["year"] = lagged["year"] + 1
    lagged = lagged.rename(
        columns={
            "market_cap_year_end": "market_cap_lag1",
            "annual_bhar_vw": "annual_bhar_vw_lag1",
            "annual_ret": "annual_ret_lag1",
        }
    )

    sample = event.merge(
        current,
        left_on=["cik", "filing_year"],
        right_on=["cik", "year"],
        how="left",
        validate="many_to_one",
    )
    sample = sample.merge(
        lagged,
        left_on=["cik", "filing_year"],
        right_on=["cik", "year"],
        how="left",
        validate="many_to_one",
        suffixes=("", "_lag1"),
    )
    sample["sic2"] = pd.to_numeric(sample["sic2"], errors="coerce").astype("Int64")
    sample["PatentMismatch"] = pd.to_numeric(sample["PatentMismatch"], errors="coerce")
    sample = sample.loc[sample["PatentMismatch"].notna()].copy()
    sample["PatentMismatch"] = sample["PatentMismatch"].astype(int)

    sample["size_bucket"] = pd.Series(pd.NA, index=sample.index, dtype="object")
    sample["mom_bucket"] = pd.Series(pd.NA, index=sample.index, dtype="object")
    for filing_year, idx in sample.groupby("filing_year").groups.items():
        del filing_year
        sample.loc[idx, "size_bucket"] = _rank_bucket(
            sample.loc[idx, "market_cap_lag1"], ["S", "M", "L"]
        )
        sample.loc[idx, "mom_bucket"] = _rank_bucket(
            sample.loc[idx, "annual_bhar_vw_lag1"], ["L", "M", "H"]
        )

    summary = {
        "ai_filing_rows": int(len(sample)),
        "unique_firms": int(sample["gvkey"].replace("", pd.NA).dropna().nunique()),
        "filing_year_min": int(sample["filing_year"].min()),
        "filing_year_max": int(sample["filing_year"].max()),
        "mismatch_share": float(sample["PatentMismatch"].mean()),
    }
    return sample, summary


def _add_peer_adjusted_outcome(
    sample: pd.DataFrame, outcome: str, group_cols: list[str], benchmark_id: str
) -> pd.DataFrame:
    use = sample[["filing_id", outcome, *group_cols]].dropna(subset=[outcome, *group_cols]).copy()
    grouped = use.groupby(group_cols)[outcome].agg(["sum", "count"]).reset_index()
    use = use.merge(grouped, on=group_cols, how="left")
    peer_sum = use["sum"] - use[outcome]
    peer_count = use["count"] - 1
    adjusted = np.where(peer_count.gt(0), use[outcome] - (peer_sum / peer_count), np.nan)
    return pd.DataFrame(
        {
            "filing_id": use["filing_id"].to_numpy(),
            f"{outcome}_{benchmark_id}": adjusted,
            f"{outcome}_{benchmark_id}_peer_count": peer_count.to_numpy(),
        }
    )


def _add_adjusted_outcomes(sample: pd.DataFrame) -> pd.DataFrame:
    out = sample.copy()
    for outcome, _, _ in OUTCOME_SPECS:
        for benchmark_id, _, group_cols in BENCHMARK_SPECS:
            if group_cols is None:
                out[f"{outcome}_{benchmark_id}"] = pd.to_numeric(out[outcome], errors="coerce")
                out[f"{outcome}_{benchmark_id}_peer_count"] = np.nan
                continue
            adjusted = _add_peer_adjusted_outcome(out, outcome, group_cols, benchmark_id)
            out = out.merge(adjusted, on="filing_id", how="left", validate="one_to_one")
    return out


def _summarize_diff(sample: pd.DataFrame, column: str, peer_col: str) -> dict[str, object]:
    use = sample[["PatentMismatch", column, peer_col]].dropna(subset=[column]).copy()
    use["PatentMismatch"] = pd.to_numeric(use["PatentMismatch"], errors="coerce").fillna(0).astype(int)
    mismatch = use.loc[use["PatentMismatch"].eq(1), column].astype(float)
    non_mismatch = use.loc[use["PatentMismatch"].eq(0), column].astype(float)
    t_stat, p_value = stats.ttest_ind(mismatch, non_mismatch, equal_var=False, nan_policy="omit")
    se_diff = np.sqrt(
        (mismatch.var(ddof=1) / len(mismatch)) + (non_mismatch.var(ddof=1) / len(non_mismatch))
    )
    median_peers = pd.to_numeric(use[peer_col], errors="coerce").median()
    return {
        "mean_mismatch": float(mismatch.mean()),
        "mean_non_mismatch": float(non_mismatch.mean()),
        "diff_mismatch_minus_non": float(mismatch.mean() - non_mismatch.mean()),
        "p_value": float(p_value),
        "se_diff": float(se_diff),
        "n_total": int(len(use)),
        "n_mismatch": int(len(mismatch)),
        "n_non_mismatch": int(len(non_mismatch)),
        "median_peer_count": None if pd.isna(median_peers) else float(median_peers),
    }


def _build_results(sample: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    table_rows: list[dict[str, object]] = []
    figure_rows: list[dict[str, object]] = []
    for outcome, outcome_label, horizon_months in OUTCOME_SPECS:
        for benchmark_id, benchmark_label, _ in BENCHMARK_SPECS:
            summary = _summarize_diff(
                sample,
                f"{outcome}_{benchmark_id}",
                f"{outcome}_{benchmark_id}_peer_count",
            )
            table_rows.append(
                {
                    "panel": outcome_label,
                    "benchmark_label": benchmark_label,
                    **summary,
                }
            )
            figure_rows.append(
                {
                    "benchmark_id": benchmark_id,
                    "benchmark_label": benchmark_label,
                    "horizon_months": horizon_months,
                    "outcome_label": outcome_label,
                    "diff_mismatch_minus_non": summary["diff_mismatch_minus_non"],
                    "ci_low": summary["diff_mismatch_minus_non"] - 1.96 * summary["se_diff"],
                    "ci_high": summary["diff_mismatch_minus_non"] + 1.96 * summary["se_diff"],
                    "p_value": summary["p_value"],
                    "n_total": summary["n_total"],
                }
            )
    return pd.DataFrame(table_rows), pd.DataFrame(figure_rows)


def _markdown_table(headers: list[str], rows: list[list[object]]) -> str:
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join(["---"] * len(headers)) + " |",
    ]
    for row in rows:
        lines.append("| " + " | ".join(str(item) for item in row) + " |")
    return "\n".join(lines)


def _render_table_outputs(table_df: pd.DataFrame) -> tuple[str, str]:
    headers = ["Row", *[label for _, label, _ in BENCHMARK_SPECS]]
    md_lines = ["# Table Main", ""]
    latex_lines = [
        "\\begin{table}[!htbp]",
        "\\centering",
        "\\caption{Post-filing return spreads under alternative benchmark adjustments}",
        "\\begin{tabular}{l" + "c" * len(BENCHMARK_SPECS) + "}",
        "\\hline",
    ]
    for panel in [label for _, label, _ in OUTCOME_SPECS]:
        subset = table_df.loc[table_df["panel"].eq(panel)].copy()
        rows = []
        for row_label, field, formatter in [
            ("Mean mismatch", "mean_mismatch", lambda x: f"{100 * x:.3f}"),
            ("Mean non-mismatch", "mean_non_mismatch", lambda x: f"{100 * x:.3f}"),
            ("Diff", "diff_mismatch_minus_non", lambda x: f"{100 * x:.3f}"),
            ("p-value", "p_value", lambda x: f"{x:.3f}"),
            ("N", "n_total", lambda x: str(int(x))),
            (
                "Median peer count",
                "median_peer_count",
                lambda x: "" if x is None or pd.isna(x) else f"{x:.1f}",
            ),
        ]:
            rendered_row = [row_label]
            for _, benchmark_label, _ in BENCHMARK_SPECS:
                value = subset.loc[subset["benchmark_label"].eq(benchmark_label), field].iloc[0]
                rendered_row.append(formatter(value))
            rows.append(rendered_row)
        md_lines.extend([f"## Panel {panel}", _markdown_table(headers, rows), ""])
        latex_lines.append(f"\\multicolumn{{{len(headers)}}}{{l}}{{\\textit{{Panel {panel}}}}} \\\\")
        latex_lines.append(" & ".join(headers) + " \\\\")
        for row in rows:
            latex_lines.append(" & ".join(str(item) for item in row) + " \\\\")
    latex_lines.extend(["\\hline", "\\end{tabular}", "\\end{table}"])
    return "\n".join(md_lines), "\n".join(latex_lines) + "\n"


def _docx_panel_rows(table_df: pd.DataFrame, panel_name: str) -> list[list[tuple[str, bool]]]:
    rows = []
    subset = table_df.loc[table_df["panel"].eq(panel_name)].copy()
    for row_label, field, formatter in [
        ("Mean mismatch", "mean_mismatch", lambda x: f"{100 * x:.3f}"),
        ("Mean non-mismatch", "mean_non_mismatch", lambda x: f"{100 * x:.3f}"),
        ("Diff", "diff_mismatch_minus_non", lambda x: f"{100 * x:.3f}"),
        ("p-value", "p_value", lambda x: f"{x:.3f}"),
        ("N", "n_total", lambda x: str(int(x))),
        (
            "Median peer count",
            "median_peer_count",
            lambda x: "" if x is None or pd.isna(x) else f"{x:.1f}",
        ),
    ]:
        rendered = [(row_label, True)]
        for _, benchmark_label, _ in BENCHMARK_SPECS:
            value = subset.loc[subset["benchmark_label"].eq(benchmark_label), field].iloc[0]
            rendered.append((formatter(value), False))
        rows.append(rendered)
    return rows


def _build_table_docx(table_df: pd.DataFrame, output_path: Path) -> None:
    document = Document()
    _set_document_defaults(document)
    _set_landscape(document)
    _add_title(document, "Table 10. Post-Filing Return Spreads Under Alternative Benchmark Adjustments")
    note = (
        "This table reports mismatch-minus-non-mismatch buy-and-hold abnormal return spreads under four benchmark constructions: raw market-adjusted BHAR, industry-adjusted BHAR, industry-year-adjusted BHAR, and a coarse characteristic-adjusted BHAR benchmarked within filing year by lagged size and lagged relative return buckets. Median peer count is not reported for the raw BHAR column because that specification does not use a peer benchmark."
    )
    _add_note(document, note)
    headers = ["", *[label for _, label, _ in BENCHMARK_SPECS]]
    for _, panel_label, _ in OUTCOME_SPECS:
        p = document.add_paragraph()
        p.add_run(f"Panel {panel_label}").bold = True
        _build_panel_table(document, headers, _docx_panel_rows(table_df, panel_label))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def _plot_spreads(figure_df: pd.DataFrame, output_base: Path) -> tuple[Path, Path]:
    _base_style()
    fig, ax = plt.subplots(figsize=(8.6, 4.8), constrained_layout=True)
    color_map = {
        "Raw BHAR": "#1d3557",
        "Industry-adjusted": "#457b9d",
        "Industry-year-adjusted": "#e76f51",
        "Characteristic-adjusted": "#2a9d8f",
    }
    horizon_labels = {1: "1m", 3: "3m", 6: "6m", 12: "12m"}
    for benchmark_label, sub in figure_df.groupby("benchmark_label", sort=False):
        ordered = sub.sort_values("horizon_months")
        ax.plot(
            ordered["horizon_months"],
            100 * ordered["diff_mismatch_minus_non"],
            marker="o",
            linewidth=2.0,
            color=color_map[benchmark_label],
            label=benchmark_label,
        )
    ax.axhline(0, color="#6c757d", linewidth=0.8)
    ax.set_xticks([1, 3, 6, 12], [horizon_labels[x] for x in [1, 3, 6, 12]])
    ax.set_xlabel("Post-filing horizon")
    ax.set_ylabel("Mismatch minus non-mismatch BHAR (pct)")
    ax.set_title("Adjusted Post-Filing Return Spreads Across Benchmark Constructions")
    ax.legend(frameon=False, ncol=2, loc="best")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    png_path = output_base.with_suffix(".png")
    pdf_path = output_base.with_suffix(".pdf")
    fig.savefig(png_path, dpi=220, bbox_inches="tight")
    fig.savefig(pdf_path, bbox_inches="tight")
    plt.close(fig)
    return png_path, pdf_path


def _result_notes(table_df: pd.DataFrame, sample_summary: dict[str, object]) -> str:
    panel_3m = table_df.loc[table_df["panel"].eq("BHAR[+2,+63]")].set_index("benchmark_label")
    panel_1m = table_df.loc[table_df["panel"].eq("BHAR[+2,+21]")].set_index("benchmark_label")
    panel_12m = table_df.loc[table_df["panel"].eq("BHAR[+2,+252]")].set_index("benchmark_label")
    return "\n".join(
        [
            f"# Result Notes: {TEST_ID}",
            "",
            f"- AI-filing sample: `{sample_summary['ai_filing_rows']}` rows across `{sample_summary['unique_firms']}` firms.",
            f"- 1m industry-adjusted diff: `{100 * panel_1m.loc['Industry-adjusted', 'diff_mismatch_minus_non']:.3f}` pct (p=`{panel_1m.loc['Industry-adjusted', 'p_value']:.3f}`).",
            f"- 3m raw diff: `{100 * panel_3m.loc['Raw BHAR', 'diff_mismatch_minus_non']:.3f}` pct (p=`{panel_3m.loc['Raw BHAR', 'p_value']:.3f}`).",
            f"- 3m industry-year-adjusted diff: `{100 * panel_3m.loc['Industry-year-adjusted', 'diff_mismatch_minus_non']:.3f}` pct (p=`{panel_3m.loc['Industry-year-adjusted', 'p_value']:.3f}`).",
            f"- 12m raw diff: `{100 * panel_12m.loc['Raw BHAR', 'diff_mismatch_minus_non']:.3f}` pct (p=`{panel_12m.loc['Raw BHAR', 'p_value']:.3f}`).",
            "- Interpretation discipline: if the spread remains weak across industry-time and characteristic benchmarks, the market-results block should not claim that the signal is robust to benchmark choice.",
            "",
        ]
    )


def _writer_packet(args: argparse.Namespace, table_df: pd.DataFrame, sample_summary: dict[str, object]) -> str:
    panel_3m = table_df.loc[table_df["panel"].eq("BHAR[+2,+63]")].set_index("benchmark_label")
    return "\n".join(
        [
            f"# Writer Packet: {TEST_ID}",
            "",
            "## Purpose",
            "- This run asks whether the post-filing spread survives alternative benchmark constructions rather than depending on the raw market-adjusted BHAR definition alone.",
            "- The benchmark ladder is designed to absorb industry AI cycles first and then a coarse size-plus-prior-return characteristic mix.",
            "",
            "## Sample Definition",
            f"- Event panel: `{args.event_panel}`",
            f"- Annual panel: `{args.annual_panel}`",
            "- Unit of observation: `AI-talking annual filing event`",
            "- Outcome horizons: `1m`, `3m`, `6m`, `12m` BHAR windows",
            "",
            "## Benchmark Ladder",
            "- Raw BHAR: event-level firm buy-and-hold return minus market buy-and-hold return",
            "- Industry-adjusted: event BHAR minus peer mean within `sic2`",
            "- Industry-year-adjusted: event BHAR minus peer mean within `filing_year × sic2`",
            "- Characteristic-adjusted: event BHAR minus peer mean within `filing_year × lagged size tercile × lagged relative-return tercile`",
            "",
            "## Sample Counts",
            f"- AI-filing rows: `{sample_summary['ai_filing_rows']}`",
            f"- Unique firms: `{sample_summary['unique_firms']}`",
            "",
            "## Main 3m Read",
            f"- Raw 3m diff: `{100 * panel_3m.loc['Raw BHAR', 'diff_mismatch_minus_non']:.3f}` pct (p=`{panel_3m.loc['Raw BHAR', 'p_value']:.3f}`)",
            f"- Industry-year-adjusted 3m diff: `{100 * panel_3m.loc['Industry-year-adjusted', 'diff_mismatch_minus_non']:.3f}` pct (p=`{panel_3m.loc['Industry-year-adjusted', 'p_value']:.3f}`)",
            f"- Characteristic-adjusted 3m diff: `{100 * panel_3m.loc['Characteristic-adjusted', 'diff_mismatch_minus_non']:.3f}` pct (p=`{panel_3m.loc['Characteristic-adjusted', 'p_value']:.3f}`)",
            "",
            "## Caption Draft",
            "This table and figure test whether the post-filing mismatch spread survives alternative benchmark constructions. Starting from raw market-adjusted BHAR, the run then removes industry means, industry-year means, and a coarse characteristic benchmark based on lagged size and lagged relative return. If the spread disappears once these benchmark adjustments are applied, the market-results block should be described as benchmark-sensitive rather than as a robust pricing anomaly.",
            "",
        ]
    )


def _dataset_summary(
    sample_summary: dict[str, object], table_df: pd.DataFrame, figure_df: pd.DataFrame, *, args: argparse.Namespace
) -> dict[str, object]:
    return {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "inputs": {
            "event_panel": str(args.event_panel),
            "annual_panel": str(args.annual_panel),
        },
        "sample_summary": sample_summary,
        "table_rows": table_df.to_dict(orient="records"),
        "figure_rows": figure_df.to_dict(orient="records"),
    }


def _copy_exports(run_dir: Path, paper_root: Path, run_id: str) -> dict[str, str]:
    exports = {
        "table_csv": paper_root / "tables" / f"{TEST_ID}_{run_id}.csv",
        "table_md": paper_root / "tables" / f"{TEST_ID}_{run_id}.md",
        "table_tex": paper_root / "latex" / f"{TEST_ID}_{run_id}.tex",
        "table_docx": paper_root / "docx" / f"{TEST_ID}_{run_id}.docx",
        "figure_png": paper_root / "figures" / f"{TEST_ID}_{run_id}.png",
        "figure_pdf": paper_root / "figures" / f"{TEST_ID}_{run_id}.pdf",
        "figure_series": paper_root / "tables" / f"{TEST_ID}_{run_id}_figure_series.csv",
        "writer_packet": paper_root / "writer_packets" / f"{TEST_ID}_{run_id}.md",
        "result_notes": paper_root / "snippets" / f"{TEST_ID}_{run_id}_result_notes.md",
    }
    for path in exports.values():
        path.parent.mkdir(parents=True, exist_ok=True)
    mapping = {
        run_dir / "table_main.csv": exports["table_csv"],
        run_dir / "table_main.md": exports["table_md"],
        run_dir / "table_main.tex": exports["table_tex"],
        run_dir / "table_main.docx": exports["table_docx"],
        run_dir / "main_figure.png": exports["figure_png"],
        run_dir / "main_figure.pdf": exports["figure_pdf"],
        run_dir / "figure_series.csv": exports["figure_series"],
        run_dir / "writer_packet.md": exports["writer_packet"],
        run_dir / "result_notes.md": exports["result_notes"],
    }
    for src, dst in mapping.items():
        shutil.copy2(src, dst)
    return {key: str(path) for key, path in exports.items()}


def main() -> None:
    args = _parse_args()
    run_dir = args.test_root / args.run_id
    run_dir.mkdir(parents=True, exist_ok=True)

    sample, sample_summary = _build_analysis_sample(args.event_panel, args.annual_panel)
    sample = _add_adjusted_outcomes(sample)
    table_df, figure_df = _build_results(sample)

    table_df.to_csv(run_dir / "table_main.csv", index=False)
    table_md, table_tex = _render_table_outputs(table_df)
    (run_dir / "table_main.md").write_text(table_md, encoding="utf-8")
    (run_dir / "table_main.tex").write_text(table_tex, encoding="utf-8")
    _build_table_docx(table_df, run_dir / "table_main.docx")
    figure_df.to_csv(run_dir / "figure_series.csv", index=False)
    fig_png, fig_pdf = _plot_spreads(figure_df, run_dir / "main_figure")
    (run_dir / "result_notes.md").write_text(_result_notes(table_df, sample_summary), encoding="utf-8")
    (run_dir / "writer_packet.md").write_text(
        _writer_packet(args, table_df, sample_summary), encoding="utf-8"
    )
    dataset_summary = _dataset_summary(sample_summary, table_df, figure_df, args=args)
    (run_dir / "dataset_summary.json").write_text(json.dumps(dataset_summary, indent=2), encoding="utf-8")

    paper_exports = _copy_exports(run_dir, args.paper_root, args.run_id)
    manifest = {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "module_path": MODULE_PATH,
        "run_dir": str(run_dir),
        "inputs": {
            "event_panel": str(args.event_panel),
            "annual_panel": str(args.annual_panel),
        },
        "outputs": {
            "dataset_summary": str(run_dir / "dataset_summary.json"),
            "table_csv": str(run_dir / "table_main.csv"),
            "table_md": str(run_dir / "table_main.md"),
            "table_tex": str(run_dir / "table_main.tex"),
            "table_docx": str(run_dir / "table_main.docx"),
            "figure_png": str(fig_png),
            "figure_pdf": str(fig_pdf),
            "figure_series": str(run_dir / "figure_series.csv"),
            "writer_packet": str(run_dir / "writer_packet.md"),
            "result_notes": str(run_dir / "result_notes.md"),
        },
        "paper_exports": paper_exports,
    }
    (run_dir / "run_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    print(f"[{TEST_ID}] wrote run bundle to {run_dir}")
    print(f"[{TEST_ID}] paper table: {paper_exports['table_docx']}")
    print(f"[{TEST_ID}] paper figure: {paper_exports['figure_png']}")


if __name__ == "__main__":
    main()
