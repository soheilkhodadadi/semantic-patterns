"""Publication run driver for Test 07: ChatGPT DID and event path."""

from __future__ import annotations

import argparse
import json
import math
import shutil
from datetime import UTC, date, datetime
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from scipy.stats import chi2
import statsmodels.api as sm

from semantic_ai_washing.analysis.delivery_table_payloads import _add_patent_mismatch
from semantic_ai_washing.analysis.publication_runs.test_03_post_filing_drift import (
    _add_note,
    _add_title,
    _set_document_defaults,
    _set_landscape,
)

REPO_ROOT = Path(__file__).resolve().parents[4]
DEFAULT_EVENT_PANEL = (
    REPO_ROOT
    / "data/processed/panel/filing_event_estimation_sample_hybrid_api_a_conf49_v1.parquet"
)
DEFAULT_ANNUAL_PANEL = (
    REPO_ROOT
    / "data/processed/panel/canonical/ever_speaker_panel_2016_2025_hybrid_api_a_conf49_v1.parquet"
)
DEFAULT_TEST_ROOT = Path(
    "/Users/soheilkhodadadi/DataWork/semantic-patterns/ai_washing/derived/test_runs/test_07_chatgpt_did"
)
DEFAULT_PAPER_ROOT = REPO_ROOT / "paper/generated"
DEFAULT_RUN_ID = f"{date.today():%Y%m%d}_hybrid_api_a_conf49_main_v1"
TEST_ID = "test_07_chatgpt_did"
MODULE_PATH = "semantic_ai_washing.analysis.publication_runs.test_07_chatgpt_did"
CONTROL_TERMS = ["ln_assets_l1", "leverage_l1", "cash_l1", "roa_l1"]
MAIN_OUTCOMES = [
    ("AI_Focus", "AI Focus"),
    ("A_S", "A/S ratio"),
    ("PatentMismatch", "PatentMismatch"),
    ("car_m1_p1", "CAR[-1,+1]"),
    ("bhar_3m", "BHAR[+2,+63]"),
]
ALT_OUTCOMES = [("AI_Focus", "AI Focus"), ("PatentMismatch", "PatentMismatch")]
EVENT_STUDY_OUTCOMES = [("PatentMismatch", "PatentMismatch"), ("car_m1_p1", "CAR[-1,+1]")]
EVENT_TIMES = [-4, -3, -2, 0, 1, 2]
PRETREND_TIMES = [-4, -3, -2]
PRE_PERIOD_START = 2018
PRE_PERIOD_END = 2022
POST_START = 2023


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--event-panel", type=Path, default=DEFAULT_EVENT_PANEL)
    parser.add_argument("--annual-panel", type=Path, default=DEFAULT_ANNUAL_PANEL)
    parser.add_argument("--test-root", type=Path, default=DEFAULT_TEST_ROOT)
    parser.add_argument("--paper-root", type=Path, default=DEFAULT_PAPER_ROOT)
    parser.add_argument("--run-id", default=DEFAULT_RUN_ID)
    return parser.parse_args()


def _base_style() -> None:
    plt.style.use("seaborn-v0_8-whitegrid")
    plt.rcParams.update(
        {
            "font.family": "serif",
            "font.serif": ["Times New Roman", "Times", "DejaVu Serif"],
            "axes.titlesize": 13,
            "axes.labelsize": 11,
            "xtick.labelsize": 10,
            "ytick.labelsize": 10,
            "legend.fontsize": 9,
        }
    )


def _ensure_sic2(df: pd.DataFrame) -> pd.DataFrame:
    panel = df.copy()
    if "sic2" not in panel.columns or panel["sic2"].isna().all():
        if "sic" in panel.columns:
            sic_raw = pd.to_numeric(panel["sic"], errors="coerce")
            panel["sic2"] = (sic_raw // 100).astype("Int64")
        else:
            panel["sic2"] = pd.Series(pd.NA, index=panel.index, dtype="Int64")
    return panel


def _load_annual_panel(annual_panel: Path, event_panel: Path) -> pd.DataFrame:
    annual = pd.read_parquet(annual_panel).copy()
    annual = _ensure_sic2(annual)
    annual = _add_patent_mismatch(annual)
    annual["cik"] = annual["cik"].astype(str)
    annual["year"] = pd.to_numeric(annual["year"], errors="coerce").astype("Int64")

    event = pd.read_parquet(event_panel).copy()
    event["cik"] = event["cik"].astype(str)
    event["filing_year"] = pd.to_numeric(event["filing_year"], errors="coerce").astype("Int64")
    event_agg = (
        event.groupby(["cik", "filing_year"], as_index=False)
        .agg(
            car_m1_p1=("car_m1_p1", "mean"),
            bhar_3m=("bhar_3m", "mean"),
        )
        .rename(columns={"filing_year": "year"})
    )

    merged = annual.merge(event_agg, on=["cik", "year"], how="left", validate="one_to_one")
    merged = merged.sort_values(["cik", "year"]).reset_index(drop=True)
    for control in ["ln_assets", "leverage", "cash", "roa"]:
        merged[f"{control}_l1"] = merged.groupby("cik")[control].shift(1)
    merged["post_chatgpt"] = merged["year"].ge(POST_START).astype(int)
    return merged


def _build_treatment_panel(panel: pd.DataFrame) -> tuple[pd.DataFrame, dict[str, object]]:
    use = panel.copy()
    pre = use.loc[use["year"].between(PRE_PERIOD_START, PRE_PERIOD_END)].copy()
    pre["talk_year"] = pd.to_numeric(pre["any_ai_talk"], errors="coerce").fillna(0).astype(int)
    pre["talk_A_S"] = pre["A_S"].where(pre["talk_year"].eq(1))
    pre["sic2"] = pd.to_numeric(pre["sic2"], errors="coerce").astype("Int64")
    pre["log_mcap"] = np.log1p(pd.to_numeric(pre["market_cap_year_end"], errors="coerce"))

    agg = pre.groupby("cik", as_index=False).agg(
        gvkey=("gvkey", "last"),
        pre_obs=("year", "count"),
        pre_patents_sum=(
            "patents_ai",
            lambda s: pd.to_numeric(s, errors="coerce").fillna(0).sum(),
        ),
        pre_a_s_mean=("talk_A_S", "mean"),
        pre_actionable_sum=("n_A", lambda s: pd.to_numeric(s, errors="coerce").fillna(0).sum()),
        pre_spec_mean=(
            "SpecShare",
            lambda s: (
                pd.to_numeric(s, errors="coerce").where(pre.loc[s.index, "talk_year"].eq(1)).mean()
            ),
        ),
        pre_size_mean=("log_mcap", "mean"),
        pre_sic2=("sic2", lambda s: s.dropna().iloc[-1] if not s.dropna().empty else pd.NA),
    )
    agg["pre_sic2"] = pd.to_numeric(agg["pre_sic2"], errors="coerce").astype("Int64")

    as_cut = agg["pre_a_s_mean"].median(skipna=True)
    agg["low_pre_as"] = agg["pre_a_s_mean"].le(as_cut)
    agg["no_pre_actionable"] = agg["pre_actionable_sum"].fillna(0).eq(0)
    agg["T1_lowcap"] = (
        agg["pre_obs"].gt(0)
        & agg["pre_patents_sum"].fillna(0).eq(0)
        & (agg["low_pre_as"].fillna(False) | agg["no_pre_actionable"])
    )

    industry_median_spec = agg.groupby("pre_sic2")["pre_spec_mean"].transform("median")
    agg["T2_highspec"] = agg["pre_spec_mean"].gt(industry_median_spec)

    size_cut = agg["pre_size_mean"].median(skipna=True)
    agg["T3_smallpre"] = agg["pre_size_mean"].le(size_cut)

    treat_cols = ["T1_lowcap", "T2_highspec", "T3_smallpre"]
    for col in treat_cols:
        agg[col] = agg[col].fillna(False).astype(int)

    use = use.merge(
        agg[["cik", *treat_cols, "pre_obs"]], on="cik", how="left", validate="many_to_one"
    )
    use = use.loc[use["pre_obs"].fillna(0).gt(0)].copy()
    for col in treat_cols:
        use[col] = use[col].fillna(0).astype(int)

    summary = {
        "pre_period_window": f"{PRE_PERIOD_START}-{PRE_PERIOD_END}",
        "post_start_year": POST_START,
        "pre_obs_firms": int(agg["pre_obs"].gt(0).sum()),
        "t1_count": int(agg["T1_lowcap"].sum()),
        "t2_count": int(agg["T2_highspec"].sum()),
        "t3_count": int(agg["T3_smallpre"].sum()),
        "as_cut": float(as_cut) if math.isfinite(float(as_cut)) else math.nan,
        "size_cut_log_mcap": float(size_cut) if math.isfinite(float(size_cut)) else math.nan,
    }
    return use, summary


def _model_sample(panel: pd.DataFrame, outcome: str, treat_col: str) -> pd.DataFrame:
    needed = [outcome, treat_col, "post_chatgpt", "year", "cik", *CONTROL_TERMS]
    use = panel.dropna(subset=needed).copy()
    use = use.loc[use["cik"].astype(str).str.len().gt(0)].copy()
    use["year"] = pd.to_numeric(use["year"], errors="coerce").astype(int)
    use["post_chatgpt"] = pd.to_numeric(use["post_chatgpt"], errors="coerce").astype(int)
    use[treat_col] = pd.to_numeric(use[treat_col], errors="coerce").astype(int)
    use["treat_post"] = use[treat_col] * use["post_chatgpt"]
    for col in CONTROL_TERMS + [outcome]:
        use[col] = pd.to_numeric(use[col], errors="coerce").astype(float)
    return use


def _two_way_absorb(
    use: pd.DataFrame, columns: list[str], *, max_iter: int = 200, tol: float = 1e-10
) -> pd.DataFrame:
    entity = use["cik"]
    year = use["year"]
    absorbed = pd.DataFrame(index=use.index)
    for col in columns:
        z = pd.to_numeric(use[col], errors="coerce").astype(float).copy()
        for _ in range(max_iter):
            old = z.to_numpy(copy=True)
            z = z - z.groupby(entity).transform("mean")
            z = z - z.groupby(year).transform("mean")
            delta = float(np.nanmax(np.abs(z.to_numpy() - old))) if len(z) else 0.0
            if delta < tol:
                break
        absorbed[col] = z
    return absorbed


def _fit_absorbed_ols(
    use: pd.DataFrame, dependent: str, rhs: list[str]
) -> sm.regression.linear_model.RegressionResultsWrapper:
    absorbed = _two_way_absorb(use, [dependent, *rhs])
    y = absorbed[dependent]
    X = absorbed[rhs]
    return sm.OLS(y, X).fit(cov_type="cluster", cov_kwds={"groups": use["cik"].astype(str)})


def _fit_did(panel: pd.DataFrame, outcome: str, treat_col: str) -> dict[str, object]:
    use = _model_sample(panel, outcome, treat_col)
    rhs = ["treat_post", *CONTROL_TERMS]
    result = _fit_absorbed_ols(use, outcome, rhs)
    coef = float(result.params.get("treat_post", math.nan))
    se = float(result.bse.get("treat_post", math.nan))
    pval = float(result.pvalues.get("treat_post", math.nan))
    return {
        "outcome": outcome,
        "treatment": treat_col,
        "coef": coef,
        "se": se,
        "p_value": pval,
        "nobs": int(result.nobs),
        "adj_r_squared": float(result.rsquared_adj),
        "formula": f"{outcome} ~ treat_post + {' + '.join(CONTROL_TERMS)} + firm FE + year FE",
    }


def _fit_event_study(panel: pd.DataFrame, outcome: str, treat_col: str) -> dict[str, object]:
    use = _model_sample(panel, outcome, treat_col).copy()
    use = use.loc[use["year"].between(2019, 2025)].copy()
    use["event_time"] = use["year"] - POST_START
    use = use.loc[use["event_time"].isin(EVENT_TIMES + [-1])].copy()

    terms = []
    for k in EVENT_TIMES:
        col = f"event_treat_{k}".replace("-", "m")
        use[col] = use[treat_col] * use["event_time"].eq(k).astype(int)
        terms.append(col)

    rhs = [*terms, *CONTROL_TERMS]
    result = _fit_absorbed_ols(use, outcome, rhs)

    rows = []
    for k in EVENT_TIMES:
        col = f"event_treat_{k}".replace("-", "m")
        coef = float(result.params.get(col, math.nan))
        se = float(result.bse.get(col, math.nan))
        pval = float(result.pvalues.get(col, math.nan))
        rows.append(
            {
                "event_time": k,
                "coef": coef,
                "se": se,
                "p_value": pval,
                "ci_low": coef - 1.96 * se,
                "ci_high": coef + 1.96 * se,
            }
        )

    test_terms = [f"event_treat_{k}".replace("-", "m") for k in PRETREND_TIMES]
    pre_beta = result.params[test_terms].to_numpy(dtype=float)
    pre_cov = result.cov_params().loc[test_terms, test_terms].to_numpy(dtype=float)
    stat = float(pre_beta.T @ np.linalg.pinv(pre_cov) @ pre_beta)
    pretrend_p = float(1.0 - chi2.cdf(stat, df=len(test_terms)))
    return {
        "outcome": outcome,
        "treatment": treat_col,
        "formula": f"{outcome} ~ event_time×treated + {' + '.join(CONTROL_TERMS)} + firm FE + year FE",
        "nobs": int(result.nobs),
        "pretrend_p_value": pretrend_p,
        "coefficients": rows,
    }


def _stars(p_value: float | None) -> str:
    if p_value is None or not math.isfinite(p_value):
        return ""
    if p_value < 0.01:
        return "***"
    if p_value < 0.05:
        return "**"
    if p_value < 0.10:
        return "*"
    return ""


def _coef_cell(coef: float, se: float, p_value: float) -> str:
    if not (math.isfinite(coef) and math.isfinite(se)):
        return ""
    return f"{coef:.4f}{_stars(p_value)} ({se:.4f})"


def _build_table_payload(
    main_results: list[dict[str, object]], alt_results: list[dict[str, object]]
) -> tuple[pd.DataFrame, dict[str, dict[str, object]], dict[tuple[str, str], dict[str, object]]]:
    main_map = {row["outcome"]: row for row in main_results}
    alt_map = {(row["treatment"], row["outcome"]): row for row in alt_results}

    rows: list[dict[str, object]] = []
    for label, title in MAIN_OUTCOMES:
        row = main_map[label]
        rows.append(
            {
                "panel": "Panel A. Main treatment T1 across outcomes",
                "row_label": title,
                "col1": _coef_cell(row["coef"], row["se"], row["p_value"]),
                "col2": str(row["nobs"]),
                "col3": f"{row['adj_r_squared']:.3f}",
            }
        )
    rows.extend(
        [
            {
                "panel": "Panel A. Main treatment T1 across outcomes",
                "row_label": "Controls / firm FE / year FE",
                "col1": "Y",
                "col2": "Y",
                "col3": "Y",
            }
        ]
    )

    for treat_col, treat_label in [
        ("T1_lowcap", "T1 LowCapability"),
        ("T2_highspec", "T2 HighSpecPre"),
        ("T3_smallpre", "T3 SmallPre"),
    ]:
        ai_row = alt_map[(treat_col, "AI_Focus")]
        mismatch_row = alt_map[(treat_col, "PatentMismatch")]
        rows.append(
            {
                "panel": "Panel B. Alternative treatments on core annual outcomes",
                "row_label": treat_label,
                "col1": _coef_cell(ai_row["coef"], ai_row["se"], ai_row["p_value"]),
                "col2": _coef_cell(
                    mismatch_row["coef"], mismatch_row["se"], mismatch_row["p_value"]
                ),
                "col3": f"{ai_row['nobs']} / {mismatch_row['nobs']}",
            }
        )

    return pd.DataFrame(rows), main_map, alt_map


def _markdown_table(headers: list[str], rows: list[list[object]]) -> str:
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join(["---"] * len(headers)) + " |",
    ]
    for row in rows:
        lines.append("| " + " | ".join(str(item) for item in row) + " |")
    return "\n".join(lines)


def _render_table_outputs(table_df: pd.DataFrame) -> tuple[str, str]:
    md_lines = ["# Table Main", ""]
    latex_lines = [
        "\\begin{table}[!htbp]",
        "\\centering",
        "\\caption{Difference-in-differences around the ChatGPT release}",
        "\\begin{tabular}{lccc}",
        "\\hline",
    ]
    panel_headers = {
        "Panel A. Main treatment T1 across outcomes": [
            "Outcome",
            "T1 DID coef.",
            "N",
            "Adj. R-squared",
        ],
        "Panel B. Alternative treatments on core annual outcomes": [
            "Treatment",
            "AI Focus DID",
            "PatentMismatch DID",
            "N (AI / mismatch)",
        ],
    }
    for panel_name, headers in panel_headers.items():
        panel_rows = table_df.loc[
            table_df["panel"].eq(panel_name), ["row_label", "col1", "col2", "col3"]
        ].values.tolist()
        md_lines.extend([f"## {panel_name}", _markdown_table(headers, panel_rows), ""])
        latex_lines.append(f"\\multicolumn{{4}}{{l}}{{\\textit{{{panel_name}}}}} \\\\")
        latex_lines.append(" & ".join(headers) + " \\\\")
        for row in panel_rows:
            latex_lines.append(" & ".join(str(item) for item in row) + " \\\\")
    latex_lines.extend(["\\hline", "\\end{tabular}", "\\end{table}"])
    return "\n".join(md_lines), "\n".join(latex_lines) + "\n"


def _docx_panel_rows(table_df: pd.DataFrame, panel_name: str) -> list[list[tuple[str, bool]]]:
    rows = []
    for row in table_df.loc[
        table_df["panel"].eq(panel_name), ["row_label", "col1", "col2", "col3"]
    ].itertuples(index=False):
        rows.append(
            [(str(row[0]), True), (str(row[1]), False), (str(row[2]), False), (str(row[3]), False)]
        )
    return rows


def _build_table_docx(table_df: pd.DataFrame, output_path: Path) -> None:
    document = Document()
    _set_document_defaults(document)
    _set_landscape(document)
    _add_title(document, "Table 7. ChatGPT DID Around AI Disclosure and Mismatch")
    note = (
        "This table reports difference-in-differences estimates around the release of ChatGPT. PostChatGPT equals one for fiscal years 2023 onward. The main treatment identifies firms with zero pre-2023 AI patents and weak pre-period disclosure credibility, defined as a low pre-period A/S ratio or no pre-period actionable disclosure. "
        "Panel A reports the main treatment DID coefficient across annual disclosure outcomes and filing-based market outcomes. Panel B reports alternative treatment definitions on the two core annual outcomes. All specifications include lagged controls, firm fixed effects, year fixed effects, and firm-clustered standard errors."
    )
    _add_note(document, note)

    panel_a_headers = ["Outcome", "T1 DID coef.", "N", "Adj. R-squared"]
    p = document.add_paragraph()
    p.add_run("Panel A. Main treatment T1 across outcomes").bold = True
    from semantic_ai_washing.analysis.publication_runs.test_03_post_filing_drift import (
        _build_panel_table,
    )

    _build_panel_table(
        document,
        panel_a_headers,
        _docx_panel_rows(table_df, "Panel A. Main treatment T1 across outcomes"),
    )

    panel_b_headers = ["Treatment", "AI Focus DID", "PatentMismatch DID", "N (AI / mismatch)"]
    p = document.add_paragraph()
    p.add_run("Panel B. Alternative treatments on core annual outcomes").bold = True
    _build_panel_table(
        document,
        panel_b_headers,
        _docx_panel_rows(table_df, "Panel B. Alternative treatments on core annual outcomes"),
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def _plot_event_study(
    event_results: list[dict[str, object]], output_base: Path
) -> tuple[Path, Path]:
    _base_style()
    fig, axes = plt.subplots(1, 2, figsize=(10.0, 4.6), constrained_layout=True, sharex=True)
    color_map = {"PatentMismatch": "#c46b48", "car_m1_p1": "#1d3557"}
    title_map = {"PatentMismatch": "PatentMismatch", "car_m1_p1": "CAR[-1,+1]"}

    for ax, outcome in zip(axes, ["PatentMismatch", "car_m1_p1"], strict=False):
        result = next(row for row in event_results if row["outcome"] == outcome)
        coeffs = pd.DataFrame(result["coefficients"])
        ax.errorbar(
            coeffs["event_time"],
            coeffs["coef"],
            yerr=1.96 * coeffs["se"],
            color=color_map[outcome],
            marker="o",
            linewidth=1.9,
            capsize=3,
        )
        ax.axhline(0, color="#6c757d", linewidth=0.8)
        ax.axvline(-1, color="#6c757d", linestyle="--", linewidth=0.8)
        ax.set_title(f"{title_map[outcome]} (pretrend p={result['pretrend_p_value']:.3f})")
        ax.set_xlabel("Event time (years; -1 = 2022 omitted)")
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)
        ax.grid(True, axis="y", color="#d9d9d9", linewidth=0.7)
        ax.grid(False, axis="x")
    axes[0].set_ylabel("DID event-time coefficient")
    fig.suptitle("Event-Study Path Around ChatGPT Release for Main Treatment", fontsize=14)

    png_path = output_base.with_suffix(".png")
    pdf_path = output_base.with_suffix(".pdf")
    fig.savefig(png_path, dpi=220, bbox_inches="tight")
    fig.savefig(pdf_path, bbox_inches="tight")
    plt.close(fig)
    return png_path, pdf_path


def _result_notes(
    main_map: dict[str, dict[str, object]],
    alt_map: dict[tuple[str, str], dict[str, object]],
    event_results: list[dict[str, object]],
    treat_summary: dict[str, object],
) -> str:
    mismatch = main_map["PatentMismatch"]
    car = main_map["car_m1_p1"]
    bhar = main_map["bhar_3m"]
    event_patent = next(row for row in event_results if row["outcome"] == "PatentMismatch")
    event_car = next(row for row in event_results if row["outcome"] == "car_m1_p1")
    return "\n".join(
        [
            "# Result Notes",
            "",
            f"- Main treatment counts: T1=`{treat_summary['t1_count']}`, T2=`{treat_summary['t2_count']}`, T3=`{treat_summary['t3_count']}` firms with pre-period observations.",
            f"- T1 DID on AI Focus: `{main_map['AI_Focus']['coef']:.4f}` with p=`{main_map['AI_Focus']['p_value']:.3f}`; on PatentMismatch: `{mismatch['coef']:.4f}` with p=`{mismatch['p_value']:.3f}`.",
            f"- T1 DID on filing outcomes: CAR=`{car['coef']:.4f}` (p=`{car['p_value']:.3f}`), BHAR[+2,+63]=`{bhar['coef']:.4f}` (p=`{bhar['p_value']:.3f}`).",
            f"- Event-study pretrend p-values for T1: PatentMismatch=`{event_patent['pretrend_p_value']:.3f}`, CAR=`{event_car['pretrend_p_value']:.3f}`.",
            f"- Alternative-treatment PatentMismatch DID: T2=`{alt_map[('T2_highspec', 'PatentMismatch')]['coef']:.4f}` (p=`{alt_map[('T2_highspec', 'PatentMismatch')]['p_value']:.3f}`), T3=`{alt_map[('T3_smallpre', 'PatentMismatch')]['coef']:.4f}` (p=`{alt_map[('T3_smallpre', 'PatentMismatch')]['p_value']:.3f}`).",
            "",
        ]
    )


def _writer_packet(
    args: argparse.Namespace,
    treat_summary: dict[str, object],
    main_map: dict[str, dict[str, object]],
    alt_map: dict[tuple[str, str], dict[str, object]],
    event_results: list[dict[str, object]],
) -> str:
    car_event = next(row for row in event_results if row["outcome"] == "car_m1_p1")
    mismatch_event = next(row for row in event_results if row["outcome"] == "PatentMismatch")
    return "\n".join(
        [
            "# Writer Packet",
            "",
            "## Metadata",
            f"- Test id: `{TEST_ID}`",
            f"- Run id: `{args.run_id}`",
            f"- Date run: `{date.today().isoformat()}`",
            f"- Script/module path: `{MODULE_PATH}`",
            f"- Input files: `{args.annual_panel}`, `{args.event_panel}`",
            "",
            "## Event Rule",
            f"- PostChatGPT = 1 for years `>= {POST_START}`",
            f"- Pre-period used for treatment construction: `{PRE_PERIOD_START}-{PRE_PERIOD_END}`",
            "",
            "## Treatment Definitions",
            "- T1 LowCapability: `zero pre-period AI patents AND (low pre-period A/S ratio OR no pre-period actionable disclosure)`",
            "- T2 HighSpecPre: `pre-period SpecShare above industry median`",
            "- T3 SmallPre: `pre-period average log market cap below sample median`",
            f"- Treated-firm counts: `T1={treat_summary['t1_count']}`, `T2={treat_summary['t2_count']}`, `T3={treat_summary['t3_count']}`",
            "",
            "## Main DID Results",
            f"- AI Focus: `{main_map['AI_Focus']['coef']:.4f}` (p=`{main_map['AI_Focus']['p_value']:.3f}`)",
            f"- A/S ratio: `{main_map['A_S']['coef']:.4f}` (p=`{main_map['A_S']['p_value']:.3f}`)",
            f"- PatentMismatch: `{main_map['PatentMismatch']['coef']:.4f}` (p=`{main_map['PatentMismatch']['p_value']:.3f}`)",
            f"- CAR[-1,+1]: `{main_map['car_m1_p1']['coef']:.4f}` (p=`{main_map['car_m1_p1']['p_value']:.3f}`)",
            f"- BHAR[+2,+63]: `{main_map['bhar_3m']['coef']:.4f}` (p=`{main_map['bhar_3m']['p_value']:.3f}`)",
            "",
            "## Event-Study Pretrends",
            f"- PatentMismatch pretrend p-value: `{mismatch_event['pretrend_p_value']:.3f}`",
            f"- CAR[-1,+1] pretrend p-value: `{car_event['pretrend_p_value']:.3f}`",
            "",
            "## Alternative Treatments",
            f"- T2 on AI Focus: `{alt_map[('T2_highspec', 'AI_Focus')]['coef']:.4f}` (p=`{alt_map[('T2_highspec', 'AI_Focus')]['p_value']:.3f}`)",
            f"- T2 on PatentMismatch: `{alt_map[('T2_highspec', 'PatentMismatch')]['coef']:.4f}` (p=`{alt_map[('T2_highspec', 'PatentMismatch')]['p_value']:.3f}`)",
            f"- T3 on AI Focus: `{alt_map[('T3_smallpre', 'AI_Focus')]['coef']:.4f}` (p=`{alt_map[('T3_smallpre', 'AI_Focus')]['p_value']:.3f}`)",
            f"- T3 on PatentMismatch: `{alt_map[('T3_smallpre', 'PatentMismatch')]['coef']:.4f}` (p=`{alt_map[('T3_smallpre', 'PatentMismatch')]['p_value']:.3f}`)",
            "",
            "## Caption Draft",
            "This table reports difference-in-differences estimates around the release of ChatGPT. Treated firms are defined using pre-period AI capability and rhetorical-tilt measures. The main treatment identifies firms with zero pre-2023 AI patents and weak pre-period disclosure credibility. The table tests whether the post-ChatGPT period increased AI disclosure, mismatch, and related market outcomes disproportionately in firms with weaker pre-period capability.",
            "",
        ]
    )


def _dataset_summary(
    args: argparse.Namespace,
    treat_summary: dict[str, object],
    main_results: list[dict[str, object]],
    alt_results: list[dict[str, object]],
    event_results: list[dict[str, object]],
) -> dict[str, object]:
    return {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "annual_panel": str(args.annual_panel),
        "event_panel": str(args.event_panel),
        "treatment_summary": treat_summary,
        "main_results": main_results,
        "alt_results": alt_results,
        "event_results": event_results,
    }


def _copy_exports(run_dir: Path, paper_root: Path, run_id: str) -> dict[str, str]:
    exports = {
        "figure_png": paper_root / "figures" / f"{TEST_ID}_{run_id}.png",
        "figure_pdf": paper_root / "figures" / f"{TEST_ID}_{run_id}.pdf",
        "table_csv": paper_root / "tables" / f"{TEST_ID}_{run_id}.csv",
        "table_md": paper_root / "tables" / f"{TEST_ID}_{run_id}.md",
        "table_tex": paper_root / "latex" / f"{TEST_ID}_{run_id}.tex",
        "table_docx": paper_root / "docx" / f"{TEST_ID}_{run_id}.docx",
        "figure_series": paper_root / "tables" / f"{TEST_ID}_{run_id}_figure_series.csv",
        "writer_packet": paper_root / "writer_packets" / f"{TEST_ID}_{run_id}.md",
        "result_notes": paper_root / "snippets" / f"{TEST_ID}_{run_id}_result_notes.md",
    }
    for path in exports.values():
        path.parent.mkdir(parents=True, exist_ok=True)
    mapping = {
        run_dir / "figure_main.png": exports["figure_png"],
        run_dir / "figure_main.pdf": exports["figure_pdf"],
        run_dir / "table_main.csv": exports["table_csv"],
        run_dir / "table_main.md": exports["table_md"],
        run_dir / "table_main.tex": exports["table_tex"],
        run_dir / "table_main.docx": exports["table_docx"],
        run_dir / "figure_series.csv": exports["figure_series"],
        run_dir / "writer_packet.md": exports["writer_packet"],
        run_dir / "result_notes.md": exports["result_notes"],
    }
    for src, dst in mapping.items():
        shutil.copy2(src, dst)
    return {key: str(path) for key, path in exports.items()}


def main() -> None:
    args = _parse_args()
    run_dir = args.test_root / args.run_id
    run_dir.mkdir(parents=True, exist_ok=True)

    panel = _load_annual_panel(args.annual_panel, args.event_panel)
    panel, treat_summary = _build_treatment_panel(panel)

    main_results = [_fit_did(panel, outcome, "T1_lowcap") for outcome, _ in MAIN_OUTCOMES]
    alt_results = []
    for treat_col in ["T1_lowcap", "T2_highspec", "T3_smallpre"]:
        for outcome, _ in ALT_OUTCOMES:
            alt_results.append(_fit_did(panel, outcome, treat_col))

    event_results = [
        _fit_event_study(panel, outcome, "T1_lowcap") for outcome, _ in EVENT_STUDY_OUTCOMES
    ]

    table_df, main_map, alt_map = _build_table_payload(main_results, alt_results)
    table_df.to_csv(run_dir / "table_main.csv", index=False)

    figure_rows = []
    for result in event_results:
        for coef_row in result["coefficients"]:
            figure_rows.append(
                {
                    "outcome": result["outcome"],
                    "pretrend_p_value": result["pretrend_p_value"],
                    **coef_row,
                }
            )
    figure_df = pd.DataFrame(figure_rows)
    figure_df.to_csv(run_dir / "figure_series.csv", index=False)

    table_md, table_tex = _render_table_outputs(table_df)
    (run_dir / "table_main.md").write_text(table_md, encoding="utf-8")
    (run_dir / "table_main.tex").write_text(table_tex, encoding="utf-8")
    _build_table_docx(table_df, run_dir / "table_main.docx")
    png_path, pdf_path = _plot_event_study(event_results, run_dir / "figure_main")

    (run_dir / "result_notes.md").write_text(
        _result_notes(main_map, alt_map, event_results, treat_summary), encoding="utf-8"
    )
    (run_dir / "writer_packet.md").write_text(
        _writer_packet(args, treat_summary, main_map, alt_map, event_results), encoding="utf-8"
    )
    summary = _dataset_summary(args, treat_summary, main_results, alt_results, event_results)
    (run_dir / "dataset_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")

    paper_exports = _copy_exports(run_dir, args.paper_root, args.run_id)
    manifest = {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "module_path": MODULE_PATH,
        "inputs": {
            "annual_panel": str(args.annual_panel),
            "event_panel": str(args.event_panel),
        },
        "run_dir": str(run_dir),
        "outputs": {
            "dataset_summary": str(run_dir / "dataset_summary.json"),
            "table_csv": str(run_dir / "table_main.csv"),
            "table_md": str(run_dir / "table_main.md"),
            "table_tex": str(run_dir / "table_main.tex"),
            "table_docx": str(run_dir / "table_main.docx"),
            "figure_series": str(run_dir / "figure_series.csv"),
            "figure_png": str(png_path),
            "figure_pdf": str(pdf_path),
            "writer_packet": str(run_dir / "writer_packet.md"),
            "result_notes": str(run_dir / "result_notes.md"),
        },
        "paper_exports": paper_exports,
    }
    (run_dir / "run_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    print(f"[{TEST_ID}] wrote run bundle to {run_dir}")
    print(f"[{TEST_ID}] paper table: {paper_exports['table_docx']}")


if __name__ == "__main__":
    main()
