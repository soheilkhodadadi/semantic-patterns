"""Publication run driver for Test 09: factor-adjusted post-filing portfolio alpha."""

from __future__ import annotations

import argparse
import json
import math
import shutil
from datetime import UTC, date, datetime
from pathlib import Path

from docx import Document
import pandas as pd
import statsmodels.api as sm

from semantic_ai_washing.analysis.delivery_table_payloads import _add_patent_mismatch
from semantic_ai_washing.analysis.ken_french_factors import (
    FF5_2X3_MONTHLY_URL,
    MOM_MONTHLY_URL,
    stage_ken_french_monthly_factors,
)
from semantic_ai_washing.analysis.publication_runs.test_03_post_filing_drift import (
    _add_note,
    _add_title,
    _build_panel_table,
    _set_document_defaults,
    _set_landscape,
)
from semantic_ai_washing.analysis.publication_runs.test_03_post_filing_portfolio_alpha import (
    HORIZONS,
    _build_positions,
    _compute_long_short_series,
    _load_market_index,
    _load_monthly_returns,
    _plot_cumulative_long_short,
)

REPO_ROOT = Path(__file__).resolve().parents[4]
DEFAULT_EVENT_PANEL = (
    REPO_ROOT
    / "data/processed/panel/filing_event_estimation_sample_hybrid_api_a_conf49_v1.parquet"
)
DEFAULT_ANNUAL_PANEL = (
    REPO_ROOT
    / "data/processed/panel/canonical/ever_speaker_panel_2016_2025_hybrid_api_a_conf49_v1.parquet"
)
DEFAULT_MONTHLY_RETURNS = REPO_ROOT / "data/interim/market/wrds_crsp_msf_full_sample_v1.parquet"
DEFAULT_MARKET_INDEX = REPO_ROOT / "data/interim/market/wrds_crsp_msi_full_sample_v1.parquet"
DEFAULT_FACTOR_ROOT = Path(
    "/Users/soheilkhodadadi/DataWork/semantic-patterns/ai_washing/derived/v3_1/factor_inputs"
)
DEFAULT_TEST_ROOT = Path(
    "/Users/soheilkhodadadi/DataWork/semantic-patterns/ai_washing/derived/v3_1/test_runs/test_09_factor_adjusted_alpha"
)
DEFAULT_PAPER_ROOT = REPO_ROOT / "paper/generated/v3_1"
DEFAULT_RUN_ID = f"{date.today():%Y%m%d}_aiw_v3_1_test_09_factor_adjusted_alpha_main_v1"
TEST_ID = "test_09_factor_adjusted_alpha"
MODULE_PATH = "semantic_ai_washing.analysis.publication_runs.test_09_factor_adjusted_alpha"
FACTOR_MODELS: list[tuple[str, str, list[str]]] = [
    ("capm", "CAPM", ["mkt_rf"]),
    ("ff3", "FF3", ["mkt_rf", "smb", "hml"]),
    ("ff5", "FF5", ["mkt_rf", "smb", "hml", "rmw", "cma"]),
    ("ff5_mom", "FF5+Mom", ["mkt_rf", "smb", "hml", "rmw", "cma", "mom"]),
]


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--event-panel", type=Path, default=DEFAULT_EVENT_PANEL)
    parser.add_argument("--annual-panel", type=Path, default=DEFAULT_ANNUAL_PANEL)
    parser.add_argument("--monthly-returns", type=Path, default=DEFAULT_MONTHLY_RETURNS)
    parser.add_argument("--market-index", type=Path, default=DEFAULT_MARKET_INDEX)
    parser.add_argument("--factor-root", type=Path, default=DEFAULT_FACTOR_ROOT)
    parser.add_argument("--test-root", type=Path, default=DEFAULT_TEST_ROOT)
    parser.add_argument("--paper-root", type=Path, default=DEFAULT_PAPER_ROOT)
    parser.add_argument("--run-id", default=DEFAULT_RUN_ID)
    parser.add_argument(
        "--size-subset",
        choices=["all", "nonbig", "big"],
        default="all",
        help="Optional yearly-median market-cap split based on annual-panel market_cap_year_end.",
    )
    parser.add_argument(
        "--refresh-factors",
        action="store_true",
        help="Redownload official Ken French factor files even if they already exist locally.",
    )
    return parser.parse_args()


def _normalize_id(series: pd.Series) -> pd.Series:
    return (
        series.astype(str)
        .str.replace(r"\.0$", "", regex=True)
        .str.strip()
        .replace({"nan": "", "None": ""})
    )


def _load_annual_signal_panel(path: Path) -> pd.DataFrame:
    panel = pd.read_parquet(path).copy()
    if "sic2" not in panel.columns or panel["sic2"].isna().all():
        if "sic" in panel.columns:
            sic_raw = pd.to_numeric(panel["sic"], errors="coerce")
            panel["sic2"] = (sic_raw // 100).astype("Int64")
        else:
            panel["sic2"] = pd.Series(pd.NA, index=panel.index, dtype="Int64")
    panel = _add_patent_mismatch(panel)
    keep = [
        "cik",
        "year",
        "permno",
        "PatentMismatch",
        "A_S",
        "AI_Focus",
        "CredAI",
        "SpecShare",
        "market_cap_year_end",
    ]
    available = [column for column in keep if column in panel.columns]
    out = panel[available].copy()
    out["cik"] = _normalize_id(out["cik"])
    out["year"] = pd.to_numeric(out["year"], errors="coerce").astype("Int64")
    if "permno" in out.columns:
        out["permno"] = _normalize_id(out["permno"])
    if "market_cap_year_end" in out.columns:
        out["market_cap_year_end"] = pd.to_numeric(
            out["market_cap_year_end"], errors="coerce"
        )
    out["PatentMismatch"] = pd.to_numeric(out["PatentMismatch"], errors="coerce")
    for column in ["A_S", "AI_Focus", "CredAI", "SpecShare"]:
        if column in out.columns:
            out[column] = pd.to_numeric(out[column], errors="coerce")
    return out


def _prepare_signal_sample(
    event_panel: Path,
    annual_panel: Path,
    *,
    size_subset: str,
) -> pd.DataFrame:
    event = pd.read_parquet(event_panel).copy()
    event["cik"] = _normalize_id(event["cik"])
    event["permno"] = _normalize_id(event["permno"])
    event["filing_year"] = pd.to_numeric(event["filing_year"], errors="coerce").astype("Int64")
    event["n_ai_total"] = pd.to_numeric(event["n_ai_total"], errors="coerce")

    annual = _load_annual_signal_panel(annual_panel)
    merged = event.merge(
        annual,
        left_on=["cik", "filing_year"],
        right_on=["cik", "year"],
        how="left",
        validate="many_to_one",
        suffixes=("", "_annual"),
    )
    merged["is_ai_filing"] = merged["n_ai_total"].fillna(0).gt(0)
    sample = merged.loc[merged["is_ai_filing"]].copy()
    sample = sample.loc[sample["PatentMismatch"].notna()].copy()
    sample["PatentMismatch"] = sample["PatentMismatch"].fillna(0).astype(int)
    sample = sample.loc[sample["permno"] != ""].copy()
    sample["filing_date"] = pd.to_datetime(sample["filing_date"])
    sample["filing_month"] = sample["filing_date"].dt.to_period("M")

    if size_subset != "all" and "market_cap_year_end" in sample.columns:
        medians = sample.groupby("filing_year")["market_cap_year_end"].median()
        sample["year_median_market_cap"] = sample["filing_year"].map(medians)
        sample["is_nonbig"] = sample["market_cap_year_end"] <= sample["year_median_market_cap"]
        if size_subset == "nonbig":
            sample = sample.loc[sample["is_nonbig"].fillna(False)].copy()
        elif size_subset == "big":
            sample = sample.loc[sample["is_nonbig"].fillna(False).eq(False)].copy()
    return sample.reset_index(drop=True)


def _fit_factor_model(series_df: pd.DataFrame, return_col: str, factor_cols: list[str]) -> dict[str, float]:
    use = series_df.dropna(subset=[return_col, *factor_cols]).copy()
    if use.empty:
        return {
            "alpha": math.nan,
            "alpha_t": math.nan,
            "alpha_p": math.nan,
            "adj_r2": math.nan,
            "n_months": 0,
        }
    X = sm.add_constant(use[factor_cols])
    result = sm.OLS(use[return_col], X).fit(cov_type="HAC", cov_kwds={"maxlags": 3})
    return {
        "alpha": float(result.params.get("const", math.nan)),
        "alpha_t": float(result.tvalues.get("const", math.nan)),
        "alpha_p": float(result.pvalues.get("const", math.nan)),
        "adj_r2": float(result.rsquared_adj),
        "n_months": int(len(use)),
    }


def _summarize_horizon(series_df: pd.DataFrame, weight_col: str, *, horizon_label: str) -> dict[str, object]:
    use = series_df.dropna(subset=[weight_col]).copy()
    mean_monthly = float(use[weight_col].mean()) if not use.empty else math.nan
    annualized_simple = float(mean_monthly * 12.0) if not math.isnan(mean_monthly) else math.nan
    row: dict[str, object] = {
        "weighting": weight_col[:2],
        "horizon": horizon_label,
        "mean_monthly_return": mean_monthly,
        "annualized_simple_return": annualized_simple,
    }
    for model_key, _, factor_cols in FACTOR_MODELS:
        fit = _fit_factor_model(series_df, weight_col, factor_cols)
        row[f"{model_key}_alpha"] = fit["alpha"]
        row[f"{model_key}_alpha_t"] = fit["alpha_t"]
        row[f"{model_key}_alpha_p"] = fit["alpha_p"]
        row[f"{model_key}_adj_r2"] = fit["adj_r2"]
        row[f"{model_key}_n_months"] = fit["n_months"]
    row["n_months"] = int(use[weight_col].notna().sum())
    return row


def _fmt_pct(value: float) -> str:
    if pd.isna(value):
        return ""
    return f"{100 * value:.3f}"


def _fmt_num(value: float) -> str:
    if pd.isna(value):
        return ""
    return f"{value:.3f}"


def _alpha_cell(row: dict[str, object], model_key: str) -> str:
    alpha = row.get(f"{model_key}_alpha", math.nan)
    p_value = row.get(f"{model_key}_alpha_p", math.nan)
    if pd.isna(alpha):
        return ""
    return f"{100 * float(alpha):.3f} [{float(p_value):.3f}]"


def _render_table_outputs(summary_rows: list[dict[str, object]]) -> tuple[str, str]:
    headers = [
        "Horizon",
        "Mean monthly",
        "Annualized simple",
        "CAPM alpha [p]",
        "FF3 alpha [p]",
        "FF5 alpha [p]",
        "FF5+Mom alpha [p]",
        "Months",
    ]
    by_key = {(row["weighting"], row["horizon"]): row for row in summary_rows}

    def panel_rows(weighting: str) -> list[list[str]]:
        rows: list[list[str]] = []
        for horizon in HORIZONS.values():
            row = by_key[(weighting, horizon)]
            rows.append(
                [
                    horizon,
                    _fmt_pct(float(row["mean_monthly_return"])),
                    _fmt_pct(float(row["annualized_simple_return"])),
                    _alpha_cell(row, "capm"),
                    _alpha_cell(row, "ff3"),
                    _alpha_cell(row, "ff5"),
                    _alpha_cell(row, "ff5_mom"),
                    str(int(row["n_months"])),
                ]
            )
        return rows

    panel_a = panel_rows("ew")
    panel_b = panel_rows("vw")

    def md_table(rows: list[list[str]]) -> str:
        lines = [
            "| " + " | ".join(headers) + " |",
            "| " + " | ".join(["---"] * len(headers)) + " |",
        ]
        lines.extend("| " + " | ".join(row) + " |" for row in rows)
        return "\n".join(lines)

    md = "\n".join(
        [
            "# Table Main",
            "",
            "## Panel A. Equal-weight long credible / short mismatch portfolios",
            md_table(panel_a),
            "",
            "## Panel B. Value-weight long credible / short mismatch portfolios",
            md_table(panel_b),
            "",
        ]
    )

    latex_lines = [
        "\\begin{table}[!htbp]",
        "\\centering",
        "\\caption{Factor-adjusted calendar-time long-short alpha after AI filing signals}",
        "\\begin{tabular}{l" + "c" * (len(headers) - 1) + "}",
        "\\hline",
        "Horizon & Mean monthly & Annualized simple & CAPM alpha [p] & FF3 alpha [p] & FF5 alpha [p] & FF5+Mom alpha [p] & Months \\",
        "\\hline",
        "\\multicolumn{8}{l}{\\textit{Panel A. Equal-weight long credible / short mismatch portfolios}} \\",
    ]
    for row in panel_a:
        latex_lines.append(" & ".join(row) + " \\")
    latex_lines.extend([
        "\\hline",
        "\\multicolumn{8}{l}{\\textit{Panel B. Value-weight long credible / short mismatch portfolios}} \\",
    ])
    for row in panel_b:
        latex_lines.append(" & ".join(row) + " \\")
    latex_lines.extend(["\\hline", "\\end{tabular}", "\\end{table}"])
    return md, "\n".join(latex_lines) + "\n"


def _build_table_docx(summary_rows: list[dict[str, object]], output_path: Path) -> None:
    document = Document()
    _set_document_defaults(document)
    _set_landscape(document)
    _add_title(document, "Table 9. Factor-Adjusted Calendar-Time Long-Short Alpha")
    note = (
        "This table reports factor-adjusted calendar-time long-short portfolio returns formed after annual AI-related filing signals. Each month, the long leg holds non-mismatch AI filers and the short leg holds mismatch AI filers that filed within the prior 1, 3, 6, or 12 months. "
        "Panel A reports equal-weight portfolios and Panel B reports value-weight portfolios using lagged market capitalization. Alpha is the intercept from monthly time-series regressions estimated with Newey-West (lag 3) standard errors. Factors come from the official Ken French Data Library monthly Fama/French 5 Factors (2x3) file and the monthly Momentum Factor file."
    )
    _add_note(document, note)
    headers = [
        "",
        "Mean monthly",
        "Annualized simple",
        "CAPM alpha [p]",
        "FF3 alpha [p]",
        "FF5 alpha [p]",
        "FF5+Mom alpha [p]",
        "Months",
    ]
    by_key = {(row["weighting"], row["horizon"]): row for row in summary_rows}

    def panel_rows(weighting: str) -> list[list[tuple[str, bool]]]:
        rows: list[list[tuple[str, bool]]] = []
        for horizon in HORIZONS.values():
            row = by_key[(weighting, horizon)]
            rows.append(
                [
                    (horizon, True),
                    (_fmt_pct(float(row["mean_monthly_return"])), False),
                    (_fmt_pct(float(row["annualized_simple_return"])), False),
                    (_alpha_cell(row, "capm"), False),
                    (_alpha_cell(row, "ff3"), False),
                    (_alpha_cell(row, "ff5"), False),
                    (_alpha_cell(row, "ff5_mom"), False),
                    (str(int(row["n_months"])), False),
                ]
            )
        return rows

    p = document.add_paragraph()
    p.add_run("Panel A. Equal-weight long credible / short mismatch portfolios").bold = True
    _build_panel_table(document, headers, panel_rows("ew"))

    p = document.add_paragraph()
    p.add_run("Panel B. Value-weight long credible / short mismatch portfolios").bold = True
    _build_panel_table(document, headers, panel_rows("vw"))
    document.save(output_path)


def _result_notes(summary_rows: list[dict[str, object]], *, size_subset: str) -> str:
    by_key = {(row["weighting"], row["horizon"]): row for row in summary_rows}
    ew_3m = by_key[("ew", "3m")]
    vw_3m = by_key[("vw", "3m")]
    return "\n".join(
        [
            f"# Result Notes: {TEST_ID}",
            "",
            f"- Sample screen: `{size_subset}`.",
            f"- Equal-weight 3-month CAPM alpha: `{100 * float(ew_3m['capm_alpha']):.3f}` pp/month with p=`{float(ew_3m['capm_alpha_p']):.3f}`.",
            f"- Equal-weight 3-month FF5+Mom alpha: `{100 * float(ew_3m['ff5_mom_alpha']):.3f}` pp/month with p=`{float(ew_3m['ff5_mom_alpha_p']):.3f}`.",
            f"- Value-weight 3-month CAPM alpha: `{100 * float(vw_3m['capm_alpha']):.3f}` pp/month with p=`{float(vw_3m['capm_alpha_p']):.3f}`.",
            f"- Value-weight 3-month FF5+Mom alpha: `{100 * float(vw_3m['ff5_mom_alpha']):.3f}` pp/month with p=`{float(vw_3m['ff5_mom_alpha_p']):.3f}`.",
            "- Interpretation discipline: this run asks whether the existing long-short spread survives richer factor benchmarking, not whether it proves a causal market-inefficiency channel by itself.",
            "",
        ]
    )


def _writer_packet(args: argparse.Namespace, summary_rows: list[dict[str, object]], factor_manifest: dict[str, object]) -> str:
    by_key = {(row["weighting"], row["horizon"]): row for row in summary_rows}
    ew_3m = by_key[("ew", "3m")]
    return "\n".join(
        [
            f"# Writer Packet: {TEST_ID}",
            "",
            "## Identification Role",
            "- This run strengthens the earlier calendar-time portfolio result by replacing the market-only benchmark with standard factor models.",
            "- The output is designed to answer whether the long-short spread survives CAPM, FF3, FF5, and FF5+Momentum benchmarking.",
            "",
            "## Sample Definition",
            f"- Event panel: `{args.event_panel}`",
            f"- Annual panel: `{args.annual_panel}`",
            f"- Monthly returns: `{args.monthly_returns}`",
            f"- Market index: `{args.market_index}`",
            f"- Size subset: `{args.size_subset}`",
            "- Unit of observation: `calendar-month long-short portfolio return`",
            "- Signal rule: `AI-talking annual filers with PatentMismatch label; one active signal per stock-month using the latest filing`",
            "",
            "## Factor Inputs",
            f"- Official FF5 monthly source: `{FF5_2X3_MONTHLY_URL}`",
            f"- Official monthly momentum source: `{MOM_MONTHLY_URL}`",
            f"- Factor month range staged locally: `{factor_manifest['month_min']}` to `{factor_manifest['month_max']}`",
            "- Inference: `Newey-West HAC with lag 3`",
            "",
            "## Headline Numbers",
            f"- EW 3m CAPM alpha: `{100 * float(ew_3m['capm_alpha']):.3f}` pp/month (p=`{float(ew_3m['capm_alpha_p']):.3f}`)",
            f"- EW 3m FF5+Mom alpha: `{100 * float(ew_3m['ff5_mom_alpha']):.3f}` pp/month (p=`{float(ew_3m['ff5_mom_alpha_p']):.3f}`)",
            "",
            "## Caption Draft",
            "This table reports factor-adjusted calendar-time long-short portfolio returns formed after annual AI-related filing signals. Each month, the long leg holds non-mismatch AI filers and the short leg holds mismatch AI filers that filed within the prior 1, 3, 6, or 12 months. Equal-weight and value-weight portfolios are benchmarked against CAPM, FF3, FF5, and FF5 plus the momentum factor using monthly data from the Ken French Data Library. Alpha is the intercept from time-series regressions estimated with Newey-West standard errors.",
            "",
        ]
    )


def _dataset_summary(
    signal_sample: pd.DataFrame,
    summary_rows: list[dict[str, object]],
    figure_series: pd.DataFrame,
    factor_manifest: dict[str, object],
    *,
    args: argparse.Namespace,
) -> dict[str, object]:
    figure_payload = figure_series.copy()
    if "month" in figure_payload.columns:
        figure_payload["month"] = figure_payload["month"].astype(str)
    return {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "inputs": {
            "event_panel": str(args.event_panel),
            "annual_panel": str(args.annual_panel),
            "monthly_returns": str(args.monthly_returns),
            "market_index": str(args.market_index),
            "factor_root": str(args.factor_root),
            "size_subset": args.size_subset,
        },
        "signal_filing_count": int(len(signal_sample)),
        "unique_permno": int(signal_sample["permno"].nunique()),
        "filing_year_min": int(signal_sample["filing_year"].min()),
        "filing_year_max": int(signal_sample["filing_year"].max()),
        "factor_month_min": factor_manifest["month_min"],
        "factor_month_max": factor_manifest["month_max"],
        "portfolio_summaries": summary_rows,
        "figure_series": figure_payload.to_dict(orient="records"),
    }


def _copy_exports(run_dir: Path, paper_root: Path, run_id: str) -> dict[str, str]:
    exports = {
        "figure_png": paper_root / "figures" / f"{TEST_ID}_{run_id}.png",
        "figure_pdf": paper_root / "figures" / f"{TEST_ID}_{run_id}.pdf",
        "table_csv": paper_root / "tables" / f"{TEST_ID}_{run_id}.csv",
        "table_md": paper_root / "tables" / f"{TEST_ID}_{run_id}.md",
        "table_tex": paper_root / "latex" / f"{TEST_ID}_{run_id}.tex",
        "table_docx": paper_root / "docx" / f"{TEST_ID}_{run_id}.docx",
        "figure_series": paper_root / "tables" / f"{TEST_ID}_{run_id}_figure_series.csv",
        "writer_packet": paper_root / "writer_packets" / f"{TEST_ID}_{run_id}.md",
        "result_notes": paper_root / "snippets" / f"{TEST_ID}_{run_id}_result_notes.md",
    }
    for path in exports.values():
        path.parent.mkdir(parents=True, exist_ok=True)
    mapping = {
        run_dir / "figure_main.png": exports["figure_png"],
        run_dir / "figure_main.pdf": exports["figure_pdf"],
        run_dir / "table_main.csv": exports["table_csv"],
        run_dir / "table_main.md": exports["table_md"],
        run_dir / "table_main.tex": exports["table_tex"],
        run_dir / "table_main.docx": exports["table_docx"],
        run_dir / "figure_series.csv": exports["figure_series"],
        run_dir / "writer_packet.md": exports["writer_packet"],
        run_dir / "result_notes.md": exports["result_notes"],
    }
    for src, dst in mapping.items():
        shutil.copy2(src, dst)
    return {key: str(path) for key, path in exports.items()}


def main() -> None:
    args = _parse_args()
    run_dir = args.test_root / args.run_id
    run_dir.mkdir(parents=True, exist_ok=True)

    signal_sample = _prepare_signal_sample(
        args.event_panel,
        args.annual_panel,
        size_subset=args.size_subset,
    )
    monthly_returns = _load_monthly_returns(args.monthly_returns)
    market_index = _load_market_index(args.market_index)
    factors, factor_manifest = stage_ken_french_monthly_factors(
        args.factor_root,
        refresh=args.refresh_factors,
    )

    summary_rows: list[dict[str, object]] = []
    figure_series = None
    for horizon_months, horizon_label in HORIZONS.items():
        positions = _build_positions(signal_sample, horizon_months)
        series = _compute_long_short_series(positions, monthly_returns, market_index)
        series = series.merge(factors, on="month", how="left", validate="many_to_one")
        for weight_col in ["ew_long_short", "vw_long_short"]:
            summary_rows.append(_summarize_horizon(series, weight_col, horizon_label=horizon_label))
        if horizon_label == "3m":
            figure_series = series.copy()

    summary_df = pd.DataFrame(summary_rows)
    summary_df.to_csv(run_dir / "table_main.csv", index=False)
    figure_save = figure_series.copy()
    figure_save["month"] = figure_save["month"].astype(str)
    figure_save.to_csv(run_dir / "figure_series.csv", index=False)

    table_md, table_tex = _render_table_outputs(summary_rows)
    (run_dir / "table_main.md").write_text(table_md, encoding="utf-8")
    (run_dir / "table_main.tex").write_text(table_tex, encoding="utf-8")
    _build_table_docx(summary_rows, run_dir / "table_main.docx")
    png_path, pdf_path = _plot_cumulative_long_short(figure_series, run_dir / "figure_main")

    (run_dir / "result_notes.md").write_text(
        _result_notes(summary_rows, size_subset=args.size_subset), encoding="utf-8"
    )
    (run_dir / "writer_packet.md").write_text(
        _writer_packet(args, summary_rows, factor_manifest), encoding="utf-8"
    )
    dataset_summary = _dataset_summary(
        signal_sample,
        summary_rows,
        figure_series,
        factor_manifest,
        args=args,
    )
    (run_dir / "dataset_summary.json").write_text(
        json.dumps(dataset_summary, indent=2), encoding="utf-8"
    )

    paper_exports = _copy_exports(run_dir, args.paper_root, args.run_id)
    manifest = {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "module_path": MODULE_PATH,
        "run_dir": str(run_dir),
        "inputs": {
            "event_panel": str(args.event_panel),
            "annual_panel": str(args.annual_panel),
            "monthly_returns": str(args.monthly_returns),
            "market_index": str(args.market_index),
            "factor_root": str(args.factor_root),
            "size_subset": args.size_subset,
            "refresh_factors": bool(args.refresh_factors),
        },
        "factor_manifest": factor_manifest,
        "outputs": {
            "dataset_summary": str(run_dir / "dataset_summary.json"),
            "table_csv": str(run_dir / "table_main.csv"),
            "table_md": str(run_dir / "table_main.md"),
            "table_tex": str(run_dir / "table_main.tex"),
            "table_docx": str(run_dir / "table_main.docx"),
            "figure_series": str(run_dir / "figure_series.csv"),
            "figure_png": str(png_path),
            "figure_pdf": str(pdf_path),
            "writer_packet": str(run_dir / "writer_packet.md"),
            "result_notes": str(run_dir / "result_notes.md"),
        },
        "paper_exports": paper_exports,
    }
    (run_dir / "run_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    print(f"[{TEST_ID}] wrote run bundle to {run_dir}")
    print(f"[{TEST_ID}] paper table: {paper_exports['table_docx']}")


if __name__ == "__main__":
    main()
