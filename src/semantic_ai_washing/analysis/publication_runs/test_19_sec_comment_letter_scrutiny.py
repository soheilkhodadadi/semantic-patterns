"""Publication run driver for Test 19: SEC comment-letter scrutiny for low-credibility AI disclosure."""

from __future__ import annotations

import argparse
import json
import math
import re
import shutil
from datetime import UTC, date, datetime
from pathlib import Path

from docx import Document
from dotenv import dotenv_values
import matplotlib.pyplot as plt
import pandas as pd
import psycopg2

from semantic_ai_washing.analysis.delivery_table_payloads import _fit_absorbed_ols, to_markdown_table
from semantic_ai_washing.analysis.publication_runs.test_03_post_filing_drift import (
    _add_note,
    _add_title,
    _build_panel_table,
    _set_document_defaults,
    _set_landscape,
)
from semantic_ai_washing.analysis.publication_runs.test_16_construct_variant_screen import (
    _add_construct_variants,
)

REPO_ROOT = Path(__file__).resolve().parents[4]
DEFAULT_ANNUAL_PANEL = (
    REPO_ROOT
    / "data/processed/panel/canonical/ever_speaker_panel_2016_2025_hybrid_api_a_conf49_v1.parquet"
)
DEFAULT_TEST_ROOT = Path(
    "/Users/soheilkhodadadi/DataWork/semantic-patterns/ai_washing/derived/v3_2/test_runs/test_19_sec_comment_letter_scrutiny"
)
DEFAULT_PAPER_ROOT = REPO_ROOT / "paper/generated/v3_2"
DEFAULT_RUN_ID = f"{date.today():%Y%m%d}_aiw_v3_2_test_19_sec_comment_letter_scrutiny_main_v1"
TEST_ID = "test_19_sec_comment_letter_scrutiny"
MODULE_PATH = "semantic_ai_washing.analysis.publication_runs.test_19_sec_comment_letter_scrutiny"

STRICT_REGEX = re.compile(
    r"\b(artificial intelligence|machine learning|generative ai|chatgpt|large language models?|"
    r"deep learning|neural networks?|foundation models?)\b",
    re.IGNORECASE,
)
ABBREV_REGEX = re.compile(r"\bA\.?I\.?\b|\bLLMs?\b")
TECH_OFFICE_REGEX = re.compile(
    r"office of technology|office of industrial applications|office of trade (and )?services",
    re.IGNORECASE,
)
CIK_REGEX = re.compile(r"\d+")

VARIANT_SPECS: list[tuple[str, str]] = [
    ("PatentMismatch", "PatentMismatch"),
    ("LowCredibility", "LowCredibility"),
    ("A_S", "A/S ratio"),
]
OUTCOME_SPECS: list[tuple[str, str]] = [
    ("any_comment_t_t1", "Any comment letter (t:t+1)"),
    ("ai_comment_narrow_t_t1", "AI scrutiny narrow (t:t+1)"),
    ("ai_comment_broad_t_t1", "AI scrutiny broad (t:t+1)"),
]
CONTROL_SPECS = ["AI_Focus", "ln_assets", "cash", "leverage", "roa"]


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--annual-panel", type=Path, default=DEFAULT_ANNUAL_PANEL)
    parser.add_argument("--test-root", type=Path, default=DEFAULT_TEST_ROOT)
    parser.add_argument("--paper-root", type=Path, default=DEFAULT_PAPER_ROOT)
    parser.add_argument("--run-id", default=DEFAULT_RUN_ID)
    parser.add_argument("--dotenv-path", type=Path, default=REPO_ROOT / ".env")
    return parser.parse_args()


def _base_style() -> None:
    plt.style.use("seaborn-v0_8-whitegrid")
    plt.rcParams.update(
        {
            "font.family": "serif",
            "font.serif": ["Times New Roman", "Times", "DejaVu Serif"],
            "axes.titlesize": 12,
            "axes.labelsize": 10,
            "xtick.labelsize": 9,
            "ytick.labelsize": 9,
        }
    )


def _normalize_cik(value: object) -> str:
    text = str(value or "").strip()
    digits = "".join(ch for ch in text if ch.isdigit())
    if not digits:
        return ""
    return digits.zfill(10)


def _normalize_ticker(value: object) -> str:
    return str(value or "").strip().upper()


def _clean_text(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _extract_ciks(value: object) -> list[str]:
    text = str(value or "")
    return [_normalize_cik(match.group(0)) for match in CIK_REGEX.finditer(text)]


def _flag_ai_scrutiny(row: pd.Series) -> tuple[bool, bool, bool]:
    issue_phrase = _clean_text(row.get("list_cl_issue_phrase", ""))
    issue_taxgrp = _clean_text(row.get("list_cl_issue_taxgrp", ""))
    body = _clean_text(row.get("cl_text", ""))
    combined_lower = f"{issue_phrase} {issue_taxgrp} {body}".lower()
    combined_raw = f"{issue_phrase} {issue_taxgrp} {body}"
    narrow = bool(STRICT_REGEX.search(combined_lower))
    broad = narrow or bool(ABBREV_REGEX.search(combined_raw))
    tech_office = bool(TECH_OFFICE_REGEX.search(combined_lower))
    return narrow, broad, tech_office


def _load_panel(path: Path) -> pd.DataFrame:
    panel = pd.read_parquet(path).copy()
    panel = _add_construct_variants(panel)
    if "sic2" not in panel.columns and "sic" in panel.columns:
        sic = pd.to_numeric(panel["sic"], errors="coerce")
        panel["sic2"] = (sic // 100).astype("Int64")
    for column in ["AI_Focus", "ln_assets", "cash", "leverage", "roa", "A_S"]:
        if column in panel.columns:
            panel[column] = pd.to_numeric(panel[column], errors="coerce")
    panel["cik"] = panel["cik"].map(_normalize_cik)
    panel["ticker"] = panel["ticker"].map(_normalize_ticker)
    panel = panel.loc[panel["any_ai_talk"].fillna(0).astype(int).eq(1)].copy()
    panel = panel.loc[panel["year"].between(2016, 2024, inclusive="both")].copy()
    return panel


def _ticker_lookup(panel: pd.DataFrame) -> dict[str, str]:
    grouped = panel.groupby("ticker")["cik"].nunique()
    unique_tickers = grouped.loc[grouped.eq(1)].index.tolist()
    mapping = (
        panel.loc[panel["ticker"].isin(unique_tickers), ["ticker", "cik"]]
        .drop_duplicates()
        .set_index("ticker")["cik"]
        .to_dict()
    )
    return mapping


def _wrds_connection(dotenv_path: Path) -> psycopg2.extensions.connection:
    values = dotenv_values(dotenv_path)
    return psycopg2.connect(
        dbname="wrds",
        user=values["WRDS_USER"],
        password=values["WRDS_PASS"],
        host=values["WRDS_DB_HOST"],
        port=int(values["WRDS_DB_PORT"]),
        connect_timeout=15,
    )


def _pull_comment_letters(panel: pd.DataFrame, *, dotenv_path: Path) -> pd.DataFrame:
    tickers = sorted({ticker for ticker in panel["ticker"].dropna().unique().tolist() if ticker})
    if not tickers:
        return pd.DataFrame()
    query = """
        select
            file_date,
            company_name,
            best_edgar_ticker,
            list_ciks_reclet,
            list_cl_issue_phrase,
            list_cl_issue_taxgrp,
            cl_text
        from audit.feed25_comment_letters
        where file_date >= date '2016-01-01'
          and best_edgar_ticker = any(%s)
    """
    conn = _wrds_connection(dotenv_path)
    try:
        frame = pd.read_sql_query(query, conn, params=(tickers,))
    finally:
        conn.close()
    return frame


def _match_comment_letters(raw: pd.DataFrame, panel: pd.DataFrame) -> tuple[pd.DataFrame, dict[str, int]]:
    if raw.empty:
        return raw.copy(), {"pulled_rows": 0, "matched_rows": 0, "matched_unique_firms": 0}

    universe_ciks = set(panel["cik"].dropna().unique().tolist())
    ticker_lookup = _ticker_lookup(panel)
    rows: list[dict[str, object]] = []
    for idx, row in raw.reset_index(drop=True).iterrows():
        ticker = _normalize_ticker(row.get("best_edgar_ticker"))
        parsed_ciks = [cik for cik in _extract_ciks(row.get("list_ciks_reclet")) if cik in universe_ciks]
        match_ciks = parsed_ciks
        match_method = "cik_list"
        if not match_ciks and ticker and ticker in ticker_lookup:
            match_ciks = [ticker_lookup[ticker]]
            match_method = "ticker_fallback"
        for matched_cik in sorted(set(match_ciks)):
            narrow, broad, tech_office = _flag_ai_scrutiny(row)
            rows.append(
                {
                    "comment_id": idx,
                    "matched_cik": matched_cik,
                    "match_method": match_method,
                    "file_date": pd.to_datetime(row.get("file_date")),
                    "event_year": int(pd.to_datetime(row.get("file_date")).year),
                    "company_name": _clean_text(row.get("company_name")),
                    "best_edgar_ticker": ticker,
                    "list_ciks_reclet": _clean_text(row.get("list_ciks_reclet")),
                    "list_cl_issue_phrase": _clean_text(row.get("list_cl_issue_phrase")),
                    "list_cl_issue_taxgrp": _clean_text(row.get("list_cl_issue_taxgrp")),
                    "ai_phrase_narrow": int(narrow),
                    "ai_phrase_broad": int(broad),
                    "tech_office_flag": int(tech_office),
                }
            )

    matched = pd.DataFrame(rows)
    summary = {
        "pulled_rows": int(len(raw)),
        "matched_rows": int(len(matched)),
        "matched_unique_firms": int(matched["matched_cik"].nunique()) if not matched.empty else 0,
        "cik_list_matches": int((matched["match_method"] == "cik_list").sum()) if not matched.empty else 0,
        "ticker_fallback_matches": int((matched["match_method"] == "ticker_fallback").sum())
        if not matched.empty
        else 0,
    }
    return matched, summary


def _aggregate_events(matched: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    if matched.empty:
        empty_events = pd.DataFrame(
            columns=["matched_cik", "year", "any_comment", "ai_comment_narrow", "ai_comment_broad"]
        )
        empty_counts = pd.DataFrame(columns=["year", "Any comment", "AI narrow", "AI broad"])
        return empty_events, empty_counts

    yearly_counts = (
        matched.groupby("event_year", as_index=False)
        .agg(
            any_comment=("comment_id", "nunique"),
            ai_comment_narrow=("ai_phrase_narrow", "sum"),
            ai_comment_broad=("ai_phrase_broad", "sum"),
        )
        .rename(
            columns={
                "event_year": "year",
                "any_comment": "Any comment",
                "ai_comment_narrow": "AI narrow",
                "ai_comment_broad": "AI broad",
            }
        )
    )
    event_panel = (
        matched.groupby(["matched_cik", "event_year"], as_index=False)
        .agg(
            any_comment=("comment_id", "nunique"),
            ai_comment_narrow=("ai_phrase_narrow", "max"),
            ai_comment_broad=("ai_phrase_broad", "max"),
        )
        .rename(columns={"matched_cik": "cik", "event_year": "year"})
    )
    event_panel["any_comment"] = (event_panel["any_comment"] > 0).astype(int)
    return event_panel, yearly_counts


def _prepare_sample(panel: pd.DataFrame, event_panel: pd.DataFrame) -> tuple[pd.DataFrame, dict[str, object]]:
    sample = panel.copy()
    for variant, _label in VARIANT_SPECS:
        if variant in sample.columns:
            sample[variant] = pd.to_numeric(sample[variant], errors="coerce")
    if event_panel.empty:
        for column in ["any_comment_t_t1", "ai_comment_narrow_t_t1", "ai_comment_broad_t_t1"]:
            sample[column] = 0.0
        return sample, {
            "sample_rows": int(len(sample)),
            "unique_firms": int(sample["cik"].nunique()),
            "years": [int(sample["year"].min()), int(sample["year"].max())],
            "outcome_counts": {column: 0 for column, _label in OUTCOME_SPECS},
        }

    current = event_panel.rename(
        columns={
            "any_comment": "any_comment_t",
            "ai_comment_narrow": "ai_comment_narrow_t",
            "ai_comment_broad": "ai_comment_broad_t",
        }
    )
    lead = event_panel.rename(
        columns={
            "year": "event_year",
            "any_comment": "any_comment_t1",
            "ai_comment_narrow": "ai_comment_narrow_t1",
            "ai_comment_broad": "ai_comment_broad_t1",
        }
    )
    sample["year_plus1"] = pd.to_numeric(sample["year"], errors="coerce") + 1
    sample = sample.merge(current, on=["cik", "year"], how="left")
    sample = sample.merge(
        lead,
        left_on=["cik", "year_plus1"],
        right_on=["cik", "event_year"],
        how="left",
        suffixes=("", ""),
    )
    sample["any_comment_t_t1"] = sample[["any_comment_t", "any_comment_t1"]].fillna(0).max(axis=1)
    sample["ai_comment_narrow_t_t1"] = (
        sample[["ai_comment_narrow_t", "ai_comment_narrow_t1"]].fillna(0).max(axis=1)
    )
    sample["ai_comment_broad_t_t1"] = (
        sample[["ai_comment_broad_t", "ai_comment_broad_t1"]].fillna(0).max(axis=1)
    )
    summary = {
        "sample_rows": int(len(sample)),
        "unique_firms": int(sample["cik"].nunique()),
        "years": [int(sample["year"].min()), int(sample["year"].max())],
        "outcome_counts": {
            outcome: int(pd.to_numeric(sample[outcome], errors="coerce").fillna(0).astype(int).sum())
            for outcome, _label in OUTCOME_SPECS
        },
    }
    return sample, summary


def _fit_model(sample: pd.DataFrame, outcome: str, variant: str) -> dict[str, object]:
    result, use, adj_r2 = _fit_absorbed_ols(
        sample,
        dependent=outcome,
        rhs_terms=[variant],
        absorb_col="sic2",
        include_year=True,
        controls=CONTROL_SPECS,
    )
    return {
        "outcome": outcome,
        "variant": variant,
        "nobs": int(result.nobs),
        "adj_r2": float(adj_r2) if adj_r2 is not None else math.nan,
        "params": result.params.to_dict(),
        "bse": result.bse.to_dict(),
        "pvalues": result.pvalues.to_dict(),
        "outcome_mean": float(use[outcome].mean()),
    }


def _stars(p_value: float | None) -> str:
    if p_value is None or not math.isfinite(p_value):
        return ""
    if p_value < 0.01:
        return "***"
    if p_value < 0.05:
        return "**"
    if p_value < 0.10:
        return "*"
    return ""


def _coef_cell(coef: float | None, se: float | None, p_value: float | None) -> str:
    if coef is None or se is None or not math.isfinite(coef) or not math.isfinite(se):
        return ""
    return f"{coef:.4f}{_stars(p_value)} ({se:.4f})"


def _build_results_table(
    rows: list[dict[str, object]],
) -> tuple[pd.DataFrame, dict[tuple[str, str], dict[str, object]]]:
    keyed = {(row["outcome"], row["variant"]): row for row in rows}
    table_rows: list[dict[str, object]] = []
    for outcome, outcome_label in OUTCOME_SPECS:
        row = {"row_label": outcome_label}
        for variant, variant_label in VARIANT_SPECS:
            result = keyed[(outcome, variant)]
            row[variant_label] = _coef_cell(
                result["params"].get(variant),
                result["bse"].get(variant),
                result["pvalues"].get(variant),
            )
        base = keyed[(outcome, VARIANT_SPECS[0][0])]
        row["Outcome mean"] = f"{base['outcome_mean']:.3f}"
        row["Observations"] = str(base["nobs"])
        table_rows.append(row)
    return pd.DataFrame(table_rows), keyed


def _render_table_outputs(table_df: pd.DataFrame) -> tuple[str, str]:
    headers = ["Outcome"] + [label for _variant, label in VARIANT_SPECS] + ["Outcome mean", "Observations"]
    md_rows = [
        [
            row["row_label"],
            *[row[label] for _variant, label in VARIANT_SPECS],
            row["Outcome mean"],
            row["Observations"],
        ]
        for _, row in table_df.iterrows()
    ]
    md = to_markdown_table(headers, [[str(value) for value in row] for row in md_rows])
    latex_lines = [
        "\\begin{table}[!htbp]",
        "\\centering",
        "\\caption{SEC comment-letter scrutiny and low-credibility AI disclosure}",
        "\\begin{tabular}{lccc rr}",
        "\\hline",
        "Outcome & PatentMismatch & LowCredibility & A/S ratio & Outcome mean & Obs. \\\\",
        "\\hline",
    ]
    for _, row in table_df.iterrows():
        latex_lines.append(
            " & ".join(
                [
                    str(row["row_label"]),
                    *[str(row[label]) for _variant, label in VARIANT_SPECS],
                    str(row["Outcome mean"]),
                    str(row["Observations"]),
                ]
            )
            + " \\\\"
        )
    latex_lines.extend(["\\hline", "\\end{tabular}", "\\end{table}"])
    return md, "\n".join(latex_lines) + "\n"


def _write_docx(
    *,
    output_path: Path,
    table_df: pd.DataFrame,
    yearly_counts: pd.DataFrame,
    match_summary: dict[str, int],
    sample_summary: dict[str, object],
) -> None:
    document = Document()
    _set_document_defaults(document)
    _set_landscape(document)
    _add_title(document, "Test 19. SEC comment-letter scrutiny")
    _add_note(
        document,
        (
            "This table links the AI-talking annual panel to Audit Analytics SEC comment letters "
            "for the same firm universe. The main outcomes are any comment-letter incidence and "
            "AI-related comment-letter incidence over a `t:t+1` window. Event definitions are "
            "reported alongside the yearly audit counts below."
        ),
    )
    _add_note(
        document,
        (
            f"Matched rows = {match_summary['matched_rows']}, matched firms = {match_summary['matched_unique_firms']}, "
            f"sample rows = {sample_summary['sample_rows']}, unique firms in sample = {sample_summary['unique_firms']}."
        ),
    )
    headers = ["Outcome"] + [label for _variant, label in VARIANT_SPECS] + ["Outcome mean", "Observations"]
    table_rows = [
        [
            (str(row["row_label"]), False),
            *[(str(row[label]), False) for _variant, label in VARIANT_SPECS],
            (str(row["Outcome mean"]), False),
            (str(row["Observations"]), False),
        ]
        for _, row in table_df.iterrows()
    ]
    _build_panel_table(document, headers, table_rows)
    document.add_paragraph("Yearly event counts")
    if not yearly_counts.empty:
        count_headers = yearly_counts.columns.tolist()
        count_rows = [
            [(str(value), False) for value in row]
            for row in yearly_counts.astype({"year": int}).values.tolist()
        ]
        _build_panel_table(document, count_headers, count_rows)
    else:
        document.add_paragraph("No matched comment-letter events found.")
    document.save(output_path)


def _write_figure(yearly_counts: pd.DataFrame, *, png_path: Path, pdf_path: Path) -> None:
    _base_style()
    fig, ax = plt.subplots(figsize=(8.5, 4.8))
    if yearly_counts.empty:
        ax.text(0.5, 0.5, "No matched comment-letter events", ha="center", va="center", fontsize=11)
        ax.set_axis_off()
    else:
        ax.plot(yearly_counts["year"], yearly_counts["AI narrow"], marker="o", label="AI narrow")
        ax.plot(yearly_counts["year"], yearly_counts["AI broad"], marker="o", label="AI broad")
        ax.plot(yearly_counts["year"], yearly_counts["Any comment"], marker="o", label="Any comment")
        ax.set_title("Matched SEC comment-letter counts by year")
        ax.set_xlabel("Year")
        ax.set_ylabel("Count")
        ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(png_path, dpi=220)
    fig.savefig(pdf_path)
    plt.close(fig)


def _write_writer_packet(
    *,
    output_path: Path,
    match_summary: dict[str, int],
    sample_summary: dict[str, object],
    yearly_counts: pd.DataFrame,
    keyed: dict[tuple[str, str], dict[str, object]],
) -> None:
    broad = keyed[("ai_comment_broad_t_t1", "PatentMismatch")]
    narrow = keyed[("ai_comment_narrow_t_t1", "PatentMismatch")]
    lines = [
        "# Writer Packet: Test 19 SEC comment-letter scrutiny",
        "",
        "## Setup",
        "",
        "- Sample: AI-talking annual panel, years 2016-2024.",
        "- SEC source: Audit Analytics `feed25_comment_letters` restricted to the panel ticker universe.",
        "- Matching order: `list_ciks_reclet` first, then unique-ticker fallback.",
        f"- Pulled rows: `{match_summary['pulled_rows']}`.",
        f"- Matched rows: `{match_summary['matched_rows']}` across `{match_summary['matched_unique_firms']}` firms.",
        f"- Event counts in sample window: any comment `{sample_summary['outcome_counts']['any_comment_t_t1']}`, narrow AI `{sample_summary['outcome_counts']['ai_comment_narrow_t_t1']}`, broad AI `{sample_summary['outcome_counts']['ai_comment_broad_t_t1']}`.",
        "",
        "## Read",
        "",
        f"- PatentMismatch on narrow AI scrutiny (`t:t+1`): `{narrow['params'].get('PatentMismatch', math.nan):.4f}` with p-value `{narrow['pvalues'].get('PatentMismatch', math.nan):.3f}`.",
        f"- PatentMismatch on broad AI scrutiny (`t:t+1`): `{broad['params'].get('PatentMismatch', math.nan):.4f}` with p-value `{broad['pvalues'].get('PatentMismatch', math.nan):.3f}`.",
        "- Use the broad definition only if the narrow definition is too sparse to stand on its own.",
        "- This run should be read as Packet D entry work: it establishes event density, matching quality, and whether a true scrutiny lane is viable.",
        "",
        "## Caution",
        "",
        "- Comment letters are matched from the panel ticker universe; rows lacking both a usable CIK list and a unique ticker fallback are excluded.",
        "- The narrow AI definition is intentionally conservative; the broad definition adds AI abbreviations but should still be checked against false positives.",
        "- A thin narrow-event count does not kill the packet; it just means Test 20 may need the broader technology/disclosure scrutiny definition.",
    ]
    if not yearly_counts.empty:
        peak = yearly_counts.sort_values("AI broad", ascending=False).iloc[0]
        lines.extend(
            [
                "",
                "## Event density",
                "",
                f"- Peak broad-AI scrutiny year in the matched sample: `{int(peak['year'])}` with `{int(peak['AI broad'])}` events.",
            ]
        )
    output_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _write_result_notes(
    *,
    output_path: Path,
    match_summary: dict[str, int],
    sample_summary: dict[str, object],
    keyed: dict[tuple[str, str], dict[str, object]],
) -> None:
    broad = keyed[("ai_comment_broad_t_t1", "PatentMismatch")]
    narrow = keyed[("ai_comment_narrow_t_t1", "PatentMismatch")]
    lines = [
        "- Packet D has started with a comment-letter scrutiny audit plus pilot incidence table.",
        f"- Matched SEC comment-letter rows: `{match_summary['matched_rows']}` across `{match_summary['matched_unique_firms']}` firms.",
        f"- Outcome counts in the AI-talking annual sample: any comment `{sample_summary['outcome_counts']['any_comment_t_t1']}`, narrow AI `{sample_summary['outcome_counts']['ai_comment_narrow_t_t1']}`, broad AI `{sample_summary['outcome_counts']['ai_comment_broad_t_t1']}`.",
        f"- PatentMismatch on narrow AI scrutiny (`t:t+1`): `{narrow['params'].get('PatentMismatch', math.nan):.4f}` (p=`{narrow['pvalues'].get('PatentMismatch', math.nan):.3f}`).",
        f"- PatentMismatch on broad AI scrutiny (`t:t+1`): `{broad['params'].get('PatentMismatch', math.nan):.4f}` (p=`{broad['pvalues'].get('PatentMismatch', math.nan):.3f}`).",
        "- Treat this as a viability result first: it tells us whether the scrutiny lane is dense enough to support the post-scrutiny cleanup design in Test 20.",
    ]
    output_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    args = _parse_args()
    run_dir = args.test_root / args.run_id
    paper_root = args.paper_root
    output_paths = {
        "table_docx": paper_root / "docx" / f"{TEST_ID}_{args.run_id}.docx",
        "table_tex": paper_root / "latex" / f"{TEST_ID}_{args.run_id}.tex",
        "table_csv": paper_root / "tables" / f"{TEST_ID}_{args.run_id}.csv",
        "counts_csv": paper_root / "tables" / f"{TEST_ID}_{args.run_id}_yearly_counts.csv",
        "figure_png": paper_root / "figures" / f"{TEST_ID}_{args.run_id}.png",
        "figure_pdf": paper_root / "figures" / f"{TEST_ID}_{args.run_id}.pdf",
        "writer_packet": paper_root / "writer_packets" / f"{TEST_ID}_{args.run_id}.md",
        "result_notes": paper_root / "snippets" / f"{TEST_ID}_{args.run_id}_result_notes.md",
    }
    run_dir.mkdir(parents=True, exist_ok=True)
    paper_root.mkdir(parents=True, exist_ok=True)
    for path in output_paths.values():
        path.parent.mkdir(parents=True, exist_ok=True)

    panel = _load_panel(args.annual_panel)
    raw_comments = _pull_comment_letters(panel, dotenv_path=args.dotenv_path)
    matched_comments, match_summary = _match_comment_letters(raw_comments, panel)
    event_panel, yearly_counts = _aggregate_events(matched_comments)
    sample, sample_summary = _prepare_sample(panel, event_panel)

    model_rows = [
        _fit_model(sample, outcome=outcome, variant=variant)
        for outcome, _label in OUTCOME_SPECS
        for variant, _variant_label in VARIANT_SPECS
    ]
    table_df, keyed = _build_results_table(model_rows)
    md_table, latex_table = _render_table_outputs(table_df)

    table_df.to_csv(output_paths["table_csv"], index=False)
    yearly_counts.to_csv(output_paths["counts_csv"], index=False)
    output_paths["table_tex"].write_text(latex_table, encoding="utf-8")
    _write_docx(
        output_path=output_paths["table_docx"],
        table_df=table_df,
        yearly_counts=yearly_counts,
        match_summary=match_summary,
        sample_summary=sample_summary,
    )
    _write_figure(yearly_counts, png_path=output_paths["figure_png"], pdf_path=output_paths["figure_pdf"])
    _write_writer_packet(
        output_path=output_paths["writer_packet"],
        match_summary=match_summary,
        sample_summary=sample_summary,
        yearly_counts=yearly_counts,
        keyed=keyed,
    )
    _write_result_notes(
        output_path=output_paths["result_notes"],
        match_summary=match_summary,
        sample_summary=sample_summary,
        keyed=keyed,
    )

    raw_export = run_dir / "matched_comment_letters.parquet"
    event_export = run_dir / "comment_letter_event_panel.parquet"
    sample_export = run_dir / "scrutiny_sample.parquet"
    matched_comments.to_parquet(raw_export, index=False)
    event_panel.to_parquet(event_export, index=False)
    sample.to_parquet(sample_export, index=False)

    for path in output_paths.values():
        shutil.copy2(path, run_dir / path.name)

    manifest = {
        "test_id": TEST_ID,
        "module_path": MODULE_PATH,
        "run_id": args.run_id,
        "run_timestamp_utc": datetime.now(UTC).isoformat(),
        "inputs": {
            "annual_panel": str(args.annual_panel),
            "dotenv_path": str(args.dotenv_path),
        },
        "intermediate_outputs": {
            "matched_comment_letters": str(raw_export),
            "event_panel": str(event_export),
            "scrutiny_sample": str(sample_export),
        },
        "paper_outputs": {key: str(path) for key, path in output_paths.items()},
        "match_summary": match_summary,
        "sample_summary": sample_summary,
        "yearly_counts": yearly_counts.to_dict(orient="records"),
        "models": model_rows,
    }
    (run_dir / "run_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
