"""Publication run driver for legacy Table B2: attrition map."""

from __future__ import annotations

import argparse
import json
import shutil
from datetime import UTC, date, datetime
from pathlib import Path

import pandas as pd
from docx import Document

from semantic_ai_washing.analysis.publication_runs.test_03_post_filing_drift import (
    _add_note,
    _add_title,
    _build_panel_table,
    _set_document_defaults,
    _set_landscape,
)

REPO_ROOT = Path(__file__).resolve().parents[4]
DEFAULT_ANNUAL_PANEL = (
    REPO_ROOT
    / "data/processed/panel/canonical/ever_speaker_panel_2016_2025_hybrid_api_a_conf49_v1.parquet"
)
DEFAULT_EVENT_PANEL = (
    REPO_ROOT
    / "data/processed/panel/filing_event_estimation_sample_hybrid_api_a_conf49_v1.parquet"
)
DEFAULT_DAILY_RETURNS = (
    REPO_ROOT / "data/interim/market/filing_event_returns_daily_hybrid_api_a_conf49_v1.parquet"
)
DEFAULT_TEST_ROOT = Path(
    "/Users/soheilkhodadadi/DataWork/semantic-patterns/ai_washing/derived/test_runs/legacy_b2_attrition_map"
)
DEFAULT_PAPER_ROOT = REPO_ROOT / "paper/generated"
DEFAULT_RUN_ID = f"{date.today():%Y%m%d}_hybrid_api_a_conf49_main_v1"
TEST_ID = "legacy_b2_attrition_map"
MODULE_PATH = "semantic_ai_washing.analysis.publication_runs.legacy_b2_attrition_map"
WINDOW_SPECS = [
    (
        -1,
        2,
        "CAR[-1,+2] complete",
        "Immediate filing-window CAR sample after daily-return matching.",
    ),
    (2, 21, "BHAR[+2,+21] complete", "Requires one trading month of post-filing returns."),
    (
        2,
        63,
        "BHAR[+2,+63] complete",
        "Three-month horizon drops later-sample filings mechanically.",
    ),
    (2, 126, "BHAR[+2,+126] complete", "Six-month horizon drops more late-sample filings."),
    (
        2,
        252,
        "BHAR[+2,+252] complete",
        "One-year horizon is the main event-study attrition screen.",
    ),
]


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--annual-panel", type=Path, default=DEFAULT_ANNUAL_PANEL)
    parser.add_argument("--event-panel", type=Path, default=DEFAULT_EVENT_PANEL)
    parser.add_argument("--daily-returns", type=Path, default=DEFAULT_DAILY_RETURNS)
    parser.add_argument("--test-root", type=Path, default=DEFAULT_TEST_ROOT)
    parser.add_argument("--paper-root", type=Path, default=DEFAULT_PAPER_ROOT)
    parser.add_argument("--run-id", default=DEFAULT_RUN_ID)
    return parser.parse_args()


def _fmt_share(numerator: int, denominator: int, label: str) -> str:
    if denominator <= 0:
        return f"n/a of {label}"
    return f"{100 * numerator / denominator:.1f}% of {label}"


def _load_attrition_inputs(args: argparse.Namespace) -> dict[str, object]:
    annual = pd.read_parquet(args.annual_panel).copy()
    annual["year"] = pd.to_numeric(annual["year"], errors="coerce")
    for column in [
        "any_ai_talk",
        "ln_assets",
        "leverage",
        "cash",
        "capx_at",
        "roa",
        "rd_intensity",
        "emp",
        "market_cap_year_end",
        "shrout",
    ]:
        annual[column] = pd.to_numeric(annual[column], errors="coerce")
    annual = annual.sort_values(["permno", "year", "cik"]).reset_index(drop=True)
    annual["share_growth_lead1"] = (
        annual.groupby("permno", sort=False)["shrout"].shift(-1).div(annual["shrout"]) - 1.0
    )
    annual["nonbig_marketcap"] = annual["market_cap_year_end"].le(
        annual.groupby("year")["market_cap_year_end"].transform("median")
    )

    event = pd.read_parquet(args.event_panel).copy()
    daily = pd.read_parquet(args.daily_returns, columns=["filing_id", "relative_day"]).copy()
    daily["relative_day"] = pd.to_numeric(daily["relative_day"], errors="coerce").astype("Int64")
    observed_days = daily.groupby("filing_id")["relative_day"].agg(
        lambda s: set(s.dropna().astype(int).tolist())
    )

    base_2016 = annual.loc[annual["year"].eq(2016)].sort_values("cik").drop_duplicates("cik")

    window_counts = []
    for start_day, end_day, label, reason in WINDOW_SPECS:
        needed_days = set(range(start_day, end_day + 1))
        n_complete = int(observed_days.apply(lambda s: needed_days.issubset(s)).sum())
        window_counts.append(
            {
                "row_label": label,
                "n": n_complete,
                "share_of_parent": _fmt_share(n_complete, len(event), "event sample"),
                "notes": reason,
            }
        )

    annual_rows = len(annual)
    event_rows = len(event)
    baseline_rows = len(base_2016)
    annual_counts = {
        "annual_rows": annual_rows,
        "annual_unique_firms": int(annual["cik"].astype(str).nunique()),
        "annual_ai_talk_rows": int(annual["any_ai_talk"].fillna(0).sum()),
        "annual_core_controls_complete": int(
            annual[["ln_assets", "leverage", "cash", "capx_at", "roa"]].notna().all(axis=1).sum()
        ),
        "annual_with_rd": int(annual["rd_intensity"].notna().sum()),
        "annual_market_linked": int(annual["market_cap_year_end"].notna().sum()),
        "annual_share_growth_complete": int(annual["share_growth_lead1"].notna().sum()),
        "annual_nonbig_market_linked": int(
            (
                annual["market_cap_year_end"].notna() & annual["nonbig_marketcap"].fillna(False)
            ).sum()
        ),
    }
    baseline_counts = {
        "baseline_rows": baseline_rows,
        "baseline_core_controls_complete": int(
            base_2016[["ln_assets", "leverage", "cash", "capx_at", "roa"]]
            .notna()
            .all(axis=1)
            .sum()
        ),
        "baseline_with_rd": int(base_2016["rd_intensity"].notna().sum()),
        "baseline_with_emp": int(base_2016["emp"].notna().sum()),
    }
    event_counts = {
        "event_rows": event_rows,
        "event_unique_filings": int(event["filing_id"].nunique()),
        "window_counts": window_counts,
    }
    return {
        "annual_counts": annual_counts,
        "baseline_counts": baseline_counts,
        "event_counts": event_counts,
    }


def _build_table_df(attrition: dict[str, object]) -> pd.DataFrame:
    annual = attrition["annual_counts"]
    baseline = attrition["baseline_counts"]
    event = attrition["event_counts"]
    rows = [
        {
            "panel": "Panel A. Expanded annual ever-speaker panel",
            "row_label": "Expanded ever-speaker backbone",
            "n": annual["annual_rows"],
            "share_of_parent": "100.0% of annual panel",
            "notes": (f"{annual['annual_unique_firms']:,} unique firms across 2016-2025."),
        },
        {
            "panel": "Panel A. Expanded annual ever-speaker panel",
            "row_label": "AI-talking firm-years",
            "n": annual["annual_ai_talk_rows"],
            "share_of_parent": _fmt_share(
                annual["annual_ai_talk_rows"], annual["annual_rows"], "annual panel"
            ),
            "notes": "Disclosure-composition and PatentMismatch are economically meaningful here.",
        },
        {
            "panel": "Panel A. Expanded annual ever-speaker panel",
            "row_label": "Rows with core annual controls",
            "n": annual["annual_core_controls_complete"],
            "share_of_parent": _fmt_share(
                annual["annual_core_controls_complete"], annual["annual_rows"], "annual panel"
            ),
            "notes": "Requires ln assets, leverage, cash, capx/assets, and ROA.",
        },
        {
            "panel": "Panel A. Expanded annual ever-speaker panel",
            "row_label": "Rows with R&D/assets",
            "n": annual["annual_with_rd"],
            "share_of_parent": _fmt_share(
                annual["annual_with_rd"], annual["annual_rows"], "annual panel"
            ),
            "notes": "Most restrictive accounting variable in the annual panel.",
        },
        {
            "panel": "Panel A. Expanded annual ever-speaker panel",
            "row_label": "Rows with CRSP market-cap linkage",
            "n": annual["annual_market_linked"],
            "share_of_parent": _fmt_share(
                annual["annual_market_linked"], annual["annual_rows"], "annual panel"
            ),
            "notes": "Needed for valuation and financing tests.",
        },
        {
            "panel": "Panel A. Expanded annual ever-speaker panel",
            "row_label": "Rows with next-year share growth",
            "n": annual["annual_share_growth_complete"],
            "share_of_parent": _fmt_share(
                annual["annual_share_growth_complete"], annual["annual_rows"], "annual panel"
            ),
            "notes": "Requires current and next-year shares outstanding.",
        },
        {
            "panel": "Panel A. Expanded annual ever-speaker panel",
            "row_label": "Non-big rows inside market-linked sample",
            "n": annual["annual_nonbig_market_linked"],
            "share_of_parent": _fmt_share(
                annual["annual_nonbig_market_linked"],
                annual["annual_market_linked"],
                "market-linked rows",
            ),
            "notes": "Matched-sample yearly median market-cap split used in the non-big refinement.",
        },
        {
            "panel": "Panel B. Filing-event backbone",
            "row_label": "Filing-event estimation sample",
            "n": event["event_rows"],
            "share_of_parent": "100.0% of event sample",
            "notes": "Annual 10-K filing events with matched daily return histories.",
        },
    ]
    for row in event["window_counts"]:
        rows.append({"panel": "Panel B. Filing-event backbone", **row})
    rows.extend(
        [
            {
                "panel": "Panel C. Baseline cross-section for legacy determinants",
                "row_label": "2016 ever-speaker baseline",
                "n": baseline["baseline_rows"],
                "share_of_parent": "100.0% of baseline cross-section",
                "notes": "One row per ever-speaker firm in 2016.",
            },
            {
                "panel": "Panel C. Baseline cross-section for legacy determinants",
                "row_label": "2016 rows with core controls",
                "n": baseline["baseline_core_controls_complete"],
                "share_of_parent": _fmt_share(
                    baseline["baseline_core_controls_complete"],
                    baseline["baseline_rows"],
                    "baseline cross-section",
                ),
                "notes": "Baseline multivariate sample with basic accounting controls present.",
            },
            {
                "panel": "Panel C. Baseline cross-section for legacy determinants",
                "row_label": "2016 rows with R&D/assets",
                "n": baseline["baseline_with_rd"],
                "share_of_parent": _fmt_share(
                    baseline["baseline_with_rd"],
                    baseline["baseline_rows"],
                    "baseline cross-section",
                ),
                "notes": "Primary pinch point for the fuller determinants appendix tables.",
            },
            {
                "panel": "Panel C. Baseline cross-section for legacy determinants",
                "row_label": "2016 rows with employment",
                "n": baseline["baseline_with_emp"],
                "share_of_parent": _fmt_share(
                    baseline["baseline_with_emp"],
                    baseline["baseline_rows"],
                    "baseline cross-section",
                ),
                "notes": "Second major source of legacy cross-sectional shrinkage.",
            },
        ]
    )
    return pd.DataFrame(rows)


def _markdown_table(headers: list[str], rows: list[list[object]]) -> str:
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join(["---"] * len(headers)) + " |",
    ]
    for row in rows:
        lines.append("| " + " | ".join(str(value) for value in row) + " |")
    return "\n".join(lines)


def _escape_tex(text: str) -> str:
    return (
        text.replace("\\", "\\textbackslash{}")
        .replace("&", "\\&")
        .replace("%", "\\%")
        .replace("_", "\\_")
    )


def _render_table_outputs(table_df: pd.DataFrame) -> tuple[str, str]:
    headers = ["Empirical block", "N", "Share of parent", "Main reason for shrinkage"]
    md_lines = ["# Table Main", ""]
    latex_lines = [
        "\\begin{table}[!htbp]",
        "\\centering",
        "\\caption{Attrition map for the expanded 2016-2025 AI-washing sample}",
        "\\begin{tabular}{p{4.0cm}cp{2.8cm}p{7.0cm}}",
        "\\hline",
    ]
    for panel_name in table_df["panel"].drop_duplicates():
        subset = table_df.loc[table_df["panel"].eq(panel_name)].copy()
        rendered = subset[["row_label", "n", "share_of_parent", "notes"]].values.tolist()
        md_lines.extend([f"## {panel_name}", _markdown_table(headers, rendered), ""])
        latex_lines.append(
            f"\\multicolumn{{4}}{{l}}{{\\textit{{{_escape_tex(panel_name)}}}}} \\\\"
        )
        latex_lines.append(" & ".join(headers) + " \\\\")
        for row in rendered:
            latex_lines.append(" & ".join(_escape_tex(str(value)) for value in row) + " \\\\")
    latex_lines.extend(["\\hline", "\\end{tabular}", "\\end{table}"])
    return "\n".join(md_lines), "\n".join(latex_lines) + "\n"


def _docx_panel_rows(table_df: pd.DataFrame, panel_name: str) -> list[list[tuple[str, bool]]]:
    rows = []
    subset = table_df.loc[table_df["panel"].eq(panel_name)].copy()
    for row in subset[["row_label", "n", "share_of_parent", "notes"]].itertuples(index=False):
        rows.append(
            [
                (str(row[0]), True),
                (str(row[1]), False),
                (str(row[2]), False),
                (str(row[3]), False),
            ]
        )
    return rows


def _build_table_docx(table_df: pd.DataFrame, output_path: Path) -> None:
    document = Document()
    _set_document_defaults(document)
    _set_landscape(document)
    _add_title(document, "Table B2. Attrition Map for the Expanded 2016-2025 Sample")
    _add_note(
        document,
        "This table maps the live sample backbone underlying the expanded 2016-2025 AI-washing design. "
        "Panel A tracks the annual ever-speaker panel, Panel B tracks the filing-event backbone used for "
        "CAR and BHAR tests, and Panel C tracks the 2016 baseline cross-section used by legacy determinants "
        "specifications. The current attrition story is driven mainly by WRDS accounting coverage, CRSP market "
        "linkage, and long-horizon event windows rather than missing patent-timing fields.",
    )
    headers = ["", "N", "Share of parent", "Main reason for shrinkage"]
    for idx, panel_name in enumerate(table_df["panel"].drop_duplicates()):
        if idx == 2:
            document.add_page_break()
        paragraph = document.add_paragraph()
        paragraph.add_run(panel_name).bold = True
        _build_panel_table(document, headers, _docx_panel_rows(table_df, panel_name))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def _result_notes(attrition: dict[str, object]) -> str:
    annual = attrition["annual_counts"]
    baseline = attrition["baseline_counts"]
    event = attrition["event_counts"]
    long_horizon = next(
        row for row in event["window_counts"] if row["row_label"] == "BHAR[+2,+252] complete"
    )
    return "\n".join(
        [
            "# Result Notes",
            "",
            (
                f"- Expanded annual backbone: `{annual['annual_rows']:,}` firm-years across "
                f"`{annual['annual_unique_firms']:,}` ever-speaker firms."
            ),
            (
                f"- Accounting coverage is the main annual-panel screen: `{annual['annual_core_controls_complete']:,}` "
                f"rows have the core annual controls, but only `{annual['annual_with_rd']:,}` retain R&D/assets."
            ),
            (
                f"- Market linkage is a second major screen: `{annual['annual_market_linked']:,}` annual rows have "
                f"year-end CRSP market cap, and `{annual['annual_share_growth_complete']:,}` retain next-year share growth."
            ),
            (
                f"- Filing-event backbone: `{event['event_rows']:,}` annual 10-K events; "
                f"`{long_horizon['n']:,}` survive to the 12-month BHAR window."
            ),
            (
                f"- Legacy determinants baseline: `{baseline['baseline_rows']:,}` 2016 firms, but only "
                f"`{baseline['baseline_core_controls_complete']:,}` retain the core baseline controls and "
                f"`{baseline['baseline_with_rd']:,}` retain R&D/assets."
            ),
            "",
        ]
    )


def _writer_packet(args: argparse.Namespace, attrition: dict[str, object]) -> str:
    annual = attrition["annual_counts"]
    baseline = attrition["baseline_counts"]
    event = attrition["event_counts"]
    return "\n".join(
        [
            "# Writer Packet",
            "",
            "## Metadata",
            f"- Test id: `{TEST_ID}`",
            f"- Run id: `{args.run_id}`",
            f"- Date run: `{date.today().isoformat()}`",
            f"- Script/module path: `{MODULE_PATH}`",
            "",
            "## Source Artifacts",
            f"- Annual panel: `{args.annual_panel}`",
            f"- Event panel: `{args.event_panel}`",
            f"- Daily returns: `{args.daily_returns}`",
            "",
            "## Main Numbers",
            f"- Expanded annual panel: `{annual['annual_rows']:,}` rows across `{annual['annual_unique_firms']:,}` firms",
            f"- AI-talking firm-years: `{annual['annual_ai_talk_rows']:,}`",
            f"- Annual rows with core controls: `{annual['annual_core_controls_complete']:,}`",
            f"- Annual rows with R&D/assets: `{annual['annual_with_rd']:,}`",
            f"- Market-linked annual rows: `{annual['annual_market_linked']:,}`",
            f"- Event-study filings: `{event['event_rows']:,}`",
            f"- Baseline 2016 firms: `{baseline['baseline_rows']:,}`",
            "",
            "## Caption Draft",
            "This appendix table summarizes sample attrition across the expanded AI-washing design. It separates the annual ever-speaker panel, the filing-event backbone, and the 2016 baseline cross-section used by the legacy determinants specifications. The current attrition pattern is driven mainly by accounting-data coverage, CRSP market linkage, and the mechanical loss of late-sample filings in longer BHAR windows.",
            "",
        ]
    )


def _dataset_summary(args: argparse.Namespace, attrition: dict[str, object]) -> dict[str, object]:
    return {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "inputs": {
            "annual_panel": str(args.annual_panel),
            "event_panel": str(args.event_panel),
            "daily_returns": str(args.daily_returns),
        },
        "attrition": attrition,
    }


def _copy_exports(run_dir: Path, paper_root: Path, run_id: str) -> dict[str, str]:
    exports = {
        "table_csv": paper_root / "tables" / f"{TEST_ID}_{run_id}.csv",
        "table_md": paper_root / "tables" / f"{TEST_ID}_{run_id}.md",
        "table_tex": paper_root / "latex" / f"{TEST_ID}_{run_id}.tex",
        "table_docx": paper_root / "docx" / f"{TEST_ID}_{run_id}.docx",
        "writer_packet": paper_root / "writer_packets" / f"{TEST_ID}_{run_id}.md",
        "result_notes": paper_root / "snippets" / f"{TEST_ID}_{run_id}_result_notes.md",
    }
    for path in exports.values():
        path.parent.mkdir(parents=True, exist_ok=True)
    mapping = {
        run_dir / "table_main.csv": exports["table_csv"],
        run_dir / "table_main.md": exports["table_md"],
        run_dir / "table_main.tex": exports["table_tex"],
        run_dir / "table_main.docx": exports["table_docx"],
        run_dir / "writer_packet.md": exports["writer_packet"],
        run_dir / "result_notes.md": exports["result_notes"],
    }
    for src, dst in mapping.items():
        shutil.copy2(src, dst)
    return {key: str(path) for key, path in exports.items()}


def main() -> None:
    args = _parse_args()
    run_dir = args.test_root / args.run_id
    run_dir.mkdir(parents=True, exist_ok=True)

    attrition = _load_attrition_inputs(args)
    table_df = _build_table_df(attrition)
    table_df.to_csv(run_dir / "table_main.csv", index=False)

    table_md, table_tex = _render_table_outputs(table_df)
    (run_dir / "table_main.md").write_text(table_md, encoding="utf-8")
    (run_dir / "table_main.tex").write_text(table_tex, encoding="utf-8")
    _build_table_docx(table_df, run_dir / "table_main.docx")
    (run_dir / "result_notes.md").write_text(_result_notes(attrition), encoding="utf-8")
    (run_dir / "writer_packet.md").write_text(_writer_packet(args, attrition), encoding="utf-8")
    dataset_summary = _dataset_summary(args, attrition)
    (run_dir / "dataset_summary.json").write_text(
        json.dumps(dataset_summary, indent=2), encoding="utf-8"
    )

    paper_exports = _copy_exports(run_dir, args.paper_root, args.run_id)
    manifest = {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "module_path": MODULE_PATH,
        "run_dir": str(run_dir),
        "outputs": {
            "dataset_summary": str(run_dir / "dataset_summary.json"),
            "table_csv": str(run_dir / "table_main.csv"),
            "table_md": str(run_dir / "table_main.md"),
            "table_tex": str(run_dir / "table_main.tex"),
            "table_docx": str(run_dir / "table_main.docx"),
            "writer_packet": str(run_dir / "writer_packet.md"),
            "result_notes": str(run_dir / "result_notes.md"),
        },
        "paper_exports": paper_exports,
    }
    (run_dir / "run_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    print(f"[{TEST_ID}] wrote run bundle to {run_dir}")
    print(f"[{TEST_ID}] paper table: {paper_exports['table_docx']}")


if __name__ == "__main__":
    main()
