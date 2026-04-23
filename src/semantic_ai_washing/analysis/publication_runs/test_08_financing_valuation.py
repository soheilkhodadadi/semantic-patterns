"""Publication run driver for Test 08: financing and valuation consequences."""

from __future__ import annotations

import argparse
import json
import math
import shutil
from datetime import UTC, date, datetime
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
import statsmodels.api as sm

from semantic_ai_washing.analysis.delivery_table_payloads import _add_patent_mismatch
from semantic_ai_washing.analysis.publication_runs.test_03_post_filing_drift import (
    _add_note,
    _add_title,
    _build_panel_table,
    _set_document_defaults,
    _set_landscape,
)

REPO_ROOT = Path(__file__).resolve().parents[4]
DEFAULT_ANNUAL_PANEL = (
    REPO_ROOT
    / "data/processed/panel/canonical/ever_speaker_panel_2016_2025_hybrid_api_a_conf49_v1.parquet"
)
DEFAULT_MARKET_FEATURES = (
    REPO_ROOT
    / "data/interim/market/annual_market_features_ever_speaker_2016_2025_hybrid_api_a_conf49_v1.csv"
)
DEFAULT_TEST_ROOT = Path(
    "/Users/soheilkhodadadi/DataWork/semantic-patterns/ai_washing/derived/test_runs/test_08_financing_valuation"
)
DEFAULT_PAPER_ROOT = REPO_ROOT / "paper/generated"
DEFAULT_RUN_ID = f"{date.today():%Y%m%d}_hybrid_api_a_conf49_main_v1"
TEST_ID = "test_08_financing_valuation"
MODULE_PATH = "semantic_ai_washing.analysis.publication_runs.test_08_financing_valuation"
FOCAL_TERMS = ["PatentMismatch", "A_S", "AI_Focus"]
VALUATION_CONTROLS = ["ln_assets", "cash", "roa"]
FINANCING_CONTROLS = ["ln_assets", "leverage", "cash", "roa"]
VALUATION_SPECS = [
    ("log_mktcap_assets", "Log MktCap/assets", VALUATION_CONTROLS),
    ("log_q_proxy", "Log Q proxy", VALUATION_CONTROLS),
    ("delta_log_q_proxy_lead1", "Delta log Q t+1", VALUATION_CONTROLS),
]
FINANCING_SPECS = [
    ("share_growth_lead1", "Delta Shares t+1", FINANCING_CONTROLS),
    ("equity_issue_lead1", "Issue>5% t+1", FINANCING_CONTROLS),
]


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--annual-panel", type=Path, default=DEFAULT_ANNUAL_PANEL)
    parser.add_argument("--market-features", type=Path, default=DEFAULT_MARKET_FEATURES)
    parser.add_argument("--test-root", type=Path, default=DEFAULT_TEST_ROOT)
    parser.add_argument("--paper-root", type=Path, default=DEFAULT_PAPER_ROOT)
    parser.add_argument("--run-id", default=DEFAULT_RUN_ID)
    parser.add_argument(
        "--size-subset",
        choices=["all", "nonbig", "big"],
        default="all",
        help="Optional size screen using the matched-sample yearly median market cap.",
    )
    parser.add_argument(
        "--include-post-chatgpt-interaction",
        action="store_true",
        help="Add PatentMismatch x PostChatGPT, with the post indicator absorbed by year fixed effects.",
    )
    return parser.parse_args()


def _base_style() -> None:
    plt.style.use("seaborn-v0_8-whitegrid")
    plt.rcParams.update(
        {
            "font.family": "serif",
            "font.serif": ["Times New Roman", "Times", "DejaVu Serif"],
            "axes.titlesize": 13,
            "axes.labelsize": 11,
            "xtick.labelsize": 10,
            "ytick.labelsize": 10,
            "legend.fontsize": 9,
        }
    )


def _normalize_id(series: pd.Series) -> pd.Series:
    return (
        series.astype(str)
        .str.replace(r"\.0$", "", regex=True)
        .str.strip()
        .replace({"nan": "", "None": ""})
    )


def _ensure_sic2(df: pd.DataFrame) -> pd.DataFrame:
    panel = df.copy()
    if "sic2" not in panel.columns or panel["sic2"].isna().all():
        if "sic" in panel.columns:
            sic_raw = pd.to_numeric(panel["sic"], errors="coerce")
            panel["sic2"] = (sic_raw // 100).astype("Int64")
        else:
            panel["sic2"] = pd.Series(pd.NA, index=panel.index, dtype="Int64")
    return panel


def _load_panel(annual_panel: Path, market_features: Path) -> pd.DataFrame:
    panel = pd.read_parquet(annual_panel).copy()
    panel = _ensure_sic2(panel)
    panel = _add_patent_mismatch(panel)
    panel["cik"] = _normalize_id(panel["cik"])
    panel["permno"] = _normalize_id(panel["permno"])
    panel["year"] = pd.to_numeric(panel["year"], errors="coerce").astype("Int64")

    market = pd.read_csv(market_features).copy()
    market["permno"] = _normalize_id(market["permno"])
    market["year"] = pd.to_numeric(market["year"], errors="coerce").astype("Int64")
    for column in ["shrout", "market_cap_year_end"]:
        market[column] = pd.to_numeric(market[column], errors="coerce")
    keep = ["permno", "year", "shrout", "market_cap_year_end"]
    panel = panel.merge(market[keep], on=["permno", "year"], how="left", suffixes=("", "_market"))
    panel = panel.sort_values(["cik", "year"]).reset_index(drop=True)
    return panel


def _winsorize(series: pd.Series, lower_q: float = 0.01, upper_q: float = 0.99) -> pd.Series:
    valid = pd.to_numeric(series, errors="coerce").replace([np.inf, -np.inf], np.nan)
    if valid.notna().sum() == 0:
        return valid
    lower = valid.quantile(lower_q)
    upper = valid.quantile(upper_q)
    return valid.clip(lower=lower, upper=upper)


def _prepare_panel(panel: pd.DataFrame) -> tuple[pd.DataFrame, dict[str, object]]:
    use = panel.copy()
    use["PatentMismatch"] = (
        pd.to_numeric(use["PatentMismatch"], errors="coerce").fillna(0).astype(int)
    )
    for column in [
        "A_S",
        "AI_Focus",
        "ln_assets",
        "leverage",
        "cash",
        "roa",
        "market_cap_year_end",
        "shrout",
    ]:
        use[column] = pd.to_numeric(use[column], errors="coerce").replace(
            [np.inf, -np.inf], np.nan
        )

    use["assets_proxy"] = np.exp(use["ln_assets"])
    use.loc[~np.isfinite(use["assets_proxy"]) | use["assets_proxy"].le(0), "assets_proxy"] = np.nan

    use["mktcap_assets"] = use["market_cap_year_end"] / use["assets_proxy"]
    use["q_proxy"] = use["mktcap_assets"] + use["leverage"]
    use["log_mktcap_assets"] = np.log1p(use["mktcap_assets"].clip(lower=0))
    use["log_q_proxy"] = np.log1p(use["q_proxy"].clip(lower=0))
    use["PostChatGPT"] = np.where(
        use["year"].notna(), use["year"].astype(int).ge(2023).astype(float), np.nan
    )
    yearly_median_market_cap = use.groupby("year")["market_cap_year_end"].transform("median")
    use["nonbig_marketcap"] = np.where(
        use["market_cap_year_end"].notna() & yearly_median_market_cap.notna(),
        use["market_cap_year_end"].le(yearly_median_market_cap),
        pd.NA,
    )

    use = use.sort_values(["permno", "year", "cik"]).reset_index(drop=True)
    shrout_lead1 = use.groupby("permno", sort=False)["shrout"].shift(-1)
    log_q_proxy_lead1 = use.groupby("cik", sort=False)["log_q_proxy"].shift(-1)

    use["share_growth_lead1"] = shrout_lead1 / use["shrout"] - 1.0
    use["equity_issue_lead1"] = np.where(
        use["share_growth_lead1"].notna(),
        (use["share_growth_lead1"] > 0.05).astype(float),
        np.nan,
    )
    use["delta_log_q_proxy_lead1"] = log_q_proxy_lead1 - use["log_q_proxy"]

    for column in [
        "log_mktcap_assets",
        "log_q_proxy",
        "delta_log_q_proxy_lead1",
        "share_growth_lead1",
    ]:
        use[column] = _winsorize(use[column])

    summary = {
        "firm_year_rows": int(len(use)),
        "unique_firms": int(use["cik"].nunique()),
        "valuation_marketcap_nonmissing": int(use["market_cap_year_end"].notna().sum()),
        "valuation_assets_nonmissing": int(use["assets_proxy"].notna().sum()),
        "share_growth_nonmissing": int(use["share_growth_lead1"].notna().sum()),
        "issue_indicator_nonmissing": int(use["equity_issue_lead1"].notna().sum()),
        "issue_rule": "Issue>5% t+1 equals one when next-year CRSP shrout growth exceeds 5%",
        "q_proxy_rule": "Q proxy equals market_cap/assets plus leverage because local WRDS pull lacks book-equity fields",
        "size_rule": "non-big uses the matched-sample yearly median market cap because NYSE breakpoints are unavailable locally",
        "nonbig_flag_nonmissing": int(pd.Series(use["nonbig_marketcap"]).notna().sum()),
    }
    return use, summary


def _apply_size_subset(
    panel: pd.DataFrame, summary: dict[str, object], size_subset: str
) -> tuple[pd.DataFrame, dict[str, object]]:
    use = panel.copy()
    subset_summary = dict(summary)
    subset_summary["size_subset"] = size_subset
    if size_subset == "all":
        return use, subset_summary
    flag = pd.Series(use["nonbig_marketcap"], index=use.index)
    if size_subset == "nonbig":
        use = use.loc[flag.fillna(False).astype(bool)].copy()
    else:
        use = use.loc[flag.notna() & ~flag.astype(bool)].copy()
    subset_summary["subset_rows"] = int(len(use))
    subset_summary["subset_unique_firms"] = int(use["cik"].nunique())
    return use, subset_summary


def _two_way_absorb(
    use: pd.DataFrame, columns: list[str], *, max_iter: int = 200, tol: float = 1e-10
) -> pd.DataFrame:
    entity = use["cik"]
    year = use["year"]
    absorbed = pd.DataFrame(index=use.index)
    for col in columns:
        z = pd.to_numeric(use[col], errors="coerce").astype(float).copy()
        for _ in range(max_iter):
            old = z.to_numpy(copy=True)
            z = z - z.groupby(entity).transform("mean")
            z = z - z.groupby(year).transform("mean")
            if len(z):
                delta_arr = np.abs(z.to_numpy() - old)
                delta = 0.0 if np.isnan(delta_arr).all() else float(np.nanmax(delta_arr))
            else:
                delta = 0.0
            if delta < tol:
                break
        absorbed[col] = z
    return absorbed


def _fit_absorbed_ols(
    use: pd.DataFrame, dependent: str, rhs: list[str]
) -> sm.regression.linear_model.RegressionResultsWrapper:
    absorbed = _two_way_absorb(use, [dependent, *rhs])
    y = absorbed[dependent]
    X = absorbed[rhs]
    return sm.OLS(y, X).fit(cov_type="cluster", cov_kwds={"groups": use["cik"].astype(str)})


def _fit_model(
    panel: pd.DataFrame,
    outcome: str,
    label: str,
    controls: list[str],
    panel_name: str,
    *,
    include_post_chatgpt_interaction: bool = False,
) -> dict[str, object]:
    rhs = [*FOCAL_TERMS]
    if include_post_chatgpt_interaction:
        interaction_name = "PatentMismatch_x_PostChatGPT"
        panel = panel.copy()
        panel[interaction_name] = panel["PatentMismatch"] * panel["PostChatGPT"]
        rhs.append(interaction_name)
    rhs.extend(controls)
    needed = [outcome, "cik", "year", *rhs]
    use = panel.dropna(subset=needed).copy()
    finite_mask = np.isfinite(use[[outcome, *rhs]].to_numpy(dtype=float)).all(axis=1)
    use = use.loc[finite_mask].copy()
    use = use.loc[use["cik"].str.len().gt(0)].copy()
    use["year"] = pd.to_numeric(use["year"], errors="coerce").astype(int)
    result = _fit_absorbed_ols(use, outcome, rhs)
    return {
        "panel_name": panel_name,
        "outcome": outcome,
        "outcome_label": label,
        "nobs": int(result.nobs),
        "adj_r_squared": float(result.rsquared_adj),
        "outcome_mean": float(use[outcome].mean()),
        "params": result.params.to_dict(),
        "bse": result.bse.to_dict(),
        "pvalues": result.pvalues.to_dict(),
    }


def _stars(p_value: float | None) -> str:
    if p_value is None or not math.isfinite(p_value):
        return ""
    if p_value < 0.01:
        return "***"
    if p_value < 0.05:
        return "**"
    if p_value < 0.10:
        return "*"
    return ""


def _coef_cell(coef: float | None, se: float | None, p_value: float | None) -> str:
    if coef is None or se is None or not math.isfinite(coef) or not math.isfinite(se):
        return ""
    return f"{coef:.4f}{_stars(p_value)} ({se:.4f})"


def _build_results_table(
    model_rows: list[dict[str, object]],
    *,
    include_post_chatgpt_interaction: bool = False,
) -> tuple[pd.DataFrame, dict[tuple[str, str], dict[str, object]]]:
    keyed = {(row["panel_name"], row["outcome"]): row for row in model_rows}
    rows: list[dict[str, object]] = []
    panel_specs = [
        ("Panel A. Valuation outcomes", VALUATION_SPECS),
        ("Panel B. Financing outcomes", FINANCING_SPECS),
    ]
    row_terms = ["PatentMismatch"]
    if include_post_chatgpt_interaction:
        row_terms.append("PatentMismatch x PostChatGPT")
    row_terms.extend(["A_S", "AI_Focus", "Outcome mean", "N", "Adj. R-squared"])
    for panel_name, specs in panel_specs:
        for term in row_terms:
            row = {"panel": panel_name, "row_label": term}
            for outcome, label, _controls in specs:
                result = keyed[(panel_name, outcome)]
                if term in FOCAL_TERMS or term == "PatentMismatch x PostChatGPT":
                    coef_term = (
                        "PatentMismatch_x_PostChatGPT"
                        if term == "PatentMismatch x PostChatGPT"
                        else term
                    )
                    row[label] = _coef_cell(
                        result["params"].get(coef_term),
                        result["bse"].get(coef_term),
                        result["pvalues"].get(coef_term),
                    )
                elif term == "Outcome mean":
                    row[label] = f"{result['outcome_mean']:.4f}"
                elif term == "N":
                    row[label] = str(result["nobs"])
                else:
                    row[label] = f"{result['adj_r_squared']:.3f}"
            rows.append(row)
    return pd.DataFrame(rows), keyed


def _markdown_table(headers: list[str], rows: list[list[object]]) -> str:
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join(["---"] * len(headers)) + " |",
    ]
    for row in rows:
        lines.append("| " + " | ".join(str(item) for item in row) + " |")
    return "\n".join(lines)


def _render_table_outputs(table_df: pd.DataFrame) -> tuple[str, str]:
    md_lines = ["# Table Main", ""]
    latex_lines = [
        "\\begin{table}[!htbp]",
        "\\centering",
        "\\caption{Financing and valuation consequences of disclosure credibility}",
        "\\begin{tabular}{lccccc}",
        "\\hline",
    ]
    for panel_name, specs in [
        ("Panel A. Valuation outcomes", VALUATION_SPECS),
        ("Panel B. Financing outcomes", FINANCING_SPECS),
    ]:
        headers = ["Row"] + [label for _outcome, label, _controls in specs]
        panel_rows = table_df.loc[table_df["panel"].eq(panel_name)].copy()
        rendered = panel_rows[
            ["row_label", *[label for _outcome, label, _controls in specs]]
        ].values.tolist()
        md_lines.extend([f"## {panel_name}", _markdown_table(headers, rendered), ""])
        latex_lines.append(
            f"\\multicolumn{{{len(headers)}}}{{l}}{{\\textit{{{panel_name}}}}} \\\\"
        )
        latex_lines.append(" & ".join(headers) + " \\\\")
        for row in rendered:
            latex_lines.append(" & ".join(str(item) for item in row) + " \\\\")
    latex_lines.extend(["\\hline", "\\end{tabular}", "\\end{table}"])
    return "\n".join(md_lines), "\n".join(latex_lines) + "\n"


def _docx_panel_rows(
    table_df: pd.DataFrame, panel_name: str, headers: list[str]
) -> list[list[tuple[str, bool]]]:
    rows = []
    subset = table_df.loc[table_df["panel"].eq(panel_name)].copy()
    for row in subset[["row_label", *headers[1:]]].itertuples(index=False):
        rows.append([(str(row[0]), True), *[(str(value), False) for value in row[1:]]])
    return rows


def _build_table_docx(table_df: pd.DataFrame, output_path: Path) -> None:
    document = Document()
    _set_document_defaults(document)
    _set_landscape(document)
    _add_title(document, "Table 8. Financing and Valuation Consequences")
    note = (
        "This table reports valuation and financing consequences associated with disclosure credibility. "
        "Panel A uses log market-capitalization-to-assets, a log Q-style proxy equal to log(1 + market-capitalization-to-assets plus leverage), "
        "and the next-year change in that log Q proxy. Panel B uses next-year change in CRSP shares outstanding and an indicator for share growth above 5%. "
        "The key regressors are PatentMismatch, the A/S ratio, and AI Focus. All specifications absorb firm and year fixed effects and use firm-clustered standard errors."
    )
    _add_note(document, note)
    for panel_name, specs in [
        ("Panel A. Valuation outcomes", VALUATION_SPECS),
        ("Panel B. Financing outcomes", FINANCING_SPECS),
    ]:
        p = document.add_paragraph()
        p.add_run(panel_name).bold = True
        headers = ["", *[label for _outcome, label, _controls in specs]]
        _build_panel_table(document, headers, _docx_panel_rows(table_df, panel_name, headers))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def _plot_patent_mismatch_coefficients(
    model_rows: list[dict[str, object]], output_base: Path
) -> tuple[Path, Path]:
    _base_style()
    fig, ax = plt.subplots(figsize=(8.4, 4.8), constrained_layout=True)
    order = [
        ("log_mktcap_assets", "Log MktCap/assets"),
        ("log_q_proxy", "Log Q proxy"),
        ("delta_log_q_proxy_lead1", "Delta log Q t+1"),
        ("share_growth_lead1", "Delta Shares t+1"),
        ("equity_issue_lead1", "Issue>5% t+1"),
    ]
    y_map = {outcome: idx for idx, (outcome, _label) in enumerate(order[::-1])}
    color_map = {
        "Panel A. Valuation outcomes": "#c46b48",
        "Panel B. Financing outcomes": "#6d597a",
    }
    for row in model_rows:
        coef = row["params"].get("PatentMismatch", math.nan)
        se = row["bse"].get("PatentMismatch", math.nan)
        if not (math.isfinite(coef) and math.isfinite(se)):
            continue
        ax.errorbar(
            coef,
            y_map[row["outcome"]],
            xerr=1.96 * se,
            fmt="o",
            color=color_map[row["panel_name"]],
            capsize=4,
            linewidth=1.6,
            label=row["panel_name"],
        )
    ax.axvline(0, color="#6c757d", linewidth=0.9)
    ax.set_yticks(range(len(order)))
    ax.set_yticklabels([label for _outcome, label in order][::-1])
    ax.set_xlabel("PatentMismatch coefficient")
    ax.set_title("PatentMismatch, Valuation, and Financing Outcomes")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.grid(True, axis="x", color="#d9d9d9", linewidth=0.7)
    ax.grid(False, axis="y")
    handles, labels = ax.get_legend_handles_labels()
    dedup: dict[str, object] = {}
    for handle, label in zip(handles, labels, strict=False):
        dedup.setdefault(label, handle)
    ax.legend(dedup.values(), dedup.keys(), frameon=False, loc="best")

    png_path = output_base.with_suffix(".png")
    pdf_path = output_base.with_suffix(".pdf")
    fig.savefig(png_path, dpi=220, bbox_inches="tight")
    fig.savefig(pdf_path, bbox_inches="tight")
    plt.close(fig)
    return png_path, pdf_path


def _result_notes(
    model_keyed: dict[tuple[str, str], dict[str, object]],
    summary: dict[str, object],
    *,
    include_post_chatgpt_interaction: bool = False,
) -> str:
    subset = summary.get("size_subset", "all")
    lines = [
        "# Result Notes",
        "",
        f"- Size subset: `{subset}`.",
        f"- Valuation (`Log MktCap/assets`): PatentMismatch `{model_keyed[('Panel A. Valuation outcomes', 'log_mktcap_assets')]['params'].get('PatentMismatch', math.nan):.4f}` (p=`{model_keyed[('Panel A. Valuation outcomes', 'log_mktcap_assets')]['pvalues'].get('PatentMismatch', math.nan):.3f}`).",
        f"- Valuation (`Log Q proxy`): PatentMismatch `{model_keyed[('Panel A. Valuation outcomes', 'log_q_proxy')]['params'].get('PatentMismatch', math.nan):.4f}` (p=`{model_keyed[('Panel A. Valuation outcomes', 'log_q_proxy')]['pvalues'].get('PatentMismatch', math.nan):.3f}`).",
        f"- Financing (`Delta Shares t+1`): PatentMismatch `{model_keyed[('Panel B. Financing outcomes', 'share_growth_lead1')]['params'].get('PatentMismatch', math.nan):.4f}` (p=`{model_keyed[('Panel B. Financing outcomes', 'share_growth_lead1')]['pvalues'].get('PatentMismatch', math.nan):.3f}`).",
        f"- Financing (`Issue>5% t+1`): PatentMismatch `{model_keyed[('Panel B. Financing outcomes', 'equity_issue_lead1')]['params'].get('PatentMismatch', math.nan):.4f}` (p=`{model_keyed[('Panel B. Financing outcomes', 'equity_issue_lead1')]['pvalues'].get('PatentMismatch', math.nan):.3f}`).",
    ]
    if include_post_chatgpt_interaction:
        lines.append(
            f"- Interaction (`PatentMismatch x PostChatGPT`) on `Log MktCap/assets`: `{model_keyed[('Panel A. Valuation outcomes', 'log_mktcap_assets')]['params'].get('PatentMismatch_x_PostChatGPT', math.nan):.4f}` (p=`{model_keyed[('Panel A. Valuation outcomes', 'log_mktcap_assets')]['pvalues'].get('PatentMismatch_x_PostChatGPT', math.nan):.3f}`)."
        )
        lines.append(
            f"- Interaction (`PatentMismatch x PostChatGPT`) on `Log Q proxy`: `{model_keyed[('Panel A. Valuation outcomes', 'log_q_proxy')]['params'].get('PatentMismatch_x_PostChatGPT', math.nan):.4f}` (p=`{model_keyed[('Panel A. Valuation outcomes', 'log_q_proxy')]['pvalues'].get('PatentMismatch_x_PostChatGPT', math.nan):.3f}`)."
        )
    lines.extend(
        [
            f"- Coverage: `{summary['firm_year_rows']}` rows, `{summary['valuation_marketcap_nonmissing']}` rows with market cap, `{summary['share_growth_nonmissing']}` rows with next-year share growth.",
            "",
        ]
    )
    return "\n".join(lines)


def _writer_packet(
    args: argparse.Namespace, summary: dict[str, object], model_rows: list[dict[str, object]]
) -> str:
    keyed = {(row["panel_name"], row["outcome"]): row for row in model_rows}
    return "\n".join(
        [
            "# Writer Packet",
            "",
            "## Metadata",
            f"- Test id: `{TEST_ID}`",
            f"- Run id: `{args.run_id}`",
            f"- Date run: `{date.today().isoformat()}`",
            f"- Script/module path: `{MODULE_PATH}`",
            f"- Input annual panel: `{args.annual_panel}`",
            f"- Input market features: `{args.market_features}`",
            "",
            "## Design",
            f"- Size subset: `{args.size_subset}`",
            f"- Post-ChatGPT interaction: `{args.include_post_chatgpt_interaction}`",
            "- Core regressors: `PatentMismatch`, `A_S`, `AI_Focus`",
            "- Panel A controls: `ln_assets`, `cash`, `roa`",
            "- Panel B controls: `ln_assets`, `leverage`, `cash`, `roa`",
            "- Fixed effects: `firm and year`",
            "- Clustering: `firm (cik)`",
            f"- Q proxy rule: `{summary['q_proxy_rule']}`",
            f"- Issuance rule: `{summary['issue_rule']}`",
            f"- Size rule: `{summary['size_rule']}`",
            "- Interaction note: `PostChatGPT is absorbed by year fixed effects; only PatentMismatch x PostChatGPT is added when requested`",
            "",
            "## Main Results",
            f"- Log MktCap/assets PatentMismatch: `{keyed[('Panel A. Valuation outcomes', 'log_mktcap_assets')]['params'].get('PatentMismatch', math.nan):.4f}` (p=`{keyed[('Panel A. Valuation outcomes', 'log_mktcap_assets')]['pvalues'].get('PatentMismatch', math.nan):.3f}`)",
            f"- Log Q proxy PatentMismatch: `{keyed[('Panel A. Valuation outcomes', 'log_q_proxy')]['params'].get('PatentMismatch', math.nan):.4f}` (p=`{keyed[('Panel A. Valuation outcomes', 'log_q_proxy')]['pvalues'].get('PatentMismatch', math.nan):.3f}`)",
            f"- Delta Shares t+1 PatentMismatch: `{keyed[('Panel B. Financing outcomes', 'share_growth_lead1')]['params'].get('PatentMismatch', math.nan):.4f}` (p=`{keyed[('Panel B. Financing outcomes', 'share_growth_lead1')]['pvalues'].get('PatentMismatch', math.nan):.3f}`)",
            f"- Issue>5% t+1 PatentMismatch: `{keyed[('Panel B. Financing outcomes', 'equity_issue_lead1')]['params'].get('PatentMismatch', math.nan):.4f}` (p=`{keyed[('Panel B. Financing outcomes', 'equity_issue_lead1')]['pvalues'].get('PatentMismatch', math.nan):.3f}`)",
            "",
            "## Caption Draft",
            "This table reports financing and valuation consequences associated with disclosure credibility. Panel A uses log market-capitalization-to-assets, a log Q-style proxy, and the next-year change in that log Q proxy. Panel B uses next-year change in CRSP shares outstanding and an indicator for substantial share growth. The key regressors are PatentMismatch, the A/S ratio, and AI Focus. All specifications absorb firm and year fixed effects and use firm-clustered standard errors.",
            "",
        ]
    )


def _dataset_summary(
    args: argparse.Namespace, summary: dict[str, object], model_rows: list[dict[str, object]]
) -> dict[str, object]:
    return {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "annual_panel": str(args.annual_panel),
        "market_features": str(args.market_features),
        "size_subset": args.size_subset,
        "include_post_chatgpt_interaction": args.include_post_chatgpt_interaction,
        "panel_summary": summary,
        "model_rows": model_rows,
    }


def _copy_exports(run_dir: Path, paper_root: Path, run_id: str) -> dict[str, str]:
    exports = {
        "figure_png": paper_root / "figures" / f"{TEST_ID}_{run_id}.png",
        "figure_pdf": paper_root / "figures" / f"{TEST_ID}_{run_id}.pdf",
        "table_csv": paper_root / "tables" / f"{TEST_ID}_{run_id}.csv",
        "table_md": paper_root / "tables" / f"{TEST_ID}_{run_id}.md",
        "table_tex": paper_root / "latex" / f"{TEST_ID}_{run_id}.tex",
        "table_docx": paper_root / "docx" / f"{TEST_ID}_{run_id}.docx",
        "figure_series": paper_root / "tables" / f"{TEST_ID}_{run_id}_figure_series.csv",
        "writer_packet": paper_root / "writer_packets" / f"{TEST_ID}_{run_id}.md",
        "result_notes": paper_root / "snippets" / f"{TEST_ID}_{run_id}_result_notes.md",
    }
    for path in exports.values():
        path.parent.mkdir(parents=True, exist_ok=True)
    mapping = {
        run_dir / "figure_main.png": exports["figure_png"],
        run_dir / "figure_main.pdf": exports["figure_pdf"],
        run_dir / "table_main.csv": exports["table_csv"],
        run_dir / "table_main.md": exports["table_md"],
        run_dir / "table_main.tex": exports["table_tex"],
        run_dir / "table_main.docx": exports["table_docx"],
        run_dir / "figure_series.csv": exports["figure_series"],
        run_dir / "writer_packet.md": exports["writer_packet"],
        run_dir / "result_notes.md": exports["result_notes"],
    }
    for src, dst in mapping.items():
        shutil.copy2(src, dst)
    return {key: str(path) for key, path in exports.items()}


def main() -> None:
    args = _parse_args()
    run_dir = args.test_root / args.run_id
    run_dir.mkdir(parents=True, exist_ok=True)

    panel = _load_panel(args.annual_panel, args.market_features)
    panel, summary = _prepare_panel(panel)
    panel, summary = _apply_size_subset(panel, summary, args.size_subset)

    model_rows = []
    for outcome, label, controls in VALUATION_SPECS:
        model_rows.append(
            _fit_model(
                panel,
                outcome,
                label,
                controls,
                "Panel A. Valuation outcomes",
                include_post_chatgpt_interaction=args.include_post_chatgpt_interaction,
            )
        )
    for outcome, label, controls in FINANCING_SPECS:
        model_rows.append(
            _fit_model(
                panel,
                outcome,
                label,
                controls,
                "Panel B. Financing outcomes",
                include_post_chatgpt_interaction=args.include_post_chatgpt_interaction,
            )
        )

    table_df, model_keyed = _build_results_table(
        model_rows,
        include_post_chatgpt_interaction=args.include_post_chatgpt_interaction,
    )
    table_df.to_csv(run_dir / "table_main.csv", index=False)
    figure_df = pd.DataFrame(
        [
            {
                "panel": row["panel_name"],
                "outcome": row["outcome_label"],
                "coef": row["params"].get("PatentMismatch", math.nan),
                "se": row["bse"].get("PatentMismatch", math.nan),
                "p_value": row["pvalues"].get("PatentMismatch", math.nan),
            }
            for row in model_rows
        ]
    )
    figure_df.to_csv(run_dir / "figure_series.csv", index=False)

    table_md, table_tex = _render_table_outputs(table_df)
    (run_dir / "table_main.md").write_text(table_md, encoding="utf-8")
    (run_dir / "table_main.tex").write_text(table_tex, encoding="utf-8")
    _build_table_docx(table_df, run_dir / "table_main.docx")
    png_path, pdf_path = _plot_patent_mismatch_coefficients(model_rows, run_dir / "figure_main")

    (run_dir / "result_notes.md").write_text(
        _result_notes(
            model_keyed,
            summary,
            include_post_chatgpt_interaction=args.include_post_chatgpt_interaction,
        ),
        encoding="utf-8",
    )
    (run_dir / "writer_packet.md").write_text(
        _writer_packet(args, summary, model_rows), encoding="utf-8"
    )
    dataset_summary = _dataset_summary(args, summary, model_rows)
    (run_dir / "dataset_summary.json").write_text(
        json.dumps(dataset_summary, indent=2), encoding="utf-8"
    )

    paper_exports = _copy_exports(run_dir, args.paper_root, args.run_id)
    manifest = {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "module_path": MODULE_PATH,
        "inputs": {
            "annual_panel": str(args.annual_panel),
            "market_features": str(args.market_features),
        },
        "run_dir": str(run_dir),
        "outputs": {
            "dataset_summary": str(run_dir / "dataset_summary.json"),
            "table_csv": str(run_dir / "table_main.csv"),
            "table_md": str(run_dir / "table_main.md"),
            "table_tex": str(run_dir / "table_main.tex"),
            "table_docx": str(run_dir / "table_main.docx"),
            "figure_series": str(run_dir / "figure_series.csv"),
            "figure_png": str(png_path),
            "figure_pdf": str(pdf_path),
            "writer_packet": str(run_dir / "writer_packet.md"),
            "result_notes": str(run_dir / "result_notes.md"),
        },
        "paper_exports": paper_exports,
    }
    (run_dir / "run_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    print(f"[{TEST_ID}] wrote run bundle to {run_dir}")
    print(f"[{TEST_ID}] paper table: {paper_exports['table_docx']}")


if __name__ == "__main__":
    main()
