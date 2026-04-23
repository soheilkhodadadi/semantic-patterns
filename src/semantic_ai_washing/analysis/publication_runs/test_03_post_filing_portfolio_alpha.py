"""Publication run driver for Test 03 companion: post-filing portfolio alpha."""

from __future__ import annotations

import argparse
import json
import math
import shutil
from datetime import UTC, date, datetime
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import statsmodels.api as sm
from docx import Document

from semantic_ai_washing.analysis.publication_runs.test_03_post_filing_drift import (
    _add_note,
    _add_title,
    _build_analysis_sample,
    _build_panel_table,
    _set_document_defaults,
    _set_landscape,
)

REPO_ROOT = Path(__file__).resolve().parents[4]
DEFAULT_EVENT_PANEL = (
    REPO_ROOT
    / "data/processed/panel/filing_event_estimation_sample_hybrid_api_a_conf49_v1.parquet"
)
DEFAULT_ANNUAL_PANEL = (
    REPO_ROOT
    / "data/processed/panel/canonical/ever_speaker_panel_2016_2025_hybrid_api_a_conf49_v1.parquet"
)
DEFAULT_MONTHLY_RETURNS = REPO_ROOT / "data/interim/market/wrds_crsp_msf_full_sample_v1.parquet"
DEFAULT_MARKET_INDEX = REPO_ROOT / "data/interim/market/wrds_crsp_msi_full_sample_v1.parquet"
DEFAULT_TEST_ROOT = Path(
    "/Users/soheilkhodadadi/DataWork/semantic-patterns/ai_washing/derived/test_runs/test_03_post_filing_portfolio_alpha"
)
DEFAULT_PAPER_ROOT = REPO_ROOT / "paper/generated"
DEFAULT_RUN_ID = f"{date.today():%Y%m%d}_hybrid_api_a_conf49_main_v1"
TEST_ID = "test_03_post_filing_portfolio_alpha"
MODULE_PATH = "semantic_ai_washing.analysis.publication_runs.test_03_post_filing_portfolio_alpha"
HORIZONS = {1: "1m", 3: "3m", 6: "6m", 12: "12m"}
WEIGHT_LABELS = {"ew": "Equal-weight", "vw": "Value-weight"}


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--event-panel", type=Path, default=DEFAULT_EVENT_PANEL)
    parser.add_argument("--annual-panel", type=Path, default=DEFAULT_ANNUAL_PANEL)
    parser.add_argument("--monthly-returns", type=Path, default=DEFAULT_MONTHLY_RETURNS)
    parser.add_argument("--market-index", type=Path, default=DEFAULT_MARKET_INDEX)
    parser.add_argument("--test-root", type=Path, default=DEFAULT_TEST_ROOT)
    parser.add_argument("--paper-root", type=Path, default=DEFAULT_PAPER_ROOT)
    parser.add_argument("--run-id", default=DEFAULT_RUN_ID)
    return parser.parse_args()


def _base_style() -> None:
    plt.style.use("seaborn-v0_8-whitegrid")
    plt.rcParams.update(
        {
            "font.family": "serif",
            "font.serif": ["Times New Roman", "Times", "DejaVu Serif"],
            "axes.titlesize": 14,
            "axes.labelsize": 11,
            "xtick.labelsize": 10,
            "ytick.labelsize": 10,
            "legend.fontsize": 10,
        }
    )


def _load_monthly_returns(path: Path) -> pd.DataFrame:
    df = pd.read_parquet(path).copy()
    df["permno"] = df["permno"].astype(str)
    df["date"] = pd.to_datetime(df["date"])
    df["month"] = df["date"].dt.to_period("M")
    for column in ["ret", "prc", "shrout"]:
        df[column] = pd.to_numeric(df[column], errors="coerce")
    df["mcap"] = (df["prc"].abs() * df["shrout"]).replace({0: np.nan})
    df = df.sort_values(["permno", "date"]).reset_index(drop=True)
    df["lag_mcap"] = df.groupby("permno")["mcap"].shift(1)
    return df[["permno", "month", "date", "ret", "lag_mcap"]]


def _load_market_index(path: Path) -> pd.DataFrame:
    df = pd.read_parquet(path).copy()
    df["date"] = pd.to_datetime(df["date"])
    df["month"] = df["date"].dt.to_period("M")
    df["vwretd"] = pd.to_numeric(df["vwretd"], errors="coerce")
    return df[["month", "vwretd"]]


def _prepare_signal_sample(event_panel: Path, annual_panel: Path) -> pd.DataFrame:
    sample = _build_analysis_sample(event_panel, annual_panel)
    sample = sample.loc[sample["is_ai_filing"]].copy()
    sample["PatentMismatch"] = pd.to_numeric(sample["PatentMismatch"], errors="coerce")
    sample = sample.loc[sample["PatentMismatch"].notna()].copy()
    sample["PatentMismatch"] = sample["PatentMismatch"].fillna(0).astype(int)
    sample["permno"] = sample["permno"].astype(str)
    sample["filing_date"] = pd.to_datetime(sample["filing_date"])
    sample["filing_month"] = sample["filing_date"].dt.to_period("M")
    return sample


def _build_positions(
    sample: pd.DataFrame, horizon_months: int, *, split_flag: int | None = None
) -> pd.DataFrame:
    use = sample.copy()
    if split_flag is not None:
        use = use.loc[use["post_chatgpt"].eq(split_flag)].copy()
    positions = []
    for row in use[["permno", "PatentMismatch", "filing_month", "filing_date"]].itertuples(
        index=False
    ):
        for k in range(1, horizon_months + 1):
            positions.append(
                (row.permno, row.PatentMismatch, row.filing_month + k, row.filing_date)
            )
    out = pd.DataFrame(positions, columns=["permno", "PatentMismatch", "month", "filing_date"])
    out = out.sort_values(["permno", "month", "filing_date"]).drop_duplicates(
        ["permno", "month"], keep="last"
    )
    return out.reset_index(drop=True)


def _weighted_average(returns: pd.Series, weights: pd.Series) -> float:
    w = pd.to_numeric(weights, errors="coerce")
    r = pd.to_numeric(returns, errors="coerce")
    if w.notna().sum() == 0 or w.fillna(0).sum() == 0:
        return float(r.mean())
    return float(np.average(r, weights=w.fillna(0)))


def _compute_long_short_series(
    positions: pd.DataFrame,
    monthly_returns: pd.DataFrame,
    market_index: pd.DataFrame,
) -> pd.DataFrame:
    merged = positions.merge(monthly_returns, on=["permno", "month"], how="left")
    merged = merged.dropna(subset=["ret"]).copy()

    ew = (
        merged.groupby(["month", "PatentMismatch"], as_index=False)["ret"]
        .mean()
        .rename(columns={"ret": "ew_ret"})
    )
    vw = (
        merged.groupby(["month", "PatentMismatch"])
        .apply(lambda g: _weighted_average(g["ret"], g["lag_mcap"]))
        .reset_index(name="vw_ret")
    )
    portfolios = ew.merge(vw, on=["month", "PatentMismatch"], how="inner")

    ew_pivot = portfolios.pivot(index="month", columns="PatentMismatch", values="ew_ret").rename(
        columns={0: "non_mismatch", 1: "mismatch"}
    )
    vw_pivot = portfolios.pivot(index="month", columns="PatentMismatch", values="vw_ret").rename(
        columns={0: "non_mismatch", 1: "mismatch"}
    )
    out = pd.DataFrame(index=sorted(set(ew_pivot.index) | set(vw_pivot.index)))
    out.index.name = "month"
    out["ew_long_short"] = ew_pivot["non_mismatch"] - ew_pivot["mismatch"]
    out["vw_long_short"] = vw_pivot["non_mismatch"] - vw_pivot["mismatch"]
    out = out.reset_index().merge(market_index, on="month", how="left")
    return out.sort_values("month").reset_index(drop=True)


def _summarize_long_short(
    series_df: pd.DataFrame, weight_col: str, *, horizon_label: str
) -> dict[str, object]:
    use = series_df.dropna(subset=[weight_col, "vwretd"]).copy()
    X = sm.add_constant(use["vwretd"])
    result = sm.OLS(use[weight_col], X).fit()
    monthly_mean = float(use[weight_col].mean())
    annualized_simple = float(monthly_mean * 12.0)
    annualized_compound = float((1.0 + monthly_mean) ** 12 - 1.0)
    return {
        "horizon": horizon_label,
        "weighting": weight_col[:2],
        "mean_monthly_return": monthly_mean,
        "annualized_simple_return": annualized_simple,
        "annualized_compound_return": annualized_compound,
        "market_alpha": float(result.params.get("const", math.nan)),
        "market_alpha_t": float(result.tvalues.get("const", math.nan)),
        "market_alpha_p": float(result.pvalues.get("const", math.nan)),
        "beta_to_market": float(result.params.get("vwretd", math.nan)),
        "n_months": int(len(use)),
    }


def _subperiod_summary(
    sample: pd.DataFrame, monthly_returns: pd.DataFrame, market_index: pd.DataFrame
) -> list[dict[str, object]]:
    rows = []
    for split_flag, split_label in [(0, "Pre-ChatGPT filings"), (1, "Post-ChatGPT filings")]:
        positions = _build_positions(sample, 3, split_flag=split_flag)
        series = _compute_long_short_series(positions, monthly_returns, market_index)
        for weight_col in ["ew_long_short", "vw_long_short"]:
            stats = _summarize_long_short(series, weight_col, horizon_label="3m")
            rows.append(
                {
                    "split": split_label,
                    "weighting": stats["weighting"],
                    "mean_monthly_return": stats["mean_monthly_return"],
                    "market_alpha": stats["market_alpha"],
                    "market_alpha_p": stats["market_alpha_p"],
                    "n_months": stats["n_months"],
                }
            )
    return rows


def _plot_cumulative_long_short(series_df: pd.DataFrame, output_base: Path) -> tuple[Path, Path]:
    _base_style()
    plot_df = series_df.dropna(subset=["ew_long_short", "vw_long_short"]).copy()
    plot_df["month_end"] = plot_df["month"].dt.to_timestamp("M")
    plot_df["cum_ew"] = (1.0 + plot_df["ew_long_short"]).cumprod() - 1.0
    plot_df["cum_vw"] = (1.0 + plot_df["vw_long_short"]).cumprod() - 1.0

    fig, ax = plt.subplots(figsize=(8.4, 4.8), constrained_layout=True)
    ax.plot(
        plot_df["month_end"],
        100 * plot_df["cum_ew"],
        color="#1d3557",
        linewidth=2.1,
        label="3m EW long credible / short mismatch",
    )
    ax.plot(
        plot_df["month_end"],
        100 * plot_df["cum_vw"],
        color="#c46b48",
        linewidth=2.1,
        label="3m VW long credible / short mismatch",
    )
    ax.axhline(0, color="#6c757d", linewidth=0.8)
    ax.set_ylabel("Cumulative portfolio return (pct)")
    ax.set_xlabel("Calendar month")
    ax.set_title("Calendar-Time Long-Short Returns After AI Filing Signal")
    ax.legend(frameon=False, loc="best")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.grid(True, axis="y", color="#d9d9d9", linewidth=0.7)
    ax.grid(False, axis="x")

    png_path = output_base.with_suffix(".png")
    pdf_path = output_base.with_suffix(".pdf")
    fig.savefig(png_path, dpi=220, bbox_inches="tight")
    fig.savefig(pdf_path, bbox_inches="tight")
    plt.close(fig)
    return png_path, pdf_path


def _markdown_table(headers: list[str], rows: list[list[object]]) -> str:
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join(["---"] * len(headers)) + " |",
    ]
    for row in rows:
        lines.append("| " + " | ".join(str(item) for item in row) + " |")
    return "\n".join(lines)


def _render_table_outputs(
    summary_rows: list[dict[str, object]], split_rows: list[dict[str, object]]
) -> tuple[str, str]:
    headers = ["Row"] + [HORIZONS[h] for h in HORIZONS]
    by_key = {(row["weighting"], row["horizon"]): row for row in summary_rows}

    def horizon_vals(weighting: str, key: str, *, pct: bool = False) -> list[str]:
        vals = []
        for horizon in HORIZONS.values():
            value = by_key[(weighting, horizon)][key]
            if pct:
                vals.append(f"{100 * value:.3f}")
            else:
                vals.append(f"{value:.4f}")
        return vals

    panel_a_rows = [
        ["Mean monthly return"] + horizon_vals("ew", "mean_monthly_return", pct=True),
        ["Annualized simple return"] + horizon_vals("ew", "annualized_simple_return", pct=True),
        ["Market alpha"] + horizon_vals("ew", "market_alpha", pct=True),
        ["Alpha p-value"] + horizon_vals("ew", "market_alpha_p"),
        ["Beta to market"] + horizon_vals("ew", "beta_to_market"),
        ["Months"] + [str(by_key[("ew", h)]["n_months"]) for h in HORIZONS.values()],
    ]
    panel_b_rows = [
        ["Mean monthly return"] + horizon_vals("vw", "mean_monthly_return", pct=True),
        ["Annualized simple return"] + horizon_vals("vw", "annualized_simple_return", pct=True),
        ["Market alpha"] + horizon_vals("vw", "market_alpha", pct=True),
        ["Alpha p-value"] + horizon_vals("vw", "market_alpha_p"),
        ["Beta to market"] + horizon_vals("vw", "beta_to_market"),
        ["Months"] + [str(by_key[("vw", h)]["n_months"]) for h in HORIZONS.values()],
    ]

    split_headers = ["Row", "Pre-ChatGPT filings", "Post-ChatGPT filings"]
    split_map = {(row["split"], row["weighting"]): row for row in split_rows}
    panel_c_rows = [
        [
            "EW mean monthly return",
            f"{100 * split_map[('Pre-ChatGPT filings', 'ew')]['mean_monthly_return']:.3f}",
            f"{100 * split_map[('Post-ChatGPT filings', 'ew')]['mean_monthly_return']:.3f}",
        ],
        [
            "EW market alpha",
            f"{100 * split_map[('Pre-ChatGPT filings', 'ew')]['market_alpha']:.3f}",
            f"{100 * split_map[('Post-ChatGPT filings', 'ew')]['market_alpha']:.3f}",
        ],
        [
            "EW alpha p-value",
            f"{split_map[('Pre-ChatGPT filings', 'ew')]['market_alpha_p']:.3f}",
            f"{split_map[('Post-ChatGPT filings', 'ew')]['market_alpha_p']:.3f}",
        ],
        [
            "VW mean monthly return",
            f"{100 * split_map[('Pre-ChatGPT filings', 'vw')]['mean_monthly_return']:.3f}",
            f"{100 * split_map[('Post-ChatGPT filings', 'vw')]['mean_monthly_return']:.3f}",
        ],
        [
            "VW market alpha",
            f"{100 * split_map[('Pre-ChatGPT filings', 'vw')]['market_alpha']:.3f}",
            f"{100 * split_map[('Post-ChatGPT filings', 'vw')]['market_alpha']:.3f}",
        ],
        [
            "VW alpha p-value",
            f"{split_map[('Pre-ChatGPT filings', 'vw')]['market_alpha_p']:.3f}",
            f"{split_map[('Post-ChatGPT filings', 'vw')]['market_alpha_p']:.3f}",
        ],
        [
            "Months",
            str(split_map[("Pre-ChatGPT filings", "ew")]["n_months"]),
            str(split_map[("Post-ChatGPT filings", "ew")]["n_months"]),
        ],
    ]

    md = "\n".join(
        [
            "# Table Main",
            "",
            "## Panel A. Equal-weight long credible / short mismatch portfolios",
            _markdown_table(headers, panel_a_rows),
            "",
            "## Panel B. Value-weight long credible / short mismatch portfolios",
            _markdown_table(headers, panel_b_rows),
            "",
            "## Panel C. Three-month subperiod split",
            _markdown_table(split_headers, panel_c_rows),
            "",
        ]
    )

    latex_lines = [
        "\\begin{table}[!htbp]",
        "\\centering",
        "\\caption{Calendar-time long-short portfolio returns after AI disclosure filing signals}",
        "\\begin{tabular}{l" + "c" * len(headers[1:]) + "}",
        "\\hline",
        " & " + " & ".join(headers[1:]) + " \\\\",
        "\\hline",
    ]
    for row in panel_a_rows + [["--"] + ["--"] * len(headers[1:])] + panel_b_rows:
        latex_lines.append(
            str(row[0]) + " & " + " & ".join(str(item) for item in row[1:]) + " \\\\"
        )
    latex_lines.extend(["\\hline", "\\end{tabular}", "\\end{table}"])
    return md, "\n".join(latex_lines) + "\n"


def _build_table_docx(
    summary_rows: list[dict[str, object]], split_rows: list[dict[str, object]], output_path: Path
) -> None:
    document = Document()
    _set_document_defaults(document)
    _set_landscape(document)
    _add_title(
        document, "Table 3. Calendar-Time Long-Short Portfolio Returns After Filing Signals"
    )
    note = (
        "This table reports calendar-time long-short portfolio returns formed after AI-related annual filing signals. Each month, the long leg holds non-mismatch AI filers and the short leg holds mismatch AI filers that filed within the prior 1, 3, 6, or 12 months. "
        "Panel A reports equal-weight portfolios and Panel B reports value-weight portfolios using lagged market capitalization. Market alpha is the intercept from a regression of the long-short monthly return on the CRSP value-weighted market return. Panel C splits the three-month portfolio on whether the underlying filing signal occurred before or after the ChatGPT release regime shift."
    )
    _add_note(document, note)

    headers = [""] + [HORIZONS[h] for h in HORIZONS]
    by_key = {(row["weighting"], row["horizon"]): row for row in summary_rows}

    def row_values(weighting: str, key: str, *, pct: bool = False) -> list[tuple[str, bool]]:
        vals: list[tuple[str, bool]] = []
        for horizon in HORIZONS.values():
            value = by_key[(weighting, horizon)][key]
            vals.append((f"{100 * value:.3f}" if pct else f"{value:.4f}", False))
        return vals

    panel_a_rows = [
        [("Mean monthly return", True)] + row_values("ew", "mean_monthly_return", pct=True),
        [("Annualized simple return", True)]
        + row_values("ew", "annualized_simple_return", pct=True),
        [("Market alpha", True)] + row_values("ew", "market_alpha", pct=True),
        [("Alpha p-value", True)] + row_values("ew", "market_alpha_p"),
        [("Beta to market", True)] + row_values("ew", "beta_to_market"),
        [("Months", True)]
        + [(str(by_key[("ew", h)]["n_months"]), False) for h in HORIZONS.values()],
    ]
    p = document.add_paragraph()
    p.add_run("Panel A. Equal-weight long credible / short mismatch portfolios").bold = True
    _build_panel_table(document, headers, panel_a_rows)

    panel_b_rows = [
        [("Mean monthly return", True)] + row_values("vw", "mean_monthly_return", pct=True),
        [("Annualized simple return", True)]
        + row_values("vw", "annualized_simple_return", pct=True),
        [("Market alpha", True)] + row_values("vw", "market_alpha", pct=True),
        [("Alpha p-value", True)] + row_values("vw", "market_alpha_p"),
        [("Beta to market", True)] + row_values("vw", "beta_to_market"),
        [("Months", True)]
        + [(str(by_key[("vw", h)]["n_months"]), False) for h in HORIZONS.values()],
    ]
    p = document.add_paragraph()
    p.add_run("Panel B. Value-weight long credible / short mismatch portfolios").bold = True
    _build_panel_table(document, headers, panel_b_rows)

    split_map = {(row["split"], row["weighting"]): row for row in split_rows}
    split_headers = ["", "Pre-ChatGPT filings", "Post-ChatGPT filings"]
    panel_c_rows = [
        [
            ("EW mean monthly return", True),
            (
                f"{100 * split_map[('Pre-ChatGPT filings', 'ew')]['mean_monthly_return']:.3f}",
                False,
            ),
            (
                f"{100 * split_map[('Post-ChatGPT filings', 'ew')]['mean_monthly_return']:.3f}",
                False,
            ),
        ],
        [
            ("EW market alpha", True),
            (f"{100 * split_map[('Pre-ChatGPT filings', 'ew')]['market_alpha']:.3f}", False),
            (f"{100 * split_map[('Post-ChatGPT filings', 'ew')]['market_alpha']:.3f}", False),
        ],
        [
            ("EW alpha p-value", True),
            (f"{split_map[('Pre-ChatGPT filings', 'ew')]['market_alpha_p']:.3f}", False),
            (f"{split_map[('Post-ChatGPT filings', 'ew')]['market_alpha_p']:.3f}", False),
        ],
        [
            ("VW mean monthly return", True),
            (
                f"{100 * split_map[('Pre-ChatGPT filings', 'vw')]['mean_monthly_return']:.3f}",
                False,
            ),
            (
                f"{100 * split_map[('Post-ChatGPT filings', 'vw')]['mean_monthly_return']:.3f}",
                False,
            ),
        ],
        [
            ("VW market alpha", True),
            (f"{100 * split_map[('Pre-ChatGPT filings', 'vw')]['market_alpha']:.3f}", False),
            (f"{100 * split_map[('Post-ChatGPT filings', 'vw')]['market_alpha']:.3f}", False),
        ],
        [
            ("VW alpha p-value", True),
            (f"{split_map[('Pre-ChatGPT filings', 'vw')]['market_alpha_p']:.3f}", False),
            (f"{split_map[('Post-ChatGPT filings', 'vw')]['market_alpha_p']:.3f}", False),
        ],
        [
            ("Months", True),
            (str(split_map[("Pre-ChatGPT filings", "ew")]["n_months"]), False),
            (str(split_map[("Post-ChatGPT filings", "ew")]["n_months"]), False),
        ],
    ]
    p = document.add_paragraph()
    p.add_run("Panel C. Three-month subperiod split").bold = True
    _build_panel_table(document, split_headers, panel_c_rows)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def _result_notes(
    summary_rows: list[dict[str, object]], split_rows: list[dict[str, object]]
) -> str:
    summary = {(row["weighting"], row["horizon"]): row for row in summary_rows}
    split_map = {(row["split"], row["weighting"]): row for row in split_rows}
    return "\n".join(
        [
            "# Result Notes",
            "",
            f"- Equal-weight 3-month long-short portfolio mean monthly return: `{100 * summary[('ew', '3m')]['mean_monthly_return']:.3f}` percentage points; market alpha = `{100 * summary[('ew', '3m')]['market_alpha']:.3f}` pp with p=`{summary[('ew', '3m')]['market_alpha_p']:.3f}`.",
            f"- Value-weight 3-month long-short portfolio mean monthly return: `{100 * summary[('vw', '3m')]['mean_monthly_return']:.3f}` percentage points; market alpha = `{100 * summary[('vw', '3m')]['market_alpha']:.3f}` pp with p=`{summary[('vw', '3m')]['market_alpha_p']:.3f}`.",
            f"- Equal-weight 12-month market alpha: `{100 * summary[('ew', '12m')]['market_alpha']:.3f}` pp with p=`{summary[('ew', '12m')]['market_alpha_p']:.3f}`; value-weight 12-month alpha = `{100 * summary[('vw', '12m')]['market_alpha']:.3f}` pp with p=`{summary[('vw', '12m')]['market_alpha_p']:.3f}`.",
            f"- Three-month equal-weight alpha is stronger in the pre-ChatGPT filing subsample (`p={split_map[('Pre-ChatGPT filings', 'ew')]['market_alpha_p']:.3f}`) than in the post-ChatGPT filing subsample (`p={split_map[('Post-ChatGPT filings', 'ew')]['market_alpha_p']:.3f}`), but the post-period has only `{split_map[('Post-ChatGPT filings', 'ew')]['n_months']}` calendar months.",
            "",
        ]
    )


def _writer_packet(
    args: argparse.Namespace,
    summary_rows: list[dict[str, object]],
    split_rows: list[dict[str, object]],
) -> str:
    summary = {(row["weighting"], row["horizon"]): row for row in summary_rows}
    return "\n".join(
        [
            "# Writer Packet",
            "",
            "## Metadata",
            f"- Test id: `{TEST_ID}`",
            f"- Run id: `{args.run_id}`",
            f"- Date run: `{date.today().isoformat()}`",
            f"- Script/module path: `{MODULE_PATH}`",
            f"- Input files: `{args.event_panel}`, `{args.annual_panel}`, `{args.monthly_returns}`, `{args.market_index}`",
            "- Unit of observation: `calendar month portfolio return`",
            "- Sample filters: `AI-talking annual filers with PatentMismatch label; one active signal per stock-month using the latest filing`",
            "",
            "## Portfolio Construction",
            "- Long leg: `non-mismatch AI filers`",
            "- Short leg: `mismatch AI filers`",
            "- Holding windows: `1, 3, 6, 12 months after filing month, starting in month +1`",
            "- Equal-weight version: `simple average across active stock-month constituents`",
            "- Value-weight version: `lagged market-cap weights from CRSP MSF`",
            "- Overlap rule: `if multiple filing signals map to the same stock-month, retain the latest filing signal`",
            "",
            "## Estimation",
            "- Main statistic: `monthly long-short return and annualized equivalent`",
            "- Alpha model: `intercept from long-short monthly return regressed on CRSP value-weighted market return`",
            "- Market benchmark: `CRSP MSI vwretd`",
            "",
            "## Results",
            f"- Strongest portfolio signal: `EW 3m` alpha = `{100 * summary[('ew', '3m')]['market_alpha']:.3f}` pp/month, p = `{summary[('ew', '3m')]['market_alpha_p']:.3f}`",
            f"- VW 3m alpha: `{100 * summary[('vw', '3m')]['market_alpha']:.3f}` pp/month, p = `{summary[('vw', '3m')]['market_alpha_p']:.3f}`",
            f"- 12m portfolio result: `EW alpha = {100 * summary[('ew', '12m')]['market_alpha']:.3f} pp/month`, `VW alpha = {100 * summary[('vw', '12m')]['market_alpha']:.3f} pp/month`",
            "- One-sentence interpretation: `The cleaner finance-style cut suggests any post-filing return differential is concentrated in equal-weight short-to-medium horizons rather than in a broad 12-month value-weighted repricing.`",
            "",
            "## Caption Draft",
            "This table reports calendar-time long-short portfolio returns formed after annual AI-related filing signals. Each month, the long leg holds non-mismatch AI filers and the short leg holds mismatch AI filers that filed within the prior 1, 3, 6, or 12 months. Panel A reports equal-weight portfolios, Panel B reports value-weight portfolios using lagged market capitalization, and Panel C splits the three-month strategy into pre- and post-ChatGPT filing cohorts. Market alpha is the intercept from a regression of the long-short monthly return on the CRSP value-weighted market return.",
            "",
        ]
    )


def _dataset_summary(
    signal_sample: pd.DataFrame,
    summary_rows: list[dict[str, object]],
    split_rows: list[dict[str, object]],
    figure_series: pd.DataFrame,
    *,
    args: argparse.Namespace,
) -> dict[str, object]:
    figure_payload = figure_series.copy()
    if "month" in figure_payload.columns:
        figure_payload["month"] = figure_payload["month"].astype(str)
    return {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "event_panel": str(args.event_panel),
        "annual_panel": str(args.annual_panel),
        "monthly_returns": str(args.monthly_returns),
        "market_index": str(args.market_index),
        "signal_filing_count": int(len(signal_sample)),
        "unique_permno": int(signal_sample["permno"].nunique()),
        "filing_year_min": int(signal_sample["filing_year"].min()),
        "filing_year_max": int(signal_sample["filing_year"].max()),
        "portfolio_summaries": summary_rows,
        "subperiod_summary": split_rows,
        "figure_series": figure_payload.to_dict(orient="records"),
    }


def _copy_exports(run_dir: Path, paper_root: Path, run_id: str) -> dict[str, str]:
    exports = {
        "figure_png": paper_root / "figures" / f"{TEST_ID}_{run_id}.png",
        "figure_pdf": paper_root / "figures" / f"{TEST_ID}_{run_id}.pdf",
        "table_csv": paper_root / "tables" / f"{TEST_ID}_{run_id}.csv",
        "table_md": paper_root / "tables" / f"{TEST_ID}_{run_id}.md",
        "table_tex": paper_root / "latex" / f"{TEST_ID}_{run_id}.tex",
        "table_docx": paper_root / "docx" / f"{TEST_ID}_{run_id}.docx",
        "figure_series": paper_root / "tables" / f"{TEST_ID}_{run_id}_figure_series.csv",
        "split_summary": paper_root / "tables" / f"{TEST_ID}_{run_id}_split_summary.csv",
        "writer_packet": paper_root / "writer_packets" / f"{TEST_ID}_{run_id}.md",
        "result_notes": paper_root / "snippets" / f"{TEST_ID}_{run_id}_result_notes.md",
    }
    for path in exports.values():
        path.parent.mkdir(parents=True, exist_ok=True)

    mapping = {
        run_dir / "figure_main.png": exports["figure_png"],
        run_dir / "figure_main.pdf": exports["figure_pdf"],
        run_dir / "table_main.csv": exports["table_csv"],
        run_dir / "table_main.md": exports["table_md"],
        run_dir / "table_main.tex": exports["table_tex"],
        run_dir / "table_main.docx": exports["table_docx"],
        run_dir / "figure_series.csv": exports["figure_series"],
        run_dir / "split_summary.csv": exports["split_summary"],
        run_dir / "writer_packet.md": exports["writer_packet"],
        run_dir / "result_notes.md": exports["result_notes"],
    }
    for src, dst in mapping.items():
        shutil.copy2(src, dst)
    return {key: str(path) for key, path in exports.items()}


def main() -> None:
    args = _parse_args()
    run_dir = args.test_root / args.run_id
    run_dir.mkdir(parents=True, exist_ok=True)

    signal_sample = _prepare_signal_sample(args.event_panel, args.annual_panel)
    monthly_returns = _load_monthly_returns(args.monthly_returns)
    market_index = _load_market_index(args.market_index)

    summary_rows: list[dict[str, object]] = []
    figure_series = None
    for horizon_months, horizon_label in HORIZONS.items():
        positions = _build_positions(signal_sample, horizon_months)
        series = _compute_long_short_series(positions, monthly_returns, market_index)
        for weight_col in ["ew_long_short", "vw_long_short"]:
            summary_rows.append(
                _summarize_long_short(series, weight_col, horizon_label=horizon_label)
            )
        if horizon_label == "3m":
            figure_series = series.copy()

    split_rows = _subperiod_summary(signal_sample, monthly_returns, market_index)
    table_df = pd.DataFrame(summary_rows)
    split_df = pd.DataFrame(split_rows)
    table_df.to_csv(run_dir / "table_main.csv", index=False)
    split_df.to_csv(run_dir / "split_summary.csv", index=False)
    figure_series.to_csv(run_dir / "figure_series.csv", index=False)

    table_md, table_tex = _render_table_outputs(summary_rows, split_rows)
    (run_dir / "table_main.md").write_text(table_md, encoding="utf-8")
    (run_dir / "table_main.tex").write_text(table_tex, encoding="utf-8")
    _build_table_docx(summary_rows, split_rows, run_dir / "table_main.docx")
    png_path, pdf_path = _plot_cumulative_long_short(figure_series, run_dir / "figure_main")

    (run_dir / "result_notes.md").write_text(
        _result_notes(summary_rows, split_rows), encoding="utf-8"
    )
    (run_dir / "writer_packet.md").write_text(
        _writer_packet(args, summary_rows, split_rows), encoding="utf-8"
    )

    dataset_summary = _dataset_summary(
        signal_sample, summary_rows, split_rows, figure_series, args=args
    )
    (run_dir / "dataset_summary.json").write_text(
        json.dumps(dataset_summary, indent=2), encoding="utf-8"
    )

    paper_exports = _copy_exports(run_dir, args.paper_root, args.run_id)
    manifest = {
        "test_id": TEST_ID,
        "run_id": args.run_id,
        "created_at_utc": datetime.now(UTC).isoformat(),
        "module_path": MODULE_PATH,
        "inputs": {
            "event_panel": str(args.event_panel),
            "annual_panel": str(args.annual_panel),
            "monthly_returns": str(args.monthly_returns),
            "market_index": str(args.market_index),
        },
        "run_dir": str(run_dir),
        "outputs": {
            "dataset_summary": str(run_dir / "dataset_summary.json"),
            "table_csv": str(run_dir / "table_main.csv"),
            "table_md": str(run_dir / "table_main.md"),
            "table_tex": str(run_dir / "table_main.tex"),
            "table_docx": str(run_dir / "table_main.docx"),
            "figure_series": str(run_dir / "figure_series.csv"),
            "split_summary": str(run_dir / "split_summary.csv"),
            "figure_png": str(png_path),
            "figure_pdf": str(pdf_path),
            "writer_packet": str(run_dir / "writer_packet.md"),
            "result_notes": str(run_dir / "result_notes.md"),
        },
        "paper_exports": paper_exports,
    }
    (run_dir / "run_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    print(f"[{TEST_ID}] wrote run bundle to {run_dir}")
    print(f"[{TEST_ID}] paper table: {paper_exports['table_docx']}")


if __name__ == "__main__":
    main()
